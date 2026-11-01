"""xlsx tool: workbooks made and changed without losing anything, with correct numbers.

``xlsx_excel_results.json`` holds what real Excel calculated for each formula; the
structural results were checked against Excel the same way. These tests pin that
behaviour without Excel.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from openpyxl import Workbook, load_workbook

from captain_claw import xlsx_formula as F
from captain_claw import xlsx_ops as X
from captain_claw.tools.document_extract import _extract_xlsx_markdown
from captain_claw.tools.edit import EditTool
from captain_claw.tools.write import WriteTool
from captain_claw.tools.xlsx_tool import XlsxTool, xlsx_file_name

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
WHOLE = F.Box(None, 1, 1, F.MAX_ROW, F.MAX_COL)
DATA_ROWS = [
    ["Name", "Qty", "Price", "Blank", None, "Flow", "When"],
    ["alpha", 10, 1, None, None, -1000, "2026-01-01"],
    ["beta", 12, 2, None, None, 600, "2026-06-30"],
    ["gamma", 15, 3, None, None, 700, "2027-03-15"],
    ["delta", 20, 4],
    ["epsilon", 30, 6],
]


def _book(*sheets, edits=None) -> X.Book:
    """A workbook saved and opened again, as the tool sees it."""
    book, _ = X.build_workbook(list(sheets), edits=edits)
    return X.Book(book.to_bytes())


def _edit(book: X.Book, *ops, sheet=None) -> tuple[X.Book, X.Report]:
    original = book.to_bytes()
    report = X.Editor(book, default_sheet=sheet).apply(list(ops))
    data = book.to_bytes()
    X.verify_saved(data, original)
    return X.Book(data), report


def _openpyxl_values(book: X.Book, sheet: str) -> list[list]:
    wb = load_workbook(io.BytesIO(book.to_bytes()), data_only=True)
    return [[c.value for c in row] for row in wb[sheet].iter_rows()]


def _parts(data: bytes) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return {n: archive.read(n) for n in archive.namelist()}


def _sparkline_workbook(tmp_path: Path) -> bytes:
    """What an openpyxl load-and-save loses: a sparkline, and a part nobody knows."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Model"
    for row in (["Item", "Q1"], ["A", 10], ["B", 20]):
        ws.append(row)
    ws["B4"] = "=SUM(B2:B3)"
    path = tmp_path / "plain.xlsx"
    wb.save(path)
    spark = ('<extLst><ext uri="{05C60535-1F16-4fd2-B633-F4F36F0B64E0}" '
             'xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main">'
             '<x14:sparklineGroups xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main">'
             '<x14:sparklineGroup><x14:sparklines><x14:sparkline><xm:f>Model!B2:B3</xm:f><xm:sqref>C2</xm:sqref>'
             '</x14:sparkline></x14:sparklines></x14:sparklineGroup></x14:sparklineGroups></ext></extLst>')
    out = io.BytesIO()
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for name in source.namelist():
            data = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                data = data.decode().replace("</worksheet>", spark + "</worksheet>").encode()
            target.writestr(name, data)
        target.writestr("xl/unknownVendorPart.bin", b"\x00\x01keep me")
    return out.getvalue()


# ── Formula engine ───────────────────────────────────────────────────


def _excel_pins() -> list[tuple[str, object]]:
    return [tuple(p) for p in json.loads((Path(__file__).parent / "xlsx_excel_results.json").read_text())]


@pytest.fixture(scope="module")
def data_book() -> X.Book:
    return _book({"name": "Data", "rows": DATA_ROWS}, {"name": "Funcs", "rows": [["x"]]})


@pytest.mark.parametrize("formula,expected", _excel_pins(), ids=[p[0][:50] for p in _excel_pins()])
def test_results_match_what_excel_calculated(data_book, formula, expected):
    stored = F.normalize_for_storage("=" + formula)
    array = F.needs_array(stored)
    calc = X.Calc(data_book)
    result = F.evaluate(F.compile_formula(stored), calc, "Funcs", 9, 2, array=array)
    if isinstance(result, F.Arr):
        result = result.rows[0][0]
    if isinstance(expected, dict):
        assert isinstance(result, F.XlError) and result.code == expected["error"]
    elif isinstance(expected, bool):
        assert result is expected
    elif isinstance(expected, float):
        assert result == pytest.approx(expected, rel=1e-6, abs=1e-9)
    else:
        assert result == expected


def test_new_functions_get_the_prefix_excel_needs():
    assert F.normalize_for_storage("=xlookup(1,A1:A2,B1:B2)") == "_xlfn.XLOOKUP(1,A1:A2,B1:B2)"
    assert F.normalize_for_storage("=STDEV.S(A1:A9)") == "_xlfn.STDEV.S(A1:A9)"
    assert F.normalize_for_storage("=SUM(A1:A2)") == "SUM(A1:A2)"
    assert F.display_formula("_xlfn.XLOOKUP(1,A1:A2,B1:B2)") == "=XLOOKUP(1,A1:A2,B1:B2)"


def test_semicolons_become_commas_but_not_inside_text_or_arrays():
    assert F.normalize_for_storage('=IF(A1>1;"a;b";2)') == 'IF(A1>1,"a;b",2)'
    assert F.normalize_for_storage("=SUM({1,2;3,4})") == "SUM({1,2;3,4})"


@pytest.mark.parametrize("formula,fragment", [
    ("=SUMM(A1)", "not an Excel function"),
    ("=SUM(A1", "too early"),
    ("=FILTER(A1:A2,B1:B2>1)", "isn't written by this tool"),
    ("=WEBSERVICE(\"http://x\")", "outside the workbook"),
    ("=[1]Sheet1!A1", "another workbook"),
    ("=cmd|'/c calc'!A1", "can't be read"),
])
def test_formulas_that_would_open_broken_are_refused(formula, fragment):
    with pytest.raises(F.FormulaError, match=fragment):
        F.normalize_for_storage(formula)


def test_operators_over_ranges_are_stored_as_array_formulas():
    assert F.needs_array("SUM(A1:A3*B1:B3)")
    assert F.needs_array("MAX(ABS(A1:A3))")
    assert F.needs_array("SUM(IF(A1:A3>0,1,0))")
    assert not F.needs_array("SUMPRODUCT(A1:A3,B1:B3)")
    assert not F.needs_array("SUMPRODUCT((A1:A3>0)*B1:B3)")
    assert not F.needs_array("SUM(A1:A3)")
    assert not F.needs_array("A1*B1")


@pytest.mark.parametrize("formula,axis,at,count,expected", [
    ("SUM(B2:B5)+$C$7", "row", 4, 2, "SUM(B2:B7)+$C$9"),
    ("SUM(B2:B5)", "row", 2, 1, "SUM(B3:B6)"),
    ("SUM(B2:B5)", "row", 6, 1, "SUM(B2:B5)"),
    ("SUM(B2:B5)+C3", "row", 3, -1, "SUM(B2:B4)+#REF!"),
    ("B3+SUM(B3:B3)", "row", 3, -1, "#REF!+SUM(#REF!)"),
    ("SUM(B2:B5)", "row", 2, -2, "SUM(B2:B3)"),
    ("SUM(A:C)+D2", "col", 2, 1, "SUM(A:D)+E2"),
    ("SUM(A:A)", "row", 1, 3, "SUM(A:A)"),
    ("SUM(2:4)", "row", 3, 1, "SUM(2:5)"),
    ("Other!B2+'My S'!B2+B2", "row", 1, 1, "Other!B2+'My S'!B2+B3"),
])
def test_references_move_the_way_excel_moves_them(formula, axis, at, count, expected):
    assert F.shift_formula(formula, "Sheet", "Sheet", axis, at, count)[0] == expected


def test_references_on_other_sheets_move_by_name():
    assert F.shift_formula("Sheet!B2+B2", "Other", "Sheet", "row", 1, 1)[0] == "Sheet!B3+B2"
    assert F.shift_formula("'Sheet'!$B$2", "Other", "SHEET", "row", 1, 1)[0] == "'Sheet'!$B$3"


def test_renaming_rewrites_only_that_sheet():
    assert F.rename_sheet_in("SUM(Old!A1:A2)+'old'!B1+Older!C1", "Old", "New Name") == \
        "SUM('New Name'!A1:A2)+'New Name'!B1+Older!C1"


@pytest.mark.parametrize("value,fmt,shown", [
    (1234.567, "#,##0.00", "1,234.57"),
    (0.256, "0%", "26%"),
    (-1200, "$#,##0;($#,##0)", "($1,200)"),
    (46000, "dd.mm.yyyy", "09.12.2025"),
    (46000, "mmm-yy", "Dec-25"),
    (12345, "0.00E+00", "1.23E+04"),
    (1234567, "#,##0,", "1,235"),
    (12.5, "[$€-407] #,##0.00", "€ 12.50"),
    (0.75, "h:mm AM/PM", "6:00 PM"),
    (0, "0.00;-0.00;\"-\"", "-"),
])
def test_number_formats_read_back(value, fmt, shown):
    assert F.format_value(value, fmt) == shown


def test_the_1900_leap_year_bug_is_kept():
    assert F.serial_to_date(60 + 1).isoformat() == "1900-03-01"
    assert F.date_to_serial(F.serial_to_date(46000)) == 46000


# ── Nothing is lost ──────────────────────────────────────────────────


def test_a_one_cell_edit_leaves_every_other_part_alone(tmp_path):
    original = _sparkline_workbook(tmp_path)
    before = _parts(original)
    book, report = _edit(X.Book(original), {"op": "set", "cell": "B2", "value": 11})
    after = _parts(book.to_bytes())
    assert b"sparkline" in after["xl/worksheets/sheet1.xml"]
    assert after["xl/unknownVendorPart.bin"] == before["xl/unknownVendorPart.bin"]
    assert after["xl/styles.xml"] == before["xl/styles.xml"]
    assert after["xl/theme/theme1.xml"] == before["xl/theme/theme1.xml"]
    assert [c.cell for c in report.written] == ["B2"]
    assert [(c.cell, c.before, c.after) for c in report.moved] == [("B4", "(no result yet)", "31")]


def test_reading_writes_nothing(tmp_path):
    original = _sparkline_workbook(tmp_path)
    book = X.Book(original)
    book.all_sheets()
    X.check_book(book)
    assert not book._dirty


def test_a_sheet_nobody_touched_keeps_its_bytes():
    book = _book({"name": "One", "rows": [[1]]}, {"name": "Two", "rows": [[2]]})
    before = _parts(book.to_bytes())
    edited, _ = _edit(book, {"op": "set", "sheet": "One", "cell": "A1", "value": 5})
    after = _parts(edited.to_bytes())
    assert after["xl/worksheets/sheet2.xml"] == before["xl/worksheets/sheet2.xml"]


def test_moving_rows_leaves_a_sheet_that_doesnt_read_them_alone():
    book = _book({"name": "One", "rows": [[1], [2]]}, {"name": "Two", "rows": [[3], [4]]})
    two = book.sheet("Two")
    first = two.cell(1, 2)
    etree.SubElement(first, X.FT, t="shared", si="0", ref="B1:B2").text = "A1*2"
    etree.SubElement(first, X.V).text = "6"
    second = two.cell(2, 2)
    etree.SubElement(second, X.FT, t="shared", si="0")
    etree.SubElement(second, X.V).text = "8"
    book.touch(two.path)
    book = X.Book(book.to_bytes())
    before = _parts(book.to_bytes())["xl/worksheets/sheet2.xml"]
    edited, _ = _edit(book, {"op": "insert_rows", "sheet": "One", "at": 1})
    assert _parts(edited.to_bytes())["xl/worksheets/sheet2.xml"] == before


def test_a_stale_calculation_order_is_dropped():
    book = _book({"name": "S", "rows": [[1, "=A1+1"]]})
    data = _parts(book.to_bytes())
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for name, raw in data.items():
            if name == "[Content_Types].xml":
                raw = raw.replace(b"</Types>", b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
            if name == "xl/_rels/workbook.xml.rels":
                raw = raw.replace(b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
            target.writestr(name, raw)
        target.writestr("xl/calcChain.xml", f'<calcChain xmlns="{MAIN}"><c r="B1" i="1"/></calcChain>')
    edited, _ = _edit(X.Book(out.getvalue()), {"op": "clear", "range": "B1"})
    parts = _parts(edited.to_bytes())
    assert "xl/calcChain.xml" not in parts
    assert b"calcChain" not in parts["xl/_rels/workbook.xml.rels"]
    assert b"calcChain" not in parts["[Content_Types].xml"]


def test_excel_recalculates_on_open():
    book = _book({"name": "S", "rows": [[1, "=A1+1"]]})
    assert b'fullCalcOnLoad="1"' in _parts(book.to_bytes())["xl/workbook.xml"]


# ── Results ──────────────────────────────────────────────────────────


def test_created_formulas_store_their_results():
    book = _book({"name": "S", "rows": [["A", 10], ["B", 20], ["Sum", "=SUM(B1:B2)"]]})
    assert _openpyxl_values(book, "S")[2] == ["Sum", 30]


def test_a_change_recalculates_everything_downstream_across_sheets():
    book = _book({"name": "In", "rows": [[0.1]]},
                 {"name": "Out", "rows": [["=In!A1*100", "=A1+1", "=SUM(A1:B1)"]]})
    edited, report = _edit(book, {"op": "set", "sheet": "In", "cell": "A1", "value": 0.2})
    assert _openpyxl_values(edited, "Out") == [[20, 21, 41]]
    assert {c.cell: c.after for c in report.moved} == {"A1": "20", "B1": "21", "C1": "41"}


def test_a_long_chain_of_running_totals_does_not_overflow():
    rows = [[1, "=A1"]] + [[1, f"=B{i}+A{i + 1}"] for i in range(1, 6000)]
    book = _book({"name": "S", "rows": rows})
    edited, _ = _edit(book, {"op": "set", "cell": "A1", "value": 2})
    assert edited.sheet("S").stored(6000, 2) == 6001


def test_what_cant_be_calculated_keeps_its_result_until_its_input_changes():
    book = _book({"name": "S", "rows": [[1, 2]]})
    sheet = book.sheet("S")
    cell = sheet.cell(1, 3)
    etree.SubElement(cell, X.FT).text = "_xlfn.LET(_xlpm.x,A1,_xlpm.x*2)"
    etree.SubElement(cell, X.V).text = "2"
    other = sheet.cell(2, 3)
    etree.SubElement(other, X.FT).text = "CELL(\"width\",B9)"
    etree.SubElement(other, X.V).text = "8"
    book.touch(sheet.path)
    book = X.Book(book.to_bytes())

    # An edit elsewhere doesn't touch the stored result.
    edited, report = _edit(book, {"op": "set", "cell": "B1", "value": 5})
    assert edited.sheet("S").stored(1, 3) == 2
    assert not report.not_calculated

    # An edit to its input does: a stale number is worse than none.
    edited, report = _edit(edited, {"op": "set", "cell": "A1", "value": 9})
    assert not edited.sheet("S").has_stored_result(1, 3)
    assert report.not_calculated[0][1] == "C1"
    assert "LET" in report.not_calculated[0][2]
    assert edited.sheet("S").formula(1, 3).startswith("_xlfn.LET")


def test_a_link_to_another_workbook_keeps_its_last_value():
    book = _book({"name": "S", "rows": [[1, 2]]})
    sheet = book.sheet("S")
    cell = sheet.cell(1, 3)
    etree.SubElement(cell, X.FT).text = "[1]Other!A1*2"
    etree.SubElement(cell, X.V).text = "42"
    book.touch(sheet.path)
    edited, report = _edit(X.Book(book.to_bytes()), {"op": "set", "cell": "A1", "value": 9})
    assert edited.sheet("S").stored(1, 3) == 42
    assert not report.not_calculated


def test_a_new_error_is_reported_with_its_formula():
    book = _book({"name": "S", "rows": [[10, 2, "=A1/B1"]]})
    edited, report = _edit(book, {"op": "set", "cell": "B1", "value": 0})
    assert report.new_errors == [("S", "C1", "#DIV/0!", "=A1/B1")]
    assert isinstance(edited.sheet("S").stored(1, 3), F.XlError)


def test_circular_references_are_named_not_hung_on():
    book, report = X.build_workbook([{"name": "S", "rows": [["=B1", "=A1"], ["=A2"]]}])
    assert any("Circular" in n for n in report.notes)
    assert not report.new_errors


def test_array_formulas_work_on_the_whole_range():
    book = _book({"name": "S", "rows": [[1, 2], [3, 4], ["=SUM(A1:A2*B1:B2)"]]})
    sheet = book.sheet("S")
    assert sheet.stored(3, 1) == 14
    assert (3, 1) in sheet.array_masters


def test_shared_formulas_are_read_and_stay_right_after_an_edit():
    book = _book({"name": "S", "rows": [[1], [2], [3]]})
    sheet = book.sheet("S")
    first = sheet.cell(1, 2)
    f = etree.SubElement(first, X.FT, t="shared", si="0", ref="B1:B3")
    f.text = "A1*2"
    for row in (2, 3):
        member = sheet.cell(row, 2)
        etree.SubElement(member, X.FT, t="shared", si="0")
    book.touch(sheet.path)
    book = X.Book(book.to_bytes())
    assert book.sheet("S").formula(3, 2) == "A3*2"
    edited, report = _edit(book, {"op": "set", "cell": "A3", "value": 10}, {"op": "insert_rows", "at": 2})
    sheet = edited.sheet("S")
    assert sheet.formula(4, 2) == "A4*2"
    assert sheet.stored(4, 2) == 20
    assert sheet.formula(1, 2) == "A1*2"


def test_dates_are_real_dates_in_a_1904_workbook_too():
    book = _book({"name": "S", "rows": [["2026-03-31", "=A1+1", "=YEAR(A1)"]]})
    assert _openpyxl_values(book, "S")[0][0].isoformat().startswith("2026-03-31")
    assert book.sheet("S").stored(1, 3) == 2026
    raw = _parts(book.to_bytes())
    workbook = raw["xl/workbook.xml"].replace(b"<workbookPr", b'<workbookPr date1904="1"', 1)
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for name, data in raw.items():
            target.writestr(name, workbook if name == "xl/workbook.xml" else data)
    edited, _ = _edit(X.Book(out.getvalue()), {"op": "set", "cell": "D1", "value": "2026-03-31"},
                      {"op": "set", "cell": "E1", "value": "=YEAR(D1)"})
    sheet = edited.sheet("S")
    assert sheet.stored(1, 4) == F.date_to_serial(F.serial_to_date(46112), True) or sheet.stored(1, 5) == 2026
    assert sheet.stored(1, 5) == 2026


# ── What a value means ───────────────────────────────────────────────


@pytest.mark.parametrize("raw,stored", [
    (5, 5), ("12", 12), ("-3.5", -3.5), ("0.5", 0.5), ("12%", 0.12), (True, True), ("TRUE", True),
    ("007", "007"), ("hello", "hello"), ("1,200", "1,200"), ("12345678901234567890", "12345678901234567890"),
])
def test_a_value_means_what_typing_it_means(raw, stored):
    book = _book({"name": "S", "rows": [[raw]]})
    assert book.sheet("S").stored(1, 1) == stored


def test_as_text_keeps_the_text():
    book = _book({"name": "S", "rows": [["x"]]})
    edited, _ = _edit(book, {"op": "set", "cell": "A1", "value": "=1+1", "as_text": True},
                      {"op": "set", "cell": "A2", "value": "1200", "as_text": True})
    sheet = edited.sheet("S")
    assert sheet.stored(1, 1) == "=1+1" and sheet.formula(1, 1) is None
    assert sheet.stored(2, 1) == "1200"


def test_control_characters_and_spaces_survive():
    book = _book({"name": "S", "rows": [["a\x01b", "  padded  ", "two\nlines"]]})
    sheet = book.sheet("S")
    assert sheet.stored(1, 1) == "a\x01b"
    assert sheet.stored(1, 2) == "  padded  "
    assert sheet.stored(1, 3) == "two\nlines"


# ── Moving rows, columns and sheets ──────────────────────────────────


def _model() -> X.Book:
    return _book(
        {"name": "In", "rows": [["Rate", 0.1]]},
        {"name": "Sales", "rows": [["R", "Q1", "Q2", "T"], ["N", 1, 2, "=SUM(B2:C2)"], ["S", 3, 4, "=SUM(B3:C3)"],
                                   ["Tot", "=SUM(B2:B3)", "=SUM(C2:C3)", "=SUM(D2:D3)"]]},
        {"name": "Sum", "rows": [["=Sales!D4", "=A1*(1+In!B1)"]]},
        edits=[{"op": "define_name", "name": "Top", "refers_to": "Sales!$D$4"},
               {"op": "merge", "sheet": "Sales", "range": "F2:G3"},
               {"op": "comment", "sheet": "Sales", "cell": "B3", "text": "check"},
               {"op": "set", "sheet": "Sum", "cell": "C1", "value": "=Top*2"}])


def test_inserting_rows_keeps_every_reference_right():
    edited, _ = _edit(_model(), {"op": "insert_rows", "sheet": "Sales", "at": 3, "count": 2},
                      {"op": "set", "sheet": "Sales", "cell": "A3", "rows": [["M", 5, 6, "=SUM(B3:C3)"]]})
    sales = edited.sheet("Sales")
    assert sales.formula(6, 2) == "SUM(B2:B5)"
    assert sales.stored(6, 4) == 21
    assert sales.formula(5, 4) == "SUM(B5:C5)"
    assert edited.sheet("Sum").formula(1, 1) == "Sales!D6"
    assert edited.sheet("Sum").stored(1, 1) == 21
    assert edited.defined_name_text("Top", None) == "Sales!$D$6"
    assert edited.sheet("Sum").stored(1, 3) == 42
    assert X.box_text(sales.merges[0]) == "F2:G5"
    assert "B5" in {c.get("ref") for c in edited.part("xl/comments1.xml").iter(f"{{{MAIN}}}comment")}


def test_deleting_rows_breaks_only_what_read_them():
    edited, report = _edit(_model(), {"op": "delete_rows", "sheet": "Sales", "at": 3})
    sales = edited.sheet("Sales")
    assert sales.formula(3, 2) == "SUM(B2:B2)"
    assert sales.stored(3, 4) == 3
    assert edited.sheet("Sum").stored(1, 1) == 3
    assert not report.new_errors
    edited, report = _edit(edited, {"op": "delete_rows", "sheet": "Sales", "rows": "3"})
    assert report.new_errors and report.new_errors[0][2] == "#REF!"
    assert "#REF!" in edited.sheet("Sum").formula(1, 1)


def test_columns_move_too_with_their_widths():
    book = _book({"name": "S", "rows": [[1, 2, "=A1+B1"]]}, edits=[{"op": "width", "sheet": "S", "columns": "B", "width": 30}])
    edited, _ = _edit(book, {"op": "insert_columns", "at": "B"}, {"op": "set", "cell": "B1", "value": 5})
    sheet = edited.sheet("S")
    assert sheet.formula(1, 4) == "A1+C1"
    assert sheet.stored(1, 4) == 3
    widths = {(c.get("min"), c.get("max")): c.get("width") for c in sheet.root.find(X._q("cols"))}
    assert widths[("3", "3")] == "30"
    edited, _ = _edit(edited, {"op": "delete_columns", "at": "B"})
    assert edited.sheet("S").formula(1, 3) == "A1+B1"


def test_formats_validations_and_filters_follow():
    book = _book({"name": "S", "rows": [[1], [2], [3], [4]]})
    sheet = book.sheet("S")
    root = sheet.root
    etree.SubElement(root, X._q("autoFilter"), ref="A1:A4")
    root.remove(root.find(X._q("autoFilter")))
    X._insert_ordered(root, etree.Element(X._q("autoFilter"), ref="A1:A4"), X._SHEET_ORDER)
    cf = etree.Element(X._q("conditionalFormatting"), sqref="A2:A4")
    rule = etree.SubElement(cf, X._q("cfRule"), type="expression", priority="1")
    etree.SubElement(rule, X._q("formula")).text = "A2>$A$3"
    X._insert_ordered(root, cf, X._SHEET_ORDER)
    dv_parent = etree.Element(X._q("dataValidations"), count="1")
    dv = etree.SubElement(dv_parent, X._q("dataValidation"), type="whole", sqref="A3")
    etree.SubElement(dv, X._q("formula1")).text = "$A$1"
    X._insert_ordered(root, dv_parent, X._SHEET_ORDER)
    book.touch(sheet.path)
    edited, _ = _edit(X.Book(book.to_bytes()), {"op": "insert_rows", "at": 2})
    root = edited.sheet("S").root
    assert root.find(X._q("autoFilter")).get("ref") == "A1:A5"
    cf = root.find(X._q("conditionalFormatting"))
    assert cf.get("sqref") == "A3:A5"
    assert next(cf.iter(X._q("formula"))).text == "A3>$A$4"
    dv = next(root.iter(X._q("dataValidation")))
    assert dv.get("sqref") == "A4" and dv.find(X._q("formula1")).text == "$A$1"


def test_a_new_row_takes_the_format_of_the_row_above():
    book = _book({"name": "S", "rows": [[1], [2]]}, edits=[{"op": "format", "sheet": "S", "range": "A1", "bold": True}])
    style = book.sheet("S").style_of(1, 1)
    edited, _ = _edit(book, {"op": "insert_rows", "at": 2})
    assert edited.sheet("S").style_of(2, 1) == style


def test_charts_and_their_anchors_follow():
    book = _book({"name": "S", "rows": [["L", "V"], ["a", 1], ["b", 2], ["c", 3]]},
                 edits=[{"op": "add_chart", "sheet": "S", "type": "column", "categories": "A2:A4", "values": "B2:B4", "anchor": "D6"}])
    edited, _ = _edit(book, {"op": "insert_rows", "at": 3}, {"op": "set", "cell": "B3", "value": 9})
    parts = _parts(edited.to_bytes())
    assert b"S!$B$2:$B$5" in parts["xl/charts/chart1.xml"]
    assert b"<xdr:row>6</xdr:row>" in parts["xl/drawings/drawing1.xml"]
    root = etree.fromstring(parts["xl/charts/chart1.xml"])
    c = f"{{{X.C_NS}}}"
    cache = next(root.iter(c + "numCache"))
    assert [p.find(c + "v").text for p in cache.findall(c + "pt")] == ["1", "9", "2", "3"]


def test_renaming_a_sheet_rewrites_formulas_names_and_charts():
    book = _book({"name": "Old", "rows": [["a", 1], ["b", 2]]}, {"name": "Use", "rows": [["=SUM(Old!B1:B2)"]]},
                 edits=[{"op": "define_name", "name": "N", "refers_to": "Old!$B$1"},
                        {"op": "add_chart", "sheet": "Old", "categories": "A1:A2", "values": "B1:B2"}])
    edited, _ = _edit(book, {"op": "rename_sheet", "sheet": "Old", "name": "New Name"})
    assert edited.sheet("Use").formula(1, 1) == "SUM('New Name'!B1:B2)"
    assert edited.defined_name_text("N", None) == "'New Name'!$B$1"
    assert b"'New Name'!$B$1:$B$2" in _parts(edited.to_bytes())["xl/charts/chart1.xml"]


def test_a_sheet_in_use_cant_be_deleted_until_nothing_reads_it():
    book = _book({"name": "A", "rows": [[1]]}, {"name": "B", "rows": [["=A!A1"]]})
    with pytest.raises(X.XlsxError, match="B!A1 read from sheet 'A'"):
        X.Editor(book).apply([{"op": "delete_sheet", "sheet": "A"}])
    book = _book({"name": "A", "rows": [[1]]}, {"name": "B", "rows": [["=A!A1"]]})
    edited, _ = _edit(book, {"op": "set", "sheet": "B", "cell": "A1", "value": 1}, {"op": "delete_sheet", "sheet": "A"})
    assert [s.name for s in edited.sheets] == ["B"]
    assert not any("sheet1.xml" in n for n in _parts(edited.to_bytes()))


def test_a_sheet_a_chart_elsewhere_reads_isnt_deleted_but_its_own_chart_goes_with_it():
    book = _book({"name": "Data", "rows": [["a", 1]]}, {"name": "View", "rows": [["x"]]},
                 edits=[{"op": "add_chart", "sheet": "View", "categories": "Data!A1", "values": "Data!B1"}])
    with pytest.raises(X.XlsxError, match="a chart on another sheet"):
        X.Editor(book).apply([{"op": "delete_sheet", "sheet": "Data"}])
    edited, _ = _edit(book, {"op": "delete_sheet", "sheet": "View"})
    assert not [n for n in _parts(edited.to_bytes()) if "chart" in n or "drawing" in n]


def test_deleting_and_moving_sheets_keeps_sheet_scoped_names_pointing_right():
    book = _book({"name": "A", "rows": [[1]]}, {"name": "B", "rows": [[2]]}, {"name": "C", "rows": [[3]]},
                 edits=[{"op": "define_name", "sheet": "C", "name": "Here", "refers_to": "C!$A$1"}])
    edited, _ = _edit(book, {"op": "move_sheet", "sheet": "C", "position": 1})
    assert [s.name for s in edited.sheets] == ["C", "A", "B"]
    assert edited.defined_name_text("Here", "C") == "C!$A$1"
    assert edited.defined_name_text("Here", "A") is None
    edited, _ = _edit(edited, {"op": "delete_sheet", "sheet": "B"})
    assert edited.defined_name_text("Here", "C") == "C!$A$1"


def test_the_last_visible_sheet_stays():
    book = _book({"name": "Only", "rows": [[1]]})
    with pytest.raises(X.XlsxError, match="at least one visible"):
        X.Editor(book).apply([{"op": "delete_sheet", "sheet": "Only"}])


# ── Tables ───────────────────────────────────────────────────────────


def _stock() -> X.Book:
    return _book({"name": "S", "rows": [["Item", "Qty", "Price", "Value"], ["Pen", 2, 1.5, "=B2*C2"], ["Ink", 3, 4, "=B3*C3"]]},
                 edits=[{"op": "add_table", "sheet": "S", "range": "A1:D3", "name": "Stock"},
                        {"op": "set", "sheet": "S", "cell": "F1", "value": "=SUM(Stock[Value])"},
                        {"op": "set", "sheet": "S", "cell": "F2", "value": "=SUM(D2:D3)"},
                        {"op": "set", "sheet": "S", "cell": "E2", "value": "=Stock[@Qty]*2"}])


def test_structured_references_calculate():
    book = _stock()
    sheet = book.sheet("S")
    assert sheet.stored(1, 6) == 15
    assert sheet.stored(2, 5) == 4
    load_workbook(io.BytesIO(book.to_bytes()))


def test_a_table_grows_and_fills_its_calculated_column():
    edited, report = _edit(_stock(), {"op": "append_rows", "table": "Stock", "rows": [["Cap", 7, 3]]})
    sheet = edited.sheet("S")
    table = edited.table_named("Stock")
    assert X.box_text(table.box) == "A1:D4"
    assert sheet.formula(4, 4) == "B4*C4" and sheet.stored(4, 4) == 21
    assert sheet.stored(1, 6) == 36
    assert sheet.formula(2, 6) == "SUM(D2:D4)", "a plain range over the whole column stretches, as in Excel"
    assert sheet.stored(2, 6) == 36
    assert edited.part(table.part).find(X._q("autoFilter")).get("ref") == "A1:D4"


def test_a_table_wont_grow_over_something():
    book = _stock()
    edited, _ = _edit(book, {"op": "set", "cell": "A4", "value": "in the way"})
    with pytest.raises(X.XlsxError, match="isn't empty"):
        X.Editor(edited).apply([{"op": "append_rows", "table": "Stock", "rows": [["x"]]}])


def test_rows_inserted_inside_a_table_stretch_it():
    edited, _ = _edit(_stock(), {"op": "insert_rows", "at": 3})
    assert X.box_text(edited.table_named("Stock").box) == "A1:D4"


def test_a_table_header_stays_in_step_with_its_cell():
    book = _book({"name": "S", "rows": [["A", "B"], [1, 2]]}, edits=[{"op": "add_table", "sheet": "S", "range": "A1:B2", "name": "T"}])
    edited, _ = _edit(book, {"op": "set", "cell": "B1", "value": "Renamed"})
    assert edited.table_named("T").columns == ["A", "Renamed"]
    with pytest.raises(X.XlsxError, match="must stay text"):
        X.Editor(edited).apply([{"op": "set", "cell": "A1", "value": "=1+1"}])
    with pytest.raises(X.XlsxError, match="already has a column"):
        X.Editor(edited).apply([{"op": "set", "cell": "A1", "value": "renamed"}])


def test_a_number_like_header_stays_text():
    book = _book({"name": "S", "rows": [["A", "B"], [1, 2]]}, edits=[{"op": "add_table", "sheet": "S", "range": "A1:B2", "name": "T"}])
    edited, _ = _edit(book, {"op": "set", "cell": "B1", "value": "2026"})
    assert edited.sheet("S").stored(1, 2) == "2026"
    assert edited.table_named("T").columns == ["A", "2026"]


def test_a_table_needs_real_headers():
    book = _book({"name": "S", "rows": [["A", 5], [1, 2]]})
    with pytest.raises(X.XlsxError, match="must hold text"):
        X.Editor(book).apply([{"op": "add_table", "range": "A1:B2"}])


# ── Guards: all or nothing, nothing silently wrong ───────────────────


def test_a_bad_edit_fails_the_whole_call():
    book = _book({"name": "S", "rows": [[1]]})
    with pytest.raises(X.XlsxError, match=r"Edit 2 \(set\).*SUMM"):
        X.Editor(book).apply([{"op": "set", "cell": "A1", "value": 2}, {"op": "set", "cell": "B1", "value": "=SUMM(A1)"}])


def test_formulas_can_name_a_sheet_made_later_in_the_call():
    book, report = X.build_workbook([{"name": "A", "rows": [["=B!A1"]]}, {"name": "B", "rows": [[5]]}])
    assert X.Book(book.to_bytes()).sheet("A").stored(1, 1) == 5
    assert not report.new_errors


@pytest.mark.parametrize("formula,message", [
    ("=Nope!A1", "no sheet 'Nope'"),
    ("=GrowthRate*2", "neither a cell address nor a defined name"),
    ("=Missing[Col]", "no table"),
])
def test_references_to_nothing_are_refused(formula, message):
    with pytest.raises(X.XlsxError, match=message):
        X.build_workbook([{"name": "S", "rows": [[formula]]}])


def test_several_sheets_need_a_sheet_name():
    book = _book({"name": "A", "rows": [[1]]}, {"name": "B", "rows": [[2]]})
    with pytest.raises(X.XlsxError, match="Say which sheet"):
        X.Editor(book).apply([{"op": "set", "cell": "A1", "value": 9}])
    edited, _ = _edit(book, {"op": "set", "cell": "B!A2", "value": 9})
    assert edited.sheet("B").stored(2, 1) == 9


def test_merged_cells_and_protected_sheets_are_respected():
    book = _book({"name": "S", "rows": [["x"]]}, edits=[{"op": "merge", "sheet": "S", "range": "B1:C1"}])
    with pytest.raises(X.XlsxError, match="merged"):
        X.Editor(book).apply([{"op": "set", "cell": "C1", "value": 1}])
    sheet = book.sheet("S")
    X._insert_ordered(sheet.root, etree.Element(X._q("sheetProtection"), sheet="1"), X._SHEET_ORDER)
    book.touch(sheet.path)
    with pytest.raises(X.XlsxError, match="protected"):
        X.Editor(X.Book(book.to_bytes())).apply([{"op": "set", "cell": "A1", "value": 1}])


def test_a_merge_wont_lose_a_value():
    book = _book({"name": "S", "rows": [["a", "b"]]})
    with pytest.raises(X.XlsxError, match="would lose"):
        X.Editor(book).apply([{"op": "merge", "range": "A1:B1"}])


def test_a_total_that_stops_short_is_pointed_out():
    book = _book({"name": "S", "rows": [[1, None, "=SUM(A1:A2)"], [2]]})
    _, report = _edit(book, {"op": "set", "cell": "A3", "value": 5})
    assert any("C1" in n and "stops just before" in n for n in report.notes)


def test_appending_below_a_totals_row_is_pointed_out():
    book = _book({"name": "S", "rows": [[1], [2], ["=SUM(A1:A2)"]]})
    _, report = _edit(book, {"op": "append_rows", "rows": [[3]]})
    assert any("totals row" in n for n in report.notes)


def test_not_a_workbook_gets_a_clear_message():
    with pytest.raises(X.XlsxError, match="password-protected"):
        X.Book(b"\xd0\xcf\x11\xe0" + b"0" * 100)
    with pytest.raises(X.XlsxError, match="Not a readable"):
        X.Book(b"plain text")


# ── Check ────────────────────────────────────────────────────────────


def _issues(book: X.Book) -> dict[str, list[X.Issue]]:
    out: dict[str, list[X.Issue]] = {}
    for issue in X.check_book(book)[0]:
        out.setdefault(issue.kind, []).append(issue)
    return out


def test_a_clean_workbook_has_nothing_to_report():
    book = _book({"name": "S", "rows": [["a", 1], ["b", 2], ["t", "=SUM(B1:B2)"]]})
    assert X.check_book(book)[0] == []


def test_check_finds_a_typed_in_number_among_formulas():
    book = _book({"name": "S", "rows": [[1, "=A1*2"], [2, 99], [3, "=A3*2"]]})
    assert _issues(book)["hardcoded_number"][0].where == "S!B2"


def test_check_finds_a_formula_that_breaks_the_pattern():
    book = _book({"name": "S", "rows": [[1, "=A1*2"], [2, "=A2*3"], [3, "=A3*2"]]})
    assert _issues(book)["inconsistent_formula"][0].where == "S!B2"


def test_check_finds_a_sum_that_stops_short():
    book = _book({"name": "S", "rows": [[1, "=SUM(A1:A2)"], [2], [3]]})
    assert _issues(book)["range_stops_short"][0].where == "S!B1"


def test_check_finds_numbers_stored_as_text():
    book = _book({"name": "S", "rows": [[1], [2]]})
    edited, _ = _edit(book, {"op": "set", "cell": "A3", "value": "3", "as_text": True})
    assert "A3" in _issues(edited)["text_number"][0].what


def test_check_finds_errors_once_at_their_source():
    book = _book({"name": "S", "rows": [[1, 0, "=A1/B1", "=C1+1", "=D1+1"]]})
    found = _issues(book)["error_value"]
    assert [i.where for i in found if i.level == "error"] == ["S!C1"]
    assert "2 more" in [i.what for i in found if i.level == "info"][0]


def test_check_finds_broken_references():
    book = _book({"name": "S", "rows": [[1, "=A1+1"], [2]]})
    edited, _ = _edit(book, {"op": "delete_columns", "at": "A"})
    assert "S!A1" in [i.where for i in _issues(edited)["broken_reference"]]


def test_check_finds_a_stale_result_a_script_left_behind():
    book = _book({"name": "S", "rows": [[1, "=A1+1"]]})
    sheet = book.sheet("S")
    sheet.cell_el(1, 2).find(X.V).text = "999"
    book.touch(sheet.path)
    found = _issues(X.Book(book.to_bytes()))["stale_result"][0]
    assert found.where == "S!B1" and "999" in found.what and "2" in found.what


def test_a_volatile_formula_is_not_called_stale():
    book = _book({"name": "S", "rows": [["=TODAY()"]]})
    sheet = book.sheet("S")
    sheet.cell_el(1, 1).find(X.V).text = "1"
    book.touch(sheet.path)
    assert "stale_result" not in _issues(X.Book(book.to_bytes()))


def test_check_warns_about_number_patterns_inside_text():
    book = _book({"name": "S", "rows": [[3.14, '=TEXT(A1,"0.00")', '=TEXT(A1,"yyyy-mm-dd")', '="a.b"&A1']]})
    found = _issues(book)["text_format"]
    assert found[0].where == "S!B1" and "1 formula(s)" in found[0].what


def test_check_flags_circular_references():
    book, _ = X.build_workbook([{"name": "S", "rows": [["=B1", "=A1"]]}])
    assert _issues(X.Book(book.to_bytes()))["circular"]


def test_check_mentions_rates_typed_into_formulas_and_hidden_rows():
    book = _book({"name": "S", "rows": [[100, "=A1*1.25"], [1]]})
    sheet = book.sheet("S")
    sheet.row_els[2].set("hidden", "1")
    book.touch(sheet.path)
    found = _issues(X.Book(book.to_bytes()))
    assert "1.25" in found["number_in_formula"][0].what
    assert "hidden" in found


# ── The tool ─────────────────────────────────────────────────────────


@pytest.fixture
def workspace(tmp_path):
    (tmp_path / "saved" / "downloads" / "s1").mkdir(parents=True)
    kw = {"_runtime_base_path": tmp_path, "_saved_base_path": tmp_path / "saved", "_session_id": "s1"}
    return tmp_path, kw


async def test_create_read_edit_check_through_the_tool(workspace):
    root, kw = workspace
    tool = XlsxTool()
    created = await tool.execute(
        action="create", path="Budget", title="Budget",
        sheets=[{"name": "Model", "header": True,
                 "rows": [["Item", "Amount"], ["Rent", 1000], ["Food", 400], ["Total", "=SUM(B2:B3)"]]}],
        edits=[{"op": "format", "sheet": "Model", "range": "B2:B4", "number_format": "#,##0.00"}], **kw)
    assert created.success, created.error
    assert "Model!B4: =SUM(B2:B3) → 1400" in created.content
    saved = next(root.rglob("Budget.xlsx"))

    read = await tool.execute(action="read", path=str(saved), **kw)
    assert 'A2 "Rent" | B2 1000 (1,000.00)' in read.content
    assert "B4 =SUM(B2:B3) → 1400 (1,400.00)" in read.content

    edited = await tool.execute(action="edit", path=str(saved), sheet="Model",
                                edits=[{"op": "set", "cell": "B2", "value": 1500}], **kw)
    assert edited.success, edited.error
    assert "Model!B2: 1000 (1,000.00) ⇒ 1500 (1,500.00)" in edited.content
    assert "Model!B4: 1400 (1,400.00) ⇒ 1900 (1,900.00)" in edited.content
    assert "Saved in place" in edited.content and "Backup:" in edited.content

    checked = await tool.execute(action="check", path=str(saved), **kw)
    assert "0 error(s), 0 warning(s), 0 note(s)" in checked.content

    found = await tool.execute(action="read", path=str(saved), find="rent", **kw)
    assert "Model!A2" in found.content


async def test_read_pages_through_a_long_sheet(workspace):
    root, kw = workspace
    tool = XlsxTool()
    await tool.execute(action="create", path="Long.xlsx", sheets=[{"name": "S", "rows": [[i] for i in range(1, 151)]}], **kw)
    saved = next(root.rglob("Long.xlsx"))
    first = await tool.execute(action="read", path=str(saved), limit=100, **kw)
    assert "A100 100" in first.content and "A101" not in first.content.split("… more")[0]
    assert "range=A101:A150" in first.content
    second = await tool.execute(action="read", path=str(saved), sheet="S", range="A101:A150", **kw)
    assert "A150 150" in second.content


async def test_an_upload_is_never_changed_and_a_copy_never_forks(workspace):
    root, kw = workspace
    upload = root / "saved" / "downloads" / "s1" / "Budget.xlsx"
    book, _ = X.build_workbook([{"name": "S", "rows": [[1, "=A1+1"]]}])
    upload.write_bytes(book.to_bytes())
    before = upload.read_bytes()
    tool = XlsxTool()
    edit = [{"op": "set", "cell": "A1", "value": 5}]
    result = await tool.execute(action="edit", path="saved/downloads/s1/Budget.xlsx", edits=edit, **kw)
    assert result.success, result.error
    assert upload.read_bytes() == before
    assert "the original Budget.xlsx is unchanged" in result.content
    copy = next(p for p in root.rglob("Budget.xlsx") if p != upload)
    assert "output" in copy.parts
    assert X.Book(copy.read_bytes()).sheet("S").stored(1, 2) == 6
    again = await tool.execute(action="edit", path="saved/downloads/s1/Budget.xlsx", edits=edit, **kw)
    assert not again.success and "already exists" in again.error
    onward = await tool.execute(action="edit", path=str(copy), edits=[{"op": "set", "cell": "A1", "value": 7}], **kw)
    assert onward.success and X.Book(copy.read_bytes()).sheet("S").stored(1, 2) == 8


async def test_create_wont_overwrite_unless_asked(workspace):
    _, kw = workspace
    tool = XlsxTool()
    sheets = [{"name": "S", "rows": [[1]]}]
    assert (await tool.execute(action="create", path="A.xlsx", sheets=sheets, **kw)).success
    again = await tool.execute(action="create", path="A.xlsx", sheets=sheets, **kw)
    assert not again.success and "already exists" in again.error
    assert (await tool.execute(action="create", path="A.xlsx", sheets=sheets, overwrite=True, **kw)).success


async def test_a_template_keeps_its_sheets_and_formats(workspace):
    root, kw = workspace
    template, _ = X.build_workbook([{"name": "House", "rows": [["Client", None], ["Total", "=SUM(B3:B9)"]]}],
                                   edits=[{"op": "format", "sheet": "House", "range": "A1:A2", "bold": True}])
    path = root / "saved" / "house.xlsx"
    path.write_bytes(template.to_bytes())
    tool = XlsxTool()
    result = await tool.execute(action="create", path="New.xlsx", template=str(path), edits=[
        {"op": "set", "sheet": "House", "cell": "B1", "value": "Acme"},
        {"op": "set", "sheet": "House", "cell": "B3", "value": 12}], **kw)
    assert result.success, result.error
    sheet = X.Book(next(root.rglob("New.xlsx")).read_bytes()).sheet("House")
    assert sheet.stored(1, 2) == "Acme" and sheet.stored(2, 2) == 12
    assert sheet.style_of(1, 1) != 0


async def test_a_macro_workbook_keeps_its_extension(workspace):
    root, kw = workspace
    book, _ = X.build_workbook([{"name": "S", "rows": [[1]]}])
    raw = _parts(book.to_bytes())
    raw["[Content_Types].xml"] = raw["[Content_Types].xml"].replace(
        X.CT_WORKBOOK.encode(), X.CT_MACRO.encode())
    raw["xl/vbaProject.bin"] = b"macros"
    path = root / "saved" / "m.xlsm"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for name, data in raw.items():
            target.writestr(name, data)
    path.write_bytes(out.getvalue())
    tool = XlsxTool()
    result = await tool.execute(action="edit", path=str(path), output_path="copy.xlsx",
                                edits=[{"op": "set", "cell": "A1", "value": 2}], **kw)
    assert result.success, result.error
    copy = next(root.rglob("copy.xlsm"))
    assert _parts(copy.read_bytes())["xl/vbaProject.bin"] == b"macros"


def test_file_names():
    assert xlsx_file_name("Budget") == "Budget.xlsx"
    assert xlsx_file_name("budget.xls") == "budget.xlsx"
    assert xlsx_file_name("data.csv") == "data.xlsx"
    assert xlsx_file_name("m.xlsm") == "m.xlsm"
    assert xlsx_file_name("q3.v2") == "q3.v2.xlsx"


async def test_write_and_edit_refuse_workbooks_and_decks(workspace):
    root, kw = workspace
    for name in ("x.xlsx", "x.xlsm", "x.xls"):
        result = await WriteTool().execute(path=name, content="not a workbook", **kw)
        assert not result.success and result.error == "xlsx_needs_xlsx_tool"
        result = await EditTool().execute(path=name, old_string="a", new_string="b", **kw)
        assert not result.success and "xlsx tool" in result.error
    result = await WriteTool().execute(path="x.pptx", content="not a deck", **kw)
    assert not result.success and result.error == "pptx_needs_script"
    result = await EditTool().execute(path="x.pptx", old_string="a", new_string="b", **kw)
    assert not result.success and "python-pptx" in result.error
    assert not list(root.rglob("x.*"))


def test_the_extractor_shows_a_formula_with_no_stored_result(tmp_path):
    wb = Workbook()
    ws = wb.active
    for row in (["Item", "Q1"], ["A", 10], [None, "=SUM(B2:B2)"]):
        ws.append(row)
    path = tmp_path / "script.xlsx"
    wb.save(path)
    assert "=SUM(B2:B2)" in _extract_xlsx_markdown(path, 50)


async def test_a_file_made_by_a_script_reads_with_calculated_results(workspace):
    root, kw = workspace
    wb = Workbook()
    ws = wb.active
    ws.title = "S"
    ws.append([10, 20, "=A1+B1"])
    path = root / "saved" / "script.xlsx"
    wb.save(path)
    read = await XlsxTool().execute(action="read", path=str(path), **kw)
    assert "C1 =A1+B1 → 30 (calculated now, not stored)" in read.content
    checked = await XlsxTool().execute(action="check", path=str(path), **kw)
    assert "1 formula(s) had no stored result" in checked.content


async def test_edits_passed_as_json_text_are_accepted(workspace):
    root, kw = workspace
    tool = XlsxTool()
    await tool.execute(action="create", path="J.xlsx", sheets=[{"name": "S", "rows": [[1]]}], **kw)
    result = await tool.execute(action="edit", path=str(next(root.rglob("J.xlsx"))),
                                edits='[{"op": "set", "cell": "A1", "value": 2}]', **kw)
    assert result.success, result.error


def test_every_agent_has_the_tool_and_eco_mode_finds_it():
    from captain_claw.agent_orchestration_mixin import _ECO_INTENT_PATTERNS
    from captain_claw.config import ToolsConfig

    assert "xlsx" in ToolsConfig().enabled
    for text in ("update Budget.xlsx", "make a spreadsheet", "fix the excel model", "napravi tablicu", "proračun za Q3"):
        assert any("xlsx" in tools for pattern, tools in _ECO_INTENT_PATTERNS if pattern.search(text)), text


def test_the_schema_suits_every_provider():
    def walk(schema, where):
        for key, value in (schema.get("properties") or {}).items():
            assert "type" in value, f"{where}.{key} has no type"
            if value["type"] == "array":
                assert value.get("items"), f"{where}.{key} has no items"
                if isinstance(value["items"], dict):
                    walk(value["items"], f"{where}.{key}[]")
            if value["type"] == "object":
                walk(value, f"{where}.{key}")

    walk(XlsxTool.parameters, "xlsx")
    json.dumps(XlsxTool.parameters)


def _with_pivot() -> X.Book:
    """A pivot table's parts, as far as this tool needs to know about them."""
    book = _book({"name": "Data", "rows": [["R", "V"], ["a", 1]]}, {"name": "View", "rows": [["pivot here"]]})
    parts = _parts(book.to_bytes())
    parts["[Content_Types].xml"] = parts["[Content_Types].xml"].replace(
        b"</Types>", b'<Override PartName="/xl/pivotTables/pivotTable1.xml" ContentType="application/'
                     b'vnd.openxmlformats-officedocument.spreadsheetml.pivotTable+xml"/></Types>')
    parts["xl/pivotTables/pivotTable1.xml"] = f'<pivotTableDefinition xmlns="{MAIN}" name="P"/>'.encode()
    parts["xl/worksheets/_rels/sheet2.xml.rels"] = (
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/pivotTable" '
        b'Target="../pivotTables/pivotTable1.xml"/></Relationships>')
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as target:
        for name, data in parts.items():
            target.writestr(name, data)
    return X.Book(out.getvalue())


def test_a_pivot_table_is_kept_and_its_old_numbers_are_pointed_out():
    book = _with_pivot()
    before = _parts(book.to_bytes())["xl/pivotTables/pivotTable1.xml"]
    edited, report = _edit(book, {"op": "set", "sheet": "Data", "cell": "B2", "value": 5})
    assert _parts(edited.to_bytes())["xl/pivotTables/pivotTable1.xml"] == before
    assert any("pivot table" in n and "Refresh" in n for n in report.notes)
    assert "pivot table" in "\n".join(X.overview(edited))


def test_rows_under_a_pivot_table_arent_moved():
    with pytest.raises(X.XlsxError, match="holds a pivot table"):
        X.Editor(_with_pivot()).apply([{"op": "insert_rows", "sheet": "View", "at": 1}])


def test_a_workbook_openpyxl_cant_read_can_still_be_edited_but_a_broken_result_cant_be_saved():
    book = _with_pivot()
    original = book.to_bytes()
    assert X._openpyxl_problem(original) is not None, "the fixture is beyond openpyxl, as some Excel files are"
    X.verify_saved(original, original)
    with pytest.raises(X.XlsxError, match="wouldn't open again"):
        X.verify_saved(original, _book({"name": "S", "rows": [[1]]}).to_bytes())
    with pytest.raises(X.XlsxError, match="Not a readable"):
        X.verify_saved(b"junk", original)
