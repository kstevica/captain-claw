"""pptx tool: decks built on layouts, kept whole, with what doesn't fit reported.

The files were opened in real PowerPoint and the text estimate compared with its
own measurements; these tests pin that behaviour without PowerPoint.
"""

from __future__ import annotations

import io
import json
import zipfile
from pathlib import Path

import pytest
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from captain_claw import pptx_ops as P
from captain_claw.tools.edit import EditTool
from captain_claw.tools.pptx_tool import PptxTool, pptx_file_name
from captain_claw.tools.write import WriteTool

P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _deck(*slides, **options) -> Presentation:
    """A deck saved and opened again, as the tool sees it."""
    prs, _ = P.build_deck(list(slides), **options)
    return P.open_deck(P.deck_bytes(prs))


def _edit(prs, *ops, image_root=None):
    report = P.Editor(prs, image_root=image_root).apply(list(ops))
    data = P.deck_bytes(prs)
    P.verify_saved(data)
    return P.open_deck(data), report


def _titles(prs) -> list[str]:
    return [s.shapes.title.text_frame.text if s.shapes.title is not None else "" for s in prs.slides]


def _parts(prs) -> dict[str, bytes]:
    with zipfile.ZipFile(io.BytesIO(P.deck_bytes(prs))) as archive:
        return {n: archive.read(n) for n in archive.namelist()}


def _picture(path: Path, size=(400, 200)) -> Path:
    Image.new("RGB", size, (200, 30, 30)).save(path)
    return path


def _finder(root: Path):
    def resolve(src: str) -> Path | None:
        found = root / src
        return found if found.is_file() else None
    return resolve


def _issues(prs) -> dict[str, list[P.Issue]]:
    out: dict[str, list[P.Issue]] = {}
    for issue in P.check_deck(prs)[0]:
        out.setdefault(issue.kind, []).append(issue)
    return out


BODY = "- First point\n- Second point\n  - A detail\n- Third point"


# ── Making a deck ────────────────────────────────────────────────────


def test_a_new_deck_is_wide_and_fills_its_layouts():
    prs = _deck({"layout": "Title Slide", "title": "Q3 review", "subtitle": "Board pack"},
                {"title": "Results", "body": BODY})
    assert P.inches(prs.slide_width) == 13.33
    assert _titles(prs) == ["Q3 review", "Results"]
    slide = prs.slides[1]
    assert slide.slide_layout.name == "Title and Content"
    body = next(s for s in slide.placeholders if P.placeholder_kind(s) == "content")
    paragraphs = body.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["First point", "Second point", "A detail", "Third point"]
    assert [p.level for p in paragraphs] == [0, 0, 1, 0]


def test_a_new_decks_layouts_agree_on_alignment_and_case():
    prs = _deck({"title": "x"})
    for layout in P.all_layouts(prs):
        assert not [d for d in layout._element.iter(qn("a:defRPr")) if d.get("cap")]
        assert not [e for e in layout._element.iter(qn("a:lvl1pPr")) if e.get("algn") == "ctr"]


def test_layouts_are_stretched_with_the_slide():
    prs = _deck({"title": "x", "body": "y"})
    for layout in P.all_layouts(prs):
        for shape in layout.placeholders:
            assert shape.left + shape.width <= prs.slide_width + 10


def test_inline_marks_become_formatting_and_real_links():
    prs = _deck({"title": "x", "body": "Plain **bold** and *italic* and [site](https://example.com)\n2*3*4 stays"})
    body = next(s for s in prs.slides[0].placeholders if P.placeholder_kind(s) == "content")
    runs = body.text_frame.paragraphs[0].runs
    assert [(r.text, r.font.bold, r.font.italic) for r in runs][:4] == [
        ("Plain ", None, None), ("bold", True, None), (" and ", None, None), ("italic", None, True)]
    assert runs[-1].hyperlink.address == "https://example.com"
    assert body.text_frame.paragraphs[1].text == "2*3*4 stays"


def test_only_web_and_mail_links_are_written():
    with pytest.raises(P.PptxError, match="http"):
        P.build_deck([{"title": "x", "body": "[bad](file:///etc/passwd)"}])


def test_a_layout_is_chosen_for_what_the_slide_holds():
    prs = _deck({"title": "Text", "body": "a"}, {"title": "Only title"},
                {"title": "Both", "body": "a", "table": [["x"], ["1"]]})
    assert [s.slide_layout.name for s in prs.slides] == ["Title and Content", "Title Only", "Two Content"]


def test_a_bad_layout_lists_the_real_ones():
    with pytest.raises(P.PptxError, match="Title and Content"):
        P.build_deck([{"layout": "Nope", "title": "x"}])


def test_unfilled_placeholders_are_removed_unless_asked_to_keep_them():
    prs = _deck({"layout": "Two Content", "title": "x", "body": "only one side"})
    assert [P.placeholder_kind(s) for s in prs.slides[0].placeholders] == ["title", "content"]
    prs = _deck({"layout": "Two Content", "title": "x", "body": "one", "keep_empty": True})
    assert len(list(prs.slides[0].placeholders)) == 3


def test_page_numbers_and_a_footer_reach_new_slides():
    prs = _deck({"title": "One"}, page_numbers=True)
    kinds = [P.placeholder_kind(s) for s in prs.slides[0].placeholders]
    assert "slide number" in kinds
    edited, _ = _edit(prs, {"op": "add_slide", "title": "Two"})
    assert "slide number" in [P.placeholder_kind(s) for s in edited.slides[1].placeholders]
    number = next(s for s in edited.slides[1].placeholders if P.placeholder_kind(s) == "slide number")
    assert number._element.find(f".//{qn('a:fld')}") is not None


def test_the_theme_takes_fonts_and_colours():
    prs = _deck({"title": "x", "body": "y"}, look={"font": "Georgia", "heading_font": "Arial",
                                                    "heading_color": "1F4E78", "accent_color": "C00000"})
    assert P.overview(prs)[0].endswith("Fonts: Arial (headings), Georgia (text).")
    theme = next(v for k, v in _parts(prs).items() if "theme" in k)
    assert b'val="C00000"' in theme
    assert b"1F4E78" in _parts(prs)["ppt/slideMasters/slideMaster1.xml"]


def test_no_stale_thumbnail_shows_the_old_deck():
    assert "docProps/thumbnail.jpeg" not in _parts(_deck({"title": "x"}))


# ── A template ───────────────────────────────────────────────────────


def _template() -> bytes:
    prs, _ = P.build_deck([{"title": "Sample one"}, {"title": "Sample two", "body": "sample"}],
                          look={"font": "Georgia"})
    prs.slide_layouts[1].name = "Brand Content"
    return P.deck_bytes(prs)


def _as_potx(data: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            raw = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                raw = raw.replace(P.CT_PRESENTATION, b"application/vnd.openxmlformats-officedocument."
                                                      b"presentationml.template.main+xml")
            target.writestr(info, raw)
    return out.getvalue()


def test_a_template_gives_its_design_and_drops_its_samples():
    prs, _ = P.build_deck([{"layout": "Brand Content", "title": "Mine", "body": "text"}], template=_template())
    prs = P.open_deck(P.deck_bytes(prs))
    assert _titles(prs) == ["Mine"]
    assert "Georgia (text)" in P.overview(prs)[0]
    assert P.inches(prs.slide_width) == 13.33


def test_a_template_can_keep_its_slides():
    prs, _ = P.build_deck([{"title": "Mine"}], template=_template(), keep_slides=True)
    assert _titles(prs) == ["Sample one", "Sample two", "Mine"]


def test_a_potx_opens_as_a_deck():
    prs = P.open_deck(_as_potx(_template()))
    assert len(prs.slides) == 2
    assert P.CT_PRESENTATION in _parts(prs)["[Content_Types].xml"]


# ── Changing text ────────────────────────────────────────────────────


def test_set_text_keeps_each_levels_formatting():
    prs = _deck({"title": "x", "body": BODY})
    body = next(s for s in prs.slides[0].placeholders if P.placeholder_kind(s) == "content")
    body.text_frame.paragraphs[0].runs[0].font.color.rgb = P._color("C00000")
    body.text_frame.paragraphs[2].runs[0].font.size = Pt(13)
    edited, _ = _edit(prs, {"op": "set_text", "slide": 1, "shape": "body", "text": "- New\n  - Sub\n- Again"})
    body = next(s for s in edited.slides[0].placeholders if P.placeholder_kind(s) == "content")
    paragraphs = body.text_frame.paragraphs
    assert [p.text for p in paragraphs] == ["New", "Sub", "Again"]
    assert str(paragraphs[0].runs[0].font.color.rgb) == "C00000"
    assert paragraphs[1].runs[0].font.size == Pt(13)
    assert str(paragraphs[2].runs[0].font.color.rgb) == "C00000"


def test_a_text_box_gets_real_bullets_and_numbers():
    prs = _deck({"title": "x"})
    edited, _ = _edit(prs, {"op": "add_text", "slide": 1, "text": "- one\n1. first\nplain", "x": 1, "y": 2, "width": 5, "height": 1})
    box = edited.slides[0].shapes[-1]
    paragraphs = box.text_frame._txBody.findall(P.A_P)
    assert paragraphs[0].find(P.A_PPR).find(qn("a:buChar")) is not None
    assert paragraphs[1].find(P.A_PPR).find(qn("a:buAutoNum")) is not None
    assert paragraphs[2].find(P.A_PPR) is None


def test_bullets_can_be_turned_off_in_a_placeholder():
    prs = _deck({"title": "x", "body": "a"})
    edited, _ = _edit(prs, {"op": "set_text", "slide": 1, "shape": "body", "text": "Just a sentence.", "bullets": False})
    body = next(s for s in edited.slides[0].placeholders if P.placeholder_kind(s) == "content")
    props = body.text_frame._txBody.findall(P.A_P)[0].find(P.A_PPR)
    assert props.find(qn("a:buNone")) is not None and props.get("marL") == "0"


def test_replace_works_across_runs_and_keeps_the_first_runs_format():
    prs = _deck({"title": "x", "body": "a"})
    body = next(s for s in prs.slides[0].placeholders if P.placeholder_kind(s) == "content")
    para = body.text_frame.paragraphs[0]
    para.clear()
    first, second, third = para.add_run(), para.add_run(), para.add_run()
    first.text, second.text, third.text = "Revenue in Q", "3 was", " good"
    first.font.bold = True
    edited, report = _edit(prs, {"op": "replace", "slide": 1, "find": "Q3 was", "replace_with": "Q4 is"})
    body = next(s for s in edited.slides[0].placeholders if P.placeholder_kind(s) == "content")
    runs = body.text_frame.paragraphs[0].runs
    assert body.text_frame.text == "Revenue in Q4 is good"
    assert runs[0].font.bold is True and runs[0].text == "Revenue in Q4 is"
    assert report.counts["replacements"] == 1


def test_replace_reaches_tables_notes_and_every_slide():
    prs = _deck({"title": "Q3 plan", "table": [["Q3", "2"], ["x", "Q3 end"]], "notes": "Say Q3."},
                {"title": "Q3 again"})
    edited, report = _edit(prs, {"op": "replace", "slide": "all", "find": "Q3", "replace_with": "Q4", "notes": True})
    assert _titles(edited) == ["Q4 plan", "Q4 again"]
    assert report.counts["replacements"] == 5
    assert P.notes_text(edited.slides[0]) == "Say Q4."
    table = next(s for s in edited.slides[0].shapes if s.has_table).table
    assert table.cell(1, 1).text == "Q4 end"


def test_replace_treats_curly_quotes_as_plain():
    prs = _deck({"title": "the client’s plan"})
    edited, _ = _edit(prs, {"op": "replace", "slide": 1, "find": "client's", "replace_with": "firm's"})
    assert _titles(edited) == ["the firm's plan"]


def test_replace_that_finds_nothing_fails_the_call():
    prs = _deck({"title": "x"})
    with pytest.raises(P.PptxError, match="isn't in those slides"):
        P.Editor(prs).apply([{"op": "replace", "slide": "all", "find": "nope", "replace_with": "y"}])


def test_format_reaches_one_word_a_paragraph_a_table_row_and_a_range_of_slides():
    prs = _deck({"title": "Big numbers here", "body": "one\ntwo"}, {"title": "Second", "table": [["a", "b"], ["1", "2"]]})
    table_id = next(s for s in prs.slides[1].shapes if s.has_table).shape_id
    edited, _ = _edit(prs,
                      {"op": "format", "slide": 1, "shape": "title", "find": "numbers", "bold": True, "color": "C00000"},
                      {"op": "format", "slide": 1, "shape": "body", "paragraph": 2, "italic": True, "font_size": 14},
                      {"op": "format", "slide": 2, "shape": table_id, "row": 2, "bold": True},
                      {"op": "format", "slide": "all", "shape": "title", "font": "Arial"})
    title = edited.slides[0].shapes.title
    runs = title.text_frame.paragraphs[0].runs
    assert [r.text for r in runs] == ["Big ", "numbers", " here"]
    assert runs[1].font.bold is True and str(runs[1].font.color.rgb) == "C00000"
    assert runs[0].font.bold is None
    assert all(r.font.name == "Arial" for r in runs)
    body = next(s for s in edited.slides[0].placeholders if P.placeholder_kind(s) == "content")
    assert body.text_frame.paragraphs[0].runs[0].font.size is None
    assert body.text_frame.paragraphs[1].runs[0].font.italic is True
    table = next(s for s in edited.slides[1].shapes if s.has_table).table
    assert table.cell(1, 0).text_frame.paragraphs[0].runs[0].font.bold is True
    assert table.cell(0, 0).text_frame.paragraphs[0].runs[0].font.bold is None
    assert edited.slides[1].shapes.title.text_frame.paragraphs[0].runs[0].font.name == "Arial"


def test_run_properties_stay_in_schema_order():
    prs = _deck({"title": "x"})
    edited, _ = _edit(prs, {"op": "format", "slide": 1, "shape": "title", "font": "Arial", "color": "C00000"},
                      {"op": "format", "slide": 1, "shape": "title", "color": "00AA00"})
    props = edited.slides[0].shapes.title.text_frame.paragraphs[0].runs[0]._r.find(P.A_RPR)
    assert [etree.QName(c).localname for c in props] == ["solidFill", "latin"]


# ── Slides ───────────────────────────────────────────────────────────


def _three() -> Presentation:
    return _deck({"title": "One"}, {"title": "Two"}, {"title": "Three"})


def test_slides_move_and_are_deleted():
    edited, _ = _edit(_three(), {"op": "move_slide", "slide": 3, "at": 1}, {"op": "delete_slide", "slide": 2})
    assert _titles(edited) == ["Three", "Two"]
    assert sorted(n for n in _parts(edited) if n.startswith("ppt/slides/slide")) == [
        "ppt/slides/slide1.xml", "ppt/slides/slide2.xml"]


def test_a_new_slide_after_a_delete_doesnt_collide_with_an_old_name():
    edited, _ = _edit(_three(), {"op": "delete_slide", "slide": 1}, {"op": "add_slide", "title": "New"},
                      {"op": "add_slide", "title": "Newer"})
    assert _titles(edited) == ["Two", "Three", "New", "Newer"]


def test_the_last_slide_stays():
    prs = _deck({"title": "Only"})
    with pytest.raises(P.PptxError, match="at least one"):
        P.Editor(prs).apply([{"op": "delete_slide", "slide": 1}])


def test_a_slide_number_outside_the_deck_is_refused():
    with pytest.raises(P.PptxError, match="no slide 9"):
        P.Editor(_three()).apply([{"op": "set_text", "slide": 9, "shape": "title", "text": "x"}])


def test_later_edits_see_earlier_moves():
    edited, _ = _edit(_three(), {"op": "add_slide", "title": "Zero", "at": 1},
                      {"op": "set_text", "slide": 2, "shape": "title", "text": "One!"})
    assert _titles(edited) == ["Zero", "One!", "Two", "Three"]


def test_duplicate_keeps_everything_and_copies_the_chart():
    prs = _deck({"title": "Chart", "chart": {"type": "column", "categories": ["a", "b"], "series": [{"name": "S", "values": [1, 2]}]},
                 "notes": "Say this."})
    edited, _ = _edit(prs, {"op": "duplicate_slide", "slide": 1},
                      {"op": "set_chart", "slide": 2, "shape": "ph:1", "categories": ["a", "b"], "series": [{"name": "S", "values": [8, 9]}]})
    assert _titles(edited) == ["Chart", "Chart"]
    first = next(s for s in edited.slides[0].shapes if s.has_chart).chart
    second = next(s for s in edited.slides[1].shapes if s.has_chart).chart
    assert list(first.plots[0].series[0].values) == [1, 2]
    assert list(second.plots[0].series[0].values) == [8, 9]
    assert first.part is not second.part
    assert P.notes_text(edited.slides[1]) == "Say this."
    parts = _parts(edited)
    assert len([n for n in parts if n.startswith("ppt/charts/chart")]) == 2
    assert len([n for n in parts if n.startswith("ppt/embeddings/")]) == 2


def test_duplicate_shares_the_picture_and_goes_after_its_source(tmp_path):
    _picture(tmp_path / "p.png")
    prs = _deck({"title": "One", "image": "p.png"}, {"title": "Two"}, image_root=_finder(tmp_path))
    edited, _ = _edit(prs, {"op": "duplicate_slide", "slide": 1})
    assert _titles(edited) == ["One", "One", "Two"]
    assert len([n for n in _parts(edited) if n.startswith("ppt/media/")]) == 1
    pictures = [s for s in edited.slides[1].shapes if s._element.tag == qn("p:pic")]
    assert pictures and pictures[0].image.size == (400, 200)


def test_smartart_isnt_duplicated_blindly():
    prs = _deck({"title": "x"})
    frame = etree.SubElement(prs.slides[0].shapes._spTree, qn("p:graphicFrame"))
    prs.slides[0].shapes._spTree.remove(frame)
    slide = prs.slides[0]
    shape = next(iter(slide.shapes))
    holder = etree.SubElement(shape._element, f"{{{P.DGM_NS}}}relIds")
    holder.set(f"{{{R}}}dm", "rId99")
    with pytest.raises(P.PptxError, match="SmartArt"):
        P.Editor(prs).apply([{"op": "duplicate_slide", "slide": 1}])


def _with_sections(prs) -> Presentation:
    ids = [el.get("id") for el in prs.slides._sldIdLst]
    root = prs.part._element
    ext_list = root.find(qn("p:extLst"))
    if ext_list is None:
        ext_list = etree.SubElement(root, qn("p:extLst"))
    ext = etree.SubElement(ext_list, qn("p:ext"), uri="{521415D9-36F7-43E2-AB2F-B90AF26B5E84}")
    sections = etree.SubElement(ext, f"{{{P14}}}sectionLst", nsmap={"p14": P14})
    for name, members in (("Intro", ids[:1]), ("Rest", ids[1:])):
        section = etree.SubElement(sections, f"{{{P14}}}section", name=name, id="{00000000-0000-0000-0000-000000000001}")
        lst = etree.SubElement(section, f"{{{P14}}}sldIdLst")
        for member in members:
            etree.SubElement(lst, f"{{{P14}}}sldId", id=member)
    return P.open_deck(P.deck_bytes(prs))


def _section_members(prs) -> dict[str, list[str]]:
    numbers = {el.get("id"): n for n, el in enumerate(prs.slides._sldIdLst, 1)}
    return {s.get("name"): [numbers[i.get("id")] for i in s.iter(f"{{{P14}}}sldId")] for s in P._sections(prs)}


def test_sections_follow_deleted_new_and_moved_slides():
    prs = _with_sections(_three())
    assert _section_members(prs) == {"Intro": [1], "Rest": [2, 3]}
    assert any("Sections" in line for line in P.overview(prs))
    edited, _ = _edit(prs, {"op": "delete_slide", "slide": 2})
    assert _section_members(edited) == {"Intro": [1], "Rest": [2]}
    edited, _ = _edit(edited, {"op": "add_slide", "title": "New"})
    assert _section_members(edited) == {"Intro": [1], "Rest": [2, 3]}
    edited, _ = _edit(edited, {"op": "move_slide", "slide": 3, "at": 2})
    assert sum(len(v) for v in _section_members(edited).values()) == 3
    edited, _ = _edit(edited, {"op": "duplicate_slide", "slide": 1})
    assert sum(len(v) for v in _section_members(edited).values()) == 4


def test_a_link_to_a_deleted_slide_goes_with_it():
    prs = _three()
    title = prs.slides[0].shapes.title
    run = title.text_frame.paragraphs[0].runs[0]
    rel = prs.slides[0].part.relate_to(prs.slides[2].part, P.RT.SLIDE)
    run._r.get_or_add_rPr().add_hlinkClick(rel)
    edited, report = _edit(prs, {"op": "delete_slide", "slide": 3})
    assert any("link" in note for note in report.notes)
    assert edited.slides[0].shapes.title._element.find(f".//{qn('a:hlinkClick')}") is None
    assert not [r for r in edited.slides[0].part.rels.values() if r.reltype == P.RT.SLIDE]


def test_hiding_a_slide():
    edited, _ = _edit(_three(), {"op": "hide_slide", "slide": "2-3"})
    assert [s._element.get("show") for s in edited.slides] == [None, "0", "0"]
    edited, _ = _edit(edited, {"op": "hide_slide", "slide": 2, "hidden": False})
    assert [s._element.get("show") for s in edited.slides] == [None, None, "0"]


def test_notes_are_set_and_read():
    edited, _ = _edit(_three(), {"op": "set_notes", "slide": 2, "text": "Talk about this."})
    assert P.notes_text(edited.slides[1]) == "Talk about this."
    assert any("notes: Talk about this." in line for line in P.describe_slide(edited, edited.slides[1], 2))


# ── Tables, charts, pictures ─────────────────────────────────────────


def test_a_table_takes_the_empty_placeholder_and_keeps_its_link_to_the_layout():
    prs = _deck({"title": "x", "table": [["Region", "Q1"], ["North", "100"], ["South", "1,200"]]})
    frame = next(s for s in prs.slides[0].shapes if s.has_table)
    layout_ph = next(p for p in prs.slides[0].slide_layout.placeholders if p.placeholder_format.idx == 1)
    assert frame.left == layout_ph.left and frame.top == layout_ph.top
    assert frame._element.xpath(".//p:nvPr/p:ph")[0].get("idx") == "1"
    table = frame.table
    assert [table.cell(1, c).text for c in range(2)] == ["North", "100"]
    assert table.first_row is True
    # Numbers line up on the right.
    assert table.cell(1, 1)._tc.txBody.findall(P.A_P)[0].find(P.A_PPR).get("algn") == "r"
    assert sum(c.width for c in table.columns) == frame.width
    assert not _issues(prs).get("overflow")


def test_a_long_table_gets_a_smaller_font_to_fit():
    rows = [["Name", "Value"]] + [[f"Row {i}", str(i)] for i in range(14)]
    prs = _deck({"title": "x", "table": rows})
    frame = next(s for s in prs.slides[0].shapes if s.has_table)
    size = frame.table.cell(0, 0)._tc.txBody.findall(P.A_P)[0].find(P.A_R).find(P.A_RPR).get("sz")
    assert int(size) < 1800
    assert frame.top + frame.height <= prs.slide_height


def test_a_table_too_long_for_the_slide_says_so():
    rows = [["Name", "Value"]] + [[f"Row {i}", str(i)] for i in range(60)]
    prs, report = P.build_deck([{"title": "x", "table": rows}])
    assert any("taller than its space" in note for note in report.notes)
    assert any(i.kind == "overflow" for i in report.fit)


def test_set_table_changes_cells_and_grows_or_shrinks_keeping_format():
    prs = _deck({"title": "x", "table": [["a", "b"], ["1", "2"]]})
    frame_id = next(s for s in prs.slides[0].shapes if s.has_table).shape_id
    edited, _ = _edit(prs, {"op": "set_table", "slide": 1, "shape": frame_id, "cells": {"r2c2": "9"}})
    table = next(s for s in edited.slides[0].shapes if s.has_table).table
    assert table.cell(1, 1).text == "9"
    edited, _ = _edit(edited, {"op": "set_table", "slide": 1, "shape": frame_id,
                               "rows": [["a", "b", "c"], ["1", "2", "3"], ["4", "5", "6"]]})
    table = next(s for s in edited.slides[0].shapes if s.has_table).table
    assert (len(table.rows), len(table.columns)) == (3, 3)
    assert table.cell(2, 2).text == "6"
    edited, _ = _edit(edited, {"op": "set_table", "slide": 1, "shape": frame_id, "rows": [["only"]]})
    table = next(s for s in edited.slides[0].shapes if s.has_table).table
    assert (len(table.rows), len(table.columns)) == (1, 1)


def test_a_table_given_a_new_column_keeps_its_width_and_rows_that_hold_their_text():
    prs = _deck({"title": "x", "table": [["a", "b"], ["1", "2"]]})
    frame = next(s for s in prs.slides[0].shapes if s.has_table)
    width = frame.width
    edited, report = _edit(prs, {"op": "set_table", "slide": 1, "shape": frame.shape_id,
                                 "rows": [["a", "b", "c", "d", "e"], ["1", "2", "3", "4", "5 " * 200]]})
    frame = next(s for s in edited.slides[0].shapes if s.has_table)
    assert sum(c.width for c in frame.table.columns) == width == frame.width
    assert frame.height == sum(r.height for r in frame.table.rows)
    assert frame.table.rows[1].height > frame.table.rows[0].height
    assert not [i for i in report.fit if "right of the slide" in i.what]


def test_check_finds_a_table_wider_than_its_frame_says():
    prs = _deck({"title": "x", "table": [["a", "b"], ["1", "2"]]})
    table = next(s for s in prs.slides[0].shapes if s.has_table).table
    table.columns[1].width = Inches(40)
    found = _issues(prs)["overflow"]
    assert any("right of the slide" in i.what for i in found)


def test_a_bad_cell_name_is_refused():
    prs = _deck({"title": "x", "table": [["a"]]})
    frame_id = next(s for s in prs.slides[0].shapes if s.has_table).shape_id
    with pytest.raises(P.PptxError, match="r2c3"):
        P.Editor(prs).apply([{"op": "set_table", "slide": 1, "shape": frame_id, "cells": {"B2": "x"}}])
    with pytest.raises(P.PptxError, match="outside the table"):
        P.Editor(prs).apply([{"op": "set_table", "slide": 1, "shape": frame_id, "cells": {"r5c5": "x"}}])


@pytest.mark.parametrize("kind", ["column", "bar", "line", "area", "pie", "doughnut"])
def test_charts_are_native_and_hold_their_data(kind):
    prs = _deck({"title": "x", "chart": {"type": kind, "title": "Sales", "categories": ["N", "S"],
                                         "series": [{"name": "Q1", "values": [1, 2]}], "labels": True}})
    frame = next(s for s in prs.slides[0].shapes if s.has_chart)
    chart = frame.chart
    assert chart.chart_title.text_frame.text == "Sales"
    assert list(chart.plots[0].categories) == ["N", "S"]
    assert list(chart.plots[0].series[0].values) == [1, 2]
    assert frame._element.xpath(".//p:nvPr/p:ph")[0].get("idx") == "1"
    assert [n for n in _parts(prs) if n.startswith("ppt/embeddings/")]


def test_a_scatter_chart_takes_points():
    prs = _deck({"title": "x", "chart": {"type": "scatter", "series": [{"name": "S", "x": [1, 2], "y": [3, 4]}]}})
    chart = next(s for s in prs.slides[0].shapes if s.has_chart).chart
    assert list(chart.plots[0].series[0].iter_values()) == [3, 4]


@pytest.mark.parametrize("chart,message", [
    ({"type": "radar", "categories": ["a"], "series": [{"values": [1]}]}, "chart type"),
    ({"type": "column", "series": [{"values": [1]}]}, "categories"),
    ({"type": "column", "categories": ["a", "b"], "series": [{"values": [1]}]}, "1 values for 2"),
    ({"type": "column", "categories": ["a"], "series": [{"values": ["x"]}]}, "isn't a number"),
])
def test_bad_chart_data_is_refused(chart, message):
    with pytest.raises(P.PptxError, match=message):
        P.build_deck([{"title": "x", "chart": chart}])


def test_set_chart_changes_numbers_and_keeps_the_chart():
    prs = _deck({"title": "x", "chart": {"type": "line", "title": "Old", "categories": ["a"], "series": [{"name": "S", "values": [1]}]}})
    frame = next(s for s in prs.slides[0].shapes if s.has_chart)
    edited, _ = _edit(prs, {"op": "set_chart", "slide": 1, "shape": frame.shape_id, "title": "New",
                            "categories": ["a", "b"], "series": [{"name": "S", "values": [5, 6]}]})
    chart = next(s for s in edited.slides[0].shapes if s.has_chart).chart
    assert chart.chart_title.text_frame.text == "New"
    assert list(chart.plots[0].series[0].values) == [5, 6]
    assert "LINE" in str(chart.chart_type)


def test_a_picture_keeps_its_shape_in_the_placeholder(tmp_path):
    _picture(tmp_path / "wide.png", (800, 200))
    prs = _deck({"title": "x", "image": {"image": "wide.png", "alt": "A red bar"}}, image_root=_finder(tmp_path))
    picture = next(s for s in prs.slides[0].shapes if s._element.tag == qn("p:pic"))
    assert abs(picture.width / picture.height - 4) < 0.01
    assert P._alt_text(picture) == "A red bar"
    assert "stretched_picture" not in _issues(prs)


def test_cover_crops_instead_of_stretching(tmp_path):
    _picture(tmp_path / "wide.png", (800, 200))
    prs = _deck({"title": "x"}, image_root=_finder(tmp_path))
    edited, _ = _edit(prs, {"op": "add_image", "slide": 1, "image": "wide.png", "fit": "cover", "x": 1, "y": 2, "width": 2, "height": 2},
                      image_root=_finder(tmp_path))
    picture = next(s for s in edited.slides[0].shapes if s._element.tag == qn("p:pic"))
    assert picture.crop_left > 0.3 and picture.width == picture.height
    assert "stretched_picture" not in _issues(edited)


def test_a_picture_it_cant_find_or_read_fails_the_edit(tmp_path):
    (tmp_path / "fake.png").write_text("not a picture")
    with pytest.raises(P.PptxError, match="Can't read the picture"):
        P.build_deck([{"title": "x", "image": "gone.png"}], image_root=_finder(tmp_path))
    with pytest.raises(P.PptxError, match="isn't a picture"):
        P.build_deck([{"title": "x", "image": "fake.png"}], image_root=_finder(tmp_path))


def test_place_and_delete_a_shape():
    prs = _deck({"title": "x", "body": "y"})
    body_id = next(s for s in prs.slides[0].placeholders if P.placeholder_kind(s) == "content").shape_id
    edited, _ = _edit(prs, {"op": "place", "slide": 1, "shape": body_id, "x": 1, "width": 6, "z": "back"})
    body = next(s for s in edited.slides[0].placeholders if P.placeholder_kind(s) == "content")
    assert P.inches(body.left) == 1 and P.inches(body.width) == 6
    assert body.height > 0 and body.top > 0, "the rest stays where the layout had it"
    assert list(edited.slides[0].shapes)[0].shape_id == body_id
    edited, _ = _edit(edited, {"op": "delete_shape", "slide": 1, "shape": body_id})
    assert not [s for s in edited.slides[0].shapes if s.shape_id == body_id]


def test_deleting_a_picture_lets_go_of_its_part(tmp_path):
    _picture(tmp_path / "p.png")
    prs = _deck({"title": "x", "image": "p.png"}, image_root=_finder(tmp_path))
    picture = next(s for s in prs.slides[0].shapes if s._element.tag == qn("p:pic"))
    edited, _ = _edit(prs, {"op": "delete_shape", "slide": 1, "shape": picture.shape_id})
    assert not [n for n in _parts(edited) if n.startswith("ppt/media/")]


def test_deleting_an_animated_shape_takes_its_animation():
    prs = _deck({"title": "x", "body": "y"})
    slide = prs.slides[0]
    body = next(s for s in slide.placeholders if P.placeholder_kind(s) == "content")
    timing = etree.fromstring(
        '<p:timing xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"><p:tnLst><p:par>'
        '<p:cTn id="1" dur="indefinite" nodeType="tmRoot"><p:childTnLst><p:seq><p:cTn id="2" nodeType="mainSeq">'
        '<p:childTnLst><p:par><p:cTn id="3"><p:childTnLst><p:set><p:cBhvr><p:cTn id="4"/><p:tgtEl>'
        f'<p:spTgt spid="{body.shape_id}"/></p:tgtEl></p:cBhvr></p:set></p:childTnLst></p:cTn></p:par></p:childTnLst>'
        '</p:cTn></p:seq></p:childTnLst></p:cTn></p:par></p:tnLst>'
        f'<p:bldLst><p:bldP spid="{body.shape_id}" grpId="0"/></p:bldLst></p:timing>')
    slide._element.append(timing)
    prs = P.open_deck(P.deck_bytes(prs))
    edited, report = _edit(prs, {"op": "delete_shape", "slide": 1, "shape": body.shape_id})
    assert edited.slides[0]._element.find(qn("p:timing")) is None
    assert any("animations" in note for note in report.notes)


def test_a_bad_edit_fails_the_whole_call():
    prs = _three()
    with pytest.raises(P.PptxError, match=r"Edit 2 \(set_text\)"):
        P.Editor(prs).apply([{"op": "add_slide", "title": "x"}, {"op": "set_text", "slide": 1, "shape": "nope", "text": "y"}])


def test_an_unknown_shape_names_whats_there():
    with pytest.raises(P.PptxError, match=r"\[2\] title"):
        P.Editor(_three()).apply([{"op": "set_text", "slide": 1, "shape": 99, "text": "y"}])


def test_not_a_deck_gets_a_clear_message():
    with pytest.raises(P.PptxError, match="password-protected"):
        P.open_deck(b"\xd0\xcf\x11\xe0" + b"0" * 100)
    with pytest.raises(P.PptxError, match="Not a readable"):
        P.open_deck(b"plain text")


# ── Nothing else is lost ─────────────────────────────────────────────


def test_a_part_nobody_knows_survives_an_edit():
    prs = _deck({"title": "x"})
    data = P.deck_bytes(prs)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            raw = source.read(info.filename)
            if info.filename == "ppt/slides/_rels/slide1.xml.rels":
                raw = raw.replace(b"</Relationships>",
                                  b'<Relationship Id="rId77" Type="http://example.com/vendor" Target="../vendor/data.bin"/></Relationships>')
            if info.filename == "[Content_Types].xml":
                raw = raw.replace(b"</Types>", b'<Default Extension="bin" ContentType="application/octet-stream"/></Types>', 1)
            target.writestr(info, raw)
        target.writestr("ppt/vendor/data.bin", b"\x00keep me")
    prs = P.open_deck(out.getvalue())
    edited, _ = _edit(prs, {"op": "set_text", "slide": 1, "shape": "title", "text": "y"})
    assert _parts(edited)["ppt/vendor/data.bin"] == b"\x00keep me"


def test_macros_are_kept():
    prs = _deck({"title": "x"})
    data = P.deck_bytes(prs)
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w") as target:
        for info in source.infolist():
            raw = source.read(info.filename)
            if info.filename == "[Content_Types].xml":
                raw = raw.replace(P.CT_PRESENTATION, P.CT_MACRO)
                raw = raw.replace(b"</Types>", b'<Default Extension="bin" ContentType="application/vnd.ms-office.vbaProject"/></Types>', 1)
            if info.filename == "ppt/_rels/presentation.xml.rels":
                raw = raw.replace(b"</Relationships>",
                                  b'<Relationship Id="rId99" Type="http://schemas.microsoft.com/office/2006/relationships/vbaProject" Target="vbaProject.bin"/></Relationships>')
            target.writestr(info, raw)
        target.writestr("ppt/vbaProject.bin", b"macros")
    prs = P.open_deck(out.getvalue())
    assert any("Macros" in line for line in P.overview(prs))
    edited, _ = _edit(prs, {"op": "set_text", "slide": 1, "shape": "title", "text": "y"})
    assert _parts(edited)["ppt/vbaProject.bin"] == b"macros"


# ── Fit ──────────────────────────────────────────────────────────────


def _box(slide, text: str, size=18, width=6.0, height=1.0, font="Arial"):
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.name = font
    return shape


def test_measuring_counts_wrapped_lines_and_their_height():
    prs = _deck({"layout": "Blank"})
    slide = prs.slides[0]
    one = P.measure_text(P.TextContext(_box(slide, "Short"), slide, prs), Inches(6))
    assert one.lines == 1
    assert one.height == pytest.approx(18 * 1.2 + 7.2, abs=0.5)
    long_text = "word " * 100
    many = P.measure_text(P.TextContext(_box(slide, long_text), slide, prs), Inches(6))
    assert many.lines > 5
    narrow = P.measure_text(P.TextContext(_box(slide, long_text), slide, prs), Inches(3))
    assert narrow.lines > many.lines


def test_a_word_wider_than_the_box_breaks():
    prs = _deck({"layout": "Blank"})
    slide = prs.slides[0]
    fit = P.measure_text(P.TextContext(_box(slide, "x" * 200, width=1), slide, prs), Inches(1))
    assert fit.lines > 3


def test_sizes_come_from_the_layout_and_master():
    prs = _deck({"title": "Title", "body": "- level one\n  - level two"})
    slide = prs.slides[0]
    title = slide.shapes.title
    ctx = P.TextContext(title, slide, prs)
    assert ctx.size(0) == 34
    assert ctx.is_title
    body = next(s for s in slide.placeholders if P.placeholder_kind(s) == "content")
    ctx = P.TextContext(body, slide, prs)
    assert ctx.size(0) == 24 and ctx.size(1) == 20
    assert ctx.has_bullet(None, 0)
    assert ctx.autofit()[0] == "shrink"


def test_a_stuffed_placeholder_is_shrunk_as_powerpoint_would_and_said_so():
    text = "\n".join(f"- Point number {i} about something that matters a good deal to the plan" for i in range(14))
    prs, report = P.build_deck([{"title": "x", "body": text}])
    prs = P.open_deck(P.deck_bytes(prs))
    body = next(s for s in prs.slides[0].placeholders if P.placeholder_kind(s) == "content")
    autofit = body.text_frame._txBody.find(qn("a:bodyPr")).find(qn("a:normAutofit"))
    assert autofit is not None and int(autofit.get("fontScale")) < 100000
    assert any("shrunk" in note for note in report.notes)
    assert not [i for i in report.fit if i.kind == "overflow"]


def test_text_that_fits_again_loses_its_shrink():
    text = "\n".join(f"- Point number {i} about something that matters a good deal to the plan" for i in range(14))
    prs, _ = P.build_deck([{"title": "x", "body": text}])
    edited, _ = _edit(prs, {"op": "set_text", "slide": 1, "shape": "body", "text": "- Short now"})
    body = next(s for s in edited.slides[0].placeholders if P.placeholder_kind(s) == "content")
    autofit = body.text_frame._txBody.find(qn("a:bodyPr")).find(qn("a:normAutofit"))
    assert autofit is None or autofit.get("fontScale") is None


def test_too_much_text_even_when_shrunk_is_reported():
    text = "\n".join("- " + "A very long point that goes on and on " * 4 for _ in range(30))
    _, report = P.build_deck([{"title": "x", "body": text}])
    assert any(i.kind == "overflow" for i in report.fit)


def test_a_text_box_grows_to_its_text():
    prs = _deck({"title": "x"})
    edited, _ = _edit(prs, {"op": "add_text", "slide": 1, "text": "word " * 80, "x": 1, "y": 2, "width": 4, "font_size": 14})
    box = edited.slides[0].shapes[-1]
    assert P.inches(box.height) > 1


# ── Check ────────────────────────────────────────────────────────────


def test_a_clean_deck_has_nothing_to_report():
    prs = _deck({"layout": "Title Slide", "title": "Hello", "subtitle": "There"}, {"title": "Points", "body": BODY})
    assert P.check_deck(prs)[0] == []


def test_check_finds_an_overflowing_box_that_doesnt_shrink():
    prs = _deck({"layout": "Blank"})
    slide = prs.slides[0]
    _box(slide, "word " * 200, height=1)
    assert "overflow" in _issues(prs)


def test_check_finds_shapes_off_the_slide():
    prs = _deck({"title": "x"})
    edited, _ = _edit(prs, {"op": "add_text", "slide": 1, "text": "gone", "x": 20, "y": 1, "width": 2, "height": 1},
                      {"op": "add_text", "slide": 1, "text": "half", "x": 12, "y": 1, "width": 3, "height": 1})
    found = _issues(edited)["off_slide"]
    assert len(found) == 2
    assert any("entirely" in i.what for i in found) and any("runs off" in i.what for i in found)


def test_check_finds_overlaps_but_not_a_background_picture(tmp_path):
    _picture(tmp_path / "p.png", (1333, 750))
    prs = _deck({"title": "x"}, image_root=_finder(tmp_path))
    edited, _ = _edit(prs, {"op": "add_text", "slide": 1, "text": "one", "x": 1, "y": 3, "width": 4, "height": 1},
                      {"op": "add_text", "slide": 1, "text": "two", "x": 2, "y": 3, "width": 4, "height": 1})
    assert _issues(edited)["overlap"][0].level == "warn"
    prs = _deck({"title": "x"}, image_root=_finder(tmp_path))
    edited, _ = _edit(prs, {"op": "add_image", "slide": 1, "image": "p.png", "x": 0, "y": 0, "width": 13.33, "height": 7.5},
                      image_root=_finder(tmp_path))
    assert "overlap" not in _issues(edited)


def test_check_finds_an_empty_placeholder_and_a_missing_title():
    prs = _deck({"layout": "Title and Content", "keep_empty": True})
    found = _issues(prs)
    assert len(found["empty_placeholder"]) == 2
    assert found["no_title"]


def test_a_blank_slide_needs_no_title():
    prs = _deck({"layout": "Blank"})
    assert "no_title" not in _issues(prs)


def test_check_finds_small_text_and_leftovers():
    prs = _deck({"title": "Lorem ipsum dolor", "body": "fine"})
    slide = prs.slides[0]
    _box(slide, "tiny", size=7)
    found = _issues(prs)
    assert found["small_text"] and found["leftover_text"]


def test_check_finds_a_stretched_picture_and_missing_alt_text(tmp_path):
    _picture(tmp_path / "p.png", (400, 400))
    prs = _deck({"title": "x"})
    prs.slides[0].shapes.add_picture(str(tmp_path / "p.png"), Inches(1), Inches(2), Inches(4), Inches(1))
    found = _issues(prs)
    assert "stretched" in found["stretched_picture"][0].what
    assert found["no_alt_text"]


def test_check_notes_wordy_slides_same_titles_and_hidden_slides():
    prs = _deck({"title": "Same", "body": "- " + "word " * 130}, {"title": "same"})
    prs.slides[1]._element.set("show", "0")
    found = _issues(prs)
    assert found["wordy"] and found["same_title"] and found["hidden"]


def test_a_tall_table_is_measured_from_its_text_not_its_stored_rows():
    prs = _deck({"title": "x"})
    frame = prs.slides[0].shapes.add_table(2, 1, Inches(1), Inches(6), Inches(2), Inches(0.5))
    frame.table.cell(0, 0).text = "word " * 120
    frame.table.rows[0].height = Inches(0.25)
    assert "overflow" in _issues(prs)


# ── The tool ─────────────────────────────────────────────────────────


@pytest.fixture
def workspace(tmp_path):
    (tmp_path / "saved" / "downloads" / "s1").mkdir(parents=True)
    kw = {"_runtime_base_path": tmp_path, "_saved_base_path": tmp_path / "saved", "_session_id": "s1"}
    return tmp_path, kw


async def test_create_read_edit_check_through_the_tool(workspace):
    root, kw = workspace
    tool = PptxTool()
    created = await tool.execute(
        action="create", path="Pitch", title="Pitch", accent_color="1F4E78",
        slides=[{"layout": "Title Slide", "title": "Hello", "subtitle": "World"},
                {"title": "Plan", "body": BODY, "notes": "Keep it short."}],
        edits=[{"op": "add_slide", "title": "Numbers", "table": [["A", "B"], ["1", "2"]]}], **kw)
    assert created.success, created.error
    assert "3 slide(s)" in created.content
    saved = next(root.rglob("Pitch.pptx"))

    read = await tool.execute(action="read", path=str(saved), **kw)
    assert read.success, read.error
    assert 'Layouts: 1 "Title Slide"' in read.content
    assert "Slide 2 · layout" in read.content
    assert "• First point" in read.content
    assert "notes: Keep it short." in read.content
    assert "r2 1 | 2" in read.content

    one = await tool.execute(action="read", path=str(saved), slide="2", **kw)
    assert "Slide 2" in one.content and "Slide 3" not in one.content

    edited = await tool.execute(action="edit", path=str(saved),
                                edits=[{"op": "set_text", "slide": 2, "shape": "title", "text": "New plan"}], **kw)
    assert edited.success, edited.error
    assert "Saved in place" in edited.content and "Backup:" in edited.content
    assert '"New plan"' in edited.content

    checked = await tool.execute(action="check", path=str(saved), **kw)
    assert "0 warning(s)" in checked.content

    found = await tool.execute(action="read", path=str(saved), find="new PLAN", **kw)
    assert "Slide 2" in found.content


async def test_read_pages_through_a_long_deck(workspace):
    root, kw = workspace
    tool = PptxTool()
    await tool.execute(action="create", path="Long.pptx", slides=[{"title": f"S{i}"} for i in range(40)], **kw)
    saved = next(root.rglob("Long.pptx"))
    first = await tool.execute(action="read", path=str(saved), **kw)
    assert "Slide 25 " in first.content and "Slide 26 " not in first.content
    assert "slide=26-40" in first.content
    second = await tool.execute(action="read", path=str(saved), slide="26-40", **kw)
    assert "Slide 40 " in second.content


async def test_a_fit_problem_comes_back_in_the_result(workspace):
    root, kw = workspace
    tool = PptxTool()
    result = await tool.execute(action="create", path="Full.pptx",
                                slides=[{"title": "x", "body": "\n".join("- " + "Long point about things " * 5 for _ in range(30))}], **kw)
    assert result.success, result.error
    assert "⚠ Fix these" in result.content and "overflow" not in result.error if result.error else True
    assert "Slide 1" in result.content


async def test_an_upload_is_never_changed_and_a_copy_never_forks(workspace):
    root, kw = workspace
    upload = root / "saved" / "downloads" / "s1" / "Pitch.pptx"
    upload.write_bytes(P.deck_bytes(_deck({"title": "Old"})))
    before = upload.read_bytes()
    tool = PptxTool()
    edit = [{"op": "set_text", "slide": 1, "shape": "title", "text": "New"}]
    result = await tool.execute(action="edit", path="saved/downloads/s1/Pitch.pptx", edits=edit, **kw)
    assert result.success, result.error
    assert upload.read_bytes() == before
    assert "the original Pitch.pptx is unchanged" in result.content
    copy = next(p for p in root.rglob("Pitch.pptx") if p != upload)
    assert "output" in copy.parts
    assert _titles(P.open_deck(copy.read_bytes())) == ["New"]
    again = await tool.execute(action="edit", path="saved/downloads/s1/Pitch.pptx", edits=edit, **kw)
    assert not again.success and "already exists" in again.error
    onward = await tool.execute(action="edit", path=str(copy),
                                edits=[{"op": "set_text", "slide": 1, "shape": "title", "text": "Newer"}], **kw)
    assert onward.success and _titles(P.open_deck(copy.read_bytes())) == ["Newer"]


async def test_create_wont_overwrite_unless_asked(workspace):
    _, kw = workspace
    tool = PptxTool()
    slides = [{"title": "x"}]
    assert (await tool.execute(action="create", path="A.pptx", slides=slides, **kw)).success
    again = await tool.execute(action="create", path="A.pptx", slides=slides, **kw)
    assert not again.success and "already exists" in again.error
    assert (await tool.execute(action="create", path="A.pptx", slides=slides, overwrite=True, **kw)).success


async def test_a_template_file_is_used_and_cant_be_edited_as_a_deck(workspace):
    root, kw = workspace
    path = root / "saved" / "house.potx"
    path.write_bytes(_as_potx(_template()))
    tool = PptxTool()
    result = await tool.execute(action="create", path="New.pptx", template=str(path),
                                slides=[{"layout": "Brand Content", "title": "Mine"}], **kw)
    assert result.success, result.error
    assert _titles(P.open_deck(next(root.rglob("New.pptx")).read_bytes())) == ["Mine"]
    read = await tool.execute(action="read", path=str(path), **kw)
    assert read.success
    edit = await tool.execute(action="edit", path=str(path), edits=[{"op": "hide_slide", "slide": 1}], **kw)
    assert not edit.success and "action=create with template" in edit.error


async def test_a_picture_goes_in_by_its_saved_path(workspace):
    root, kw = workspace
    _picture(root / "saved" / "chart.png")
    tool = PptxTool()
    result = await tool.execute(action="create", path="Pic.pptx",
                                slides=[{"title": "x", "image": "saved/chart.png"}], **kw)
    assert result.success, result.error
    prs = P.open_deck(next(root.rglob("Pic.pptx")).read_bytes())
    assert [s for s in prs.slides[0].shapes if s._element.tag == qn("p:pic")]


async def test_render_says_plainly_when_libreoffice_is_missing(workspace, monkeypatch):
    root, kw = workspace
    monkeypatch.setattr(P, "find_renderer", lambda: None)
    tool = PptxTool()
    await tool.execute(action="create", path="R.pptx", slides=[{"title": "x"}], **kw)
    result = await tool.execute(action="render", path=str(next(root.rglob("R.pptx"))), **kw)
    assert not result.success and "LibreOffice" in result.error and "check" in result.error


async def test_render_saves_a_picture_per_slide(workspace, monkeypatch):
    root, kw = workspace
    monkeypatch.setattr(P, "find_renderer", lambda: "soffice")
    monkeypatch.setattr(P, "render_pages", lambda data, renderer: [b"page1", b"page2"])
    tool = PptxTool()
    await tool.execute(action="create", path="R.pptx", slides=[{"title": "a"}, {"title": "b"}], **kw)
    result = await tool.execute(action="render", path=str(next(root.rglob("R.pptx"))), slide="2", **kw)
    assert result.success, result.error
    picture = next(root.rglob("R-slide-2.png"))
    assert picture.read_bytes() == b"page2"
    assert "Slide 2" in result.content and "Slide 1" not in result.content


def test_rendering_a_deck_never_fetches_what_it_links_to(tmp_path):
    prs = _deck({"title": "x"})
    prs.slides[0].part.relate_to("http://169.254.169.254/latest", P.RT.IMAGE, is_external=True)
    prs.slides[0].part.relate_to("https://example.com", P.RT.HYPERLINK, is_external=True)
    clean = P.without_outside_links(P.deck_bytes(prs))
    with zipfile.ZipFile(io.BytesIO(clean)) as archive:
        rels = archive.read("ppt/slides/_rels/slide1.xml.rels")
    assert b"169.254" not in rels and b"about:blank" in rels
    assert b"https://example.com" in rels


def test_rendering_runs_libreoffice_on_a_clean_copy_with_its_own_profile(tmp_path):
    fake = tmp_path / "soffice"
    log = tmp_path / "args.txt"
    fake.write_text(
        "#!/bin/sh\n"
        f'echo "$@" > "{log}"\n'
        'outdir=""\n'
        'while [ $# -gt 0 ]; do if [ "$1" = "--outdir" ]; then outdir="$2"; fi; shift; done\n'
        'printf "%%PDF-1.4\\n" > "$outdir/deck.pdf"\n')
    fake.chmod(0o755)
    with pytest.raises(Exception):
        # The fake PDF has no pages to draw; what matters here is how it was called.
        P.render_pages(P.deck_bytes(_deck({"title": "x"})), str(fake))
    args = log.read_text()
    assert "--headless" in args and "--convert-to pdf" in args
    assert "-env:UserInstallation=file://" in args


def test_file_names():
    assert pptx_file_name("Pitch") == "Pitch.pptx"
    assert pptx_file_name("pitch.ppt") == "pitch.pptx"
    assert pptx_file_name("house.potx") == "house.pptx"
    assert pptx_file_name("m.pptm") == "m.pptm"
    assert pptx_file_name("q3.v2") == "q3.v2.pptx"


async def test_write_and_edit_point_to_the_tool(workspace):
    root, kw = workspace
    result = await WriteTool().execute(path="x.pptx", content="not a deck", **kw)
    assert not result.success and result.error == "pptx_needs_pptx_tool"
    assert "pptx tool" in result.content
    result = await EditTool().execute(path="x.pptx", old_string="a", new_string="b", **kw)
    assert not result.success and "pptx tool" in result.error


async def test_edits_passed_as_json_text_are_accepted(workspace):
    root, kw = workspace
    tool = PptxTool()
    await tool.execute(action="create", path="J.pptx", slides=[{"title": "a"}], **kw)
    result = await tool.execute(action="edit", path=str(next(root.rglob("J.pptx"))),
                                edits='[{"op": "set_text", "slide": 1, "shape": "title", "text": "b"}]', **kw)
    assert result.success, result.error


def test_every_agent_has_the_tool_and_eco_mode_finds_it():
    from captain_claw.agent_orchestration_mixin import _ECO_INTENT_PATTERNS
    from captain_claw.config import ToolsConfig

    assert "pptx" in ToolsConfig().enabled
    for text in ("update Pitch.pptx", "make a slide deck", "fix the powerpoint", "napravi prezentaciju", "dodaj slajd"):
        assert any("pptx" in tools for pattern, tools in _ECO_INTENT_PATTERNS if pattern.search(text)), text


def test_the_schema_suits_every_provider():
    def walk(schema, where):
        for key, value in (schema.get("properties") or {}).items():
            assert "type" in value, f"{where}.{key} has no type"
            if value["type"] == "array":
                assert value.get("items"), f"{where}.{key} has no items"
                if isinstance(value["items"], dict):
                    walk(value["items"], f"{where}.{key}[]")
            if value["type"] == "object":
                walk(value, f"{where}.{key}")

    walk(PptxTool.parameters, "pptx")
    json.dumps(PptxTool.parameters)
