"""Drive mount — placeholder tree, manifest, lazy listing, refresh pruning.

Runs against a fake DriveClient (an in-memory Drive tree), so no Google, no
network, no tokens. Verifies the mount is a real filesystem tree of honest
placeholders and that a refresh reflects upstream adds/removes.
"""

import tempfile
from pathlib import Path

import pytest

from captain_claw import vfs_drive
from captain_claw.drive_client import DriveFile, FOLDER_MIME


class FakeDrive:
    """An in-memory Drive: folder_id -> list of child DriveFiles."""

    def __init__(self, tree: dict[str, list[DriveFile]], content: dict | None = None):
        self.tree = tree
        # file_id -> (bytes, ext) returned by fetch()
        self.content = content or {}
        self.list_calls: list[str] = []
        self.drive_ids: list[str] = []  # the shared-drive corpus of each list_folder call
        self.fetch_calls: list[str] = []

    async def list_folder(self, folder_id, *, drive_id="", order_by="folder,name",
                          max_files=None, sleep=None):
        self.list_calls.append(folder_id)
        self.drive_ids.append(drive_id)
        children = self.tree.get(folder_id, [])
        if max_files is not None and len(children) > max_files:
            return children[:max_files], True
        return list(children), False

    async def fetch(self, f, *, sleep=None):
        self.fetch_calls.append(f.id)
        return self.content[f.id]

    async def close(self):
        pass


def _folder(fid, name):
    return DriveFile(id=fid, name=name, mime_type=FOLDER_MIME)


def _file(fid, name, *, size=100, modified="2026-07-20T10:00:00Z", mime="text/plain"):
    return DriveFile(id=fid, name=name, mime_type=mime, size=size, modified_time=modified)


@pytest.fixture()
def mount(monkeypatch):
    with tempfile.TemporaryDirectory() as tmp:
        monkeypatch.setattr(vfs_drive, "user_root", lambda uid: Path(tmp) / uid)
        yield Path(tmp)


def _sample_drive():
    # root/ ├ report.pdf  ├ notes.txt  └ sub/ └ deep.md
    return FakeDrive({
        "ROOT": [_folder("F_SUB", "sub"), _file("f1", "report.pdf", mime="application/pdf"),
                 _file("f2", "notes.txt")],
        "F_SUB": [_file("f3", "deep.md", mime="text/markdown")],
    })


class TestCreateMount:
    async def test_builds_the_real_tree(self, mount):
        drive = _sample_drive()
        summary = await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        assert (root / "report.pdf").is_file()
        assert (root / "notes.txt").is_file()
        assert (root / "sub").is_dir()
        assert (root / "sub" / "deep.md").is_file()
        assert summary["files"] == 3 and summary["dirs"] == 1
        assert summary["truncated"] is False

    async def test_placeholder_is_an_honest_marker_not_empty(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        text = (vfs_drive.mount_root("alice", "acme") / "report.pdf").read_text()
        assert "Google Drive" in text
        assert "report.pdf" in text
        assert "id f1" in text
        assert "not been downloaded" in text.lower()

    async def test_manifest_records_every_file_as_placeholder(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        man = vfs_drive.Manifest.load(vfs_drive.mount_root("alice", "acme"))
        assert set(man.files) == {"report.pdf", "notes.txt", "sub/deep.md"}
        assert all(e["state"] == "placeholder" for e in man.files.values())
        assert man.dirs[""] == "ROOT" and man.dirs["sub"] == "F_SUB"
        assert man.files["report.pdf"]["file_id"] == "f1"

    async def test_is_placeholder_reads_the_manifest(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        assert vfs_drive.is_placeholder(root, "report.pdf") is True
        assert vfs_drive.is_placeholder(root, "does/not/exist") is False


class TestCap:
    async def test_cap_truncates_and_reports(self, mount):
        many = FakeDrive({"ROOT": [_file(f"f{i}", f"{i}.txt") for i in range(10)]})
        summary = await vfs_drive.create_mount(many, "alice", "big", "ROOT", max_files=4)
        assert summary["files"] == 4
        assert summary["truncated"] is True


class TestRefresh:
    async def test_new_upstream_file_appears(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        drive.tree["ROOT"].append(_file("f9", "added.txt"))
        root = vfs_drive.mount_root("alice", "acme")
        await vfs_drive.sync(drive, root)
        assert (root / "added.txt").is_file()
        assert "added.txt" in vfs_drive.Manifest.load(root).files

    async def test_vanished_upstream_file_is_pruned_locally(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        assert (root / "notes.txt").is_file()
        drive.tree["ROOT"] = [f for f in drive.tree["ROOT"] if f.name != "notes.txt"]
        await vfs_drive.sync(drive, root)
        assert not (root / "notes.txt").exists()  # no ghost left behind
        assert "notes.txt" not in vfs_drive.Manifest.load(root).files

    async def test_cloned_state_survives_refresh_when_unchanged(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        man = vfs_drive.Manifest.load(root)
        man.files["report.pdf"]["state"] = "cloned"
        man.save(root)
        await vfs_drive.sync(drive, root)  # same modifiedTime upstream
        assert vfs_drive.Manifest.load(root).files["report.pdf"]["state"] == "cloned"

    async def test_changed_file_reverts_to_placeholder(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        man = vfs_drive.Manifest.load(root)
        man.files["report.pdf"]["state"] = "cloned"
        man.save(root)
        # upstream edit bumps modifiedTime
        for f in drive.tree["ROOT"]:
            if f.name == "report.pdf":
                f.modified_time = "2026-07-25T12:00:00Z"
        await vfs_drive.sync(drive, root)
        assert vfs_drive.Manifest.load(root).files["report.pdf"]["state"] == "placeholder"


class TestLazyListDir:
    async def test_list_dir_requires_a_known_directory(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        with pytest.raises(ValueError):
            await vfs_drive.list_dir(drive, root, "unknown-subdir")

    async def test_list_dir_of_known_subdir(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        out = await vfs_drive.list_dir(drive, root, "sub")
        assert out["files"] == 1  # the listing reports sub/'s one child
        assert (root / "sub" / "deep.md").is_file()


class TestLinkEntry:
    def test_link_entry_is_a_readonly_gdrive_link(self):
        ent = vfs_drive.link_entry("alice", "acme", "ROOT", clonemd=True)
        assert ent["mode"] == "ro"
        assert ent["kind"] == "gdrive"
        assert ent["drive"]["folder_id"] == "ROOT"
        assert ent["drive"]["clonemd"] is True
        assert Path(ent["path"]).name == "acme"

    def test_is_drive_link(self):
        assert vfs_drive.is_drive_link({"kind": "gdrive"}) is True
        assert vfs_drive.is_drive_link({"mode": "rw"}) is False
        assert vfs_drive.is_drive_link(None) is False


class TestRemoveMount:
    async def test_keep_cloned_preserves_cloned_files(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        man = vfs_drive.Manifest.load(root)
        man.files["report.pdf"]["state"] = "cloned"
        man.save(root)
        result = vfs_drive.remove_mount("alice", "acme", keep_cloned=True)
        assert result["kept"] == 1
        assert (root / "report.pdf").is_file()   # cloned file survives
        assert not (root / "notes.txt").exists()  # placeholder removed

    async def test_full_removal(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        vfs_drive.remove_mount("alice", "acme", keep_cloned=False)
        assert not root.exists()


# ── Phase 2: reading remote files as if local ─────────────────────────

from captain_claw.drive_client import DriveError


class TestBytesToText:
    def test_plain_text_decodes(self):
        assert vfs_drive.bytes_to_text(b"hello\nworld", ".txt", "a.txt") == "hello\nworld"

    def test_unknown_binary_is_not_text(self):
        assert vfs_drive.bytes_to_text(b"\x89PNG\r\n", ".png", "img.png") is None

    def test_xlsx_converts_to_markdown(self):
        import io
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.append(["Name", "Score"])
        ws.append(["Ada", 99])
        buf = io.BytesIO()
        wb.save(buf)
        out = vfs_drive.bytes_to_text(buf.getvalue(), ".xlsx", "s.xlsx")
        # Routing is what matters here: .xlsx bytes go through the xlsx
        # extractor (which emits a markdown doc), not decoded as raw text.
        assert out is not None and out.lstrip().startswith("#")


class TestFindMount:
    async def test_finds_mount_and_rel(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        found = vfs_drive.find_mount(root / "sub" / "deep.md")
        assert found == (root, "sub/deep.md")

    async def test_outside_any_mount_is_none(self, mount):
        assert vfs_drive.find_mount(mount / "nowhere.txt") is None


def _text_drive():
    """A mount whose files carry real fetchable text."""
    drive = FakeDrive(
        {"ROOT": [_file("f1", "notes.txt"), _file("f2", "data.csv")]},
        content={"f1": (b"the quarterly revenue rose sharply", ".txt"),
                 "f2": (b"a,b\n1,2", ".csv")},
    )
    return drive


class TestHydrate:
    async def test_hydrate_fetches_caches_and_marks(self, mount):
        drive = _text_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        text = await vfs_drive.hydrate(drive, root, "notes.txt")
        assert "quarterly revenue" in text
        man = vfs_drive.Manifest.load(root)
        assert man.files["notes.txt"]["state"] == "hydrated"
        # cache file written under .drive-cache/<file_id>
        assert (root / ".drive-cache" / "f1").is_file()

    async def test_second_hydrate_uses_cache_no_refetch(self, mount):
        drive = _text_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        await vfs_drive.hydrate(drive, root, "notes.txt")
        drive.fetch_calls.clear()
        await vfs_drive.hydrate(drive, root, "notes.txt")
        assert drive.fetch_calls == []  # served from cache

    async def test_changed_upstream_refetches(self, mount):
        drive = _text_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        await vfs_drive.hydrate(drive, root, "notes.txt")
        # simulate a refresh that bumped modifiedTime for this file
        man = vfs_drive.Manifest.load(root)
        man.files["notes.txt"]["modified"] = "2026-08-01T00:00:00Z"
        man.save(root)
        drive.fetch_calls.clear()
        await vfs_drive.hydrate(drive, root, "notes.txt")
        assert drive.fetch_calls == ["f1"]  # cache was stale

    async def test_binary_file_raises_readable_error(self, mount):
        drive = FakeDrive({"ROOT": [_file("f1", "pic.png", mime="image/png")]},
                          content={"f1": (b"\x89PNG", ".png")})
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        with pytest.raises(DriveError):
            await vfs_drive.hydrate(drive, root, "pic.png")


class TestReadThrough:
    async def test_placeholder_returns_content(self, mount, monkeypatch):
        drive = _text_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        monkeypatch.setattr("captain_claw.drive_client.make_client", lambda: drive)
        out = await vfs_drive.read_through(root / "notes.txt")
        assert out is not None and "quarterly revenue" in out

    async def test_non_mount_path_returns_none(self, mount):
        assert await vfs_drive.read_through(mount / "loose.txt") is None

    async def test_cloned_file_reads_normally(self, mount, monkeypatch):
        drive = _text_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        man = vfs_drive.Manifest.load(root)
        man.files["notes.txt"]["state"] = "cloned"
        man.save(root)
        monkeypatch.setattr("captain_claw.drive_client.make_client", lambda: drive)
        # cloned => read the real on-disk file, not a fetch
        assert await vfs_drive.read_through(root / "notes.txt") is None

    async def test_drive_error_falls_back_to_marker(self, mount, monkeypatch):
        class Boom(FakeDrive):
            async def fetch(self, f, *, sleep=None):
                raise DriveError("network down")

        drive = Boom({"ROOT": [_file("f1", "notes.txt")]}, content={})
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        monkeypatch.setattr("captain_claw.drive_client.make_client", lambda: drive)
        out = await vfs_drive.read_through(root / "notes.txt")
        assert "could not fetch" in out and "network down" in out


class TestFilterSearchable:
    async def test_placeholders_skipped_and_counted(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        paths = [root / "report.pdf", root / "notes.txt", root / "sub" / "deep.md"]
        searchable, skipped = vfs_drive.filter_searchable(paths)
        assert searchable == []  # all placeholders
        assert skipped == 3

    async def test_cloned_is_searchable(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        man = vfs_drive.Manifest.load(root)
        man.files["notes.txt"]["state"] = "cloned"
        man.save(root)
        searchable, skipped = vfs_drive.filter_searchable([root / "notes.txt", root / "report.pdf"])
        assert searchable == [root / "notes.txt"]
        assert skipped == 1

    async def test_non_mount_files_pass_through(self, mount):
        loose = mount / "x.txt"
        searchable, skipped = vfs_drive.filter_searchable([loose])
        assert searchable == [loose] and skipped == 0

    async def test_manifest_and_cache_are_excluded_not_counted(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT")
        root = vfs_drive.mount_root("alice", "acme")
        paths = [root / ".drive-manifest.json", root / ".drive-cache" / "f1"]
        searchable, skipped = vfs_drive.filter_searchable(paths)
        assert searchable == [] and skipped == 0  # internal, neither searched nor counted


class TestReadOnlyWrites:
    """Agent-side write/edit must refuse a mode:ro link, not just the FD panel."""

    def _ro_link(self, tmp, monkeypatch):
        import json as _json

        monkeypatch.setenv("CLAW_VFS_ROOT", str(tmp))
        monkeypatch.setenv("CLAW_VFS_USER", "local")
        monkeypatch.delenv("FD_OWNER_ID", raising=False)
        uroot = tmp / "local"
        uroot.mkdir(parents=True, exist_ok=True)
        (uroot / ".drive" / "acme").mkdir(parents=True, exist_ok=True)
        (uroot / ".vfs-links.json").write_text(_json.dumps({
            "acme": {"path": str(uroot / ".drive" / "acme"), "mode": "ro", "kind": "gdrive"}
        }))

    def test_project_is_readonly_reads_the_link(self, tmp_path, monkeypatch):
        from captain_claw import vfs

        self._ro_link(tmp_path, monkeypatch)
        assert vfs.project_is_readonly("acme") is True
        assert vfs.project_is_readonly("not-a-link") is False

    async def test_write_tool_refuses_ro_mount(self, tmp_path, monkeypatch):
        from captain_claw.tools.write import WriteTool

        self._ro_link(tmp_path, monkeypatch)
        r = await WriteTool().execute(path="vfs:acme/x.txt", content="nope")
        assert r.success is False
        assert "read-only" in r.error.lower()

    async def test_edit_tool_refuses_ro_mount(self, tmp_path, monkeypatch):
        from captain_claw.tools.edit import EditTool

        self._ro_link(tmp_path, monkeypatch)
        r = await EditTool().execute(
            path="vfs:acme/x.txt", action="replace_string",
            old_string="a", new_string="b",
        )
        assert r.success is False
        assert "read-only" in r.error.lower()


class TestReadToolIntegration:
    """The read tool itself returns Drive content for a placeholder path —
    the actual 'read remote files as if local' behaviour, end to end."""

    async def test_read_tool_hydrates_placeholder(self, tmp_path, monkeypatch):
        import json as _json

        from captain_claw.tools.read import ReadTool

        # Align every root: the agent resolver (vfs.user_root via CLAW_VFS_ROOT)
        # and vfs_drive.user_root must land on the same tree.
        monkeypatch.setenv("CLAW_VFS_ROOT", str(tmp_path))
        monkeypatch.setenv("CLAW_VFS_USER", "local")
        monkeypatch.delenv("FD_OWNER_ID", raising=False)
        monkeypatch.setattr(vfs_drive, "user_root", lambda uid: tmp_path / uid)

        drive = _text_drive()
        await vfs_drive.create_mount(drive, "local", "acme", "ROOT")
        uroot = tmp_path / "local"
        (uroot / ".vfs-links.json").write_text(_json.dumps({
            "acme": vfs_drive.link_entry("local", "acme", "ROOT"),
        }))
        monkeypatch.setattr("captain_claw.drive_client.make_client", lambda: drive)

        r = await ReadTool().execute(path="vfs:acme/notes.txt")
        assert r.success is True
        assert "quarterly revenue" in r.content  # fetched, not the marker

    async def test_read_tool_on_missing_mount_file_is_clean_404(self, tmp_path, monkeypatch):
        import json as _json

        from captain_claw.tools.read import ReadTool

        monkeypatch.setenv("CLAW_VFS_ROOT", str(tmp_path))
        monkeypatch.setenv("CLAW_VFS_USER", "local")
        monkeypatch.delenv("FD_OWNER_ID", raising=False)
        monkeypatch.setattr(vfs_drive, "user_root", lambda uid: tmp_path / uid)
        drive = _text_drive()
        await vfs_drive.create_mount(drive, "local", "acme", "ROOT")
        uroot = tmp_path / "local"
        (uroot / ".vfs-links.json").write_text(_json.dumps({
            "acme": vfs_drive.link_entry("local", "acme", "ROOT"),
        }))
        r = await ReadTool().execute(path="vfs:acme/nope.txt")
        assert r.success is False and "not found" in r.error.lower()


# ── Phase 3: clonemd (convert to real local Markdown) ─────────────────


def _clonable_drive():
    """Deterministic clone content: a Google Doc, a plain-text note, another
    Google Doc, and an image that stays a placeholder. Native docs export to
    markdown bytes that decode directly, so no real PDF/Office parsing is
    involved — clone ROUTING and naming are what these tests pin."""
    return FakeDrive(
        {"ROOT": [
            _file("g1", "Report", mime="application/vnd.google-apps.document"),
            _file("f2", "notes.txt"),
            _file("g3", "Quarterly Review",
                  mime="application/vnd.google-apps.document"),
            _file("i1", "logo.png", mime="image/png"),
        ]},
        content={
            "g1": (b"# Report\n\nThe merger closed.", ".md"),
            "f2": (b"the merger closed in Q3", ".txt"),
            "g3": (b"# Quarterly Review\n\nRevenue up.", ".md"),
            "i1": (b"\x89PNG", ".png"),
        },
    )


class TestClonemd:
    async def test_clone_creates_real_markdown_and_drops_placeholder(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        # "Report" (Google Doc) -> Report.md with real content.
        assert (root / "Report.md").is_file()
        assert "merger closed" in (root / "Report.md").read_text()
        man = vfs_drive.Manifest.load(root)
        assert man.files["Report"]["state"] == "cloned"
        assert man.files["Report"]["cloned_path"] == "Report.md"

    async def test_docx_style_rename_drops_original(self, mount):
        # A .docx-named source (with markdown-export bytes so no real parsing):
        # report.docx -> report.md, original placeholder path gone.
        drive = FakeDrive(
            {"ROOT": [_file("d1", "report.docx", mime="text/markdown")]},
            content={"d1": (b"converted body", ".md")},
        )
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        assert (root / "report.md").is_file()
        assert not (root / "report.docx").exists()

    async def test_plain_text_cloned_verbatim_keeps_name(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        assert (root / "notes.txt").read_text() == "the merger closed in Q3"
        assert vfs_drive.Manifest.load(root).files["notes.txt"]["state"] == "cloned"

    async def test_unconvertible_stays_placeholder(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        assert vfs_drive.Manifest.load(root).files["logo.png"]["state"] == "placeholder"
        assert "Google Drive" in (root / "logo.png").read_text()

    async def test_corrupt_convertible_falls_back_to_placeholder(self, mount):
        # A .pdf whose bytes aren't a real PDF: conversion fails, and clone must
        # degrade to a placeholder rather than crash.
        drive = FakeDrive(
            {"ROOT": [_file("f1", "broken.pdf", mime="application/pdf")]},
            content={"f1": (b"not really a pdf", ".pdf")},
        )
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        assert vfs_drive.Manifest.load(root).files["broken.pdf"]["state"] == "placeholder"
        assert (root / "broken.pdf").is_file()

    async def test_collision_keeps_distinct_names(self, mount):
        # Two same-stem sources both want the same .md name; the second falls
        # back to keeping its full name.
        drive = FakeDrive(
            {"ROOT": [
                _file("g1", "Report", mime="application/vnd.google-apps.document"),
                _file("d2", "Report.docx", mime="text/markdown"),
            ]},
            content={"g1": (b"# doc", ".md"), "d2": (b"converted", ".md")},
        )
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        files = {p.name for p in root.iterdir() if p.is_file() and not p.name.startswith(".")}
        assert "Report.md" in files
        assert "Report.docx.md" in files  # collision fallback keeps the ext

    async def test_cloned_files_are_grep_searchable(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        paths = [root / "notes.txt", root / "Report.md", root / "logo.png"]
        searchable, skipped = vfs_drive.filter_searchable(paths)
        assert (root / "notes.txt") in searchable
        assert (root / "Report.md") in searchable
        assert skipped == 1  # only the image placeholder

    async def test_unchanged_reclone_skips_refetch(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        drive.fetch_calls.clear()
        await vfs_drive.sync(drive, root)
        assert drive.fetch_calls == []  # nothing changed → no refetch

    async def test_changed_upstream_reconverts(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        for f in drive.tree["ROOT"]:
            if f.name == "notes.txt":
                f.modified_time = "2026-09-01T00:00:00Z"
        drive.content["f2"] = (b"restated: closed in Q4", ".txt")
        drive.fetch_calls.clear()
        await vfs_drive.sync(drive, root)
        assert "f2" in drive.fetch_calls
        assert (root / "notes.txt").read_text() == "restated: closed in Q4"

    async def test_vanished_cloned_output_is_pruned(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        assert (root / "Report.md").is_file()
        drive.tree["ROOT"] = [f for f in drive.tree["ROOT"] if f.name != "Report"]
        await vfs_drive.sync(drive, root)
        assert not (root / "Report.md").exists()  # cloned output removed too
        assert "Report" not in vfs_drive.Manifest.load(root).files

    async def test_disabling_clonemd_is_nondestructive(self, mount):
        drive = _clonable_drive()
        await vfs_drive.create_mount(drive, "alice", "acme", "ROOT", clonemd=True)
        root = vfs_drive.mount_root("alice", "acme")
        man = vfs_drive.Manifest.load(root)
        man.clonemd = False
        man.save(root)
        await vfs_drive.sync(drive, root)
        assert (root / "Report.md").is_file()  # left in place
        assert vfs_drive.Manifest.load(root).files["Report"]["state"] == "cloned"


class TestSharedDriveMount:
    """A mount rooted in a Shared Drive threads its drive id into every listing,
    so folders read back with the right corpus (not the empty default)."""

    async def test_shared_drive_id_flows_into_listings(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "team", "ROOT", shared_drive_id="DRIVE-X")
        # Every list_folder during the sync carried the shared-drive corpus.
        assert drive.drive_ids and all(d == "DRIVE-X" for d in drive.drive_ids)
        man = vfs_drive.Manifest.load(vfs_drive.mount_root("alice", "team"))
        assert man.shared_drive_id == "DRIVE-X"

    async def test_my_drive_mount_uses_no_corpus(self, mount):
        drive = _sample_drive()
        await vfs_drive.create_mount(drive, "alice", "mine", "ROOT")
        assert all(d == "" for d in drive.drive_ids)  # default corpus
        assert vfs_drive.Manifest.load(vfs_drive.mount_root("alice", "mine")).shared_drive_id == ""

    def test_link_entry_carries_shared_drive_id(self):
        ent = vfs_drive.link_entry("alice", "team", "ROOT", shared_drive_id="DRIVE-X")
        assert ent["drive"]["shared_drive_id"] == "DRIVE-X"
