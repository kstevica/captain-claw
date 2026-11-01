"""Excel formulas for agents: parse them, calculate them, keep them correct.

A workbook made or changed without a calculation engine stores formulas with no
result, so previews, phones and our own readers show blanks. This module gives
:mod:`captain_claw.xlsx_ops` what it needs to store real results:

* :func:`compile_formula` parses a formula once (openpyxl's tokenizer, a small
  precedence parser here) and lists what it depends on;
* :func:`evaluate` calculates it against a model of the workbook;
* :func:`normalize_for_storage` turns what a model typed into what the file must
  hold (``XLOOKUP`` is stored as ``_xlfn.XLOOKUP``, or Excel shows ``#NAME?``);
* :func:`shift_formula` and :func:`rename_sheet_in` keep references right when
  rows, columns or sheets move, the way Excel itself does.

What isn't implemented raises :class:`NotComputable` instead of guessing: the
caller leaves that cell for Excel to calculate on open and says so. Number
results were compared with real Excel (tests/test_tools/test_xlsx_formula.py
pins them without it).
"""

from __future__ import annotations

import datetime as dt
import math
import re
from dataclasses import dataclass, field
from decimal import ROUND_DOWN, ROUND_HALF_UP, ROUND_UP, Decimal, InvalidOperation
from functools import lru_cache
from typing import Any, Callable, Iterable, Iterator

from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import Token, TokenizerError
from openpyxl.formula.translate import Translator, TranslatorError
from openpyxl.utils import FORMULAE, column_index_from_string, get_column_letter

MAX_ROW = 1_048_576
MAX_COL = 16_384


class XlError:
    """An Excel error value (a value in a cell, not a Python exception)."""

    __slots__ = ("code",)

    def __init__(self, code: str) -> None:
        self.code = code

    def __repr__(self) -> str:
        return self.code

    def __eq__(self, other: object) -> bool:
        return isinstance(other, XlError) and other.code == self.code

    def __hash__(self) -> int:
        return hash(self.code)


DIV0, NA, NAME, NULL, NUM, REF, VALUE = (
    XlError("#DIV/0!"), XlError("#N/A"), XlError("#NAME?"), XlError("#NULL!"),
    XlError("#NUM!"), XlError("#REF!"), XlError("#VALUE!"),
)
ERRORS = {e.code: e for e in (DIV0, NA, NAME, NULL, NUM, REF, VALUE)}
ERRORS["#GETTING_DATA"] = NA


class FormulaError(ValueError):
    """A formula that can't be stored; the message is for the model."""


class NotComputable(Exception):
    """Excel can calculate this, we can't: the cell is left for Excel."""


class NeedCell(Exception):
    """A cell this formula reads hasn't been calculated yet (dynamic references)."""

    def __init__(self, sheet: str, row: int, col: int) -> None:
        super().__init__(sheet, row, col)
        self.key = (sheet, row, col)


class _Fail(Exception):
    """Early exit from a function with an error value."""

    def __init__(self, error: XlError) -> None:
        super().__init__(error.code)
        self.error = error


# ── Names Excel knows ────────────────────────────────────────────────

# Functions newer than the file format: stored with a prefix, or Excel shows
# #NAME? until someone retypes the formula.
_FUTURE = """ACOT ACOTH AGGREGATE ARABIC BASE BETA.DIST BETA.INV BINOM.DIST BINOM.DIST.RANGE BINOM.INV
BITAND BITLSHIFT BITOR BITRSHIFT BITXOR CEILING.MATH CEILING.PRECISE CHISQ.DIST CHISQ.DIST.RT CHISQ.INV
CHISQ.INV.RT CHISQ.TEST COMBINA CONCAT CONFIDENCE.NORM CONFIDENCE.T COT COTH COVARIANCE.P COVARIANCE.S CSC
CSCH DAYS DECIMAL ERF.PRECISE ERFC.PRECISE EXPON.DIST F.DIST F.DIST.RT F.INV F.INV.RT F.TEST FILTERXML
FLOOR.MATH FLOOR.PRECISE FORECAST.ETS FORECAST.ETS.CONFINT FORECAST.ETS.SEASONALITY FORECAST.ETS.STAT
FORECAST.LINEAR FORMULATEXT GAMMA GAMMA.DIST GAMMA.INV GAMMALN.PRECISE GAUSS HYPGEOM.DIST IFNA IFS IMCOSH
IMCOT IMCSC IMCSCH IMSEC IMSECH IMSINH IMTAN ISFORMULA ISOWEEKNUM LOGNORM.DIST LOGNORM.INV MAXIFS MINIFS
MODE.MULT MODE.SNGL MUNIT NEGBINOM.DIST NORM.DIST NORM.INV NORM.S.DIST NORM.S.INV NUMBERVALUE PDURATION
PERCENTILE.EXC PERCENTILE.INC PERCENTRANK.EXC PERCENTRANK.INC PERMUTATIONA PHI POISSON.DIST QUARTILE.EXC
QUARTILE.INC RANK.AVG RANK.EQ RRI SEC SECH SHEET SHEETS SKEW.P STDEV.P STDEV.S SWITCH T.DIST T.DIST.2T
T.DIST.RT T.INV T.INV.2T T.TEST TEXTJOIN UNICHAR UNICODE VAR.P VAR.S WEBSERVICE WEIBULL.DIST XOR Z.TEST
XLOOKUP XMATCH TEXTBEFORE TEXTAFTER ARRAYTOTEXT VALUETOTEXT""".split()
FUTURE_FUNCTIONS = frozenset(_FUTURE)
# These spill over neighbouring cells or need metadata this tool doesn't write:
# stored as a plain formula Excel would show them as `=@FILTER(…)` with one value.
SPILLING_FUNCTIONS = frozenset("""FILTER SORT SORTBY UNIQUE SEQUENCE RANDARRAY LET LAMBDA TEXTSPLIT VSTACK
HSTACK TOCOL TOROW WRAPROWS WRAPCOLS TAKE DROP CHOOSEROWS CHOOSECOLS EXPAND MAP REDUCE SCAN BYROW BYCOL
MAKEARRAY ANCHORARRAY SINGLE""".split())
_SPILL_PREFIX = {"FILTER": "_xlfn._xlws.", "SORT": "_xlfn._xlws."}
# They reach outside the workbook when the file is opened.
FORBIDDEN_FUNCTIONS = frozenset({"WEBSERVICE", "FILTERXML", "CALL", "REGISTER", "REGISTER.ID", "RTD",
                                 "SQL.REQUEST", "EXEC", "IMAGE"})
KNOWN_FUNCTIONS = frozenset(FORMULAE) | FUTURE_FUNCTIONS | SPILLING_FUNCTIONS | frozenset({
    "NETWORKDAYS.INTL", "WORKDAY.INTL", "ISO.CEILING", "ECMA.CEILING", "HYPERLINK", "SUBTOTAL",
    "GETPIVOTDATA", "CUBEVALUE", "CUBEMEMBER", "CUBESET", "ERROR.TYPE", "PHONETIC", "BAHTTEXT", "DATEDIF",
})
_PREFIX_RE = re.compile(r"_xlfn\._xlws\.|_xlfn\.|_xlws\.|_xlpm\.", re.I)


def bare_function_name(token_value: str) -> str:
    """``_xlfn.XLOOKUP(`` → ``XLOOKUP``."""
    return _PREFIX_RE.sub("", token_value.rstrip("(")).upper()


def display_formula(stored: str) -> str:
    """A stored formula as a person types it: no ``_xlfn.`` and a leading ``=``."""
    text = _PREFIX_RE.sub("", stored)
    return text if text.startswith("=") else "=" + text


# ── References ───────────────────────────────────────────────────────

_CELL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)(\d{1,7})$")
_COLS_RE = re.compile(r"^(\$?)([A-Za-z]{1,3}):(\$?)([A-Za-z]{1,3})$")
_ROWS_RE = re.compile(r"^(\$?)(\d{1,7}):(\$?)(\d{1,7})$")


@dataclass(frozen=True, slots=True)
class Box:
    """A block of cells, 1-based and inclusive."""

    sheet: str | None
    r1: int
    c1: int
    r2: int
    c2: int

    @property
    def is_cell(self) -> bool:
        return self.r1 == self.r2 and self.c1 == self.c2

    @property
    def area(self) -> int:
        return (self.r2 - self.r1 + 1) * (self.c2 - self.c1 + 1)

    def contains(self, row: int, col: int) -> bool:
        return self.r1 <= row <= self.r2 and self.c1 <= col <= self.c2


@dataclass(frozen=True, slots=True)
class NameOperand:
    sheet: str | None
    name: str


@dataclass(frozen=True, slots=True)
class TableOperand:
    table: str  # "" = the table this formula sits in
    special: tuple[str, ...]
    first_col: str | None
    last_col: str | None


@dataclass(frozen=True, slots=True)
class OpaqueOperand:
    """Valid in Excel, not followed here: other workbooks, 3-D references."""

    reason: str


def split_sheet_prefix(text: str) -> tuple[str | None, str, bool]:
    """``'My Sheet'!A1`` → (``My Sheet``, ``A1``, True)."""
    if text.startswith("'"):
        i = 1
        while i < len(text):
            if text[i] == "'":
                if i + 1 < len(text) and text[i + 1] == "'":
                    i += 2
                    continue
                break
            i += 1
        if i + 1 < len(text) and text[i + 1] == "!":
            return text[1:i].replace("''", "'"), text[i + 2:], True
        return None, text, False
    # The first ! : the reference after it may be Sheet1!#REF!
    bang = text.find("!")
    bracket = text.find("[")
    if bang > 0 and (bracket < 0 or bang < bracket or text.startswith("[")):
        return text[:bang], text[bang + 1:], True
    return None, text, False


def quote_sheet(name: str) -> str:
    """A sheet name as a formula needs it."""
    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", name) and not _CELL_RE.match(name) \
            and not re.fullmatch(r"[Rr]\d*[Cc]\d*", name) and name.upper() not in ("TRUE", "FALSE"):
        return name
    return "'" + name.replace("'", "''") + "'"


def parse_box(body: str, sheet: str | None = None) -> Box | None:
    """``A1``, ``$A$1:B2``, ``A:C`` or ``2:5`` → a :class:`Box`; None if it isn't one."""
    match = _CELL_RE.match(body)
    if match:
        col = column_index_from_string(match.group(2).upper())
        row = int(match.group(4))
        if col <= MAX_COL and 1 <= row <= MAX_ROW:
            return Box(sheet, row, col, row, col)
        return None
    match = _COLS_RE.match(body)
    if match:
        c1 = column_index_from_string(match.group(2).upper())
        c2 = column_index_from_string(match.group(4).upper())
        if c1 > MAX_COL or c2 > MAX_COL:
            return None
        return Box(sheet, 1, min(c1, c2), MAX_ROW, max(c1, c2))
    match = _ROWS_RE.match(body)
    if match:
        r1, r2 = int(match.group(2)), int(match.group(4))
        if not (1 <= r1 <= MAX_ROW and 1 <= r2 <= MAX_ROW):
            return None
        return Box(sheet, min(r1, r2), 1, max(r1, r2), MAX_COL)
    if body.count(":") == 1:
        left, right = body.split(":")
        a, b = _CELL_RE.match(left), _CELL_RE.match(right)
        if a and b:
            c1 = column_index_from_string(a.group(2).upper())
            c2 = column_index_from_string(b.group(2).upper())
            r1, r2 = int(a.group(4)), int(b.group(4))
            if max(c1, c2) <= MAX_COL and 1 <= min(r1, r2) and max(r1, r2) <= MAX_ROW:
                return Box(sheet, min(r1, r2), min(c1, c2), max(r1, r2), max(c1, c2))
    return None


def _parse_structured(text: str) -> TableOperand | None:
    """``Sales[Amount]``, ``[@Amount]``, ``Sales[[#Totals],[Amount]]``, ``T[[A]:[B]]``."""
    match = re.fullmatch(r"([^\[\]]*)\[(.*)\]", text, re.S)
    if match is None:
        return None
    table, inner = match.group(1), match.group(2)
    parts: list[str] = []
    depth, start = 0, 0
    if "[" not in inner:
        parts = [inner]
    else:
        i = 0
        while i < len(inner):
            ch = inner[i]
            if ch == "'":
                i += 2
                continue
            if ch == "[":
                depth += 1
            elif ch == "]":
                depth -= 1
            elif ch == "," and depth == 0:
                parts.append(inner[start:i])
                start = i + 1
            i += 1
        parts.append(inner[start:])
    special: list[str] = []
    first: str | None = None
    last: str | None = None

    def unescape(name: str) -> str:
        return re.sub(r"'(.)", r"\1", name)

    for raw in parts:
        part = raw.strip()
        this_row = part.startswith("@")
        if this_row:
            special.append("#This Row")
            part = part[1:].strip()
        if not part:
            continue
        columns = re.findall(r"\[((?:[^\[\]']|'.)*)\]", part) if part.startswith("[") else [part]
        for name in columns:
            if name.startswith("#"):
                special.append(name)
            elif first is None:
                first = last = unescape(name)
            else:
                last = unescape(name)
    return TableOperand(table, tuple(special), first, last)


def parse_operand(text: str) -> Box | NameOperand | TableOperand | OpaqueOperand | XlError:
    """What a RANGE token refers to."""
    if text.startswith("[") and "]" in text and "!" in text.split("]", 1)[1][:200]:
        return OpaqueOperand("it refers to another workbook")
    sheet, body, _ = split_sheet_prefix(text)
    if sheet is not None:
        if "[" in sheet:
            return OpaqueOperand("it refers to another workbook")
        if ":" in sheet:
            return OpaqueOperand("it uses a 3-D reference")
    if body == "#REF!":
        return REF
    if "[" in body:
        structured = _parse_structured(body)
        return structured if structured is not None else OpaqueOperand("unreadable table reference")
    box = parse_box(body, sheet)
    if box is not None:
        return box
    if re.fullmatch(r"[\w.\\?]+", body, re.U):
        return NameOperand(sheet, body)
    return OpaqueOperand(f"unreadable reference {text}")


# ── Parse tree ───────────────────────────────────────────────────────


class Ref:
    """A reference as a value: functions see where it points, not just what's there."""

    __slots__ = ("sheet", "r1", "c1", "r2", "c2")

    def __init__(self, sheet: str, r1: int, c1: int, r2: int, c2: int) -> None:
        self.sheet, self.r1, self.c1, self.r2, self.c2 = sheet, r1, c1, r2, c2

    @property
    def is_cell(self) -> bool:
        return self.r1 == self.r2 and self.c1 == self.c2

    @property
    def height(self) -> int:
        return self.r2 - self.r1 + 1

    @property
    def width(self) -> int:
        return self.c2 - self.c1 + 1

    def single(self, ctx: "Ctx") -> Any:
        return ctx.model.value(self.sheet, self.r1, self.c1)

    def _clip(self, ctx: "Ctx") -> tuple[int, int]:
        """Whole columns and rows stop at the used part of the sheet."""
        r2, c2 = self.r2, self.c2
        rows, cols = ctx.model.used(self.sheet)
        if self.height > 100_000:
            r2 = max(self.r1, min(r2, rows))
        if self.width > 2_000:
            c2 = max(self.c1, min(c2, cols))
        return r2, c2

    def matrix(self, ctx: "Ctx") -> list[list[Any]]:
        r2, c2 = self._clip(ctx)
        value = ctx.model.value
        return [[value(self.sheet, r, c) for c in range(self.c1, c2 + 1)] for r in range(self.r1, r2 + 1)]

    def present(self, ctx: "Ctx") -> Iterator[Any]:
        """Values of the cells that aren't blank (fast on sparse whole columns)."""
        for _, _, val in ctx.model.values_in(self.sheet, self.r1, self.c1, self.r2, self.c2):
            yield val

    def arr(self, ctx: "Ctx") -> "Arr":
        return Arr(self.matrix(ctx))


class Arr:
    """An array value (a literal, or a calculation over a range)."""

    __slots__ = ("rows",)

    def __init__(self, rows: list[list[Any]]) -> None:
        self.rows = rows

    @property
    def height(self) -> int:
        return len(self.rows)

    @property
    def width(self) -> int:
        return len(self.rows[0]) if self.rows else 0

    def flat(self) -> Iterator[Any]:
        for row in self.rows:
            yield from row

    def at(self, r: int, c: int) -> Any:
        """Element (r, c); a single row or column stretches, as in Excel; else #N/A."""
        if self.height == 1:
            r = 0
        if self.width == 1:
            c = 0
        if r < self.height and c < self.width:
            return self.rows[r][c]
        return NA


class Ctx:
    """Where a formula is being calculated."""

    __slots__ = ("model", "sheet", "row", "col", "array")

    def __init__(self, model: Any, sheet: str, row: int, col: int, array: bool = False) -> None:
        self.model, self.sheet, self.row, self.col, self.array = model, sheet, row, col, array

    def with_array(self) -> "Ctx":
        return self if self.array else Ctx(self.model, self.sheet, self.row, self.col, True)


class Node:
    __slots__ = ()

    def ev(self, ctx: Ctx) -> Any:
        raise NotImplementedError


class Lit(Node):
    __slots__ = ("value",)

    def __init__(self, value: Any) -> None:
        self.value = value

    def ev(self, ctx: Ctx) -> Any:
        return self.value


class Blank(Node):
    """An argument left out: ``IF(A1,,2)``."""

    __slots__ = ()

    def ev(self, ctx: Ctx) -> Any:
        return None


class Operand(Node):
    __slots__ = ("op",)

    def __init__(self, op: Box | NameOperand | TableOperand | OpaqueOperand | XlError) -> None:
        self.op = op

    def ev(self, ctx: Ctx) -> Any:
        return resolve_operand(self.op, ctx)


class Unary(Node):
    __slots__ = ("op", "x")

    def __init__(self, op: str, x: Node) -> None:
        self.op, self.x = op, x

    def ev(self, ctx: Ctx) -> Any:
        return _map1(self.x.ev(ctx), ctx, lambda v: _unary(self.op, v))


class Percent(Node):
    __slots__ = ("x",)

    def __init__(self, x: Node) -> None:
        self.x = x

    def ev(self, ctx: Ctx) -> Any:
        return _map1(self.x.ev(ctx), ctx, _percent)


class Binary(Node):
    __slots__ = ("op", "a", "b")

    def __init__(self, op: str, a: Node, b: Node) -> None:
        self.op, self.a, self.b = op, a, b

    def ev(self, ctx: Ctx) -> Any:
        a = _spread(self.a.ev(ctx), ctx)
        b = _spread(self.b.ev(ctx), ctx)
        if isinstance(a, Arr) or isinstance(b, Arr):
            left = a if isinstance(a, Arr) else Arr([[a]])
            right = b if isinstance(b, Arr) else Arr([[b]])
            height = max(left.height, right.height)
            width = max(left.width, right.width)
            return Arr([[_binary(self.op, left.at(r, c), right.at(r, c)) for c in range(width)]
                        for r in range(height)])
        return _binary(self.op, a, b)


class Call(Node):
    __slots__ = ("name", "args")

    def __init__(self, name: str, args: list[Node]) -> None:
        self.name, self.args = name, args

    def ev(self, ctx: Ctx) -> Any:
        spec = FUNCTIONS.get(self.name)
        if spec is None:
            if self.name in KNOWN_FUNCTIONS:
                raise NotComputable(f"{self.name} isn't calculated here")
            raise NotComputable(f"{self.name} is not a built-in function")
        try:
            if spec.lazy:
                return spec.fn(ctx, *self.args)
            arg_ctx = ctx.with_array() if spec.array_args else ctx
            values = [a.ev(arg_ctx) for a in self.args]
            return spec.fn(ctx, *values)
        except _Fail as fail:
            return fail.error
        except ZeroDivisionError:
            return DIV0
        except (OverflowError, ValueError):
            return NUM
        except TypeError as exc:
            # Too many or too few arguments.
            if "argument" in str(exc):
                raise NotComputable(f"{self.name}: wrong number of arguments") from None
            raise


class Unsupported(Node):
    """Parsed tokens this module has no meaning for (``A1 B1``, ``X:INDEX(…)``)."""

    __slots__ = ("why",)

    def __init__(self, why: str) -> None:
        self.why = why

    def ev(self, ctx: Ctx) -> Any:
        raise NotComputable(self.why)


_PRECEDENCE = {"=": 1, "<>": 1, "<": 1, "<=": 1, ">": 1, ">=": 1, "&": 2, "+": 3, "-": 3,
               "*": 4, "/": 4, "^": 5}


class _SyntaxProblem(Exception):
    pass


class _Parser:
    def __init__(self, items: list[Token]) -> None:
        self.items = [t for t in items if t.type != Token.WSPACE]
        self.i = 0

    def peek(self) -> Token | None:
        return self.items[self.i] if self.i < len(self.items) else None

    def next(self) -> Token:
        token = self.peek()
        if token is None:
            raise _SyntaxProblem("the formula ends too early (a missing bracket or value?)")
        self.i += 1
        return token

    def parse(self) -> Node:
        node = self.expression(0)
        left = self.peek()
        if left is not None:
            raise _SyntaxProblem(f"unexpected {left.value!r}")
        return node

    def expression(self, min_power: int) -> Node:
        left = self.prefix()
        while True:
            token = self.peek()
            if token is None:
                return left
            if token.type == Token.OP_POST:
                if 6 < min_power:
                    return left
                self.next()
                left = Percent(left)
                continue
            if token.type == Token.OP_IN:
                power = _PRECEDENCE.get(token.value)
                if power is None:
                    raise _SyntaxProblem(f"unknown operator {token.value!r}")
                if power < min_power:
                    return left
                self.next()
                left = Binary(token.value, left, self.expression(power + 1))
                continue
            return left

    def prefix(self) -> Node:
        token = self.next()
        if token.type == Token.OP_PRE:
            # Excel binds a leading minus tighter than ^: -2^2 is 4.
            return Unary(token.value, self.expression(7))
        if token.type == Token.OPERAND:
            return self.operand(token)
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            return self.call(token)
        if token.type == Token.PAREN and token.subtype == Token.OPEN:
            node = self.expression(0)
            closing = self.next()
            if closing.type != Token.PAREN or closing.subtype != Token.CLOSE:
                raise _SyntaxProblem(f"expected ) but found {closing.value!r}")
            return node
        if token.type == Token.ARRAY and token.subtype == Token.OPEN:
            return self.array()
        raise _SyntaxProblem(f"unexpected {token.value!r}")

    def operand(self, token: Token) -> Node:
        if token.subtype == Token.NUMBER:
            return Lit(float(token.value))
        if token.subtype == Token.TEXT:
            return Lit(token.value[1:-1].replace('""', '"'))
        if token.subtype == Token.LOGICAL:
            return Lit(token.value.upper() == "TRUE")
        if token.subtype == Token.ERROR:
            return Lit(ERRORS.get(token.value.upper(), VALUE))
        return Operand(parse_operand(token.value))

    def call(self, token: Token) -> Node:
        raw = token.value
        if raw.startswith(":"):
            raise _SyntaxProblem("a range built from functions")
        name = bare_function_name(raw)
        args: list[Node] = []
        if self.peek() is not None and self.peek().type == Token.FUNC and self.peek().subtype == Token.CLOSE:
            self.next()
            return Call(name, args)
        while True:
            token = self.peek()
            if token is None:
                raise _SyntaxProblem(f"{name}( has no closing bracket")
            if token.type == Token.SEP or (token.type == Token.FUNC and token.subtype == Token.CLOSE):
                args.append(Blank())
            else:
                args.append(self.expression(0))
            token = self.next()
            if token.type == Token.FUNC and token.subtype == Token.CLOSE:
                return Call(name, args)
            if token.type != Token.SEP:
                raise _SyntaxProblem(f"expected , or ) but found {token.value!r}")

    def array(self) -> Node:
        rows: list[list[Any]] = [[]]
        while True:
            token = self.next()
            if token.type == Token.ARRAY and token.subtype == Token.CLOSE:
                return Lit(Arr(rows))
            if token.type == Token.SEP:
                if token.subtype == Token.ROW:
                    rows.append([])
                continue
            negative = False
            if token.type == Token.OP_PRE:
                negative = token.value == "-"
                token = self.next()
            node = self.operand(token)
            if not isinstance(node, Lit):
                raise _SyntaxProblem("arrays hold constants only")
            value = node.value
            rows[-1].append(-value if negative else value)


@dataclass(slots=True)
class Compiled:
    """A formula, parsed once."""

    text: str
    tree: Node
    operands: list[Box | NameOperand | TableOperand | OpaqueOperand | XlError] = field(default_factory=list)
    functions: frozenset[str] = frozenset()
    syntax_problem: str = ""

    @property
    def dynamic(self) -> bool:
        """Could it read anything? Another workbook's numbers don't change with an
        edit here; what INDIRECT or a 3-D reference reads can't be known."""
        return bool(self.functions & _DYNAMIC) or any(
            isinstance(o, OpaqueOperand) and "3-D" in o.reason for o in self.operands)

    @property
    def position_sensitive(self) -> bool:
        return bool(self.functions & _POSITIONAL)


_DYNAMIC = frozenset({"OFFSET", "INDIRECT"})
_POSITIONAL = frozenset({"ROW", "COLUMN", "OFFSET", "INDIRECT", "CELL", "ADDRESS"})
VOLATILE = frozenset({"NOW", "TODAY", "RAND", "RANDBETWEEN", "RANDARRAY", "OFFSET", "INDIRECT", "INFO", "CELL"})


def _tokens(text: str) -> list[Token]:
    return Tokenizer(text if text.startswith("=") else "=" + text).items


@lru_cache(maxsize=200_000)
def compile_formula(stored: str) -> Compiled:
    """Parse a formula from a file. Never raises: what can't be parsed can't be
    calculated, and is still listed with whatever it refers to."""
    try:
        items = _tokens(stored)
    except (TokenizerError, ValueError) as exc:
        return Compiled(stored, Unsupported(f"unreadable formula: {exc}"), syntax_problem=str(exc))
    operands = []
    functions = set()
    for token in items:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            operands.append(parse_operand(token.value))
        elif token.type == Token.FUNC and token.subtype == Token.OPEN:
            functions.add(bare_function_name(token.value.lstrip(":")))
    try:
        tree: Node = _Parser(items).parse()
        problem = ""
    except _SyntaxProblem as exc:
        tree = Unsupported(f"unreadable formula: {exc}")
        problem = str(exc)
    return Compiled(stored, tree, operands, frozenset(functions), problem)


def _semicolons_to_commas(text: str) -> str:
    """Argument separators in a local Excel's style become the file's commas."""
    out: list[str] = []
    quote = ""
    braces = brackets = 0
    for ch in text:
        if quote:
            if ch == quote:
                quote = ""
        elif ch in "\"'":
            quote = ch
        elif ch == "{":
            braces += 1
        elif ch == "}":
            braces -= 1
        elif ch == "[":
            brackets += 1
        elif ch == "]":
            brackets -= 1
        elif ch == ";" and braces == 0 and brackets == 0:
            ch = ","
        out.append(ch)
    return "".join(out)


def normalize_for_storage(formula: str, *, allow_custom: bool = False) -> str:
    """What a model typed → the text the file must hold (no leading ``=``).

    Refuses what would open as an error or reach outside the file, so a bad
    formula fails the edit instead of turning up as ``#NAME?`` in front of a client.
    """
    text = formula.strip()
    if text.startswith("="):
        text = text[1:].strip()
    if not text:
        raise FormulaError("an empty formula")
    if len(text) > 8192:
        raise FormulaError("a formula can have at most 8,192 characters")
    text = _semicolons_to_commas(text)
    try:
        items = _tokens(text)
    except (TokenizerError, ValueError) as exc:
        raise FormulaError(f"=`{text[:120]}` can't be read: {exc}") from None
    for token in items:
        if token.type == Token.FUNC and token.subtype == Token.OPEN:
            raw = token.value.lstrip(":")
            name = bare_function_name(raw)
            if name in FORBIDDEN_FUNCTIONS:
                raise FormulaError(f"{name} reaches outside the workbook and isn't allowed here.")
            if name in SPILLING_FUNCTIONS:
                raise FormulaError(
                    f"{name} isn't written by this tool (it spills over cells or needs newer Excel "
                    "metadata). Work out the result and set the values, or use single-cell formulas "
                    "such as SUMIFS, INDEX/MATCH or XLOOKUP.")
            if name not in KNOWN_FUNCTIONS and not allow_custom:
                raise FormulaError(f"{name} is not an Excel function (in =`{text[:120]}`).")
            prefix = "_xlfn." if name in FUTURE_FUNCTIONS else ""
            token.value = (":" if token.value.startswith(":") else "") + prefix + name + "("
        elif token.type == Token.OPERAND and token.subtype == Token.RANGE:
            operand = parse_operand(token.value)
            if isinstance(operand, OpaqueOperand) and "workbook" in operand.reason:
                raise FormulaError(f"{token.value} refers to another workbook; links to other files "
                                   "aren't written by this tool.")
    try:
        _Parser(items).parse()
    except _SyntaxProblem as exc:
        raise FormulaError(f"=`{text[:120]}`: {exc}") from None
    return "".join(t.value for t in items)


def needs_array(stored: str) -> bool:
    """Does the formula work on whole ranges with operators or plain functions?

    Saved as an ordinary formula Excel takes one value from each range; saved as
    an array formula it does what the formula says. SUMPRODUCT and the like work
    on arrays by themselves.
    """
    return _wants_array(compile_formula(stored).tree)


def _is_block(node: Node) -> bool:
    return isinstance(node, Operand) and isinstance(node.op, Box) and not node.op.is_cell


def _wants_array(node: Node) -> bool:
    if isinstance(node, Binary):
        return _is_block(node.a) or _is_block(node.b) or _wants_array(node.a) or _wants_array(node.b)
    if isinstance(node, (Unary, Percent)):
        return _is_block(node.x) or _wants_array(node.x)
    if isinstance(node, Call):
        spec = FUNCTIONS.get(node.name)
        if spec is not None and spec.array_args:
            return False
        if spec is not None and spec.scalar and any(_is_block(a) for a in node.args):
            return True
        return any(_wants_array(a) for a in node.args)
    return False


def translate(stored: str, origin: str, target: str) -> str:
    """The formula as it reads when copied from *origin* to *target*."""
    try:
        moved = Translator("=" + stored.lstrip("="), origin=origin).translate_formula(target)
    except TranslatorError as exc:
        raise FormulaError(f"copying the formula to {target} leaves the sheet ({exc})") from None
    return moved[1:] if moved.startswith("=") else moved


# ── Rewriting references when things move ────────────────────────────


def _shift_edge(value: int, at: int, count: int) -> int | None:
    """Where row or column *value* goes; None when it's deleted."""
    if count >= 0:
        return value + count if value >= at else value
    gone = -count
    if value < at:
        return value
    if value < at + gone:
        return None
    return value - gone


def shift_span(first: int, last: int, at: int, count: int) -> tuple[int, int] | None:
    """A run of rows/columns after an insert (*count* > 0) or delete; None if it all went."""
    if count >= 0:
        return (first + count if first >= at else first, last + count if last >= at else last)
    gone = -count
    end = at + gone - 1
    if first >= at and last <= end:
        return None
    new_first = at if at <= first <= end else (first - gone if first > end else first)
    new_last = at - 1 if at <= last <= end else (last - gone if last > end else last)
    return new_first, new_last


def _shift_body(body: str, axis: str, at: int, count: int, limit: int) -> tuple[str, bool]:
    """A reference after a shift, ``$`` kept. Returns (new text, its area changed)."""
    cell = _CELL_RE.match(body)
    if cell:
        col_abs, col, row_abs, row = cell.groups()
        if axis == "row":
            moved = _shift_edge(int(row), at, count)
            if moved is None or moved > limit:
                return "#REF!", True
            return f"{col_abs}{col}{row_abs}{moved}", False
        moved = _shift_edge(column_index_from_string(col.upper()), at, count)
        if moved is None or moved > limit:
            return "#REF!", True
        return f"{col_abs}{get_column_letter(moved)}{row_abs}{row}", False
    cols = _COLS_RE.match(body)
    if cols:
        if axis == "row":
            return body, False
        a_abs, a, b_abs, b = cols.groups()
        first, last = column_index_from_string(a.upper()), column_index_from_string(b.upper())
        span = shift_span(first, last, at, count)
        if span is None or span[1] > limit:
            return "#REF!", True
        return (f"{a_abs}{get_column_letter(span[0])}:{b_abs}{get_column_letter(span[1])}",
                span[1] - span[0] != last - first)
    rows = _ROWS_RE.match(body)
    if rows:
        if axis == "col":
            return body, False
        a_abs, a, b_abs, b = rows.groups()
        first, last = int(a), int(b)
        span = shift_span(first, last, at, count)
        if span is None or span[1] > limit:
            return "#REF!", True
        return f"{a_abs}{span[0]}:{b_abs}{span[1]}", span[1] - span[0] != last - first
    if body.count(":") == 1:
        left, right = body.split(":")
        a, b = _CELL_RE.match(left), _CELL_RE.match(right)
        if a and b:
            ca_abs, ca, ra_abs, ra = a.groups()
            cb_abs, cb, rb_abs, rb = b.groups()
            if axis == "row":
                first, last = int(ra), int(rb)
                span = shift_span(first, last, at, count)
                if span is None or span[1] > limit:
                    return "#REF!", True
                return (f"{ca_abs}{ca}{ra_abs}{span[0]}:{cb_abs}{cb}{rb_abs}{span[1]}",
                        span[1] - span[0] != last - first)
            first, last = column_index_from_string(ca.upper()), column_index_from_string(cb.upper())
            span = shift_span(first, last, at, count)
            if span is None or span[1] > limit:
                return "#REF!", True
            return (f"{ca_abs}{get_column_letter(span[0])}{ra_abs}{ra}:"
                    f"{cb_abs}{get_column_letter(span[1])}{rb_abs}{rb}",
                    span[1] - span[0] != last - first)
    return body, False


def shift_formula(stored: str, own_sheet: str | None, target_sheet: str, axis: str,
                  at: int, count: int) -> tuple[str, bool]:
    """A formula after rows or columns moved on *target_sheet*, the way Excel rewrites it.

    Returns (new text, whether a block it reads changed size or broke). *axis* is
    ``row`` or ``col``; *count* is positive for an insert at *at*, negative for a
    delete starting at *at*.
    """
    try:
        items = _tokens(stored)
    except (TokenizerError, ValueError):
        return stored, False
    limit = MAX_ROW if axis == "row" else MAX_COL
    target = target_sheet.casefold()
    changed = False
    smaller = False
    for token in items:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            continue
        sheet, body, has_prefix = split_sheet_prefix(token.value)
        where = sheet if has_prefix else own_sheet
        if where is None or where.casefold() != target or "[" in body:
            continue
        new_body, area_changed = _shift_body(body, axis, at, count, limit)
        if new_body != body:
            prefix = token.value[: len(token.value) - len(body)] if has_prefix else ""
            token.value = prefix + new_body
            changed = True
        smaller = smaller or area_changed
    if not changed:
        return stored, smaller
    text = "".join(t.value for t in items)
    return (("=" + text) if stored.startswith("=") else text), smaller


def rename_sheet_in(stored: str, old: str, new: str) -> str:
    """A formula after sheet *old* was renamed to *new*."""
    if old.casefold() not in stored.casefold():
        return stored
    try:
        items = _tokens(stored)
    except (TokenizerError, ValueError):
        return stored
    changed = False
    for token in items:
        if token.type != Token.OPERAND or token.subtype != Token.RANGE:
            continue
        sheet, body, has_prefix = split_sheet_prefix(token.value)
        if has_prefix and sheet is not None and sheet.casefold() == old.casefold():
            token.value = f"{quote_sheet(new)}!{body}"
            changed = True
    if not changed:
        return stored
    text = "".join(t.value for t in items)
    return ("=" + text) if stored.startswith("=") else text


def mentions_sheet(stored: str, sheet_name: str) -> bool:
    """Does the formula read *sheet_name* by name?"""
    if sheet_name.casefold() not in stored.casefold():
        return False
    for operand in compile_formula(stored).operands:
        if isinstance(operand, (Box, NameOperand)) and operand.sheet is not None \
                and operand.sheet.casefold() == sheet_name.casefold():
            return True
    return False


# ── Evaluation ───────────────────────────────────────────────────────


def resolve_operand(op: Any, ctx: Ctx) -> Any:
    if isinstance(op, Box):
        sheet = ctx.sheet
        if op.sheet is not None:
            sheet = ctx.model.sheet_name(op.sheet)
            if sheet is None:
                return REF
        return Ref(sheet, op.r1, op.c1, op.r2, op.c2)
    if isinstance(op, XlError):
        return op
    if isinstance(op, NameOperand):
        text = ctx.model.defined_name(op.name, op.sheet or ctx.sheet)
        if text is None:
            return NAME
        compiled = compile_formula(text)
        return compiled.tree.ev(ctx)
    if isinstance(op, TableOperand):
        return _resolve_table(op, ctx)
    raise NotComputable(op.reason)


def _resolve_table(op: TableOperand, ctx: Ctx) -> Any:
    table = (ctx.model.table_named(op.table) if op.table
             else ctx.model.table_at(ctx.sheet, ctx.row, ctx.col))
    if table is None:
        return REF if not op.table else NAME
    c1, c2 = table.c1, table.c2
    if op.first_col is not None:
        columns = [c.casefold() for c in table.columns]
        try:
            first = columns.index(op.first_col.casefold())
            last = columns.index((op.last_col or op.first_col).casefold())
        except ValueError:
            return REF
        c1, c2 = table.c1 + first, table.c1 + last
    data_first = table.r1 + table.header_rows
    data_last = table.r2 - table.totals_rows
    special = {s.casefold() for s in op.special}
    if "#this row" in special:
        if not data_first <= ctx.row <= data_last:
            return VALUE
        r1 = r2 = ctx.row
    elif "#all" in special:
        r1, r2 = table.r1, table.r2
    elif "#headers" in special and "#data" in special:
        r1, r2 = table.r1, data_last
    elif "#data" in special and "#totals" in special:
        r1, r2 = data_first, table.r2
    elif "#headers" in special:
        if not table.header_rows:
            return REF
        r1, r2 = table.r1, table.r1
    elif "#totals" in special:
        if not table.totals_rows:
            return REF
        r1, r2 = table.r2, table.r2
    else:
        r1, r2 = data_first, data_last
    return Ref(table.sheet, r1, c1, r2, c2)


def evaluate(compiled: Compiled, model: Any, sheet: str, row: int, col: int,
             *, array: bool = False) -> Any:
    """The formula's value. Raises NotComputable or NeedCell. Array formulas
    return an :class:`Arr`; everything else one value."""
    ctx = Ctx(model, sheet, row, col, array)
    try:
        result = compiled.tree.ev(ctx)
    except _Fail as fail:
        return fail.error
    except ZeroDivisionError:
        return DIV0
    if array:
        if isinstance(result, Ref):
            return result.arr(ctx)
        return result
    return _final(result, ctx)


def _final(result: Any, ctx: Ctx) -> Any:
    if isinstance(result, Ref):
        if result.is_cell:
            value = result.single(ctx)
        else:
            value = _intersect(result, ctx)
        return 0.0 if value is None else value
    if isinstance(result, Arr):
        if result.width == 1 and result.height == 1:
            value = result.rows[0][0]
            return 0.0 if value is None else value
        # A plain formula that makes an array spills in new Excel; not stored here.
        raise NotComputable("the formula makes an array")
    return 0.0 if result is None else result


def _intersect(ref: Ref, ctx: Ctx) -> Any:
    """A range where one value is wanted: the cell in this row or column."""
    if ref.width == 1 and ref.r1 <= ctx.row <= ref.r2 and ctx.sheet == ref.sheet:
        return ctx.model.value(ref.sheet, ctx.row, ref.c1)
    if ref.height == 1 and ref.c1 <= ctx.col <= ref.c2 and ctx.sheet == ref.sheet:
        return ctx.model.value(ref.sheet, ref.r1, ctx.col)
    if ref.is_cell:
        return ref.single(ctx)
    return VALUE


def _spread(value: Any, ctx: Ctx) -> Any:
    """A Ref operand of an operator: one value, or an array in array formulas."""
    if isinstance(value, Ref):
        if value.is_cell:
            return value.single(ctx)
        return value.arr(ctx) if ctx.array else _intersect(value, ctx)
    return value


def _map1(value: Any, ctx: Ctx, fn: Callable[[Any], Any]) -> Any:
    value = _spread(value, ctx)
    if isinstance(value, Arr):
        return Arr([[fn(v) for v in row] for row in value.rows])
    return fn(value)


def _to_number(value: Any) -> float | XlError:
    if value is None:
        return 0.0
    if isinstance(value, bool):
        return 1.0 if value else 0.0
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, XlError):
        return value
    parsed = parse_number_text(str(value))
    return VALUE if parsed is None else parsed


_NUMBER_TEXT = re.compile(r"^\s*([-+]?)\s*([$€£¥]?)\s*(\d{1,3}(?:,\d{3})+|\d*)(\.\d*)?(?:[eE]([-+]?\d+))?\s*(%?)\s*$")


def parse_number_text(text: str) -> float | None:
    """A text Excel would read as a number: ``12``, ``1,200.50``, ``12%``, ``$5``, an ISO date."""
    match = _NUMBER_TEXT.match(text)
    if match and (match.group(3) or (match.group(4) and len(match.group(4)) > 1)):
        sign, _, whole, frac, exp, pct = match.groups()
        try:
            number = float((whole or "0").replace(",", "") + (frac or "") + (f"e{exp}" if exp else ""))
        except ValueError:
            return None
        if pct:
            number /= 100
        return -number if sign == "-" else number
    iso = re.fullmatch(r"\s*(\d{4})-(\d{1,2})-(\d{1,2})\s*", text)
    if iso:
        try:
            return float(date_to_serial(dt.date(int(iso.group(1)), int(iso.group(2)), int(iso.group(3)))))
        except ValueError:
            return None
    return None


def _unary(op: str, value: Any) -> Any:
    number = _to_number(value)
    if isinstance(number, XlError):
        return number
    return -number if op == "-" else number


def _percent(value: Any) -> Any:
    number = _to_number(value)
    return number if isinstance(number, XlError) else number / 100


def _binary(op: str, a: Any, b: Any) -> Any:
    if isinstance(a, XlError):
        return a
    if isinstance(b, XlError):
        return b
    if op == "&":
        return to_text(a) + to_text(b)
    if op in ("=", "<>", "<", "<=", ">", ">="):
        order = compare(a, b)
        return {"=": order == 0, "<>": order != 0, "<": order < 0, "<=": order <= 0,
                ">": order > 0, ">=": order >= 0}[op]
    x, y = _to_number(a), _to_number(b)
    if isinstance(x, XlError):
        return x
    if isinstance(y, XlError):
        return y
    try:
        if op == "+":
            return _snap(x + y, x, y)
        if op == "-":
            return _snap(x - y, x, y)
        if op == "*":
            return x * y
        if op == "/":
            return DIV0 if y == 0 else x / y
        if x == 0 and y == 0:
            return NUM
        if x == 0 and y < 0:
            return DIV0
        result = x ** y
        return NUM if isinstance(result, complex) or math.isinf(result) else result
    except OverflowError:
        return NUM


def _snap(result: float, x: float, y: float) -> float:
    """Excel zeroes what's left of a subtraction when it's only binary noise."""
    if result != 0 and abs(result) < 2e-15 * max(abs(x), abs(y)):
        return 0.0
    return result


def _rank(value: Any) -> int:
    if isinstance(value, bool):
        return 2
    if isinstance(value, str):
        return 1
    return 0


def compare(a: Any, b: Any) -> int:
    """Excel's ordering: numbers < text < logical; a blank takes the other's type."""
    if a is None and b is None:
        return 0
    if a is None:
        a = "" if isinstance(b, str) else (False if isinstance(b, bool) else 0.0)
    if b is None:
        b = "" if isinstance(a, str) else (False if isinstance(a, bool) else 0.0)
    ra, rb = _rank(a), _rank(b)
    if ra != rb:
        return -1 if ra < rb else 1
    if ra == 1:
        x, y = a.casefold(), b.casefold()
        return 0 if x == y else (-1 if x < y else 1)
    if ra == 2:
        return 0 if a == b else (-1 if not a else 1)
    if a == b or abs(a - b) <= 1e-15 * max(abs(a), abs(b)):
        return 0
    return -1 if a < b else 1


def number_to_text(value: float) -> str:
    """A number in text, as ``&`` and General show it (15 significant digits)."""
    if value == 0:
        return "0"
    if float(value).is_integer() and abs(value) < 1e15:
        return str(int(value))
    text = f"{value:.15g}"
    if "e" in text:
        mantissa, exponent = text.split("e")
        power = int(exponent)
        if -10 < power < 0:
            return f"{value:.{min(15 - power - 1, 30)}f}".rstrip("0")
        return f"{mantissa}E{'+' if power >= 0 else '-'}{abs(power):02d}"
    return text


def to_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        return number_to_text(float(value))
    if isinstance(value, XlError):
        raise _Fail(value)
    return str(value)


# ── Dates ────────────────────────────────────────────────────────────

_EPOCH = dt.date(1899, 12, 30)
_EPOCH_1904 = dt.date(1904, 1, 1)


def date_to_serial(day: dt.date, date1904: bool = False) -> int:
    if date1904:
        return (day - _EPOCH_1904).days
    serial = (day - _EPOCH).days
    # Excel counts a 29 Feb 1900 that never existed.
    return serial - 1 if serial < 61 else serial


def serial_to_date(serial: float, date1904: bool = False) -> dt.date:
    whole = math.floor(serial)
    if date1904:
        return _EPOCH_1904 + dt.timedelta(days=whole)
    if whole < 0:
        raise ValueError("a date before 1900")
    if whole == 0:
        return dt.date(1899, 12, 31)
    return _EPOCH + dt.timedelta(days=whole + (1 if whole < 61 else 0))


def _serial(ctx: Ctx, value: Any) -> dt.date:
    number = _n(value)
    if number < 0 or number > 2_958_465:
        raise _Fail(NUM)
    return serial_to_date(number, ctx.model.date1904)


def _make_serial(ctx: Ctx, day: dt.date) -> float:
    return float(date_to_serial(day, ctx.model.date1904))


def _add_months(day: dt.date, months: int) -> dt.date:
    index = day.year * 12 + day.month - 1 + months
    year, month = divmod(index, 12)
    month += 1
    if year < 1900 or year > 9999:
        raise _Fail(NUM)
    last = _month_end(year, month)
    return dt.date(year, month, min(day.day, last))


def _month_end(year: int, month: int) -> int:
    following = dt.date(year + (month == 12), (month % 12) + 1, 1)
    return (following - dt.timedelta(days=1)).day


# ── Number formats (TEXT, and how cells read back) ───────────────────

_MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
           "September", "October", "November", "December"]
_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]


def _split_sections(fmt: str) -> list[str]:
    sections: list[str] = []
    current: list[str] = []
    i = 0
    while i < len(fmt):
        ch = fmt[i]
        if ch == '"':
            end = fmt.find('"', i + 1)
            end = len(fmt) - 1 if end < 0 else end
            current.append(fmt[i:end + 1])
            i = end + 1
            continue
        if ch == "\\" and i + 1 < len(fmt):
            current.append(fmt[i:i + 2])
            i += 2
            continue
        if ch == "[":
            end = fmt.find("]", i)
            end = len(fmt) - 1 if end < 0 else end
            current.append(fmt[i:end + 1])
            i = end + 1
            continue
        if ch == ";":
            sections.append("".join(current))
            current = []
        else:
            current.append(ch)
        i += 1
    sections.append("".join(current))
    return sections


def _strip_literals(section: str) -> str:
    return re.sub(r'"[^"]*"|\\.|_.|\*.|\[[^\]]*\]', "", section)


def is_date_format(fmt: str) -> bool:
    if not fmt or fmt.lower() == "general":
        return False
    first = _strip_literals(_split_sections(fmt)[0]).lower().replace("am/pm", "").replace("a/p", "")
    return bool(re.search(r"[ymdhs]", first))


def _literal_parts(section: str) -> Iterator[tuple[str, str]]:
    """Split a section into ("lit", text) and ("code", chars)."""
    i = 0
    code: list[str] = []

    def flush() -> Iterator[tuple[str, str]]:
        if code:
            yield "code", "".join(code)
            code.clear()

    while i < len(section):
        ch = section[i]
        if ch == '"':
            yield from flush()
            end = section.find('"', i + 1)
            end = len(section) if end < 0 else end
            yield "lit", section[i + 1:end]
            i = end + 1
        elif ch == "\\" and i + 1 < len(section):
            yield from flush()
            yield "lit", section[i + 1]
            i += 2
        elif ch == "_" and i + 1 < len(section):
            yield from flush()
            yield "lit", " "
            i += 2
        elif ch == "*" and i + 1 < len(section):
            i += 2
        elif ch == "[":
            yield from flush()
            end = section.find("]", i)
            end = len(section) - 1 if end < 0 else end
            inner = section[i + 1:end]
            if inner.startswith("$"):
                yield "lit", inner[1:].split("-")[0]
            elif re.fullmatch(r"h+|m+|s+", inner, re.I):
                yield "elapsed", inner.lower()
            elif re.match(r"[<>=]", inner):
                yield "condition", inner
            i = end + 1
        else:
            code.append(ch)
            i += 1
    yield from flush()


def _format_date(section: str, serial: float, date1904: bool) -> str | None:
    try:
        day = serial_to_date(serial, date1904)
    except (ValueError, OverflowError):
        return None
    seconds = round((serial - math.floor(serial)) * 86400)
    hour, minute, second = seconds // 3600 % 24, seconds // 60 % 60, seconds % 60
    parts = list(_literal_parts(section))
    if any(kind == "elapsed" or kind == "condition" for kind, _ in parts):
        return None
    text = "".join(v for k, v in parts if k == "code").lower()
    twelve = "am/pm" in text or "a/p" in text
    out: list[str] = []
    # A month is a minute when it follows an hour or comes before seconds.
    tokens: list[tuple[str, str]] = []
    for kind, value in parts:
        if kind != "code":
            tokens.append(("lit", value))
            continue
        for tok in re.findall(r"am/pm|a/p|y+|m+|d+|h+|s+|\.0+|.", value, re.I):
            tokens.append(("tok", tok))
    for index, (kind, value) in enumerate(tokens):
        if kind == "lit":
            out.append(value)
            continue
        low = value.lower()
        if low.startswith("y"):
            out.append(str(day.year) if len(low) > 2 else f"{day.year % 100:02d}")
        elif low.startswith("m") and low != "am/pm":
            before = [t.lower() for k, t in tokens[:index] if k == "tok" and re.match(r"[a-z]", t, re.I)]
            after = [t.lower() for k, t in tokens[index + 1:] if k == "tok" and re.match(r"[a-z]", t, re.I)]
            minutes = (before and before[-1].startswith("h")) or (after and after[0].startswith("s"))
            if minutes and len(low) <= 2:
                out.append(f"{minute:02d}" if len(low) == 2 else str(minute))
            elif len(low) == 1:
                out.append(str(day.month))
            elif len(low) == 2:
                out.append(f"{day.month:02d}")
            elif len(low) == 3:
                out.append(_MONTHS[day.month - 1][:3])
            elif len(low) == 5:
                out.append(_MONTHS[day.month - 1][0])
            else:
                out.append(_MONTHS[day.month - 1])
        elif low.startswith("d"):
            if len(low) == 1:
                out.append(str(day.day))
            elif len(low) == 2:
                out.append(f"{day.day:02d}")
            elif len(low) == 3:
                out.append(_DAYS[day.weekday()][:3])
            else:
                out.append(_DAYS[day.weekday()])
        elif low.startswith("h"):
            shown = hour % 12 or 12 if twelve else hour
            out.append(f"{shown:02d}" if len(low) == 2 else str(shown))
        elif low.startswith("s"):
            out.append(f"{second:02d}" if len(low) == 2 else str(second))
        elif low == "am/pm":
            am = hour < 12
            out.append(("AM" if am else "PM") if value[0].isupper() else ("am" if am else "pm"))
        elif low == "a/p":
            out.append("a" if hour < 12 else "p")
        else:
            out.append(value)
    return "".join(out)


def _round_decimal(value: float, places: int) -> Decimal:
    quantum = Decimal(1).scaleb(-places)
    return Decimal(repr(abs(value))).quantize(quantum, rounding=ROUND_HALF_UP)


def _format_number(section: str, value: float) -> str | None:
    parts = list(_literal_parts(section))
    if any(kind in ("elapsed", "condition") for kind, _ in parts):
        return None
    code = "".join(v for k, v in parts if k == "code")
    if "/" in code:
        return None  # fractions
    if re.search(r"[eE][+-]", code):
        match = re.search(r"([0#?.,]*)[eE]([+-])(0+)", code)
        if match is None:
            return None
        mantissa, sign, digits = match.groups()
        places = len(mantissa.split(".")[1]) if "." in mantissa else 0
        text = f"{abs(value):.{places}e}"
        head, exponent = text.split("e")
        power = int(exponent)
        marker = "+" if power >= 0 else "-"
        if sign == "-" and power >= 0:
            marker = ""
        return head + "E" + marker + str(abs(power)).zfill(len(digits))
    percents = code.count("%")
    scaled = value * (100 ** percents)
    placeholders = re.findall(r"[0#?.,]+", code)
    if not placeholders:
        return "".join(v for _, v in parts)
    number_code = max(placeholders, key=len)
    trailing = re.search(r",+$", number_code)
    if trailing:
        scaled /= 1000 ** len(trailing.group())
    body = number_code.rstrip(",")
    whole_code, _, frac_code = body.partition(".")
    thousands = "," in whole_code
    whole_code = whole_code.replace(",", "")
    rounded = _round_decimal(scaled, len(frac_code))
    text = f"{rounded:f}"
    whole, _, frac = text.partition(".")
    if whole == "0" and "0" not in whole_code:
        whole = ""
    whole = whole.rjust(whole_code.count("0"), "0") if whole_code else whole
    if thousands and whole:
        whole = f"{int(whole):,}".rjust(len(whole), "0") if whole.lstrip("0") else whole
    if frac_code:
        # A trailing # shows nothing when the digit is zero.
        keep = len(frac_code)
        while keep > 0 and frac_code[keep - 1] != "0" and frac[keep - 1] == "0":
            keep -= 1
        frac = frac[:keep]
    number = whole + ("." + frac if frac or (frac_code and "." in body) else "")
    out: list[str] = []
    placed = False
    for kind, chunk in parts:
        if kind != "code":
            out.append(chunk)
            continue
        if not placed and number_code in chunk:
            before, after = chunk.split(number_code, 1)
            out.append(before + number + after)
            placed = True
        else:
            out.append(chunk)
    return "".join(out)


def format_value(value: Any, fmt: str, date1904: bool = False) -> str | None:
    """A value as the number format *fmt* shows it; None when the format is beyond us."""
    if value is None:
        return ""
    if isinstance(value, XlError):
        return value.code
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if not fmt or fmt.lower() == "general":
        return number_to_text(float(value)) if isinstance(value, (int, float)) else str(value)
    sections = _split_sections(fmt)
    if isinstance(value, str):
        if len(sections) >= 4:
            section = sections[3]
        elif "@" in fmt:
            section = next((s for s in sections if "@" in s), "@")
        else:
            return value
        return "".join(v for _, v in _literal_parts(section)).replace("@", value)
    number = float(value)
    if len(sections) == 1:
        section, negative_sign = sections[0], number < 0
    elif number > 0 or len(sections) == 2 and number == 0:
        section, negative_sign = sections[0], False
    elif number < 0:
        section, negative_sign = sections[1], False
    else:
        section, negative_sign = sections[2], False
    if section.lower() == "general":
        text = number_to_text(abs(number))
        return ("-" + text) if negative_sign else text
    if is_date_format(section):
        if number < 0:
            return None
        return _format_date(section, number, date1904)
    text = _format_number(section, abs(number))
    if text is None:
        return None
    is_zero = all(ch in "0.,%" or not ch.isdigit() for ch in text) and not re.search(r"[1-9]", text)
    return ("-" + text) if negative_sign and not is_zero else text


# ── Argument helpers ─────────────────────────────────────────────────


def _one(value: Any, ctx: Ctx | None = None) -> Any:
    if isinstance(value, Ref):
        if value.is_cell and ctx is not None:
            return value.single(ctx)
        raise _Fail(VALUE)
    if isinstance(value, Arr):
        return value.rows[0][0] if value.rows and value.rows[0] else None
    return value


def _n(value: Any) -> float:
    number = _to_number(value)
    if isinstance(number, XlError):
        raise _Fail(number)
    return number


def _s(value: Any) -> str:
    return to_text(value)


def _b(value: Any) -> bool:
    if isinstance(value, XlError):
        raise _Fail(value)
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    if isinstance(value, (int, float)):
        return value != 0
    text = str(value).strip().upper()
    if text == "TRUE":
        return True
    if text == "FALSE":
        return False
    raise _Fail(VALUE)


def _i(value: Any) -> int:
    return math.trunc(_n(value))


def _opt(args: tuple[Any, ...], index: int, default: Any) -> Any:
    if index < len(args) and args[index] is not None:
        return args[index]
    return default


def _matrix(value: Any, ctx: Ctx) -> list[list[Any]]:
    if isinstance(value, Ref):
        return value.matrix(ctx)
    if isinstance(value, Arr):
        return value.rows
    return [[value]]


def _flat(value: Any, ctx: Ctx) -> Iterator[Any]:
    for row in _matrix(value, ctx):
        yield from row


def _numbers(args: Iterable[Any], ctx: Ctx, *, count_logical_refs: bool = False) -> Iterator[float]:
    """Numbers an aggregate sees: in ranges only real numbers; typed-in arguments
    are converted, and a bad one is an error."""
    for arg in args:
        if isinstance(arg, Ref):
            source: Iterable[Any] = arg.present(ctx)
        elif isinstance(arg, Arr):
            source = arg.flat()
        else:
            if arg is None:
                yield 0.0
                continue
            yield _n(arg)
            continue
        for value in source:
            if isinstance(value, XlError):
                raise _Fail(value)
            if isinstance(value, bool):
                if count_logical_refs:
                    yield 1.0 if value else 0.0
                continue
            if isinstance(value, (int, float)):
                yield float(value)
            elif count_logical_refs and isinstance(value, str):
                yield 0.0


FUNCTIONS: dict[str, "_Spec"] = {}


@dataclass(slots=True)
class _Spec:
    fn: Callable[..., Any]
    lazy: bool = False
    array_args: bool = False
    scalar: bool = False


def _fn(*names: str, lazy: bool = False, array_args: bool = False) -> Callable[[Callable[..., Any]], Callable[..., Any]]:
    def register(fn: Callable[..., Any]) -> Callable[..., Any]:
        for name in names:
            FUNCTIONS[name] = _Spec(fn, lazy, array_args)
        return fn
    return register


def _scalar(*names: str) -> Callable[[Callable[..., Any]], Callable[..., Any]]:
    """A function of plain values: over a range it runs per cell in array
    formulas, and takes this row's cell otherwise."""
    def register(fn: Callable[..., Any]) -> Callable[..., Any]:
        def wrapped(ctx: Ctx, *args: Any) -> Any:
            values = [_spread(a, ctx) for a in args]
            arrays = [v for v in values if isinstance(v, Arr)]
            if not arrays:
                return fn(ctx, *values)
            height = max(a.height for a in arrays)
            width = max(a.width for a in arrays)

            def one(r: int, c: int) -> Any:
                try:
                    return fn(ctx, *[v.at(r, c) if isinstance(v, Arr) else v for v in values])
                except _Fail as fail:
                    return fail.error
                except ZeroDivisionError:
                    return DIV0
                except (OverflowError, ValueError):
                    return NUM

            return Arr([[one(r, c) for c in range(width)] for r in range(height)])

        for name in names:
            FUNCTIONS[name] = _Spec(wrapped, scalar=True)
        return fn
    return register


# ── Logic ────────────────────────────────────────────────────────────


@_fn("IF", lazy=True)
def _if(ctx: Ctx, test: Node, yes: Node | None = None, no: Node | None = None) -> Any:
    value = _spread(test.ev(ctx), ctx)
    if isinstance(value, Arr):
        then_v = _spread(yes.ev(ctx), ctx) if yes is not None else True
        else_v = _spread(no.ev(ctx), ctx) if no is not None else False

        def pick(r: int, c: int, item: Any) -> Any:
            if isinstance(item, XlError):
                return item
            chosen = then_v if _b(item) else else_v
            return chosen.at(r, c) if isinstance(chosen, Arr) else chosen

        return Arr([[pick(r, c, item) for c, item in enumerate(row)] for r, row in enumerate(value.rows)])
    if isinstance(value, XlError):
        return value
    if _b(value):
        return yes.ev(ctx) if yes is not None else True
    if no is None:
        return False
    return no.ev(ctx)


@_fn("IFERROR", lazy=True)
def _iferror(ctx: Ctx, value: Node, fallback: Node) -> Any:
    result = _spread(value.ev(ctx), ctx)
    if isinstance(result, Arr):
        other = _spread(fallback.ev(ctx), ctx)
        return Arr([[other if isinstance(v, XlError) else v for v in row] for row in result.rows])
    return fallback.ev(ctx) if isinstance(result, XlError) else result


@_fn("IFNA", lazy=True)
def _ifna(ctx: Ctx, value: Node, fallback: Node) -> Any:
    result = _spread(value.ev(ctx), ctx)
    return fallback.ev(ctx) if result == NA else result


@_fn("IFS", lazy=True)
def _ifs(ctx: Ctx, *pairs: Node) -> Any:
    if len(pairs) % 2:
        raise NotComputable("IFS: odd number of arguments")
    for test, value in zip(pairs[::2], pairs[1::2]):
        checked = _spread(test.ev(ctx), ctx)
        if isinstance(checked, XlError):
            return checked
        if _b(checked):
            return value.ev(ctx)
    return NA


@_fn("SWITCH", lazy=True)
def _switch(ctx: Ctx, expression: Node, *rest: Node) -> Any:
    value = _spread(expression.ev(ctx), ctx)
    if isinstance(value, XlError):
        return value
    pairs = len(rest) // 2
    for i in range(pairs):
        candidate = _spread(rest[i * 2].ev(ctx), ctx)
        if not isinstance(candidate, XlError) and compare(value, candidate) == 0:
            return rest[i * 2 + 1].ev(ctx)
    return rest[-1].ev(ctx) if len(rest) % 2 else NA


@_fn("CHOOSE", lazy=True)
def _choose(ctx: Ctx, index: Node, *options: Node) -> Any:
    chosen = _spread(index.ev(ctx), ctx)
    if isinstance(chosen, XlError):
        return chosen
    i = _i(chosen)
    if not 1 <= i <= len(options):
        return VALUE
    return options[i - 1].ev(ctx)


def _logicals(args: Iterable[Any], ctx: Ctx) -> list[bool]:
    out: list[bool] = []
    for arg in args:
        if isinstance(arg, (Ref, Arr)):
            for value in (arg.present(ctx) if isinstance(arg, Ref) else arg.flat()):
                if isinstance(value, XlError):
                    raise _Fail(value)
                if isinstance(value, bool):
                    out.append(value)
                elif isinstance(value, (int, float)):
                    out.append(value != 0)
        else:
            out.append(_b(arg))
    return out


@_fn("AND")
def _and(ctx: Ctx, *args: Any) -> Any:
    values = _logicals(args, ctx)
    return all(values) if values else VALUE


@_fn("OR")
def _or(ctx: Ctx, *args: Any) -> Any:
    values = _logicals(args, ctx)
    return any(values) if values else VALUE


@_fn("XOR")
def _xor(ctx: Ctx, *args: Any) -> Any:
    values = _logicals(args, ctx)
    return (sum(values) % 2 == 1) if values else VALUE


@_scalar("NOT")
def _not(ctx: Ctx, value: Any) -> Any:
    return not _b(value)


@_fn("TRUE")
def _true(ctx: Ctx) -> bool:
    return True


@_fn("FALSE")
def _false(ctx: Ctx) -> bool:
    return False


@_fn("NA")
def _na(ctx: Ctx) -> XlError:
    return NA


@_fn("ISBLANK")
def _isblank(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: v is None)


@_fn("ISNUMBER")
def _isnumber(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: isinstance(v, (int, float)) and not isinstance(v, bool))


@_fn("ISTEXT")
def _istext(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: isinstance(v, str))


@_fn("ISNONTEXT")
def _isnontext(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: not isinstance(v, str))


@_fn("ISLOGICAL")
def _islogical(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: isinstance(v, bool))


@_fn("ISERROR")
def _iserror(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: isinstance(v, XlError))


@_fn("ISERR")
def _iserr(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: isinstance(v, XlError) and v != NA)


@_fn("ISNA")
def _isna(ctx: Ctx, value: Any) -> Any:
    return _map1(value, ctx, lambda v: v == NA)


@_scalar("ISEVEN")
def _iseven(ctx: Ctx, value: Any) -> bool:
    return _i(value) % 2 == 0


@_scalar("ISODD")
def _isodd(ctx: Ctx, value: Any) -> bool:
    return _i(value) % 2 == 1


@_fn("ISFORMULA")
def _isformula(ctx: Ctx, ref: Any) -> Any:
    if not isinstance(ref, Ref):
        return VALUE
    return ctx.model.formula_at(ref.sheet, ref.r1, ref.c1) is not None


@_fn("ERROR.TYPE")
def _error_type(ctx: Ctx, value: Any) -> Any:
    item = _one(value, ctx)
    codes = {NULL: 1, DIV0: 2, VALUE: 3, REF: 4, NAME: 5, NUM: 6, NA: 7}
    return float(codes[item]) if isinstance(item, XlError) and item in codes else NA


@_fn("TYPE")
def _type(ctx: Ctx, value: Any) -> float:
    if isinstance(value, (Arr,)) or (isinstance(value, Ref) and not value.is_cell):
        return 64.0
    item = _one(value, ctx)
    if isinstance(item, XlError):
        return 16.0
    if isinstance(item, bool):
        return 4.0
    if isinstance(item, str):
        return 2.0
    return 1.0


@_fn("N")
def _ntype(ctx: Ctx, value: Any) -> Any:
    item = _one(value, ctx)
    if isinstance(item, XlError):
        return item
    if isinstance(item, bool):
        return 1.0 if item else 0.0
    return float(item) if isinstance(item, (int, float)) else 0.0


@_fn("T")
def _ttype(ctx: Ctx, value: Any) -> Any:
    item = _one(value, ctx)
    if isinstance(item, XlError):
        return item
    return item if isinstance(item, str) else ""


# ── Math ─────────────────────────────────────────────────────────────


@_fn("SUM")
def _sum(ctx: Ctx, *args: Any) -> float:
    return math.fsum(_numbers(args, ctx))


@_fn("PRODUCT")
def _product(ctx: Ctx, *args: Any) -> float:
    numbers = list(_numbers(args, ctx))
    return math.prod(numbers) if numbers else 0.0


@_fn("SUMPRODUCT", array_args=True)
def _sumproduct(ctx: Ctx, *args: Any) -> Any:
    blocks = [_matrix(a, ctx.with_array()) for a in args]
    if not blocks:
        return VALUE
    height, width = len(blocks[0]), len(blocks[0][0]) if blocks[0] else 0
    if any(len(b) != height or (b and len(b[0]) != width) for b in blocks):
        return VALUE
    total = 0.0
    for r in range(height):
        for c in range(width):
            product = 1.0
            for block in blocks:
                value = block[r][c]
                if isinstance(value, XlError):
                    return value
                product *= float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else 0.0
            total += product
    return total


@_scalar("ABS")
def _abs(ctx: Ctx, value: Any) -> float:
    return abs(_n(value))


@_scalar("SIGN")
def _sign(ctx: Ctx, value: Any) -> float:
    number = _n(value)
    return float((number > 0) - (number < 0))


def _round_with(value: float, digits: int, rounding: str) -> float:
    if value == 0 or math.isinf(value):
        return value
    sign = -1.0 if value < 0 else 1.0
    try:
        amount = Decimal(repr(abs(value))).quantize(Decimal(1).scaleb(-digits), rounding=rounding)
    except InvalidOperation:
        return value
    return sign * float(amount)


@_scalar("ROUND")
def _round(ctx: Ctx, value: Any, digits: Any = 0) -> float:
    return _round_with(_n(value), _i(digits), ROUND_HALF_UP)


@_scalar("ROUNDUP")
def _roundup(ctx: Ctx, value: Any, digits: Any = 0) -> float:
    return _round_with(_n(value), _i(digits), ROUND_UP)


@_scalar("ROUNDDOWN")
def _rounddown(ctx: Ctx, value: Any, digits: Any = 0) -> float:
    return _round_with(_n(value), _i(digits), ROUND_DOWN)


@_scalar("TRUNC")
def _trunc(ctx: Ctx, value: Any, digits: Any = 0) -> float:
    return _round_with(_n(value), _i(digits), ROUND_DOWN)


@_scalar("INT")
def _int(ctx: Ctx, value: Any) -> float:
    return float(math.floor(_n(value)))


@_scalar("MOD")
def _mod(ctx: Ctx, value: Any, divisor: Any) -> Any:
    n, d = _n(value), _n(divisor)
    if d == 0:
        return DIV0
    return n - d * math.floor(n / d)


@_scalar("POWER")
def _power(ctx: Ctx, base: Any, exponent: Any) -> Any:
    return _binary("^", base, exponent)


@_scalar("SQRT")
def _sqrt(ctx: Ctx, value: Any) -> Any:
    number = _n(value)
    return NUM if number < 0 else math.sqrt(number)


@_scalar("EXP")
def _exp(ctx: Ctx, value: Any) -> float:
    return math.exp(_n(value))


@_scalar("LN")
def _ln(ctx: Ctx, value: Any) -> Any:
    number = _n(value)
    return NUM if number <= 0 else math.log(number)


@_scalar("LOG10")
def _log10(ctx: Ctx, value: Any) -> Any:
    number = _n(value)
    return NUM if number <= 0 else math.log10(number)


@_scalar("LOG")
def _log(ctx: Ctx, value: Any, base: Any = 10) -> Any:
    number, b = _n(value), _n(base)
    if number <= 0 or b <= 0:
        return NUM
    if b == 1:
        return DIV0
    return math.log10(number) if b == 10 else math.log(number) / math.log(b)


@_fn("PI")
def _pi(ctx: Ctx) -> float:
    return math.pi


for _name, _math in (("SIN", math.sin), ("COS", math.cos), ("TAN", math.tan), ("ASIN", math.asin),
                     ("ACOS", math.acos), ("ATAN", math.atan), ("SINH", math.sinh), ("COSH", math.cosh),
                     ("TANH", math.tanh), ("DEGREES", math.degrees), ("RADIANS", math.radians)):
    _scalar(_name)((lambda f: lambda ctx, value: f(_n(value)))(_math))


@_scalar("ATAN2")
def _atan2(ctx: Ctx, x: Any, y: Any) -> Any:
    a, b = _n(x), _n(y)
    return DIV0 if a == 0 and b == 0 else math.atan2(b, a)


@_scalar("FACT")
def _fact(ctx: Ctx, value: Any) -> Any:
    number = _i(value)
    return NUM if number < 0 or number > 170 else float(math.factorial(number))


@_scalar("COMBIN")
def _combin(ctx: Ctx, n: Any, k: Any) -> Any:
    a, b = _i(n), _i(k)
    return NUM if a < 0 or b < 0 or b > a else float(math.comb(a, b))


@_scalar("PERMUT")
def _permut(ctx: Ctx, n: Any, k: Any) -> Any:
    a, b = _i(n), _i(k)
    return NUM if a < 0 or b < 0 or b > a else float(math.perm(a, b))


@_fn("GCD")
def _gcd(ctx: Ctx, *args: Any) -> Any:
    numbers = [math.trunc(n) for n in _numbers(args, ctx)]
    return NUM if any(n < 0 for n in numbers) else float(math.gcd(*numbers))


@_fn("LCM")
def _lcm(ctx: Ctx, *args: Any) -> Any:
    numbers = [math.trunc(n) for n in _numbers(args, ctx)]
    return NUM if any(n < 0 for n in numbers) else float(math.lcm(*numbers))


def _to_multiple(number: float, significance: float, rounder: Callable[[float], float]) -> float:
    if significance == 0:
        return 0.0
    # Divide as decimals: 4.42/0.05 is 88.4 in binary noise otherwise.
    return float(Decimal(repr(significance)) * Decimal(rounder(float(Decimal(repr(number)) / Decimal(repr(significance))))))


@_scalar("MROUND")
def _mround(ctx: Ctx, value: Any, multiple: Any) -> Any:
    n, m = _n(value), _n(multiple)
    if m == 0:
        return 0.0
    if n * m < 0:
        return NUM
    quotient = Decimal(repr(abs(n))) / Decimal(repr(abs(m)))
    count = quotient.quantize(Decimal(1), rounding=ROUND_HALF_UP)
    result = float(count * Decimal(repr(abs(m))))
    return -result if n < 0 else result


@_scalar("CEILING")
def _ceiling(ctx: Ctx, value: Any, significance: Any = 1) -> Any:
    n, s = _n(value), _n(significance)
    if n > 0 and s < 0:
        return NUM
    return _to_multiple(n, s, math.ceil)


@_scalar("FLOOR")
def _floor(ctx: Ctx, value: Any, significance: Any = 1) -> Any:
    n, s = _n(value), _n(significance)
    if s == 0:
        return DIV0 if n else 0.0
    if n > 0 and s < 0:
        return NUM
    return _to_multiple(n, s, math.floor)


@_scalar("CEILING.MATH")
def _ceiling_math(ctx: Ctx, value: Any, significance: Any = 1, mode: Any = 0) -> float:
    n, s = _n(value), abs(_n(significance if significance is not None else 1))
    if n < 0 and _n(mode) != 0:
        return _to_multiple(n, s, math.floor)
    return _to_multiple(n, s, math.ceil)


@_scalar("FLOOR.MATH")
def _floor_math(ctx: Ctx, value: Any, significance: Any = 1, mode: Any = 0) -> float:
    n, s = _n(value), abs(_n(significance if significance is not None else 1))
    if n < 0 and _n(mode) != 0:
        return _to_multiple(n, s, math.ceil)
    return _to_multiple(n, s, math.floor)


@_scalar("EVEN")
def _even(ctx: Ctx, value: Any) -> float:
    n = _n(value)
    result = math.ceil(abs(n) / 2) * 2
    return float(-result if n < 0 else result)


@_scalar("ODD")
def _odd(ctx: Ctx, value: Any) -> float:
    n = _n(value)
    result = math.ceil((abs(n) + 1) / 2) * 2 - 1
    return float(-result if n < 0 else result)


_SUBTOTALS = {1: "AVERAGE", 2: "COUNT", 3: "COUNTA", 4: "MAX", 5: "MIN", 6: "PRODUCT",
              7: "STDEV", 8: "STDEVP", 9: "SUM", 10: "VAR", 11: "VARP"}


@_fn("SUBTOTAL")
def _subtotal(ctx: Ctx, code: Any, *refs: Any) -> Any:
    number = _i(_one(code, ctx))
    base = number - 100 if number > 100 else number
    name = _SUBTOTALS.get(base)
    if name is None:
        return VALUE
    kept: list[Any] = []
    for ref in refs:
        if not isinstance(ref, Ref):
            return VALUE
        rows: list[list[Any]] = []
        last_row, last_col = ref._clip(ctx)
        for r in range(ref.r1, last_row + 1):
            if ctx.model.row_hidden(ref.sheet, r):
                # Filtered rows are always left out, hidden ones only by 101-111;
                # the file doesn't tell them apart reliably.
                raise NotComputable("SUBTOTAL over hidden rows")
            row: list[Any] = []
            for c in range(ref.c1, last_col + 1):
                formula = ctx.model.formula_at(ref.sheet, r, c)
                # A subtotal skips other subtotals in its range.
                skip = formula is not None and "SUBTOTAL(" in formula.upper()
                row.append(None if skip else ctx.model.value(ref.sheet, r, c))
            rows.append(row)
        kept.append(Arr(rows))
    return FUNCTIONS[name].fn(ctx, *kept)


# ── Counting and statistics ──────────────────────────────────────────


@_fn("AVERAGE")
def _average(ctx: Ctx, *args: Any) -> Any:
    numbers = list(_numbers(args, ctx))
    return math.fsum(numbers) / len(numbers) if numbers else DIV0


@_fn("AVERAGEA")
def _averagea(ctx: Ctx, *args: Any) -> Any:
    numbers = list(_numbers(args, ctx, count_logical_refs=True))
    return math.fsum(numbers) / len(numbers) if numbers else DIV0


@_fn("COUNT")
def _count(ctx: Ctx, *args: Any) -> float:
    total = 0
    for arg in args:
        if isinstance(arg, (Ref, Arr)):
            for value in (arg.present(ctx) if isinstance(arg, Ref) else arg.flat()):
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    total += 1
        elif isinstance(arg, (int, float)) or arg is None:
            total += 1
        elif isinstance(arg, str) and parse_number_text(arg) is not None:
            total += 1
    return float(total)


@_fn("COUNTA")
def _counta(ctx: Ctx, *args: Any) -> float:
    total = 0
    for arg in args:
        if isinstance(arg, Ref):
            total += sum(1 for _ in arg.present(ctx))
        elif isinstance(arg, Arr):
            total += sum(1 for v in arg.flat() if v is not None)
        else:
            total += 1
    return float(total)


@_fn("COUNTBLANK")
def _countblank(ctx: Ctx, ref: Any) -> Any:
    if not isinstance(ref, Ref):
        return VALUE
    filled = sum(1 for v in ref.present(ctx) if v != "")
    return float(ref.height * ref.width - filled)


def _extreme(args: Iterable[Any], ctx: Ctx, pick: Callable[..., float], logical: bool = False) -> float:
    numbers = list(_numbers(args, ctx, count_logical_refs=logical))
    return pick(numbers) if numbers else 0.0


@_fn("MAX")
def _max(ctx: Ctx, *args: Any) -> float:
    return _extreme(args, ctx, max)


@_fn("MIN")
def _min(ctx: Ctx, *args: Any) -> float:
    return _extreme(args, ctx, min)


@_fn("MAXA")
def _maxa(ctx: Ctx, *args: Any) -> float:
    return _extreme(args, ctx, max, True)


@_fn("MINA")
def _mina(ctx: Ctx, *args: Any) -> float:
    return _extreme(args, ctx, min, True)


@_fn("MEDIAN")
def _median(ctx: Ctx, *args: Any) -> Any:
    numbers = sorted(_numbers(args, ctx))
    if not numbers:
        return NUM
    middle = len(numbers) // 2
    return numbers[middle] if len(numbers) % 2 else (numbers[middle - 1] + numbers[middle]) / 2


@_fn("MODE", "MODE.SNGL")
def _mode(ctx: Ctx, *args: Any) -> Any:
    counts: dict[float, int] = {}
    for number in _numbers(args, ctx):
        counts[number] = counts.get(number, 0) + 1
    best = max(counts.values(), default=0)
    if best < 2:
        return NA
    return next(k for k, v in counts.items() if v == best)


def _variance(numbers: list[float], sample: bool) -> Any:
    size = len(numbers)
    if size < (2 if sample else 1):
        return DIV0
    mean = math.fsum(numbers) / size
    return math.fsum((x - mean) ** 2 for x in numbers) / (size - 1 if sample else size)


def _deviation(numbers: list[float], sample: bool) -> Any:
    variance = _variance(numbers, sample)
    return variance if isinstance(variance, XlError) else math.sqrt(variance)


@_fn("VAR", "VAR.S")
def _var(ctx: Ctx, *args: Any) -> Any:
    return _variance(list(_numbers(args, ctx)), True)


@_fn("VARP", "VAR.P")
def _varp(ctx: Ctx, *args: Any) -> Any:
    return _variance(list(_numbers(args, ctx)), False)


@_fn("STDEV", "STDEV.S")
def _stdev(ctx: Ctx, *args: Any) -> Any:
    return _deviation(list(_numbers(args, ctx)), True)


@_fn("STDEVP", "STDEV.P")
def _stdevp(ctx: Ctx, *args: Any) -> Any:
    return _deviation(list(_numbers(args, ctx)), False)


@_fn("LARGE")
def _large(ctx: Ctx, values: Any, k: Any) -> Any:
    numbers = sorted(_numbers([values], ctx), reverse=True)
    index = math.ceil(_n(_one(k, ctx)))
    return numbers[index - 1] if 1 <= index <= len(numbers) else NUM


@_fn("SMALL")
def _small(ctx: Ctx, values: Any, k: Any) -> Any:
    numbers = sorted(_numbers([values], ctx))
    index = math.ceil(_n(_one(k, ctx)))
    return numbers[index - 1] if 1 <= index <= len(numbers) else NUM


@_fn("RANK", "RANK.EQ")
def _rank_eq(ctx: Ctx, value: Any, ref: Any, order: Any = 0) -> Any:
    number = _n(_one(value, ctx))
    numbers = list(_numbers([ref], ctx))
    if number not in numbers:
        return NA
    ascending = _n(order if order is not None else 0) != 0
    return float(1 + sum(1 for x in numbers if (x < number if ascending else x > number)))


@_fn("RANK.AVG")
def _rank_avg(ctx: Ctx, value: Any, ref: Any, order: Any = 0) -> Any:
    number = _n(_one(value, ctx))
    numbers = list(_numbers([ref], ctx))
    if number not in numbers:
        return NA
    ascending = _n(order if order is not None else 0) != 0
    ahead = sum(1 for x in numbers if (x < number if ascending else x > number))
    return ahead + (numbers.count(number) + 1) / 2


def _percentile_inc(numbers: list[float], p: float) -> Any:
    if not numbers or p < 0 or p > 1:
        return NUM
    position = p * (len(numbers) - 1)
    low = math.floor(position)
    high = min(low + 1, len(numbers) - 1)
    return numbers[low] + (numbers[high] - numbers[low]) * (position - low)


@_fn("PERCENTILE", "PERCENTILE.INC")
def _percentile(ctx: Ctx, values: Any, p: Any) -> Any:
    return _percentile_inc(sorted(_numbers([values], ctx)), _n(_one(p, ctx)))


@_fn("PERCENTILE.EXC")
def _percentile_exc(ctx: Ctx, values: Any, p: Any) -> Any:
    numbers = sorted(_numbers([values], ctx))
    fraction = _n(_one(p, ctx))
    position = fraction * (len(numbers) + 1)
    if not numbers or position < 1 or position > len(numbers):
        return NUM
    low = math.floor(position)
    high = min(low + 1, len(numbers))
    return numbers[low - 1] + (numbers[high - 1] - numbers[low - 1]) * (position - low)


@_fn("QUARTILE", "QUARTILE.INC")
def _quartile(ctx: Ctx, values: Any, q: Any) -> Any:
    part = math.trunc(_n(_one(q, ctx)))
    if not 0 <= part <= 4:
        return NUM
    return _percentile_inc(sorted(_numbers([values], ctx)), part / 4)


def _pairs(xs: Any, ys: Any, ctx: Ctx) -> list[tuple[float, float]]:
    left, right = list(_flat(xs, ctx)), list(_flat(ys, ctx))
    if len(left) != len(right):
        raise _Fail(NA)
    pairs = []
    for a, b in zip(left, right):
        if isinstance(a, XlError):
            raise _Fail(a)
        if isinstance(b, XlError):
            raise _Fail(b)
        if isinstance(a, (int, float)) and isinstance(b, (int, float)) \
                and not isinstance(a, bool) and not isinstance(b, bool):
            pairs.append((float(a), float(b)))
    return pairs


def _line(xs: Any, ys: Any, ctx: Ctx) -> tuple[float, float, float, float]:
    """Sums for a straight-line fit of known ys on known xs."""
    pairs = _pairs(xs, ys, ctx)
    if not pairs:
        raise _Fail(DIV0)
    mean_x = math.fsum(p[0] for p in pairs) / len(pairs)
    mean_y = math.fsum(p[1] for p in pairs) / len(pairs)
    sxy = math.fsum((a - mean_x) * (b - mean_y) for a, b in pairs)
    sxx = math.fsum((a - mean_x) ** 2 for a, _ in pairs)
    syy = math.fsum((b - mean_y) ** 2 for _, b in pairs)
    return mean_x, mean_y, sxy, math.sqrt(sxx * syy) if sxx * syy > 0 else 0.0


@_fn("CORREL", "PEARSON")
def _correl(ctx: Ctx, xs: Any, ys: Any) -> Any:
    _, _, sxy, root = _line(xs, ys, ctx)
    return DIV0 if root == 0 else sxy / root


@_fn("SLOPE")
def _slope(ctx: Ctx, ys: Any, xs: Any) -> Any:
    pairs = _pairs(xs, ys, ctx)
    if not pairs:
        return DIV0
    mean_x = math.fsum(p[0] for p in pairs) / len(pairs)
    sxx = math.fsum((a - mean_x) ** 2 for a, _ in pairs)
    if sxx == 0:
        return DIV0
    mean_y = math.fsum(p[1] for p in pairs) / len(pairs)
    return math.fsum((a - mean_x) * (b - mean_y) for a, b in pairs) / sxx


@_fn("INTERCEPT")
def _intercept(ctx: Ctx, ys: Any, xs: Any) -> Any:
    pairs = _pairs(xs, ys, ctx)
    if not pairs:
        return DIV0
    slope = _slope(ctx, ys, xs)
    if isinstance(slope, XlError):
        return slope
    mean_x = math.fsum(p[0] for p in pairs) / len(pairs)
    mean_y = math.fsum(p[1] for p in pairs) / len(pairs)
    return mean_y - slope * mean_x


@_fn("FORECAST", "FORECAST.LINEAR")
def _forecast(ctx: Ctx, x: Any, ys: Any, xs: Any) -> Any:
    slope = _slope(ctx, ys, xs)
    if isinstance(slope, XlError):
        return slope
    return _intercept(ctx, ys, xs) + slope * _n(_one(x, ctx))


def _criterion(criterion: Any) -> Callable[[Any], bool]:
    """COUNTIF's matching rule."""
    if isinstance(criterion, XlError):
        return lambda v: v == criterion
    if isinstance(criterion, bool):
        return lambda v: isinstance(v, bool) and v == criterion
    if criterion is None:
        criterion = 0.0
    if isinstance(criterion, (int, float)):
        target = float(criterion)
        return lambda v: (isinstance(v, (int, float)) and not isinstance(v, bool) and float(v) == target) \
            or (isinstance(v, str) and parse_number_text(v) == target)
    text = str(criterion)
    op = "="
    for candidate in (">=", "<=", "<>", ">", "<", "="):
        if text.startswith(candidate):
            op, text = candidate, text[len(candidate):]
            break
    number = parse_number_text(text) if text.strip() else None
    if number is not None:
        def numeric(v: Any) -> bool:
            if isinstance(v, bool) or v is None or isinstance(v, XlError):
                return op == "<>" and not isinstance(v, XlError)
            if isinstance(v, str):
                parsed = parse_number_text(v)
                if parsed is None:
                    return op == "<>"
                v = parsed if op in ("=", "<>") else None
                if v is None:
                    return False
            return {"=": v == number, "<>": v != number, ">": v > number, "<": v < number,
                    ">=": v >= number, "<=": v <= number}[op]
        return numeric
    logical = text.upper()
    if logical in ("TRUE", "FALSE") and op in ("=", "<>"):
        wanted = logical == "TRUE"
        return lambda v: (isinstance(v, bool) and v == wanted) == (op == "=")
    if op in ("=", "<>"):
        if text == "":
            # "=" means truly blank; "" matches blank and empty text; "<>" not blank.
            if criterion == "":
                return lambda v: v is None or v == ""
            if op == "=":
                return lambda v: v is None
            return lambda v: v is not None
        pattern = _wildcard(text)

        def textual(v: Any) -> bool:
            hit = isinstance(v, str) and pattern.fullmatch(v) is not None
            return hit if op == "=" else not hit
        return textual

    def ordered(v: Any) -> bool:
        if not isinstance(v, str):
            return False
        order = compare(v, text)
        return {">": order > 0, "<": order < 0, ">=": order >= 0, "<=": order <= 0}[op]
    return ordered


def _wildcard(text: str) -> re.Pattern[str]:
    out: list[str] = []
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "~" and i + 1 < len(text) and text[i + 1] in "*?~":
            out.append(re.escape(text[i + 1]))
            i += 2
            continue
        out.append(".*" if ch == "*" else "." if ch == "?" else re.escape(ch))
        i += 1
    return re.compile("".join(out), re.I | re.S)


def _criteria_hits(ctx: Ctx, pairs: tuple[Any, ...]) -> tuple[list[tuple[int, int]], tuple[int, int]]:
    """Positions (row, col within the block) where every range meets its criterion."""
    if len(pairs) % 2 or not pairs:
        raise _Fail(VALUE)
    first = _matrix(pairs[0], ctx)
    shape = (len(first), len(first[0]) if first else 0)
    hits = [(r, c) for r in range(shape[0]) for c in range(shape[1])]
    for block, criterion in zip(pairs[::2], pairs[1::2]):
        rows = _matrix(block, ctx)
        if (len(rows), len(rows[0]) if rows else 0) != shape:
            raise _Fail(VALUE)
        match = _criterion(_one(criterion, ctx) if not isinstance(criterion, (Ref, Arr)) or
                           (isinstance(criterion, Ref) and criterion.is_cell)
                           else _intersect_value(criterion, ctx))
        hits = [(r, c) for r, c in hits if match(rows[r][c])]
    return hits, shape


def _intersect_value(value: Any, ctx: Ctx) -> Any:
    if isinstance(value, Ref):
        return _intersect(value, ctx)
    if isinstance(value, Arr):
        return value.rows[0][0]
    return value


def _sum_block(values: Any, hits: list[tuple[int, int]], ctx: Ctx) -> Iterator[Any]:
    if isinstance(values, Ref):
        for r, c in hits:
            yield ctx.model.value(values.sheet, values.r1 + r, values.c1 + c)
    else:
        rows = _matrix(values, ctx)
        for r, c in hits:
            yield rows[r][c] if r < len(rows) and c < len(rows[r]) else None


def _real_numbers(values: Iterable[Any]) -> list[float]:
    out: list[float] = []
    for value in values:
        if isinstance(value, XlError):
            raise _Fail(value)
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            out.append(float(value))
    return out


@_fn("COUNTIF")
def _countif(ctx: Ctx, block: Any, criterion: Any) -> float:
    return float(len(_criteria_hits(ctx, (block, criterion))[0]))


@_fn("COUNTIFS")
def _countifs(ctx: Ctx, *pairs: Any) -> float:
    return float(len(_criteria_hits(ctx, pairs)[0]))


@_fn("SUMIF")
def _sumif(ctx: Ctx, block: Any, criterion: Any, values: Any = None) -> float:
    hits, _ = _criteria_hits(ctx, (block, criterion))
    return math.fsum(_real_numbers(_sum_block(values if values is not None else block, hits, ctx)))


@_fn("SUMIFS")
def _sumifs(ctx: Ctx, values: Any, *pairs: Any) -> Any:
    hits, shape = _criteria_hits(ctx, pairs)
    rows = _matrix(values, ctx)
    if (len(rows), len(rows[0]) if rows else 0) != shape:
        return VALUE
    return math.fsum(_real_numbers(_sum_block(values, hits, ctx)))


@_fn("AVERAGEIF")
def _averageif(ctx: Ctx, block: Any, criterion: Any, values: Any = None) -> Any:
    hits, _ = _criteria_hits(ctx, (block, criterion))
    numbers = _real_numbers(_sum_block(values if values is not None else block, hits, ctx))
    return math.fsum(numbers) / len(numbers) if numbers else DIV0


@_fn("AVERAGEIFS")
def _averageifs(ctx: Ctx, values: Any, *pairs: Any) -> Any:
    hits, _ = _criteria_hits(ctx, pairs)
    numbers = _real_numbers(_sum_block(values, hits, ctx))
    return math.fsum(numbers) / len(numbers) if numbers else DIV0


@_fn("MAXIFS")
def _maxifs(ctx: Ctx, values: Any, *pairs: Any) -> float:
    hits, _ = _criteria_hits(ctx, pairs)
    return max(_real_numbers(_sum_block(values, hits, ctx)), default=0.0)


@_fn("MINIFS")
def _minifs(ctx: Ctx, values: Any, *pairs: Any) -> float:
    hits, _ = _criteria_hits(ctx, pairs)
    return min(_real_numbers(_sum_block(values, hits, ctx)), default=0.0)


# ── Lookup and reference ─────────────────────────────────────────────


def _same(a: Any, b: Any) -> bool:
    if isinstance(a, XlError) or isinstance(b, XlError) or a is None or b is None:
        return False
    return _rank(a) == _rank(b) and compare(a, b) == 0


def _vector(value: Any, ctx: Ctx) -> list[Any]:
    rows = _matrix(value, ctx)
    if len(rows) == 1:
        return list(rows[0])
    if rows and len(rows[0]) == 1:
        return [row[0] for row in rows]
    raise _Fail(NA)


def _approximate(items: list[Any], target: Any, *, descending: bool = False) -> int | None:
    """Index of the last value ≤ target (first ≥ target when descending) in sorted data."""
    found = None
    for index, item in enumerate(items):
        if item is None or isinstance(item, XlError) or _rank(item) != _rank(target):
            continue
        order = compare(item, target)
        if descending:
            if order >= 0:
                found = index
            else:
                break
        elif order <= 0:
            found = index
        else:
            break
    return found


def _find_exact(items: list[Any], target: Any) -> int | None:
    if isinstance(target, str) and any(ch in target for ch in "*?~"):
        pattern = _wildcard(target)
        for index, item in enumerate(items):
            if isinstance(item, str) and pattern.fullmatch(item):
                return index
        return None
    for index, item in enumerate(items):
        if _same(item, target):
            return index
    return None


@_fn("MATCH")
def _match(ctx: Ctx, target: Any, block: Any, kind: Any = 1) -> Any:
    wanted = _intersect_value(target, ctx)
    if isinstance(wanted, XlError):
        return wanted
    items = _vector(block, ctx)
    mode = _i(kind if kind is not None else 0)
    if mode == 0:
        found = _find_exact(items, wanted)
    else:
        found = _approximate(items, wanted, descending=mode < 0)
    return NA if found is None else float(found + 1)


@_fn("VLOOKUP")
def _vlookup(ctx: Ctx, target: Any, block: Any, column: Any, approximate: Any = True) -> Any:
    wanted = _intersect_value(target, ctx)
    if isinstance(wanted, XlError):
        return wanted
    rows = _matrix(block, ctx)
    index = _i(_one(column, ctx))
    if index < 1:
        return VALUE
    if rows and index > len(rows[0]):
        return REF
    first = [row[0] for row in rows]
    near = _b(_one(approximate, ctx)) if approximate is not None else False
    found = _approximate(first, wanted) if near else _find_exact(first, wanted)
    if found is None:
        return NA
    return rows[found][index - 1]


@_fn("HLOOKUP")
def _hlookup(ctx: Ctx, target: Any, block: Any, row: Any, approximate: Any = True) -> Any:
    wanted = _intersect_value(target, ctx)
    if isinstance(wanted, XlError):
        return wanted
    rows = _matrix(block, ctx)
    index = _i(_one(row, ctx))
    if index < 1:
        return VALUE
    if index > len(rows):
        return REF
    near = _b(_one(approximate, ctx)) if approximate is not None else False
    found = _approximate(rows[0], wanted) if near else _find_exact(rows[0], wanted)
    if found is None:
        return NA
    return rows[index - 1][found]


@_fn("XLOOKUP")
def _xlookup(ctx: Ctx, target: Any, block: Any, results: Any, missing: Any = None,
             match_mode: Any = 0, search_mode: Any = 1) -> Any:
    wanted = _intersect_value(target, ctx)
    if isinstance(wanted, XlError):
        return wanted
    items = _vector(block, ctx)
    mode = _i(match_mode if match_mode is not None else 0)
    search = _i(search_mode if search_mode is not None else 1)
    order = list(range(len(items))) if search >= 0 else list(range(len(items) - 1, -1, -1))
    if abs(search) == 2:
        raise NotComputable("XLOOKUP binary search")
    found = None
    if mode == 2:
        pattern = _wildcard(str(wanted))
        found = next((i for i in order if isinstance(items[i], str) and pattern.fullmatch(items[i])), None)
    else:
        found = next((i for i in order if _same(items[i], wanted)), None)
        if found is None and mode in (-1, 1):
            best = None
            for i in order:
                item = items[i]
                if item is None or isinstance(item, XlError) or _rank(item) != _rank(wanted):
                    continue
                gap = compare(item, wanted)
                if (mode == -1 and gap < 0 and (best is None or compare(item, items[best]) > 0)) or \
                        (mode == 1 and gap > 0 and (best is None or compare(item, items[best]) < 0)):
                    best = i
            found = best
    if found is None:
        return NA if missing is None else _intersect_value(missing, ctx)
    rows = _matrix(results, ctx)
    source_rows = _matrix(block, ctx)
    across = len(source_rows) == 1
    if across:
        if found >= (len(rows[0]) if rows else 0):
            return VALUE
        if len(rows) == 1:
            return rows[0][found]
    else:
        if found >= len(rows):
            return VALUE
        if rows and len(rows[0]) == 1:
            return rows[found][0]
    raise NotComputable("XLOOKUP returning several cells")


@_fn("LOOKUP")
def _lookup(ctx: Ctx, target: Any, block: Any, results: Any = None) -> Any:
    wanted = _intersect_value(target, ctx)
    if isinstance(wanted, XlError):
        return wanted
    if results is None:
        rows = _matrix(block, ctx)
        if rows and len(rows[0]) > len(rows):
            items, out = rows[0], rows[-1]
        else:
            items, out = [r[0] for r in rows], [r[-1] for r in rows]
    else:
        items, out = _vector(block, ctx), _vector(results, ctx)
    found = _approximate(items, wanted)
    if found is None or found >= len(out):
        return NA
    return out[found]


@_fn("INDEX")
def _index(ctx: Ctx, block: Any, row: Any = None, col: Any = None) -> Any:
    r = _i(_one(row, ctx)) if row is not None else 0
    c = _i(_one(col, ctx)) if col is not None else 0
    if isinstance(block, Ref):
        height, width = block.height, block.width
        if col is None and height == 1 and width > 1:
            r, c = 1, r
        if r < 0 or c < 0 or r > height or c > width:
            return REF
        top, bottom = (block.r1 + r - 1,) * 2 if r else (block.r1, block.r2)
        left, right = (block.c1 + c - 1,) * 2 if c else (block.c1, block.c2)
        if col is None and width > 1 and height > 1 and r:
            return Ref(block.sheet, top, block.c1, bottom, block.c2)
        return Ref(block.sheet, top, left, bottom, right)
    rows = _matrix(block, ctx)
    height, width = len(rows), len(rows[0]) if rows else 0
    if col is None and height == 1:
        r, c = 1, r
    if r < 0 or c < 0 or r > height or c > width:
        return REF
    if r and c:
        return rows[r - 1][c - 1]
    if r:
        return Arr([rows[r - 1]])
    if c:
        return Arr([[row[c - 1]] for row in rows])
    return Arr(rows)


@_fn("ROW")
def _row(ctx: Ctx, ref: Any = None) -> Any:
    if ref is None:
        return float(ctx.row)
    if not isinstance(ref, Ref):
        return VALUE
    if ref.height > 1 and ctx.array:
        return Arr([[float(r)] for r in range(ref.r1, ref.r2 + 1)])
    return float(ref.r1)


@_fn("COLUMN")
def _column(ctx: Ctx, ref: Any = None) -> Any:
    if ref is None:
        return float(ctx.col)
    if not isinstance(ref, Ref):
        return VALUE
    if ref.width > 1 and ctx.array:
        return Arr([[float(c) for c in range(ref.c1, ref.c2 + 1)]])
    return float(ref.c1)


@_fn("ROWS")
def _rows(ctx: Ctx, block: Any) -> Any:
    if isinstance(block, (Ref, Arr)):
        return float(block.height)
    return 1.0


@_fn("COLUMNS")
def _columns(ctx: Ctx, block: Any) -> Any:
    if isinstance(block, (Ref, Arr)):
        return float(block.width)
    return 1.0


@_fn("OFFSET")
def _offset(ctx: Ctx, ref: Any, rows: Any, cols: Any, height: Any = None, width: Any = None) -> Any:
    if not isinstance(ref, Ref):
        return VALUE
    r1 = ref.r1 + _i(_one(rows, ctx))
    c1 = ref.c1 + _i(_one(cols, ctx))
    h = _i(_one(height, ctx)) if height is not None else ref.height
    w = _i(_one(width, ctx)) if width is not None else ref.width
    if h == 0 or w == 0:
        return REF
    if h < 0:
        r1, h = r1 + h + 1, -h
    if w < 0:
        c1, w = c1 + w + 1, -w
    if r1 < 1 or c1 < 1 or r1 + h - 1 > MAX_ROW or c1 + w - 1 > MAX_COL:
        return REF
    return Ref(ref.sheet, r1, c1, r1 + h - 1, c1 + w - 1)


@_fn("INDIRECT")
def _indirect(ctx: Ctx, text: Any, a1: Any = True) -> Any:
    if a1 is not None and not _b(_one(a1, ctx)):
        raise NotComputable("INDIRECT with R1C1 references")
    operand = parse_operand(_s(_one(text, ctx)).strip())
    if isinstance(operand, (Box, NameOperand, TableOperand)):
        result = resolve_operand(operand, ctx)
        return result if isinstance(result, Ref) else REF
    return REF


@_fn("ADDRESS")
def _address(ctx: Ctx, row: Any, col: Any, kind: Any = 1, a1: Any = True, sheet: Any = None) -> Any:
    r, c = _i(_one(row, ctx)), _i(_one(col, ctx))
    mode = _i(kind if kind is not None else 1)
    if r < 1 or c < 1 or r > MAX_ROW or c > MAX_COL or not 1 <= mode <= 4:
        return VALUE
    if a1 is not None and not _b(_one(a1, ctx)):
        raise NotComputable("ADDRESS in R1C1 style")
    text = f"{'$' if mode in (1, 3) else ''}{get_column_letter(c)}{'$' if mode in (1, 2) else ''}{r}"
    return f"{quote_sheet(_s(_one(sheet, ctx)))}!{text}" if sheet is not None else text


# ── Text ─────────────────────────────────────────────────────────────


@_scalar("LEFT")
def _left(ctx: Ctx, text: Any, count: Any = 1) -> Any:
    n = _i(count if count is not None else 1)
    return VALUE if n < 0 else _s(text)[:n]


@_scalar("RIGHT")
def _right(ctx: Ctx, text: Any, count: Any = 1) -> Any:
    n = _i(count if count is not None else 1)
    if n < 0:
        return VALUE
    value = _s(text)
    return value[-n:] if n else ""


@_scalar("MID")
def _mid(ctx: Ctx, text: Any, start: Any, count: Any) -> Any:
    first, n = _i(start), _i(count)
    if first < 1 or n < 0:
        return VALUE
    return _s(text)[first - 1:first - 1 + n]


@_scalar("LEN")
def _len(ctx: Ctx, text: Any) -> float:
    return float(len(_s(text)))


@_scalar("TRIM")
def _trim(ctx: Ctx, text: Any) -> str:
    return re.sub(" +", " ", _s(text)).strip(" ")


@_scalar("UPPER")
def _upper(ctx: Ctx, text: Any) -> str:
    return _s(text).upper()


@_scalar("LOWER")
def _lower(ctx: Ctx, text: Any) -> str:
    return _s(text).lower()


@_scalar("PROPER")
def _proper(ctx: Ctx, text: Any) -> str:
    return re.sub(r"[^\W\d_]+", lambda m: m.group().capitalize(), _s(text).lower())


@_scalar("CLEAN")
def _clean(ctx: Ctx, text: Any) -> str:
    return "".join(ch for ch in _s(text) if ord(ch) >= 32)


@_scalar("EXACT")
def _exact(ctx: Ctx, a: Any, b: Any) -> bool:
    return _s(a) == _s(b)


@_scalar("REPT")
def _rept(ctx: Ctx, text: Any, count: Any) -> Any:
    n = _i(count)
    return VALUE if n < 0 else _s(text) * n


@_scalar("CHAR", "UNICHAR")
def _char(ctx: Ctx, code: Any) -> Any:
    number = _i(code)
    return VALUE if number < 1 or number > 0x10FFFF else chr(number)


@_scalar("CODE", "UNICODE")
def _code(ctx: Ctx, text: Any) -> Any:
    value = _s(text)
    return float(ord(value[0])) if value else VALUE


@_fn("CONCATENATE")
def _concatenate(ctx: Ctx, *args: Any) -> str:
    return "".join(_s(_intersect_value(a, ctx) if isinstance(a, (Ref, Arr)) else a) for a in args)


@_fn("CONCAT")
def _concat(ctx: Ctx, *args: Any) -> str:
    return "".join(_s(v) for a in args for v in _flat(a, ctx))


@_fn("TEXTJOIN")
def _textjoin(ctx: Ctx, separator: Any, ignore_empty: Any, *args: Any) -> str:
    sep = _s(_one(separator, ctx))
    skip = _b(_one(ignore_empty, ctx))
    parts = [_s(v) for a in args for v in _flat(a, ctx)]
    return sep.join(p for p in parts if p != "" or not skip)


@_scalar("FIND")
def _find(ctx: Ctx, needle: Any, text: Any, start: Any = 1) -> Any:
    first = _i(start if start is not None else 1)
    value = _s(text)
    if first < 1 or first > len(value) + 1:
        return VALUE
    found = value.find(_s(needle), first - 1)
    return VALUE if found < 0 else float(found + 1)


@_scalar("SEARCH")
def _search(ctx: Ctx, needle: Any, text: Any, start: Any = 1) -> Any:
    first = _i(start if start is not None else 1)
    value = _s(text)
    if first < 1 or first > len(value) + 1:
        return VALUE
    match = _wildcard(_s(needle)).search(value, first - 1)
    return VALUE if match is None else float(match.start() + 1)


@_scalar("SUBSTITUTE")
def _substitute(ctx: Ctx, text: Any, old: Any, new: Any, which: Any = None) -> Any:
    value, before, after = _s(text), _s(old), _s(new)
    if not before:
        return value
    if which is None:
        return value.replace(before, after)
    nth = _i(which)
    if nth < 1:
        return VALUE
    position = -1
    for _ in range(nth):
        position = value.find(before, position + 1)
        if position < 0:
            return value
    return value[:position] + after + value[position + len(before):]


@_scalar("REPLACE")
def _replace(ctx: Ctx, text: Any, start: Any, count: Any, new: Any) -> Any:
    first, n = _i(start), _i(count)
    if first < 1 or n < 0:
        return VALUE
    value = _s(text)
    return value[:first - 1] + _s(new) + value[first - 1 + n:]


@_scalar("TEXT")
def _text(ctx: Ctx, value: Any, fmt: Any) -> Any:
    if isinstance(value, XlError):
        return value
    text_fmt = _s(fmt)
    item = 0.0 if value is None else value
    if isinstance(item, str):
        parsed = parse_number_text(item)
        if parsed is not None:
            item = parsed
    shown = format_value(item, text_fmt, ctx.model.date1904)
    if shown is None:
        raise NotComputable(f"TEXT format {text_fmt!r}")
    return shown


@_scalar("VALUE")
def _value(ctx: Ctx, text: Any) -> Any:
    if isinstance(text, (int, float)) and not isinstance(text, bool):
        return float(text)
    if text is None:
        return 0.0
    parsed = parse_number_text(_s(text))
    return VALUE if parsed is None else parsed


@_scalar("NUMBERVALUE")
def _numbervalue(ctx: Ctx, text: Any, decimal: Any = ".", group: Any = ",") -> Any:
    value = _s(text).replace(" ", "")
    point = _s(decimal if decimal is not None else ".")[:1]
    thousands = _s(group if group is not None else ",")[:1]
    if not value:
        return 0.0
    percents = len(value) - len(value.rstrip("%"))
    value = value.rstrip("%").replace(thousands, "").replace(point, ".")
    try:
        return float(value) / (100 ** percents)
    except ValueError:
        return VALUE


@_scalar("FIXED")
def _fixed(ctx: Ctx, value: Any, digits: Any = 2, no_commas: Any = False) -> Any:
    places = _i(digits if digits is not None else 2)
    number = _round_with(_n(value), places, ROUND_HALF_UP)
    commas = not (_b(no_commas) if no_commas is not None else False)
    shown = max(places, 0)
    return f"{number:,.{shown}f}" if commas else f"{number:.{shown}f}"


@_scalar("TEXTBEFORE")
def _textbefore(ctx: Ctx, text: Any, delimiter: Any) -> Any:
    value, mark = _s(text), _s(delimiter)
    found = value.find(mark)
    return NA if found < 0 else value[:found]


@_scalar("TEXTAFTER")
def _textafter(ctx: Ctx, text: Any, delimiter: Any) -> Any:
    value, mark = _s(text), _s(delimiter)
    found = value.find(mark)
    return NA if found < 0 else value[found + len(mark):]


# ── Dates and times ──────────────────────────────────────────────────


@_scalar("DATE")
def _date(ctx: Ctx, year: Any, month: Any, day: Any) -> Any:
    y, m, d = _i(year), _i(month), _i(day)
    if 0 <= y < 1900:
        y += 1900
    if y < 0 or y > 9999:
        return NUM
    first = _add_months(dt.date(max(y, 1900), 1, 1), m - 1)
    result = first + dt.timedelta(days=d - 1)
    if result.year > 9999:
        return NUM
    return _make_serial(ctx, result)


@_scalar("YEAR")
def _year(ctx: Ctx, value: Any) -> float:
    return float(_serial(ctx, value).year)


@_scalar("MONTH")
def _month(ctx: Ctx, value: Any) -> float:
    return float(_serial(ctx, value).month)


@_scalar("DAY")
def _day(ctx: Ctx, value: Any) -> float:
    return float(_serial(ctx, value).day)


@_fn("TODAY")
def _today(ctx: Ctx) -> float:
    return _make_serial(ctx, dt.date.today())


@_fn("NOW")
def _now(ctx: Ctx) -> float:
    moment = dt.datetime.now()
    seconds = moment.hour * 3600 + moment.minute * 60 + moment.second
    return _make_serial(ctx, moment.date()) + seconds / 86400


@_scalar("TIME")
def _time(ctx: Ctx, hour: Any, minute: Any, second: Any) -> Any:
    total = _i(hour) * 3600 + _i(minute) * 60 + _i(second)
    return NUM if total < 0 else (total % 86400) / 86400


def _day_seconds(value: float) -> int:
    return round((value - math.floor(value)) * 86400)


@_scalar("HOUR")
def _hour(ctx: Ctx, value: Any) -> float:
    return float(_day_seconds(_n(value)) // 3600 % 24)


@_scalar("MINUTE")
def _minute(ctx: Ctx, value: Any) -> float:
    return float(_day_seconds(_n(value)) // 60 % 60)


@_scalar("SECOND")
def _second(ctx: Ctx, value: Any) -> float:
    return float(_day_seconds(_n(value)) % 60)


@_scalar("EDATE")
def _edate(ctx: Ctx, start: Any, months: Any) -> Any:
    return _make_serial(ctx, _add_months(_serial(ctx, start), _i(months)))


@_scalar("EOMONTH")
def _eomonth(ctx: Ctx, start: Any, months: Any) -> Any:
    moved = _add_months(_serial(ctx, start).replace(day=1), _i(months))
    return _make_serial(ctx, moved.replace(day=_month_end(moved.year, moved.month)))


@_scalar("WEEKDAY")
def _weekday(ctx: Ctx, value: Any, kind: Any = 1) -> Any:
    mode = _i(kind if kind is not None else 1)
    # Python: Monday = 0.
    weekday = _serial(ctx, value).weekday()
    if mode == 1 or mode == 17:
        return float((weekday + 1) % 7 + 1)
    if mode == 2 or mode == 11:
        return float(weekday + 1)
    if mode == 3:
        return float(weekday)
    if 12 <= mode <= 16:
        return float((weekday - (mode - 11)) % 7 + 1)
    return NUM


@_scalar("WEEKNUM")
def _weeknum(ctx: Ctx, value: Any, kind: Any = 1) -> Any:
    day = _serial(ctx, value)
    mode = _i(kind if kind is not None else 1)
    if mode == 21:
        return float(day.isocalendar()[1])
    start = {1: 6, 2: 0, 11: 0, 12: 1, 13: 2, 14: 3, 15: 4, 16: 5, 17: 6}.get(mode)
    if start is None:
        return NUM
    first = dt.date(day.year, 1, 1)
    offset = (first.weekday() - start) % 7
    return float(((day - first).days + offset) // 7 + 1)


@_scalar("ISOWEEKNUM")
def _isoweeknum(ctx: Ctx, value: Any) -> float:
    return float(_serial(ctx, value).isocalendar()[1])


@_scalar("DAYS")
def _days(ctx: Ctx, end: Any, start: Any) -> float:
    return float(math.floor(_n(end)) - math.floor(_n(start)))


def _days360(a: dt.date, b: dt.date, european: bool) -> int:
    d1, d2 = a.day, b.day
    if european:
        d1, d2 = min(d1, 30), min(d2, 30)
    else:
        a_last = a.day == _month_end(a.year, a.month)
        b_last = b.day == _month_end(b.year, b.month)
        if a.month == 2 and a_last:
            if b.month == 2 and b_last:
                d2 = 30
            d1 = 30
        if d2 == 31 and d1 >= 30:
            d2 = 30
        if d1 == 31:
            d1 = 30
    return (b.year - a.year) * 360 + (b.month - a.month) * 30 + d2 - d1


@_scalar("DAYS360")
def _days360_fn(ctx: Ctx, start: Any, end: Any, method: Any = False) -> float:
    return float(_days360(_serial(ctx, start), _serial(ctx, end),
                          _b(method) if method is not None else False))


@_scalar("DATEDIF")
def _datedif(ctx: Ctx, start: Any, end: Any, unit: Any) -> Any:
    a, b = _serial(ctx, start), _serial(ctx, end)
    if a > b:
        return NUM
    kind = _s(unit).upper()
    months = (b.year - a.year) * 12 + b.month - a.month
    if b.day < a.day:
        months -= 1
    if kind == "D":
        return float((b - a).days)
    if kind == "M":
        return float(months)
    if kind == "Y":
        return float(months // 12)
    if kind == "YM":
        return float(months % 12)
    if kind == "MD":
        anchor = _add_months(a, months)
        return float((b - anchor).days)
    if kind == "YD":
        anchor = _add_months(a, (months // 12) * 12)
        return float((b - anchor).days)
    return NUM


@_scalar("DATEVALUE")
def _datevalue(ctx: Ctx, text: Any) -> Any:
    value = _s(text).strip()
    match = re.fullmatch(r"(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T].*)?", value)
    if match:
        return _make_serial(ctx, dt.date(int(match.group(1)), int(match.group(2)), int(match.group(3))))
    # 3/4/2026 reads differently per locale; Excel knows the user's, we don't.
    raise NotComputable("DATEVALUE of a local date text")


@_scalar("TIMEVALUE")
def _timevalue(ctx: Ctx, text: Any) -> Any:
    match = re.fullmatch(r"\s*(\d{1,2}):(\d{2})(?::(\d{2}))?\s*", _s(text))
    if match is None:
        raise NotComputable("TIMEVALUE of a local time text")
    hour, minute, second = int(match.group(1)), int(match.group(2)), int(match.group(3) or 0)
    return (hour * 3600 + minute * 60 + second) / 86400


@_scalar("YEARFRAC")
def _yearfrac(ctx: Ctx, start: Any, end: Any, basis: Any = 0) -> Any:
    a, b = _serial(ctx, start), _serial(ctx, end)
    if a > b:
        a, b = b, a
    mode = _i(basis if basis is not None else 0)
    if mode == 0:
        return _days360(a, b, False) / 360
    if mode == 2:
        return (b - a).days / 360
    if mode == 3:
        return (b - a).days / 365
    if mode == 4:
        return _days360(a, b, True) / 360
    if mode != 1:
        return NUM
    if a == b:
        return 0.0
    if (b.year == a.year) or (b.year == a.year + 1 and (b.month, b.day) <= (a.month, a.day)):
        leap = (_is_leap(a.year) and (a.month, a.day) <= (2, 29)) or \
            (_is_leap(b.year) and (b.month, b.day) >= (2, 29))
        if b.year == a.year:
            leap = _is_leap(a.year)
        return (b - a).days / (366 if leap else 365)
    years = range(a.year, b.year + 1)
    average = sum(366 if _is_leap(y) else 365 for y in years) / len(years)
    return (b - a).days / average


def _is_leap(year: int) -> bool:
    return year % 4 == 0 and (year % 100 != 0 or year % 400 == 0)


def _holidays(value: Any, ctx: Ctx) -> set[int]:
    if value is None:
        return set()
    return {math.floor(v) for v in _numbers([value], ctx)}


@_fn("NETWORKDAYS")
def _networkdays(ctx: Ctx, start: Any, end: Any, holidays: Any = None) -> float:
    a, b = math.floor(_n(_one(start, ctx))), math.floor(_n(_one(end, ctx)))
    sign = 1.0
    if a > b:
        a, b, sign = b, a, -1.0
    skip = _holidays(holidays, ctx)
    total = 0
    for serial in range(a, b + 1):
        if serial_to_date(serial, ctx.model.date1904).weekday() < 5 and serial not in skip:
            total += 1
    return sign * total


@_fn("WORKDAY")
def _workday(ctx: Ctx, start: Any, days: Any, holidays: Any = None) -> float:
    serial = math.floor(_n(_one(start, ctx)))
    left = _i(_one(days, ctx))
    skip = _holidays(holidays, ctx)
    step = 1 if left >= 0 else -1
    left = abs(left)
    while left:
        serial += step
        if serial_to_date(serial, ctx.model.date1904).weekday() < 5 and serial not in skip:
            left -= 1
    return float(serial)


# ── Money ────────────────────────────────────────────────────────────


def _pmt(rate: float, nper: float, pv: float, fv: float, kind: float) -> float:
    if nper == 0:
        raise _Fail(NUM)
    if rate == 0:
        return -(pv + fv) / nper
    growth = (1 + rate) ** nper
    return -(fv + pv * growth) * rate / ((1 + rate * kind) * (growth - 1))


def _fv(rate: float, nper: float, pmt: float, pv: float, kind: float) -> float:
    if rate == 0:
        return -(pv + pmt * nper)
    growth = (1 + rate) ** nper
    return -(pv * growth + pmt * (1 + rate * kind) * (growth - 1) / rate)


def _ipmt(rate: float, per: float, nper: float, pv: float, fv: float, kind: float) -> float:
    if per < 1 or per > nper:
        raise _Fail(NUM)
    payment = _pmt(rate, nper, pv, fv, kind)
    if kind and per == 1:
        return 0.0
    interest = _fv(rate, per - 1, payment, pv, kind) * rate
    return interest / (1 + rate) if kind else interest


def _optional(value: Any, default: float = 0.0) -> float:
    return default if value is None else _n(value)


@_scalar("PMT")
def _pmt_fn(ctx: Ctx, rate: Any, nper: Any, pv: Any, fv: Any = None, kind: Any = None) -> float:
    return _pmt(_n(rate), _n(nper), _n(pv), _optional(fv), _optional(kind))


@_scalar("FV")
def _fv_fn(ctx: Ctx, rate: Any, nper: Any, pmt: Any, pv: Any = None, kind: Any = None) -> float:
    return _fv(_n(rate), _n(nper), _optional(pmt), _optional(pv), _optional(kind))


@_scalar("PV")
def _pv_fn(ctx: Ctx, rate: Any, nper: Any, pmt: Any, fv: Any = None, kind: Any = None) -> float:
    r, n, p, f, t = _n(rate), _n(nper), _optional(pmt), _optional(fv), _optional(kind)
    if r == 0:
        return -(f + p * n)
    growth = (1 + r) ** n
    return -(f + p * (1 + r * t) * (growth - 1) / r) / growth


@_scalar("NPER")
def _nper_fn(ctx: Ctx, rate: Any, pmt: Any, pv: Any, fv: Any = None, kind: Any = None) -> Any:
    r, p, present, f, t = _n(rate), _n(pmt), _n(pv), _optional(fv), _optional(kind)
    if r == 0:
        return DIV0 if p == 0 else -(present + f) / p
    numerator = p * (1 + r * t) - f * r
    denominator = p * (1 + r * t) + present * r
    if denominator == 0 or numerator / denominator <= 0:
        return NUM
    return math.log(numerator / denominator) / math.log(1 + r)


@_scalar("IPMT")
def _ipmt_fn(ctx: Ctx, rate: Any, per: Any, nper: Any, pv: Any, fv: Any = None, kind: Any = None) -> float:
    return _ipmt(_n(rate), _n(per), _n(nper), _n(pv), _optional(fv), _optional(kind))


@_scalar("PPMT")
def _ppmt_fn(ctx: Ctx, rate: Any, per: Any, nper: Any, pv: Any, fv: Any = None, kind: Any = None) -> float:
    r, p, n, present, f, t = _n(rate), _n(per), _n(nper), _n(pv), _optional(fv), _optional(kind)
    return _pmt(r, n, present, f, t) - _ipmt(r, p, n, present, f, t)


@_scalar("CUMIPMT")
def _cumipmt(ctx: Ctx, rate: Any, nper: Any, pv: Any, first: Any, last: Any, kind: Any) -> Any:
    r, n, present, a, b, t = _n(rate), _n(nper), _n(pv), _i(first), _i(last), _i(kind)
    if r <= 0 or n <= 0 or present <= 0 or a < 1 or b < a or t not in (0, 1):
        return NUM
    return math.fsum(_ipmt(r, p, n, present, 0.0, t) for p in range(a, b + 1))


@_scalar("CUMPRINC")
def _cumprinc(ctx: Ctx, rate: Any, nper: Any, pv: Any, first: Any, last: Any, kind: Any) -> Any:
    r, n, present, a, b, t = _n(rate), _n(nper), _n(pv), _i(first), _i(last), _i(kind)
    if r <= 0 or n <= 0 or present <= 0 or a < 1 or b < a or t not in (0, 1):
        return NUM
    payment = _pmt(r, n, present, 0.0, t)
    return math.fsum(payment - _ipmt(r, p, n, present, 0.0, t) for p in range(a, b + 1))


@_fn("NPV")
def _npv(ctx: Ctx, rate: Any, *args: Any) -> Any:
    r = _n(_one(rate, ctx))
    if r == -1:
        return DIV0
    return math.fsum(v / (1 + r) ** (i + 1) for i, v in enumerate(_numbers(args, ctx)))


def _solve(fn: Callable[[float], float], guess: float) -> float:
    """A root by Newton's method, falling back to bisection over a wide bracket."""
    rate = guess
    for _ in range(100):
        try:
            value = fn(rate)
            step = (fn(rate + 1e-7) - value) / 1e-7
        except (OverflowError, ZeroDivisionError, ValueError):
            break
        if isinstance(value, complex) or isinstance(step, complex):
            break
        if abs(value) < 1e-9:
            return rate
        if step == 0:
            break
        rate -= value / step
        if rate <= -1:
            rate = (rate + value / step - 1) / 2
    points = [-0.99 + 0.01 * i for i in range(0, 100)] + [0.1 * i for i in range(10, 101)] + [10.0, 100.0, 1000.0]
    last_x: float | None = None
    last_y: float | None = None
    for x in points:
        try:
            y = fn(x)
        except (OverflowError, ZeroDivisionError, ValueError):
            continue
        if isinstance(y, complex):
            continue
        if last_y is not None and last_x is not None and y * last_y <= 0:
            low, high = last_x, x
            for _ in range(200):
                middle = (low + high) / 2
                y_mid = fn(middle)
                if y_mid * last_y <= 0:
                    high = middle
                else:
                    low, last_y = middle, y_mid
            return (low + high) / 2
        last_x, last_y = x, y
    raise _Fail(NUM)


@_fn("IRR")
def _irr(ctx: Ctx, values: Any, guess: Any = None) -> Any:
    flows = list(_numbers([values], ctx))
    if not any(v > 0 for v in flows) or not any(v < 0 for v in flows):
        return NUM

    def npv(rate: float) -> float:
        return math.fsum(v / (1 + rate) ** i for i, v in enumerate(flows))

    return _solve(npv, _n(_one(guess, ctx)) if guess is not None else 0.1)


@_fn("XNPV")
def _xnpv(ctx: Ctx, rate: Any, values: Any, dates: Any) -> Any:
    r = _n(_one(rate, ctx))
    flows, when = list(_numbers([values], ctx)), list(_numbers([dates], ctx))
    if len(flows) != len(when) or not flows:
        return NUM
    if r <= -1:
        return NUM
    return math.fsum(v / (1 + r) ** ((math.floor(d) - math.floor(when[0])) / 365) for v, d in zip(flows, when))


@_fn("XIRR")
def _xirr(ctx: Ctx, values: Any, dates: Any, guess: Any = None) -> Any:
    flows, when = list(_numbers([values], ctx)), list(_numbers([dates], ctx))
    if len(flows) != len(when) or not any(v > 0 for v in flows) or not any(v < 0 for v in flows):
        return NUM

    def npv(rate: float) -> float:
        return math.fsum(v / (1 + rate) ** ((math.floor(d) - math.floor(when[0])) / 365)
                         for v, d in zip(flows, when))

    return _solve(npv, _n(_one(guess, ctx)) if guess is not None else 0.1)


@_fn("MIRR")
def _mirr(ctx: Ctx, values: Any, finance: Any, reinvest: Any) -> Any:
    flows = list(_numbers([values], ctx))
    f, r = _n(_one(finance, ctx)), _n(_one(reinvest, ctx))
    n = len(flows)
    negative = math.fsum(v / (1 + f) ** i for i, v in enumerate(flows) if v < 0)
    positive = math.fsum(v * (1 + r) ** (n - i - 1) for i, v in enumerate(flows) if v > 0)
    if negative == 0 or positive == 0 or n < 2:
        return DIV0
    return (positive / -negative) ** (1 / (n - 1)) - 1


@_fn("RATE")
def _rate(ctx: Ctx, nper: Any, pmt: Any, pv: Any, fv: Any = None, kind: Any = None, guess: Any = None) -> Any:
    n, p, present = _n(_one(nper, ctx)), _n(_one(pmt, ctx)), _n(_one(pv, ctx))
    f, t = _optional(_one(fv, ctx)), _optional(_one(kind, ctx))

    def balance(rate: float) -> float:
        if abs(rate) < 1e-12:
            return present + p * n + f
        growth = (1 + rate) ** n
        return present * growth + p * (1 + rate * t) * (growth - 1) / rate + f

    return _solve(balance, _n(_one(guess, ctx)) if guess is not None else 0.1)


@_scalar("EFFECT")
def _effect(ctx: Ctx, rate: Any, periods: Any) -> Any:
    r, n = _n(rate), math.trunc(_n(periods))
    return NUM if r <= 0 or n < 1 else (1 + r / n) ** n - 1


@_scalar("NOMINAL")
def _nominal(ctx: Ctx, rate: Any, periods: Any) -> Any:
    r, n = _n(rate), math.trunc(_n(periods))
    return NUM if r <= 0 or n < 1 else n * ((1 + r) ** (1 / n) - 1)


@_scalar("SLN")
def _sln(ctx: Ctx, cost: Any, salvage: Any, life: Any) -> Any:
    years = _n(life)
    return DIV0 if years == 0 else (_n(cost) - _n(salvage)) / years


@_scalar("SYD")
def _syd(ctx: Ctx, cost: Any, salvage: Any, life: Any, period: Any) -> Any:
    c, s, n, p = _n(cost), _n(salvage), _n(life), _n(period)
    if n <= 0 or p < 1 or p > n:
        return NUM
    return (c - s) * (n - p + 1) * 2 / (n * (n + 1))
