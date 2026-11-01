"""PowerPoint decks (.pptx) for agents, without scripts.

The promise is slides that fit, in the firm's own template:

* **The template does the design.** Slides come from the deck's layouts and fill
  their placeholders; a table, chart or picture takes the empty placeholder that
  suits it. Text stays in the master's fonts and sizes unless asked otherwise.
* **Nothing overflows unnoticed.** There is no layout engine in a .pptx library,
  so text is measured here (font metrics, what the shape inherits from layout and
  master, wrapping) and each result says what doesn't fit. Where a placeholder
  shrinks text on overflow, the shrink PowerPoint would apply is stored, because
  PowerPoint only works it out itself when someone edits the text.
* **Nothing else is lost.** python-pptx keeps every part it doesn't model.
  Deleting, moving and duplicating slides, which it has no calls for, are done on
  the XML here, along with what else points at a slide: sections, custom shows and
  links from other slides.

Nothing here decides where files live; ``tools/pptx_tool.py`` owns paths, backups
and saving.
"""

from __future__ import annotations

import copy
import io
import re
import shutil
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable, Iterator

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.chart.plot import XyPlot
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.package import XmlPart
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.presentation import Presentation as DeckClass
from pptx.util import Emu, Inches, Pt

from captain_claw.docx_ops import _FOLD

EMU_PER_INCH = 914400
EMU_PER_PT = 12700
MAX_UNZIPPED = 1024 * 1024 * 1024

CT_PRESENTATION = b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml"
CT_MACRO = b"application/vnd.ms-powerpoint.presentation.macroEnabled.main+xml"
_OTHER_MAIN_TYPES = {
    b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml": CT_PRESENTATION,
    b"application/vnd.openxmlformats-officedocument.presentationml.slideshow.main+xml": CT_PRESENTATION,
    b"application/vnd.ms-powerpoint.template.macroEnabled.main+xml": CT_MACRO,
    b"application/vnd.ms-powerpoint.slideshow.macroEnabled.main+xml": CT_MACRO,
}
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
C_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
DGM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
A_P, A_R, A_T, A_BR, A_FLD = qn("a:p"), qn("a:r"), qn("a:t"), qn("a:br"), qn("a:fld")
A_RPR, A_PPR, A_END = qn("a:rPr"), qn("a:pPr"), qn("a:endParaRPr")
_BULLET_TAGS = tuple(qn(f"a:{t}") for t in ("buNone", "buAutoNum", "buChar", "buBlip"))
_FOOTER_TYPES = (PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER)


class PptxError(ValueError):
    """A request the deck can't satisfy; the message is for the model."""


# ── Opening and saving ───────────────────────────────────────────────


def open_deck(data: bytes) -> DeckClass:
    """A deck from its bytes. A template or slide show opens as an ordinary deck."""
    if data[:4] == b"\xd0\xcf\x11\xe0":
        raise PptxError("This is a password-protected deck or an old .ppt file. "
                        "Ask for an unprotected copy saved as .pptx.")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNZIPPED:
                raise PptxError("This deck is too large to open here (over 1 GB unpacked).")
            types = archive.read("[Content_Types].xml")
    except (zipfile.BadZipFile, KeyError) as exc:
        raise PptxError(f"Not a readable PowerPoint deck ({exc}).") from None
    for other, plain in _OTHER_MAIN_TYPES.items():
        if other in types:
            data = _with_content_types(data, types.replace(other, plain))
            break
    try:
        return Presentation(io.BytesIO(data))
    except PptxError:
        raise
    except Exception as exc:
        raise PptxError(f"Not a readable PowerPoint deck ({type(exc).__name__}: {exc}).") from None


def _with_content_types(data: bytes, types: bytes) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            target.writestr(info, types if info.filename == "[Content_Types].xml" else source.read(info.filename))
    return out.getvalue()


def deck_bytes(prs: DeckClass) -> bytes:
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()


def verify_saved(data: bytes) -> None:
    """The bytes must open again, every slide readable, before they replace a file."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)):
            raise PptxError("The result has two parts with one name, so nothing was saved.")
    prs = open_deck(data)
    for slide in prs.slides:
        list(slide.shapes)
        _ = slide.slide_layout.name


# ── Addresses ────────────────────────────────────────────────────────


_KIND_NAMES = {
    PP_PLACEHOLDER.TITLE: "title", PP_PLACEHOLDER.CENTER_TITLE: "title", PP_PLACEHOLDER.SUBTITLE: "subtitle",
    PP_PLACEHOLDER.BODY: "body", PP_PLACEHOLDER.OBJECT: "content", PP_PLACEHOLDER.CHART: "chart",
    PP_PLACEHOLDER.TABLE: "table", PP_PLACEHOLDER.PICTURE: "picture", PP_PLACEHOLDER.BITMAP: "picture",
    PP_PLACEHOLDER.MEDIA_CLIP: "media", PP_PLACEHOLDER.ORG_CHART: "diagram",
    PP_PLACEHOLDER.DATE: "date", PP_PLACEHOLDER.FOOTER: "footer", PP_PLACEHOLDER.SLIDE_NUMBER: "slide number",
    PP_PLACEHOLDER.VERTICAL_BODY: "body", PP_PLACEHOLDER.VERTICAL_OBJECT: "content",
    PP_PLACEHOLDER.VERTICAL_TITLE: "title",
}


def placeholder_kind(shape: Any) -> str | None:
    if not shape.is_placeholder:
        return None
    return _KIND_NAMES.get(shape.placeholder_format.type, "content")


def _walk(shapes: Iterable[Any]) -> Iterator[Any]:
    """Shapes, including those inside groups."""
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _walk(shape.shapes)


def inches(emu: Any) -> float:
    return round(int(emu or 0) / EMU_PER_INCH, 2)


def all_layouts(prs: DeckClass) -> list[Any]:
    return [layout for master in prs.slide_masters for layout in master.slide_layouts]


def _layout_description(layout: Any) -> str:
    parts = []
    for ph in layout.placeholders:
        if ph.placeholder_format.type in _FOOTER_TYPES:
            continue
        kind = _KIND_NAMES.get(ph.placeholder_format.type, "content")
        parts.append(f"{kind}" if kind == "title" else f"{kind}:{ph.placeholder_format.idx}")
    return ", ".join(parts) or "nothing to fill"


def find_layout(prs: DeckClass, ref: Any) -> Any:
    layouts = all_layouts(prs)
    listing = "; ".join(f'{i} "{layout.name}"' for i, layout in enumerate(layouts, 1))
    if isinstance(ref, (int, float)) or str(ref).strip().isdigit():
        index = int(ref)
        if 1 <= index <= len(layouts):
            return layouts[index - 1]
        raise PptxError(f"There is no layout {index}. Layouts: {listing}.")
    wanted = str(ref).strip().casefold()
    for layout in layouts:
        if layout.name.casefold() == wanted:
            return layout
    near = [layout for layout in layouts if wanted in layout.name.casefold()]
    if len(near) == 1:
        return near[0]
    raise PptxError(f"There is no layout {ref!r}. Layouts: {listing}.")


def _content_placeholders(layout: Any) -> list[Any]:
    return [ph for ph in layout.placeholders if ph.placeholder_format.type in (
        PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.BODY)]


def _default_layout(prs: DeckClass, spec: dict[str, Any]) -> Any:
    """The layout a slide's content asks for when it names none."""
    visuals = sum(1 for key in ("table", "chart", "image") if spec.get(key))
    text = 1 if spec.get("body") else 0
    wanted = visuals + text
    layouts = all_layouts(prs)
    titled = [la for la in layouts if any(
        ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE) for ph in la.placeholders)]
    if wanted == 0 and spec.get("title") and not spec.get("subtitle"):
        for layout in titled:
            if not [ph for ph in layout.placeholders if ph.placeholder_format.type not in _FOOTER_TYPES
                    and ph.placeholder_format.type not in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)]:
                return layout
    # A layout with room for each piece, and no more than that.
    for count in range(max(wanted, 1), 5):
        for layout in titled:
            if len(_content_placeholders(layout)) == count and "caption" not in layout.name.casefold() \
                    and "vertical" not in layout.name.casefold():
                return layout
    return titled[0] if titled else layouts[0]


# ── What a shape's text inherits ─────────────────────────────────────

_THEME_FONTS: dict[int, tuple[str, str]] = {}


def _theme_fonts(master: Any) -> tuple[str, str]:
    key = id(master.part)
    found = _THEME_FONTS.get(key)
    if found is None:
        major, minor = "Calibri", "Calibri"
        try:
            root = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
            major = root.find(f".//{qn('a:majorFont')}/{qn('a:latin')}").get("typeface") or major
            minor = root.find(f".//{qn('a:minorFont')}/{qn('a:latin')}").get("typeface") or minor
        except Exception:
            pass
        found = _THEME_FONTS[key] = (major, minor)
    return found


class TextContext:
    """Where a shape's (or table cell's) text gets what it doesn't set itself:
    its own list style, the layout's and master's placeholders, the master's
    text styles, then the deck's defaults."""

    def __init__(self, shape: Any, slide: Any, prs: DeckClass, *, tx_body: Any = None) -> None:
        self.shape = shape
        self.slide = slide
        self.master = slide.slide_layout.slide_master
        self.tx_body = tx_body
        self.chain: list[Any] = []
        kind = None
        if tx_body is None:
            self.tx_body = shape._element.find(qn("p:txBody"))
            element_chain = [shape._element]
            if shape.is_placeholder and shape.shape_type != MSO_SHAPE_TYPE.GROUP:
                kind = placeholder_kind(shape)
                current = shape
                for _ in range(2):
                    try:
                        current = current._base_placeholder
                    except Exception:
                        current = None
                    if current is None:
                        break
                    element_chain.append(current._element)
            self.chain = [el.find(qn("p:txBody")) for el in element_chain]
            self.chain = [el for el in self.chain if el is not None]
        else:
            self.chain = [tx_body]
        self.is_title = kind == "title"
        styles = self.master._element.find(qn("p:txStyles"))
        self.base: list[Any] = []
        if styles is not None:
            if kind == "title":
                self.base.append(styles.find(qn("p:titleStyle")))
            elif kind in ("body", "content", "subtitle", "chart", "table", "picture", "media", "diagram"):
                self.base.append(styles.find(qn("p:bodyStyle")))
            elif kind is not None:
                self.base.append(styles.find(qn("p:otherStyle")))
        self.base.append(prs.part._element.find(qn("p:defaultTextStyle")))
        self.base = [el for el in self.base if el is not None]

    def _levels(self, level: int) -> Iterator[Any]:
        tag = qn(f"a:lvl{level + 1}pPr")
        for body in self.chain:
            style = body.find(qn("a:lstStyle"))
            if style is not None:
                found = style.find(tag)
                if found is not None:
                    yield found
        for base in self.base:
            found = base.find(tag)
            if found is not None:
                yield found

    def _body_prs(self) -> Iterator[Any]:
        for body in self.chain:
            found = body.find(qn("a:bodyPr"))
            if found is not None:
                yield found

    def body_attr(self, name: str, default: str | None = None) -> str | None:
        for props in self._body_prs():
            value = props.get(name)
            if value is not None:
                return value
        return default

    def inset(self, name: str, default: int) -> int:
        value = self.body_attr(name)
        return int(value) if value is not None else default

    def autofit(self) -> tuple[str, float, float]:
        """("shrink" | "grow" | "none", stored font scale, stored line reduction)."""
        for props in self._body_prs():
            shrink = props.find(qn("a:normAutofit"))
            if shrink is not None:
                return ("shrink", int(shrink.get("fontScale", "100000")) / 100000,
                        int(shrink.get("lnSpcReduction", "0")) / 100000)
            if props.find(qn("a:spAutoFit")) is not None:
                return "grow", 1.0, 0.0
            if props.find(qn("a:noAutofit")) is not None:
                return "none", 1.0, 0.0
        return "none", 1.0, 0.0

    def size(self, level: int, r_pr: Any = None) -> float:
        if r_pr is not None and r_pr.get("sz"):
            return int(r_pr.get("sz")) / 100
        for props in self._levels(level):
            defaults = props.find(qn("a:defRPr"))
            if defaults is not None and defaults.get("sz"):
                return int(defaults.get("sz")) / 100
        return 18.0

    def is_bold(self, level: int, r_pr: Any = None) -> bool:
        if r_pr is not None and r_pr.get("b") is not None:
            return r_pr.get("b") in ("1", "true")
        for props in self._levels(level):
            defaults = props.find(qn("a:defRPr"))
            if defaults is not None and defaults.get("b") is not None:
                return defaults.get("b") in ("1", "true")
        return False

    def font(self, level: int, r_pr: Any = None) -> str:
        typeface = None
        if r_pr is not None:
            latin = r_pr.find(qn("a:latin"))
            if latin is not None:
                typeface = latin.get("typeface")
        if typeface is None:
            for props in self._levels(level):
                defaults = props.find(qn("a:defRPr"))
                latin = defaults.find(qn("a:latin")) if defaults is not None else None
                if latin is not None and latin.get("typeface"):
                    typeface = latin.get("typeface")
                    break
        major, minor = _theme_fonts(self.master)
        if typeface is None or typeface == "+mn-lt":
            return major if (typeface is None and self.is_title) else minor
        if typeface == "+mj-lt":
            return major
        return typeface

    def para_attr(self, p_pr: Any, level: int, name: str, default: int = 0) -> int:
        if p_pr is not None and p_pr.get(name) is not None:
            return int(p_pr.get(name))
        for props in self._levels(level):
            if props.get(name) is not None:
                return int(props.get(name))
        return default

    def spacing(self, p_pr: Any, level: int, tag: str) -> tuple[str, float] | None:
        """("pct", 0.9) or ("pts", 12.0) for lnSpc, spcBef or spcAft."""
        for props in ([p_pr] if p_pr is not None else []) + list(self._levels(level)):
            holder = props.find(qn(f"a:{tag}"))
            if holder is None:
                continue
            pct = holder.find(qn("a:spcPct"))
            if pct is not None:
                return "pct", int(pct.get("val", "100000")) / 100000
            pts = holder.find(qn("a:spcPts"))
            if pts is not None:
                return "pts", int(pts.get("val", "0")) / 100
        return None

    def has_bullet(self, p_pr: Any, level: int) -> bool:
        for props in ([p_pr] if p_pr is not None else []) + list(self._levels(level)):
            for child in props:
                if child.tag in _BULLET_TAGS:
                    return child.tag != _BULLET_TAGS[0]
        return False


# ── Measuring text ───────────────────────────────────────────────────

# Helvetica's advance widths (the standard AFM table), per 1000 of the font size.
_WIDTHS = dict(zip(
    " !\"#$%&'()*+,-./0123456789:;<=>?@ABCDEFGHIJKLMNOPQRSTUVWXYZ[\\]^_`abcdefghijklmnopqrstuvwxyz{|}~",
    [278, 278, 355, 556, 556, 889, 667, 191, 333, 333, 389, 584, 278, 333, 278, 278, 556, 556, 556, 556, 556, 556,
     556, 556, 556, 556, 278, 278, 584, 584, 584, 556, 1015, 667, 667, 722, 722, 667, 611, 778, 722, 278, 500, 667,
     556, 833, 722, 778, 667, 778, 722, 667, 611, 722, 667, 944, 667, 667, 611, 278, 278, 278, 469, 556, 333, 556,
     556, 500, 556, 556, 278, 556, 556, 222, 222, 500, 222, 833, 556, 556, 556, 556, 333, 500, 278, 556, 500, 722,
     500, 500, 500, 334, 260, 334, 584]))
# Other fonts against Helvetica, (regular, bold). Measured in PowerPoint for the
# fonts it had; Lato, Montserrat and Roboto are their published widths.
_FONT_WIDTHS = {
    "calibri": (0.92, 0.94), "calibri light": (0.91, 0.91), "arial": (1.0, 1.07), "helvetica": (1.0, 1.07),
    "helvetica neue": (1.01, 1.07), "aptos": (0.96, 1.01), "aptos display": (0.9, 0.9), "cambria": (0.96, 1.03),
    "times new roman": (0.92, 0.97), "georgia": (1.0, 1.16), "verdana": (1.16, 1.29), "tahoma": (1.02, 1.15),
    "segoe ui": (0.99, 1.0), "trebuchet ms": (1.03, 1.07), "garamond": (0.88, 0.96), "century gothic": (1.09, 1.09),
    "gill sans mt": (0.92, 1.05), "franklin gothic book": (0.95, 0.95), "open sans": (1.05, 1.1), "roboto": (0.98, 1.02),
    "lato": (0.97, 1.02), "montserrat": (1.14, 1.2),
}
_DEFAULT_WIDTHS = (1.0, 1.07)
# PowerPoint set every font it measured at 1.2 of its size (Windows varies a little by font).
LINE_HEIGHT = 1.2


def text_width(text: str, size: float, font: str = "", bold: bool = False) -> float:
    """Width of *text* in points."""
    total = 0
    for ch in text:
        width = _WIDTHS.get(ch)
        if width is None:
            width = 1000 if ord(ch) >= 0x2E80 else 556
        total += width
    regular, heavy = _FONT_WIDTHS.get(font.casefold(), _DEFAULT_WIDTHS)
    return total / 1000 * size * (heavy if bold else regular)


_SPLIT_WORDS = re.compile(r"(\s+)")


def _wrap_count(pieces: list[tuple[str, float, str, bool]], first_width: float, next_width: float) -> int:
    """Lines needed for (text, size, font, bold) pieces at these widths, in points."""
    lines = 1
    room = max(first_width, 1.0)
    used = 0.0
    pending_space = 0.0
    for text, size, font, bold in pieces:
        for token in _SPLIT_WORDS.split(text):
            if not token:
                continue
            width = text_width(token, size, font, bold)
            if token.isspace():
                pending_space += width
                continue
            if used and used + pending_space + width > room:
                lines += 1
                room = max(next_width, 1.0)
                used = 0.0
            elif used:
                used += pending_space
            pending_space = 0.0
            # A word wider than the line breaks across lines.
            while width > room:
                lines += 1
                width -= room
                room = max(next_width, 1.0)
            used += width
    return lines


@dataclass(slots=True)
class TextFit:
    height: float  # points, insets included
    lines: int
    largest: float
    smallest: float


def measure_text(ctx: TextContext, width_emu: int, *, scale: float = 1.0, reduction: float = 0.0) -> TextFit:
    """How tall the text comes out in a box this wide."""
    wrap = ctx.body_attr("wrap", "square") != "none"
    width = (width_emu - ctx.inset("lIns", 91440) - ctx.inset("rIns", 91440)) / EMU_PER_PT
    total = (ctx.inset("tIns", 45720) + ctx.inset("bIns", 45720)) / EMU_PER_PT
    lines_total = 0
    largest = 0.0
    smallest = 0.0
    previous_after = 0.0
    paragraphs = ctx.tx_body.findall(A_P) if ctx.tx_body is not None else []
    for index, para in enumerate(paragraphs):
        p_pr = para.find(A_PPR)
        level = int(p_pr.get("lvl", "0")) if p_pr is not None else 0
        segments: list[list[tuple[str, float, str, bool]]] = [[]]
        para_size = 0.0
        for child in para:
            if child.tag in (A_R, A_FLD):
                r_pr = child.find(A_RPR)
                size = ctx.size(level, r_pr) * scale
                font = ctx.font(level, r_pr)
                t = child.find(A_T)
                text = t.text if t is not None and t.text else ""
                segments[-1].append((text, size, font, ctx.is_bold(level, r_pr)))
                if text.strip():
                    para_size = max(para_size, size)
                    smallest = size if not smallest else min(smallest, size)
            elif child.tag == A_BR:
                segments.append([])
        if not para_size:
            # An empty line is as tall as its own mark says.
            end = para.find(A_END)
            para_size = ctx.size(level, end) * scale
        largest = max(largest, para_size)
        mar_l = ctx.para_attr(p_pr, level, "marL")
        indent = ctx.para_attr(p_pr, level, "indent")
        next_width = width - mar_l / EMU_PER_PT
        first_width = width - max(mar_l + indent, mar_l if indent < 0 else 0) / EMU_PER_PT
        count = 0
        for pieces in segments:
            count += 1 if not wrap else _wrap_count(pieces, first_width, next_width)
        spacing = ctx.spacing(p_pr, level, "lnSpc")
        if spacing is not None and spacing[0] == "pts":
            line_height = spacing[1]
        else:
            factor = spacing[1] if spacing is not None else 1.0
            line_height = para_size * LINE_HEIGHT * max(factor - reduction, 0.1)
        before = ctx.spacing(p_pr, level, "spcBef")
        after = ctx.spacing(p_pr, level, "spcAft")
        gap_before = 0.0
        if index and before is not None:
            # Space above the first paragraph isn't drawn.
            gap_before = before[1] if before[0] == "pts" else before[1] * para_size * LINE_HEIGHT
        gap_after = 0.0
        if after is not None:
            gap_after = after[1] if after[0] == "pts" else after[1] * para_size * LINE_HEIGHT
        total += previous_after + gap_before + count * line_height
        previous_after = gap_after
        lines_total += count
    return TextFit(total, lines_total, largest, smallest or largest)


_SHRINK_STEPS = ((1.0, 0.1), (0.925, 0.1), (0.925, 0.2), (0.85, 0.2), (0.775, 0.2), (0.7, 0.2), (0.625, 0.2), (0.55, 0.2))


def apply_autofit(shape: Any, slide: Any, prs: DeckClass) -> float | None:
    """Store what PowerPoint would work out for shrink-on-overflow text. Returns
    the font scale when it had to shrink, None when the text fits as it is."""
    ctx = TextContext(shape, slide, prs)
    kind, _, _ = ctx.autofit()
    if kind != "shrink" or ctx.tx_body is None:
        return None
    box = int(shape.height or 0) / EMU_PER_PT
    width = int(shape.width or 0)
    own = ctx.tx_body.find(qn("a:bodyPr"))

    def store(scale: float, reduction: float) -> None:
        holder = own.find(qn("a:normAutofit"))
        if scale >= 1.0 and not reduction:
            if holder is not None:
                holder.attrib.pop("fontScale", None)
                holder.attrib.pop("lnSpcReduction", None)
            return
        if holder is None:
            for tag in ("a:noAutofit", "a:spAutoFit"):
                old = own.find(qn(tag))
                if old is not None:
                    own.remove(old)
            holder = etree.SubElement(own, qn("a:normAutofit"))
            # The schema wants autofit before anything else that can be in bodyPr.
            own.remove(holder)
            own.insert(0 if own.find(qn("a:prstTxWarp")) is None else 1, holder)
        if scale < 1.0:
            holder.set("fontScale", str(int(scale * 100000)))
        else:
            holder.attrib.pop("fontScale", None)
        holder.set("lnSpcReduction", str(int(reduction * 100000)))

    if measure_text(ctx, width).height <= box:
        store(1.0, 0.0)
        return None
    for scale, reduction in _SHRINK_STEPS:
        if measure_text(ctx, width, scale=scale, reduction=reduction).height <= box:
            store(scale, reduction)
            return scale
    store(*_SHRINK_STEPS[-1])
    return _SHRINK_STEPS[-1][0]


# ── Reading a deck out ───────────────────────────────────────────────


def paragraph_text(para: Any) -> str:
    parts = []
    for child in para:
        if child.tag in (A_R, A_FLD):
            t = child.find(A_T)
            parts.append(t.text if t is not None and t.text else "")
        elif child.tag == A_BR:
            parts.append("\n")
    return "".join(parts)


def _shape_text(shape: Any) -> str:
    if not shape.has_text_frame:
        return ""
    return "\n".join(paragraph_text(p) for p in shape.text_frame._txBody.findall(A_P)).strip()


def _table_cells(shape: Any) -> list[list[str]]:
    table = shape.table
    return [[table.cell(r, c).text.replace("\n", " / ") for c in range(len(table.columns))]
            for r in range(len(table.rows))]


def _chart_summary(chart: Any) -> str:
    kind = str(chart.chart_type).split(" ")[0].split(".")[-1].replace("_", " ").lower()
    title = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else ""
    parts = [f'{kind} "{title}"' if title else kind]
    try:
        plot = chart.plots[0]
        if isinstance(plot, XyPlot):
            for series in list(plot.series)[:12]:
                points = list(series.iter_values())[:40]
                parts.append(f"{series.name}: {_compact_list(points)}")
        else:
            parts.append("categories: " + _compact_list(list(plot.categories)[:40]))
            for series in list(plot.series)[:12]:
                parts.append(f"{series.name}: {_compact_list(list(series.values)[:40])}")
    except Exception:
        parts.append("(its data can't be read here)")
    return "; ".join(parts)


def _compact_list(values: list[Any]) -> str:
    shown = []
    for value in values:
        if isinstance(value, float) and value.is_integer():
            shown.append(str(int(value)))
        else:
            shown.append(str(value))
    return ", ".join(shown)


def _alt_text(shape: Any) -> str:
    """Alt text worth the name: a bare file name, as scripts leave it, isn't."""
    props = shape._element.xpath(".//p:cNvPr")
    text = (props[0].get("descr") or "") if props else ""
    return "" if re.fullmatch(r"\S+\.(png|jpe?g|gif|bmp|tiff?|emf|wmf|svg|webp)", text, re.I) else text


def _smartart_text(shape: Any) -> str:
    """The words in a SmartArt diagram: read here, changed only in PowerPoint."""
    ids = shape._element.findall(f".//{{{DGM_NS}}}relIds")
    if not ids:
        return ""
    data_id = ids[0].get(f"{{{R_NS}}}dm")
    try:
        root = etree.fromstring(shape.part.related_part(data_id).blob)
    except Exception:
        return ""
    return " / ".join(t.text for t in root.iter(A_T) if t.text and t.text.strip())


def is_smartart(shape: Any) -> bool:
    return bool(shape._element.findall(f".//{{{DGM_NS}}}relIds"))


def _geometry(shape: Any) -> str:
    try:
        return f"at {inches(shape.left)}, {inches(shape.top)} size {inches(shape.width)} × {inches(shape.height)} in"
    except Exception:
        return ""


def describe_shape(shape: Any, slide: Any, prs: DeckClass, *, indent: str = "  ", depth: int = 0) -> list[str]:
    """A shape as text, addressed by its id: what `edit` needs to aim at it."""
    head = f"{indent}[{shape.shape_id}] "
    kind = placeholder_kind(shape)
    ph = f":{shape.placeholder_format.idx}" if kind and kind != "title" else ""
    label = f"{kind}{ph}" if kind else None
    where = _geometry(shape)
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        lines = [f"{head}group {shape.name!r} {where}"]
        if depth < 3:
            for child in shape.shapes:
                lines += describe_shape(child, slide, prs, indent=indent + "  ", depth=depth + 1)
        return lines
    if getattr(shape, "has_table", False) and shape.has_table:
        cells = _table_cells(shape)
        lines = [f"{head}{label or 'table'} table {len(cells)} rows × {len(cells[0]) if cells else 0} columns {where}"]
        lines += [f"{indent}    r{r} " + " | ".join(row) for r, row in enumerate(cells[:40], 1)]
        if len(cells) > 40:
            lines.append(f"{indent}    … and {len(cells) - 40} more rows")
        return lines
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return [f"{head}{label or 'chart'} chart {where}: {_chart_summary(shape.chart)}"]
    if is_smartart(shape):
        return [f"{head}SmartArt {where} (change it in PowerPoint): {_smartart_text(shape)[:400]}"]
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE or shape._element.tag == qn("p:pic"):
        alt = _alt_text(shape)
        return [f"{head}{label or 'picture'} picture {shape.name!r} {where}" + (f', alt: "{alt}"' if alt else ", no alt text")]
    if not shape.has_text_frame:
        return [f"{head}{str(shape.shape_type).split(' ')[0].lower()} {shape.name!r} {where}"]
    ctx = TextContext(shape, slide, prs)
    paragraphs = ctx.tx_body.findall(A_P) if ctx.tx_body is not None else []
    text = "\n".join(paragraph_text(p) for p in paragraphs).strip()
    label = label or f"text box {shape.name!r}"
    if not text:
        return [f"{head}{label} {where}: (empty)"]
    lines_out: list[str] = []
    size_shown = None
    for para in paragraphs:
        p_pr = para.find(A_PPR)
        level = int(p_pr.get("lvl", "0")) if p_pr is not None else 0
        body = paragraph_text(para).replace("\n", " ⏎ ")
        run = para.find(A_R)
        if size_shown is None and run is not None:
            size_shown = ctx.size(level, run.find(A_RPR))
        mark = "• " if ctx.has_bullet(p_pr, level) else ""
        lines_out.append("  " * level + mark + body)
    scale = ctx.autofit()[1]
    size_text = f", {size_shown * scale:g}pt" if size_shown else ""
    if len(lines_out) == 1 and kind in ("title", "subtitle", "slide number", "date", "footer"):
        return [f'{head}{label}{size_text}: "{lines_out[0]}"']
    out = [f"{head}{label} {where}{size_text}:"]
    out += [f"{indent}    {line}" for line in lines_out[:60]]
    if len(lines_out) > 60:
        out.append(f"{indent}    … and {len(lines_out) - 60} more lines")
    return out


def notes_text(slide: Any) -> str:
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    return frame.text.strip() if frame is not None else ""


def describe_slide(prs: DeckClass, slide: Any, number: int) -> list[str]:
    hidden = " (hidden)" if slide._element.get("show") == "0" else ""
    lines = [f'Slide {number}{hidden} · layout "{slide.slide_layout.name}"']
    for shape in slide.shapes:
        lines += describe_shape(shape, slide, prs)
    notes = notes_text(slide)
    if notes:
        shown = notes if len(notes) <= 800 else notes[:800] + "…"
        lines.append("  notes: " + shown.replace("\n", " ⏎ "))
    return lines


def overview(prs: DeckClass) -> list[str]:
    width, height = prs.slide_width, prs.slide_height
    major, minor = _theme_fonts(prs.slide_masters[0])
    lines = [f"{len(prs.slides)} slide(s), {inches(width)} × {inches(height)} in. Fonts: {major} (headings), {minor} (text)."]
    described = "; ".join(f'{i} "{la.name}" [{_layout_description(la)}]' for i, la in enumerate(all_layouts(prs), 1))
    lines.append(f"Layouts: {described}")
    sections = _sections(prs)
    if sections:
        numbers = {el.get("id"): n for n, el in enumerate(prs.slides._sldIdLst, 1)}
        parts = []
        for section in sections:
            ids = [numbers.get(s.get("id")) for s in section.iter(f"{{{P14}}}sldId")]
            ids = [i for i in ids if i]
            span = f"{min(ids)}–{max(ids)}" if ids else "empty"
            parts.append(f'"{section.get("name")}" ({span})')
        lines.append("Sections: " + ", ".join(parts))
    if any(s._element.get("show") == "0" for s in prs.slides):
        hidden = [str(n) for n, s in enumerate(prs.slides, 1) if s._element.get("show") == "0"]
        lines.append("Hidden slides: " + ", ".join(hidden))
    if _has_macros(prs):
        lines.append("Macros: yes (kept as they are, never run).")
    return lines


def _has_macros(prs: DeckClass) -> bool:
    return prs.part.content_type == CT_MACRO.decode()


def find_text(prs: DeckClass, needle: str, *, limit: int = 100) -> list[str]:
    wanted = needle.translate(_FOLD).casefold()
    out: list[str] = []
    for number, slide in enumerate(prs.slides, 1):
        for shape in _walk(slide.shapes):
            texts: list[str] = []
            if shape.has_text_frame:
                texts = [paragraph_text(p) for p in shape.text_frame._txBody.findall(A_P)]
            elif getattr(shape, "has_table", False) and shape.has_table:
                texts = [" | ".join(row) for row in _table_cells(shape)]
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                texts = [_chart_summary(shape.chart)]
            for text in texts:
                if wanted in text.translate(_FOLD).casefold():
                    out.append(f"Slide {number} [{shape.shape_id}]: {text.strip()[:240]}")
                    if len(out) >= limit:
                        return out
        notes = notes_text(slide)
        if notes and wanted in notes.translate(_FOLD).casefold():
            out.append(f"Slide {number} notes: {notes[:240]}")
    return out


# ── Writing text ─────────────────────────────────────────────────────

_LINE = re.compile(r"^(\s*)(?:([-*•–])\s+|(\d+)[.)]\s+)?(.*)$", re.S)
_INLINE = re.compile(
    r"\*\*(?P<bold>.+?)\*\*"
    r"|(?<![\w*])\*(?=\S)(?P<italic>.+?)(?<=\S)\*(?![\w*])"
    r"|\[(?P<link>[^\]]+)\]\((?P<url>[^)\s]+)\)")
_ILLEGAL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f￾￿]")


@dataclass(slots=True)
class _Line:
    level: int
    marker: str  # "", "bullet" or "number"
    body: str


def _parse_lines(text: str) -> list[_Line]:
    lines: list[_Line] = []
    for raw in _ILLEGAL.sub("", str(text)).split("\n"):
        match = _LINE.match(raw.rstrip())
        indent, bullet, number, body = match.groups()
        width = len(indent.expandtabs(2))
        marker = "bullet" if bullet else ("number" if number else "")
        lines.append(_Line(min(width // 2, 8), marker, body))
    return lines


def _segments(body: str) -> list[tuple[str, bool, bool, str | None]]:
    """(text, bold, italic, link) pieces of **bold**, *italic* and [links](url)."""
    out: list[tuple[str, bool, bool, str | None]] = []
    position = 0
    for match in _INLINE.finditer(body):
        if match.start() > position:
            out.append((body[position:match.start()], False, False, None))
        if match.group("bold") is not None:
            out.append((match.group("bold"), True, False, None))
        elif match.group("italic") is not None:
            out.append((match.group("italic"), False, True, None))
        else:
            out.append((match.group("link"), False, False, match.group("url")))
        position = match.end()
    if position < len(body):
        out.append((body[position:], False, False, None))
    return out


def _set_bullet(p: Any, kind: str, level: int) -> None:
    """An explicit bullet (or none) on a paragraph, in schema order."""
    p_pr = p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag in _BULLET_TAGS:
            p_pr.remove(child)
    if kind == "none":
        element = etree.Element(qn("a:buNone"))
        p_pr.set("marL", "0")
        p_pr.set("indent", "0")
    elif kind == "number":
        element = etree.Element(qn("a:buAutoNum"), type="arabicPeriod")
        p_pr.set("marL", str(342900 * (level + 1)))
        p_pr.set("indent", "-342900")
    else:
        element = etree.Element(qn("a:buChar"), char="•")
        p_pr.set("marL", str(285750 * (level + 1)))
        p_pr.set("indent", "-285750")
    before = (qn("a:lnSpc"), qn("a:spcBef"), qn("a:spcAft"), qn("a:buClrTx"), qn("a:buClr"), qn("a:buSzTx"),
              qn("a:buSzPct"), qn("a:buSzPts"), qn("a:buFontTx"), qn("a:buFont"))
    position = 0
    for index, child in enumerate(p_pr):
        if child.tag in before:
            position = index + 1
    p_pr.insert(position, element)


def write_text(tx_body: Any, part: Any, text: str, *, ctx: TextContext | None = None,
               bullets: bool | None = None) -> None:
    """Replace a text body's words, keeping each level's formatting: the paragraph
    and run properties already there, per level, or those of its first paragraph."""
    old = list(tx_body.findall(A_P))
    templates: dict[int, tuple[Any, Any, Any]] = {}
    for para in old:
        p_pr = para.find(A_PPR)
        level = int(p_pr.get("lvl", "0")) if p_pr is not None else 0
        run = para.find(A_R)
        if run is None:
            run = para.find(A_FLD)
        r_pr = run.find(A_RPR) if run is not None else None
        if level not in templates:
            templates[level] = (p_pr, r_pr, para.find(A_END))
        elif templates[level][1] is None and r_pr is not None:
            templates[level] = (templates[level][0], r_pr, templates[level][2])
    fallback = templates.get(min(templates)) if templates else (None, None, None)
    for para in old:
        tx_body.remove(para)
    lines = _parse_lines(text)
    for line in lines:
        p = tx_body.add_p()
        same_level = templates.get(line.level)
        p_pr, r_pr, end = same_level or (None, fallback[1], fallback[2])
        if p_pr is not None:
            p.insert(0, copy.deepcopy(p_pr))
        if line.level:
            p.get_or_add_pPr().set("lvl", str(line.level))
        inherited_bullet = ctx.has_bullet(p.find(A_PPR), line.level) if ctx is not None else False
        if line.marker == "number":
            _set_bullet(p, "number", line.level)
        elif line.marker == "bullet" and not inherited_bullet:
            _set_bullet(p, "bullet", line.level)
        elif bullets is False:
            _set_bullet(p, "none", line.level)
        elif bullets is True and not inherited_bullet:
            _set_bullet(p, "bullet", line.level)
        for piece, bold, italic, link in _segments(line.body) or [("", False, False, None)]:
            if not piece and r_pr is None and not bold and not italic:
                continue
            r = p.add_r()
            if r_pr is not None:
                own = copy.deepcopy(r_pr)
                for stale in own.findall(qn("a:hlinkClick")):
                    own.remove(stale)
                r.insert(0, own)
            r.text = piece
            if bold or italic or link:
                props = r.get_or_add_rPr()
                if bold:
                    props.set("b", "1")
                if italic:
                    props.set("i", "1")
                if link:
                    if not re.match(r"(https?://|mailto:)", link, re.I):
                        raise PptxError(f"{link!r}: links here are http(s) or mailto addresses.")
                    props.add_hlinkClick(part.relate_to(link, RT.HYPERLINK, is_external=True))
        if end is not None:
            p.append(copy.deepcopy(end))


def replace_in_paragraph(para: Any, find: str, new: str) -> int:
    """Replace every *find* in a paragraph, across runs, keeping the formatting of
    the run each match starts in. Typographic quotes and dashes match plain ones."""
    runs = [(child, child.find(A_T)) for child in para if child.tag == A_R]
    runs = [(r, t) for r, t in runs if t is not None]
    pieces = [t.text or "" for _, t in runs]
    full = "".join(pieces)
    target = find.translate(_FOLD)
    folded = full.translate(_FOLD)
    positions = []
    start = folded.find(target)
    while start >= 0 and target:
        positions.append(start)
        start = folded.find(target, start + len(target))
    if not positions:
        return 0
    for match_start in reversed(positions):
        match_end = match_start + len(find)
        offset = 0
        inserted = False
        for index, piece in enumerate(pieces):
            piece_end = offset + len(piece)
            if piece_end > match_start and offset < match_end:
                keep_before = piece[: max(match_start - offset, 0)]
                keep_after = piece[max(match_end - offset, 0):] if match_end < piece_end else ""
                pieces[index] = keep_before + (new if not inserted else "") + keep_after
                inserted = True
            offset = piece_end
        # Pieces after this match have moved, but the next match is earlier: its offsets still hold.
    for (r, t), piece in zip(runs, pieces):
        if piece:
            t.text = piece
        else:
            para.remove(r)
    return len(positions)


def _isolate(para: Any, find: str) -> list[Any]:
    """The runs that hold exactly *find*, splitting runs at its edges."""
    target = find.translate(_FOLD)
    full = "".join((c.find(A_T).text or "") for c in para if c.tag == A_R and c.find(A_T) is not None)
    start = full.translate(_FOLD).find(target)
    if start < 0 or not target:
        return []
    end = start + len(find)
    offset = 0
    covered: list[Any] = []
    for run in [c for c in para if c.tag == A_R]:
        t = run.find(A_T)
        text = t.text or "" if t is not None else ""
        run_start, run_end = offset, offset + len(text)
        offset = run_end
        if run_end <= start or run_start >= end or not text:
            continue
        cut_from = max(start - run_start, 0)
        cut_to = min(end - run_start, len(text))
        middle = run
        if cut_from:
            left = copy.deepcopy(run)
            left.find(A_T).text = text[:cut_from]
            run.addprevious(left)
        if cut_to < len(text):
            right = copy.deepcopy(run)
            right.find(A_T).text = text[cut_to:]
            run.addnext(right)
        middle.find(A_T).text = text[cut_from:cut_to]
        covered.append(middle)
    return covered


# ── Sections and what else points at a slide ─────────────────────────


def _sections(prs: DeckClass) -> list[Any]:
    return list(prs.part._element.iter(f"{{{P14}}}section"))


def _section_forget(prs: DeckClass, slide_id: str) -> None:
    for section in _sections(prs):
        for item in section.iter(f"{{{P14}}}sldId"):
            if item.get("id") == slide_id:
                item.getparent().remove(item)
                break


def _section_join(prs: DeckClass, slide_id: str, neighbour_id: str | None) -> None:
    """A slide goes in its neighbour's section (or the first): PowerPoint expects
    every slide of a sectioned deck to be listed in one."""
    sections = _sections(prs)
    if not sections:
        return
    for section in sections:
        lst = section.find(f"{{{P14}}}sldIdLst")
        if lst is None:
            continue
        for index, item in enumerate(lst):
            if item.get("id") == neighbour_id:
                lst.insert(index + 1, etree.Element(f"{{{P14}}}sldId", id=slide_id))
                return
    lst = sections[0].find(f"{{{P14}}}sldIdLst")
    if lst is None:
        lst = etree.SubElement(sections[0], f"{{{P14}}}sldIdLst")
    lst.insert(0, etree.Element(f"{{{P14}}}sldId", id=slide_id))


def _slide_id_of(prs: DeckClass, slide: Any) -> Any:
    return next(el for el in prs.slides._sldIdLst if prs.part.related_part(el.rId) is slide.part)


def _position_of(prs: DeckClass, slide: Any) -> int:
    for number, el in enumerate(prs.slides._sldIdLst, 1):
        if prs.part.related_part(el.rId) is slide.part:
            return number
    raise PptxError("That slide is no longer in the deck.")


def _move_to(prs: DeckClass, slide: Any, position: int) -> None:
    ids = prs.slides._sldIdLst
    element = _slide_id_of(prs, slide)
    ids.remove(element)
    ids.insert(max(min(position - 1, len(ids)), 0), element)
    if _sections(prs):
        slide_id = element.get("id")
        _section_forget(prs, slide_id)
        members = list(ids)
        index = members.index(element)
        _section_join(prs, slide_id, members[index - 1].get("id") if index else None)
    prs.part.rename_slide_parts([el.rId for el in ids])


def _delete_slide(prs: DeckClass, slide: Any) -> int:
    """Remove a slide and everything that names it. Returns the links it broke."""
    element = _slide_id_of(prs, slide)
    r_id, slide_id = element.rId, element.get("id")
    prs.slides._sldIdLst.remove(element)
    _section_forget(prs, slide_id)
    for show in list(prs.part._element.iter(qn("p:custShow"))):
        for item in list(show.iter(qn("p:sld"))):
            if item.get(qn("r:id")) == r_id:
                item.getparent().remove(item)
    broken = 0
    for other in prs.slides:
        for rel_id, rel in list(other.part.rels.items()):
            if rel.reltype == RT.SLIDE and not rel.is_external and rel.target_part is slide.part:
                for link in list(other._element.iter(qn("a:hlinkClick"), qn("a:hlinkHover"))):
                    if link.get(qn("r:id")) == rel_id:
                        link.getparent().remove(link)
                        broken += 1
                other.part.rels.pop(rel_id)
    # The next new slide takes a name from how many there are: no gaps, or two
    # slides would share one.
    prs.part.drop_rel(r_id)
    prs.part.rename_slide_parts([el.rId for el in prs.slides._sldIdLst])
    return broken


_SHARED_PART_PREFIXES = ("/ppt/media/", "/ppt/slideLayouts/", "/ppt/slideMasters/", "/ppt/theme/", "/ppt/tags/",
                         "/ppt/slides/", "/ppt/notesSlides/")


def _clone_part(part: Any, package: Any, clones: dict[int, Any]) -> Any:
    """A separate copy of a chart (or the like) and what it owns."""
    found = clones.get(id(part))
    if found is not None:
        return found
    template = re.sub(r"\d*(\.\w+)$", r"%d\1", str(part.partname))
    partname = PackURI(package.next_partname(template))
    if isinstance(part, XmlPart):
        new = part.__class__(partname, part.content_type, package, copy.deepcopy(part._element))
    else:
        new = part.__class__(partname, part.content_type, package, part.blob)
    clones[id(part)] = new
    id_map: dict[str, str] = {}
    for rel_id, rel in part.rels.items():
        if rel.is_external:
            id_map[rel_id] = new.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif str(rel.target_part.partname).startswith(_SHARED_PART_PREFIXES):
            id_map[rel_id] = new.relate_to(rel.target_part, rel.reltype)
        else:
            id_map[rel_id] = new.relate_to(_clone_part(rel.target_part, package, clones), rel.reltype)
    if isinstance(new, XmlPart):
        _rewrite_rel_ids(new._element, id_map)
    return new


def _rewrite_rel_ids(root: Any, id_map: dict[str, str]) -> None:
    prefix = f"{{{R_NS}}}"
    for el in root.iter():
        if not isinstance(el.tag, str):
            continue
        for name, value in el.attrib.items():
            if name.startswith(prefix) and value in id_map:
                el.set(name, id_map[value])


def _duplicate_slide(prs: DeckClass, slide: Any) -> Any:
    for shape in _walk(slide.shapes):
        if is_smartart(shape):
            raise PptxError("That slide has SmartArt, which can't be duplicated here. Duplicate it in PowerPoint.")
    new = prs.slides.add_slide(slide.slide_layout)
    tree = new.shapes._spTree
    for child in list(tree):
        if child.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.remove(child)
    package = prs.part.package
    clones: dict[int, Any] = {}
    id_map: dict[str, str] = {}
    for rel_id, rel in slide.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            id_map[rel_id] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif str(rel.target_part.partname).startswith(_SHARED_PART_PREFIXES):
            id_map[rel_id] = new.part.relate_to(rel.target_part, rel.reltype)
        else:
            id_map[rel_id] = new.part.relate_to(_clone_part(rel.target_part, package, clones), rel.reltype)
    source_data = slide._element.find(qn("p:cSld"))
    target_data = new._element.find(qn("p:cSld"))
    for child in source_data.find(qn("p:spTree")):
        if child.tag not in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            tree.append(copy.deepcopy(child))
    background = source_data.find(qn("p:bg"))
    if background is not None:
        target_data.insert(0, copy.deepcopy(background))
    if source_data.get("name"):
        target_data.set("name", source_data.get("name"))
    for child in slide._element:
        if child.tag == qn("p:cSld"):
            continue
        if child.tag == qn("p:extLst"):
            continue  # holds the slide's own creation id
        existing = new._element.find(child.tag)
        if existing is not None:
            new._element.remove(existing)
        new._element.append(copy.deepcopy(child))
    if slide._element.get("show") is not None:
        new._element.set("show", slide._element.get("show"))
    _rewrite_rel_ids(new._element, id_map)
    notes = notes_text(slide)
    if notes:
        frame = new.notes_slide.notes_text_frame
        if frame is not None:
            frame.text = notes
    return new


# ── Editing ──────────────────────────────────────────────────────────

OPS = ("set_text", "replace", "format", "add_slide", "duplicate_slide", "delete_slide", "move_slide", "hide_slide",
       "set_notes", "add_text", "add_image", "add_table", "set_table", "add_chart", "set_chart", "place",
       "delete_shape", "set_alt", "set_theme")
_CHART_TYPES = {
    ("column", False): XL_CHART_TYPE.COLUMN_CLUSTERED, ("column", True): XL_CHART_TYPE.COLUMN_STACKED,
    ("bar", False): XL_CHART_TYPE.BAR_CLUSTERED, ("bar", True): XL_CHART_TYPE.BAR_STACKED,
    ("line", False): XL_CHART_TYPE.LINE, ("line", True): XL_CHART_TYPE.LINE_STACKED,
    ("area", False): XL_CHART_TYPE.AREA, ("area", True): XL_CHART_TYPE.AREA_STACKED,
    ("pie", False): XL_CHART_TYPE.PIE, ("doughnut", False): XL_CHART_TYPE.DOUGHNUT,
    ("scatter", False): XL_CHART_TYPE.XY_SCATTER_LINES_NO_MARKERS,
}
_TABLE_SIZES = (18, 16, 14, 12, 11, 10)


@dataclass(slots=True)
class Issue:
    level: str  # error, warn, info
    kind: str
    where: str
    what: str


@dataclass(slots=True)
class Report:
    counts: dict[str, int] = field(default_factory=dict)
    changed: list[str] = field(default_factory=list)
    fit: list[Issue] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


def _color(value: Any) -> RGBColor:
    text = str(value or "").strip().lstrip("#")
    if re.fullmatch(r"[0-9A-Fa-f]{6}", text):
        return RGBColor.from_string(text.upper())
    raise PptxError(f"{value!r} is not a hex colour (use e.g. C00000).")


def _number(op: dict[str, Any], key: str) -> float | None:
    value = op.get(key)
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        raise PptxError(f"`{key}` must be a number.") from None


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT, "justify": PP_ALIGN.JUSTIFY}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "center": MSO_ANCHOR.MIDDLE, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


class Editor:
    """Applies a list of changes to a deck in memory: all of them, or none."""

    def __init__(self, prs: DeckClass, *, image_root: Callable[[str], Path | None] | None = None) -> None:
        self.prs = prs
        self.image_root = image_root
        self.report = Report()
        self.touched: dict[int, Any] = {}
        self.touched_shapes: set[tuple[int, int]] = set()
        self.created: dict[int, Any] = {}
        self.keep_empty: set[int] = set()
        self.page_numbers = False

    # Addressing

    def _count(self, name: str, number: int = 1) -> None:
        self.report.counts[name] = self.report.counts.get(name, 0) + number

    def _touch(self, slide: Any, shape: Any = None) -> None:
        self.touched[id(slide.part)] = slide
        if shape is not None:
            self.touched_shapes.add((id(slide.part), shape.shape_id))

    def _slide_at(self, number: Any) -> Any:
        count = len(self.prs.slides)
        try:
            index = int(number)
        except (TypeError, ValueError):
            raise PptxError(f"`slide` must be a slide number (1–{count}), got {number!r}.") from None
        if not 1 <= index <= count:
            raise PptxError(f"There is no slide {index}: the deck has {count}.")
        return self.prs.slides[index - 1]

    def _one_slide(self, op: dict[str, Any]) -> Any:
        if op.get("slide") in (None, ""):
            raise PptxError("needs `slide`: its number as read shows it.")
        return self._slide_at(op["slide"])

    def _slide_range(self, op: dict[str, Any]) -> list[Any]:
        raw = op.get("slide")
        if raw in (None, ""):
            raise PptxError('needs `slide`: a number, a range like "2-5", or "all".')
        text = str(raw).strip().lower()
        if text == "all":
            return list(self.prs.slides)
        match = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", text)
        if match:
            first, last = sorted((int(match.group(1)), int(match.group(2))))
            return [self._slide_at(n) for n in range(first, last + 1)]
        return [self._slide_at(text)]

    def _find_shape(self, slide: Any, ref: Any) -> Any | None:
        """A shape by id, by name, or by what it is: title, subtitle, body, or a placeholder number."""
        if ref is None or ref == "":
            return None
        text = str(ref).strip()
        if re.fullmatch(r"\d+", text):
            wanted = int(text)
            for shape in _walk(slide.shapes):
                if shape.shape_id == wanted:
                    return shape
            return None
        low = text.casefold()
        if low.startswith("ph:") or low.startswith("placeholder:"):
            idx = low.split(":", 1)[1].strip()
            for shape in slide.placeholders:
                if str(shape.placeholder_format.idx) == idx:
                    return shape
            return None
        if low in ("title", "subtitle", "body", "content"):
            for shape in slide.placeholders:
                kind = placeholder_kind(shape)
                if kind == low or (low in ("body", "content") and kind in ("body", "content")):
                    return shape
            return None
        for shape in _walk(slide.shapes):
            if shape.name.casefold() == low:
                return shape
        return None

    def _shape(self, slide: Any, ref: Any) -> Any:
        if ref in (None, ""):
            raise PptxError("needs `shape`: its id from read, or title, subtitle, body, ph:<number>.")
        shape = self._find_shape(slide, ref)
        if shape is None:
            known = ", ".join(f"[{s.shape_id}] {placeholder_kind(s) or s.name}" for s in slide.shapes)
            raise PptxError(f"Slide {_position_of(self.prs, slide)} has no shape {ref!r}. It has: {known or 'none'}.")
        return shape

    # Operations

    def apply(self, ops: list[dict[str, Any]]) -> Report:
        for index, op in enumerate(ops, 1):
            if not isinstance(op, dict):
                raise PptxError(f"Edit {index} is not an object with an `op`.")
            name = str(op.get("op") or "").strip().lower().replace("-", "_")
            handler = getattr(self, f"_op_{name}", None) if name in OPS else None
            if handler is None:
                raise PptxError(f"Edit {index}: unknown op {op.get('op')!r}; use one of: {', '.join(OPS)}.")
            try:
                handler(op)
            except PptxError as exc:
                raise PptxError(f"Edit {index} ({name}): {exc}") from None
        return self.finish()

    def _op_set_text(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        shape = self._shape(slide, op.get("shape"))
        if "text" not in op:
            raise PptxError("needs `text`.")
        if not shape.has_text_frame:
            raise PptxError(f"[{shape.shape_id}] holds no text. For a table use set_table, for a chart set_chart.")
        bullets = op.get("bullets")
        self._write(slide, shape, str(op["text"]), bullets=bullets if isinstance(bullets, bool) else None)
        self._count("texts set")

    def _write(self, slide: Any, shape: Any, text: str, *, bullets: bool | None = None) -> None:
        ctx = TextContext(shape, slide, self.prs)
        write_text(shape.text_frame._txBody, slide.part, text, ctx=ctx, bullets=bullets)
        self._touch(slide, shape)

    def _op_replace(self, op: dict[str, Any]) -> None:
        find = str(op.get("find") or "")
        if not find or "replace_with" not in op:
            raise PptxError("needs `find` and `replace_with`.")
        new = str(op["replace_with"])
        total = 0
        for slide in self._slide_range(op):
            hits = 0
            for shape in _walk(slide.shapes):
                bodies = []
                if shape.has_text_frame:
                    bodies.append(shape.text_frame._txBody)
                elif getattr(shape, "has_table", False) and shape.has_table:
                    bodies += [cell._tc.txBody for row in shape.table.rows for cell in row.cells]
                elif getattr(shape, "has_chart", False) and shape.has_chart and shape.chart.has_title \
                        and shape.chart.chart_title.has_text_frame:
                    bodies.append(shape.chart.chart_title.text_frame._txBody)
                for body in bodies:
                    for para in body.findall(A_P):
                        done = replace_in_paragraph(para, find, new)
                        if done:
                            hits += done
                            self._touch(slide, shape)
            if op.get("notes") and slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                for para in slide.notes_slide.notes_text_frame._txBody.findall(A_P):
                    hits += replace_in_paragraph(para, find, new)
                self.touched[id(slide.part)] = slide
            total += hits
        if not total:
            raise PptxError(f"{find!r} isn't in those slides. `read` with find searches the deck.")
        self._count("replacements", total)

    def _op_format(self, op: dict[str, Any]) -> None:
        slides = self._slide_range(op)
        done = 0
        for slide in slides:
            shape = self._find_shape(slide, op.get("shape"))
            if shape is None:
                if len(slides) == 1:
                    self._shape(slide, op.get("shape"))
                continue
            self._format_shape(slide, shape, op)
            done += 1
        if not done:
            raise PptxError(f"No shape {op.get('shape')!r} on those slides.")
        self._count("shapes formatted", done)

    def _format_shape(self, slide: Any, shape: Any, op: dict[str, Any]) -> None:
        if op.get("fill") not in (None, ""):
            if not hasattr(shape, "fill"):
                raise PptxError(f"[{shape.shape_id}] can't take a fill.")
            if str(op["fill"]).lower() == "none":
                shape.fill.background()
            else:
                shape.fill.solid()
                shape.fill.fore_color.rgb = _color(op["fill"])
        if op.get("valign"):
            if not shape.has_text_frame:
                raise PptxError(f"[{shape.shape_id}] holds no text.")
            anchor = _ANCHOR.get(str(op["valign"]).lower())
            if anchor is None:
                raise PptxError(f"valign {op['valign']!r}: use top, center or bottom.")
            shape.text_frame.vertical_anchor = anchor
        paragraphs = self._paragraphs_for(shape, op)
        for para in paragraphs:
            self._format_paragraph(para, op)
        self._touch(slide, shape)

    def _paragraphs_for(self, shape: Any, op: dict[str, Any]) -> list[Any]:
        bodies: list[Any] = []
        if getattr(shape, "has_table", False) and shape.has_table:
            table = shape.table
            row, column = op.get("row"), op.get("column")
            for r in range(len(table.rows)):
                for c in range(len(table.columns)):
                    if (row in (None, "") or int(row) == r + 1) and (column in (None, "") or int(column) == c + 1):
                        bodies.append(table.cell(r, c)._tc.txBody)
        elif shape.has_text_frame:
            bodies.append(shape.text_frame._txBody)
        elif any(key in op for key in ("bold", "italic", "underline", "color", "font", "font_size", "align")):
            raise PptxError(f"[{shape.shape_id}] holds no text to format.")
        out: list[Any] = []
        for body in bodies:
            out += body.findall(A_P)
        if op.get("paragraph") not in (None, ""):
            index = int(op["paragraph"])
            if not 1 <= index <= len(out):
                raise PptxError(f"[{shape.shape_id}] has {len(out)} paragraph(s), no {index}.")
            out = [out[index - 1]]
        return out

    def _format_paragraph(self, para: Any, op: dict[str, Any]) -> None:
        if op.get("align"):
            alignment = _ALIGN.get(str(op["align"]).lower())
            if alignment is None:
                raise PptxError(f"align {op['align']!r}: use left, center or right.")
            para.get_or_add_pPr().algn = alignment
        if op.get("find"):
            runs = _isolate(para, str(op["find"]))
        else:
            runs = [child for child in para if child.tag in (A_R, A_FLD)]
        for run in runs:
            props = run.find(A_RPR)
            if props is None:
                props = run.get_or_add_rPr() if hasattr(run, "get_or_add_rPr") else None
                if props is None:
                    props = etree.Element(A_RPR)
                    run.insert(0, props)
            self._format_run(props, op)

    def _format_run(self, props: Any, op: dict[str, Any]) -> None:
        for key, attribute in (("bold", "b"), ("italic", "i")):
            if key in op and op[key] is not None:
                props.set(attribute, "1" if op[key] else "0")
        if op.get("underline") is not None:
            props.set("u", "sng" if op["underline"] else "none")
        if op.get("font_size") not in (None, ""):
            size = float(op["font_size"])
            if not 1 <= size <= 400:
                raise PptxError("font_size is in points (1–400).")
            props.set("sz", str(int(size * 100)))
        if op.get("color"):
            _color(op["color"])
            for tag in ("a:solidFill", "a:noFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill"):
                for old in props.findall(qn(tag)):
                    props.remove(old)
            fill = etree.Element(qn("a:solidFill"))
            etree.SubElement(fill, qn("a:srgbClr"), val=str(op["color"]).strip().lstrip("#").upper())
            line = props.find(qn("a:ln"))
            props.insert(0 if line is None else 1, fill)
        if op.get("font"):
            for tag in ("a:latin", "a:ea", "a:cs"):
                for old in props.findall(qn(tag)):
                    props.remove(old)
            latin = etree.Element(qn("a:latin"), typeface=str(op["font"]))
            position = 0
            for index, child in enumerate(props):
                if child.tag in (qn("a:ln"), qn("a:solidFill"), qn("a:noFill"), qn("a:gradFill"), qn("a:blipFill"),
                                 qn("a:pattFill"), qn("a:grpFill"), qn("a:effectLst"), qn("a:effectDag"),
                                 qn("a:highlight"), qn("a:uLnTx"), qn("a:uLn"), qn("a:uFillTx"), qn("a:uFill")):
                    position = index + 1
            props.insert(position, latin)

    # Slides

    def _op_add_slide(self, op: dict[str, Any]) -> None:
        self.add_slide(op)
        self._count("slides added")

    def add_slide(self, spec: dict[str, Any]) -> Any:
        prs = self.prs
        layout = find_layout(prs, spec["layout"]) if spec.get("layout") not in (None, "") \
            else _default_layout(prs, spec)
        slide = prs.slides.add_slide(layout)
        element = _slide_id_of(prs, slide)
        count = len(prs.slides._sldIdLst)
        position = spec.get("at")
        if position not in (None, ""):
            _move_to(prs, slide, int(position))
        elif _sections(prs):
            members = list(prs.slides._sldIdLst)
            _section_join(prs, element.get("id"), members[-2].get("id") if count > 1 else None)
        self.created[id(slide.part)] = slide
        if spec.get("keep_empty"):
            self.keep_empty.add(id(slide.part))
        self._copy_footers(slide)
        self._fill_slide(slide, spec)
        self._touch(slide)
        return slide

    def _fill_slide(self, slide: Any, spec: dict[str, Any]) -> None:
        for key in ("title", "subtitle"):
            if spec.get(key) in (None, ""):
                continue
            shape = self._find_shape(slide, key)
            if shape is None and key == "subtitle":
                shape = self._find_shape(slide, "body")
            if shape is None:
                raise PptxError(f'Layout "{slide.slide_layout.name}" has no {key}: {_layout_description(slide.slide_layout)}.')
            self._write(slide, shape, str(spec[key]))
        if spec.get("body") not in (None, ""):
            shape = self._empty_placeholder(slide, ("content", "body"))
            if shape is None:
                raise PptxError(f'Layout "{slide.slide_layout.name}" has nowhere for body text: '
                                f"{_layout_description(slide.slide_layout)}.")
            body = spec["body"]
            text = "\n".join(str(item) for item in body) if isinstance(body, list) else str(body)
            self._write(slide, shape, text)
        placeholders = spec.get("placeholders")
        if isinstance(placeholders, dict):
            for ref, value in placeholders.items():
                shape = self._find_shape(slide, f"ph:{ref}" if str(ref).isdigit() else ref)
                if shape is None or not shape.has_text_frame:
                    raise PptxError(f'Layout "{slide.slide_layout.name}" has no text placeholder {ref!r}: '
                                    f"{_layout_description(slide.slide_layout)}.")
                self._write(slide, shape, str(value))
        if spec.get("table"):
            table = spec["table"]
            self.add_table(slide, table if isinstance(table, dict) else {"rows": table})
        if spec.get("chart"):
            self.add_chart(slide, spec["chart"])
        if spec.get("image"):
            image = spec["image"]
            self.add_image(slide, image if isinstance(image, dict) else {"image": image})
        if spec.get("notes") not in (None, ""):
            self._set_notes(slide, str(spec["notes"]))

    def _copy_footers(self, slide: Any) -> None:
        """New slides carry the deck's slide numbers and footer: python-pptx leaves
        them off, and a deck with them on every slide but the new one looks wrong."""
        wanted: dict[Any, str] = {}
        for other in self.prs.slides:
            if other.part is slide.part:
                continue
            for ph in other.placeholders:
                kind = ph.placeholder_format.type
                if kind in _FOOTER_TYPES:
                    wanted.setdefault(kind, _shape_text(ph))
        if self.page_numbers:
            wanted.setdefault(PP_PLACEHOLDER.SLIDE_NUMBER, "")
        if not wanted:
            return
        tree = slide.shapes._spTree
        for ph in slide.slide_layout.placeholders:
            kind = ph.placeholder_format.type
            if kind not in wanted:
                continue
            new = copy.deepcopy(ph._element)
            props = new.find(qn("p:spPr"))
            for child in list(props):
                props.remove(child)
            new.xpath(".//p:cNvPr")[0].set("id", str(slide.shapes._next_shape_id))
            tree.append(new)
            if kind == PP_PLACEHOLDER.FOOTER:
                body = new.find(qn("p:txBody"))
                write_text(body, slide.part, wanted[kind])
            elif kind == PP_PLACEHOLDER.DATE and wanted[kind]:
                body = new.find(qn("p:txBody"))
                write_text(body, slide.part, wanted[kind])

    def _op_duplicate_slide(self, op: dict[str, Any]) -> None:
        source = self._one_slide(op)
        position = _position_of(self.prs, source)
        new = _duplicate_slide(self.prs, source)
        at = int(op["at"]) if op.get("at") not in (None, "") else position + 1
        _move_to(self.prs, new, at)
        self._touch(new)
        self._count("slides duplicated")

    def _op_delete_slide(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        if len(self.prs.slides) == 1:
            raise PptxError("A deck needs at least one slide.")
        broken = _delete_slide(self.prs, slide)
        if broken:
            self.report.notes.append(f"{broken} link(s) from other slides pointed at that slide and were removed.")
        self.touched.pop(id(slide.part), None)
        self.created.pop(id(slide.part), None)
        self._count("slides deleted")

    def _op_move_slide(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        if op.get("at") in (None, ""):
            raise PptxError("needs `at`: where it ends up (1 = first).")
        _move_to(self.prs, slide, int(op["at"]))
        self._count("slides moved")

    def _op_hide_slide(self, op: dict[str, Any]) -> None:
        hidden = op.get("hidden", True)
        for slide in self._slide_range(op):
            if hidden:
                slide._element.set("show", "0")
            else:
                slide._element.attrib.pop("show", None)
            self._count("slides hidden" if hidden else "slides shown")

    def _op_set_notes(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        if "text" not in op:
            raise PptxError("needs `text`.")
        self._set_notes(slide, str(op["text"]))
        self._touch(slide)
        self._count("notes set")

    def _set_notes(self, slide: Any, text: str) -> None:
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            raise PptxError("This deck's notes page has nowhere for text.")
        frame.text = text

    # Boxes

    def _box(self, slide: Any, op: dict[str, Any], *, default_height: float | None = None) -> tuple[int, int, int, int] | None:
        """(x, y, width, height) in EMU when the op says where; every one, or none."""
        given = {key: _number(op, key) for key in ("x", "y", "width", "height")}
        if all(value is None for value in given.values()):
            return None
        if given["height"] is None:
            given["height"] = default_height
        missing = [key for key, value in given.items() if value is None]
        if missing:
            raise PptxError(f"say x, y, width and height in inches, or none of them (missing {', '.join(missing)}).")
        if given["width"] <= 0 or given["height"] <= 0:
            raise PptxError("width and height must be above zero.")
        return (Inches(given["x"]), Inches(given["y"]), Inches(given["width"]), Inches(given["height"]))

    def _free_area(self, slide: Any) -> tuple[int, int, int, int]:
        """The open part of a slide: under whatever's already there."""
        prs = self.prs
        margin = Inches(0.5)
        left, width = margin, prs.slide_width - 2 * margin
        top = margin
        title = self._find_shape(slide, "title")
        if title is not None:
            left = max(title.left, margin)
            width = min(title.width, prs.slide_width - 2 * left)
            top = title.top + title.height
        for shape in slide.shapes:
            if shape.is_placeholder and placeholder_kind(shape) in ("title", "date", "footer", "slide number"):
                continue
            if shape.is_placeholder and shape.has_text_frame and not _shape_text(shape):
                continue
            top = max(top, (shape.top or 0) + (shape.height or 0))
        top += Inches(0.1)
        bottom = prs.slide_height - Inches(0.45)
        for ph in slide.placeholders:
            if ph.placeholder_format.type in _FOOTER_TYPES:
                bottom = min(bottom, ph.top)
        if bottom - top < Inches(1.0):
            top = max(bottom - Inches(1.0), 0)
        return left, top, width, bottom - top

    def _empty_placeholder(self, slide: Any, kinds: tuple[str, ...]) -> Any | None:
        for ph in slide.placeholders:
            if placeholder_kind(ph) in kinds and ph.has_text_frame and not _shape_text(ph) \
                    and (id(slide.part), ph.shape_id) not in self.touched_shapes:
                return ph
        return None

    def _place_visual(self, slide: Any, op: dict[str, Any], kinds: tuple[str, ...]) -> tuple[tuple[int, int, int, int], int | None]:
        """Where a table, chart or picture goes: where it was told, or the empty
        placeholder that suits it (which it replaces), or the open area."""
        box = self._box(slide, op)
        if box is not None:
            return box, None
        ph = self._empty_placeholder(slide, kinds)
        if ph is not None:
            box = (ph.left, ph.top, ph.width, ph.height)
            idx = ph.placeholder_format.idx
            ph._element.getparent().remove(ph._element)
            return box, idx
        return self._free_area(slide), None

    @staticmethod
    def _link_placeholder(shape: Any, idx: int | None) -> None:
        """What PowerPoint writes when content sits in a placeholder, so a change of
        layout carries it along."""
        if idx is None:
            return
        props = shape._element.xpath(".//p:nvPr")[0]
        ph = etree.Element(qn("p:ph"), idx=str(idx))
        props.insert(0, ph)

    @staticmethod
    def _set_alt(shape: Any, text: str) -> None:
        shape._element.xpath(".//p:cNvPr")[0].set("descr", text)

    def _op_add_text(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        if "text" not in op:
            raise PptxError("needs `text`.")
        # A text box is as tall as its text: the height is worked out below.
        box = self._box(slide, op, default_height=0.4) or self._free_area(slide)
        shape = slide.shapes.add_textbox(box[0], box[1], box[2], box[3])
        frame = shape.text_frame
        frame.word_wrap = True
        frame.auto_size = MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT
        self._write(slide, shape, str(op["text"]))
        if any(key in op for key in ("bold", "italic", "underline", "color", "font", "font_size", "align", "fill", "valign")):
            self._format_shape(slide, shape, op)
        ctx = TextContext(shape, slide, self.prs)
        shape.height = Emu(int(measure_text(ctx, shape.width).height * EMU_PER_PT))
        self._count("text boxes added")

    def _op_add_image(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        self.add_image(slide, op)
        self._count("images added")

    def add_image(self, slide: Any, op: dict[str, Any]) -> Any:
        src = str(op.get("image") or "").strip()
        if not src:
            raise PptxError("needs `image`: a saved picture's path.")
        found = self.image_root(src) if self.image_root is not None else None
        if found is None:
            raise PptxError(f"Can't read the picture {src!r}. Use a path from this workspace.")
        box, idx = self._place_visual(slide, op, ("picture", "content", "body"))
        try:
            shape = slide.shapes.add_picture(str(found), box[0], box[1])
        except Exception as exc:
            raise PptxError(f"{Path(src).name} isn't a picture PowerPoint can show ({type(exc).__name__}).") from None
        px_w, px_h = shape.image.size
        ratio = px_w / px_h if px_h else 1.0
        if str(op.get("fit") or "contain").lower() == "cover":
            frame_ratio = box[2] / box[3]
            if ratio > frame_ratio:
                crop = (1 - frame_ratio / ratio) / 2
                shape.crop_left = shape.crop_right = crop
            else:
                crop = (1 - ratio / frame_ratio) / 2
                shape.crop_top = shape.crop_bottom = crop
            shape.left, shape.top, shape.width, shape.height = box
        else:
            width = box[2]
            height = int(width / ratio)
            if height > box[3]:
                height = box[3]
                width = int(height * ratio)
            shape.left = Emu(box[0] + (box[2] - width) // 2)
            shape.top = Emu(box[1] + (box[3] - height) // 2)
            shape.width, shape.height = Emu(width), Emu(height)
        self._link_placeholder(shape, idx)
        # python-pptx files the picture's name as its alt text, which helps nobody.
        self._set_alt(shape, str(op.get("alt") or ""))
        self._touch(slide, shape)
        return shape

    def _op_add_table(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        self.add_table(slide, op)
        self._count("tables added")

    def add_table(self, slide: Any, op: dict[str, Any]) -> Any:
        rows = op.get("rows")
        if not isinstance(rows, list) or not rows:
            raise PptxError("needs `rows`: a list of rows, each a list of values.")
        rows = [row if isinstance(row, list) else [row] for row in rows]
        width = max(len(row) for row in rows)
        if not width:
            raise PptxError("`rows` has no values.")
        text = [["" if v is None else str(v) for v in row] + [""] * (width - len(row)) for row in rows]
        box, idx = self._place_visual(slide, op, ("table", "content", "body"))
        frame = slide.shapes.add_table(len(text), width, box[0], box[1], box[2], box[3])
        table = frame.table
        table.first_row = op.get("header", True) is not False
        columns = _column_widths(text, box[2])
        for index, value in enumerate(columns):
            table.columns[index].width = Emu(value)
        size = float(op["font_size"]) if op.get("font_size") not in (None, "") else None
        chosen = size or _table_font_size(text, columns, box[3], self._table_font(slide))
        heights = _row_heights(text, columns, chosen, self._table_font(slide))
        for r, row in enumerate(text):
            table.rows[r].height = Emu(heights[r])
            for c, value in enumerate(row):
                cell = table.cell(r, c)
                write_text(cell._tc.txBody, slide.part, value)
                self._size_cell(cell, chosen)
                if r and _looks_numeric(value):
                    for para in cell._tc.txBody.findall(A_P):
                        para.get_or_add_pPr().algn = PP_ALIGN.RIGHT
        frame.height = Emu(sum(heights))
        self._link_placeholder(frame, idx)
        if op.get("alt"):
            self._set_alt(frame, str(op["alt"]))
        if sum(heights) > box[3]:
            self.report.notes.append(f"Slide {_position_of(self.prs, slide)}: the table is taller than its space even at "
                                     f"{chosen:g}pt. Split it across slides or cut rows.")
        self._touch(slide, frame)
        return frame

    def _table_font(self, slide: Any) -> str:
        return _theme_fonts(slide.slide_layout.slide_master)[1]

    @staticmethod
    def _size_cell(cell: Any, size: float) -> None:
        for para in cell._tc.txBody.findall(A_P):
            for run in para.findall(A_R):
                run.get_or_add_rPr().set("sz", str(int(size * 100)))
            end = para.find(A_END)
            if end is None:
                end = etree.SubElement(para, A_END)
            end.set("sz", str(int(size * 100)))

    def _op_set_table(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        shape = self._shape(slide, op.get("shape"))
        if not (getattr(shape, "has_table", False) and shape.has_table):
            raise PptxError(f"[{shape.shape_id}] isn't a table.")
        table = shape.table
        if isinstance(op.get("rows"), list):
            rows = [row if isinstance(row, list) else [row] for row in op["rows"]]
            width = sum(int(column.width) for column in table.columns)
            _resize_table(table, len(rows), max(len(row) for row in rows))
            text: list[list[str]] = []
            for r, row in enumerate(rows):
                cells = ["" if v is None else str(v) for v in row]
                text.append(cells + [""] * (len(table.columns) - len(cells)))
                for c in range(len(table.columns)):
                    self._set_cell(slide, table, r, c, text[r][c])
            # The same width as before, shared out afresh: a new column would
            # otherwise push the table off the slide.
            columns = _column_widths(text, width)
            for index, value in enumerate(columns):
                table.columns[index].width = Emu(value)
            heights = _row_heights(text, columns, _table_text_size(table), self._table_font(slide))
            for r, height in enumerate(heights):
                table.rows[r].height = Emu(height)
            shape.width, shape.height = Emu(width), Emu(sum(heights))
        elif isinstance(op.get("cells"), dict):
            for ref, value in op["cells"].items():
                match = re.fullmatch(r"r(\d+)c(\d+)", str(ref).strip().lower())
                if match is None:
                    raise PptxError(f'{ref!r}: name a cell like "r2c3" (row 2, column 3).')
                r, c = int(match.group(1)) - 1, int(match.group(2)) - 1
                if not (0 <= r < len(table.rows) and 0 <= c < len(table.columns)):
                    raise PptxError(f"{ref} is outside the table ({len(table.rows)} × {len(table.columns)}).")
                self._set_cell(slide, table, r, c, value)
        else:
            raise PptxError('needs `rows` (the whole table) or `cells` ({"r2c3": "new"}).')
        self._touch(slide, shape)
        self._count("tables set")

    @staticmethod
    def _set_cell(slide: Any, table: Any, r: int, c: int, value: Any) -> None:
        cell = table.cell(r, c)
        write_text(cell._tc.txBody, slide.part, "" if value is None else str(value))

    def _op_add_chart(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        self.add_chart(slide, op)
        self._count("charts added")

    def add_chart(self, slide: Any, op: dict[str, Any]) -> Any:
        kind = str(op.get("type") or "column").lower()
        stacked = bool(op.get("stacked"))
        chart_type = _CHART_TYPES.get((kind, stacked)) or _CHART_TYPES.get((kind, False))
        if chart_type is None:
            raise PptxError(f"chart type {kind!r}: use column, bar, line, area, pie, doughnut or scatter.")
        data = _chart_data(op, kind)
        box, idx = self._place_visual(slide, op, ("chart", "content", "body"))
        frame = slide.shapes.add_chart(chart_type, Emu(box[0]), Emu(box[1]), Emu(box[2]), Emu(box[3]), data)
        chart = frame.chart
        chart.font.size = Pt(float(op.get("font_size") or 12))
        title = str(op.get("title") or "")
        chart.has_title = bool(title)
        if title:
            chart.chart_title.text_frame.text = title
        series_count = len(op.get("series") or [])
        wants_legend = op.get("legend")
        if wants_legend is None:
            wants_legend = series_count > 1 or kind in ("pie", "doughnut")
        chart.has_legend = bool(wants_legend)
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
        if kind not in ("pie", "doughnut"):
            axis = chart.value_axis
            axis.has_major_gridlines = True
            axis.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
            axis.format.line.fill.background()
        if op.get("labels"):
            plot = chart.plots[0]
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.show_value = kind not in ("pie", "doughnut")
            labels.show_percentage = kind in ("pie", "doughnut")
            if kind == "column" and not stacked:
                labels.position = XL_LABEL_POSITION.OUTSIDE_END
        self._link_placeholder(frame, idx)
        if op.get("alt"):
            self._set_alt(frame, str(op["alt"]))
        self._touch(slide, frame)
        return frame

    def _op_set_chart(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        shape = self._shape(slide, op.get("shape"))
        if not (getattr(shape, "has_chart", False) and shape.has_chart):
            raise PptxError(f"[{shape.shape_id}] isn't a chart.")
        chart = shape.chart
        kind = "scatter" if isinstance(chart.plots[0], XyPlot) else "column"
        if op.get("series"):
            data = _chart_data(op, kind)
            try:
                chart.replace_data(data)
            except Exception as exc:
                raise PptxError(f"Its data couldn't be replaced ({type(exc).__name__}: {exc}).") from None
        if "title" in op:
            chart.has_title = bool(op["title"])
            if op["title"]:
                chart.chart_title.text_frame.text = str(op["title"])
        self._touch(slide, shape)
        self._count("charts set")

    def _op_place(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        shape = self._shape(slide, op.get("shape"))
        for key, attribute in (("x", "left"), ("y", "top"), ("width", "width"), ("height", "height")):
            value = _number(op, key)
            if value is not None:
                # A placeholder takes all four, or it would keep nothing from its layout.
                current = [shape.left, shape.top, shape.width, shape.height]
                if shape._element.xpath(".//a:xfrm") == [] and shape.is_placeholder:
                    shape.left, shape.top, shape.width, shape.height = current
                setattr(shape, attribute, Inches(value))
        depth = str(op.get("z") or "").lower()
        if depth:
            tree = shape._element.getparent()
            tree.remove(shape._element)
            if depth == "front":
                tree.append(shape._element)
            elif depth == "back":
                # After the group's own two header elements.
                tree.insert(2, shape._element)
            else:
                raise PptxError(f"z {op['z']!r}: use front or back.")
        self._touch(slide, shape)
        self._count("shapes placed")

    def _op_delete_shape(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        shape = self._shape(slide, op.get("shape"))
        shape_id = shape.shape_id
        rel_ids = [v for el in shape._element.iter() if isinstance(el.tag, str)
                   for k, v in el.attrib.items() if k.startswith(f"{{{R_NS}}}")]
        shape._element.getparent().remove(shape._element)
        for rel_id in set(rel_ids):
            if rel_id in slide.part.rels:
                slide.part.drop_rel(rel_id)
        if _drop_animations_of(slide, shape_id):
            self.report.notes.append(f"Slide {_position_of(self.prs, slide)}: animations on that shape went with it.")
        self._touch(slide)
        self._count("shapes deleted")

    def _op_set_alt(self, op: dict[str, Any]) -> None:
        slide = self._one_slide(op)
        shape = self._shape(slide, op.get("shape"))
        if "text" not in op:
            raise PptxError("needs `text`.")
        self._set_alt(shape, str(op["text"]))
        self._count("alt texts set")

    def _op_set_theme(self, op: dict[str, Any]) -> None:
        apply_look(self.prs, font=op.get("font"), heading_font=op.get("heading_font"),
                   heading_color=op.get("heading_color"), accent_color=op.get("accent_color"))
        for slide in self.prs.slides:
            self.touched[id(slide.part)] = slide
        self._count("theme changes")

    # Finishing

    def finish(self) -> Report:
        prs = self.prs
        for key, slide in list(self.created.items()):
            if key in self.keep_empty:
                continue
            for ph in list(slide.placeholders):
                if placeholder_kind(ph) in ("date", "footer", "slide number"):
                    continue
                if ph.has_text_frame and not _shape_text(ph) and (key, ph.shape_id) not in self.touched_shapes:
                    # An empty placeholder reads "Click to add text" to everyone who opens the file.
                    ph._element.getparent().remove(ph._element)
        shrunk: list[str] = []
        for key, slide in self.touched.items():
            if not any(el.rId and prs.part.related_part(el.rId) is slide.part for el in prs.slides._sldIdLst):
                continue
            number = _position_of(prs, slide)
            for shape in list(slide.shapes):
                if (key, shape.shape_id) not in self.touched_shapes or not shape.has_text_frame:
                    continue
                scale = apply_autofit(shape, slide, prs)
                if scale is not None:
                    shrunk.append(f"Slide {number} [{shape.shape_id}]: text shrunk to {scale * 100:g}% to fit, "
                                  "as PowerPoint does.")
        ordered = sorted(self.touched.values(), key=lambda s: _position_of(prs, s)) if self.touched else []
        for slide in ordered:
            number = _position_of(prs, slide)
            self.report.changed += describe_slide(prs, slide, number)
            self.report.fit += fit_issues(prs, slide, number)
        self.report.notes += shrunk
        if self.touched or self.report.counts:
            _drop_thumbnail(prs)
        return self.report


def _drop_thumbnail(prs: DeckClass) -> None:
    """The stored picture of the first slide would show the old deck in file browsers."""
    package = prs.part.package
    for rel_id, rel in list(package._rels.items()):
        if rel.reltype == RT.THUMBNAIL:
            package._rels.pop(rel_id)


def _drop_animations_of(slide: Any, shape_id: int) -> bool:
    """Animation that targets a shape no longer there makes PowerPoint repair the file."""
    timing = slide._element.find(qn("p:timing"))
    if timing is None:
        return False
    found = False
    for target in list(timing.iter(qn("p:spTgt"))):
        if target.get("spid") != str(shape_id):
            continue
        found = True
        # The whole click step it belongs to: a step that names no shape is broken too.
        step = target
        while step.getparent() is not None:
            parent = step.getparent()
            if parent.tag == qn("p:childTnLst") and parent.getparent().get("nodeType") == "mainSeq":
                parent.remove(step)
                break
            step = parent
    for build in list(timing.iter(qn("p:bldP"), qn("p:bldGraphic"), qn("p:bldDgm"), qn("p:bldOleChart"))):
        if build.get("spid") == str(shape_id):
            build.getparent().remove(build)
    if found and not list(timing.iter(qn("p:spTgt"))):
        slide._element.remove(timing)
    return found


# ── Tables and charts ────────────────────────────────────────────────


def _looks_numeric(text: str) -> bool:
    return bool(re.fullmatch(r"\s*[-+(]?[$€£]?\s*\d[\d.,]*\s*[%)kKmMbB]?\s*", text))


def _column_widths(rows: list[list[str]], total: int) -> list[int]:
    """Column widths in EMU, in step with what they hold."""
    columns = len(rows[0])
    weights = []
    for c in range(columns):
        longest = max((len(row[c]) for row in rows), default=1)
        weights.append(min(max(longest, 5), 40) ** 0.7)
    scale = total / sum(weights)
    widths = [int(w * scale) for w in weights]
    widths[-1] += total - sum(widths)
    return widths


_CELL_MARGIN_X = 91440 * 2
_CELL_MARGIN_Y = 45720 * 2


def _row_heights(rows: list[list[str]], columns: list[int], size: float, font: str) -> list[int]:
    line = size * LINE_HEIGHT * EMU_PER_PT
    heights = []
    for row in rows:
        lines = 1
        for c, value in enumerate(row):
            room = (columns[c] - _CELL_MARGIN_X) / EMU_PER_PT
            count = sum(_wrap_count([(part, size, font, False)], room, room) for part in value.split("\n"))
            lines = max(lines, count)
        heights.append(int(lines * line + _CELL_MARGIN_Y))
    return heights


def _table_text_size(table: Any) -> float:
    for run in table._tbl.iter(A_R):
        props = run.find(A_RPR)
        if props is not None and props.get("sz"):
            return int(props.get("sz")) / 100
    return 18.0


def _table_font_size(rows: list[list[str]], columns: list[int], room: int, font: str) -> float:
    for size in _TABLE_SIZES:
        if sum(_row_heights(rows, columns, size, font)) <= room:
            return float(size)
    return float(_TABLE_SIZES[-1])


def _resize_table(table: Any, rows: int, columns: int) -> None:
    """Rows and columns added by copying the last, so they keep their formatting."""
    element = table._tbl
    grid = element.find(qn("a:tblGrid"))
    while len(grid) < columns:
        last = grid[-1]
        new = copy.deepcopy(last)
        for ext in new.findall(qn("a:extLst")):
            new.remove(ext)
        grid.append(new)
        for tr in element.findall(qn("a:tr")):
            cells = tr.findall(qn("a:tc"))
            cell = copy.deepcopy(cells[-1])
            for attr in ("gridSpan", "hMerge", "rowSpan", "vMerge"):
                cell.attrib.pop(attr, None)
            cells[-1].addnext(cell)
    while len(grid) > columns:
        grid.remove(grid[-1])
        for tr in element.findall(qn("a:tr")):
            tr.remove(tr.findall(qn("a:tc"))[-1])
    rows_now = element.findall(qn("a:tr"))
    while len(rows_now) < rows:
        new = copy.deepcopy(rows_now[-1])
        for ext in new.findall(qn("a:extLst")):
            new.remove(ext)
        rows_now[-1].addnext(new)
        rows_now = element.findall(qn("a:tr"))
    while len(rows_now) > rows:
        element.remove(rows_now[-1])
        rows_now = element.findall(qn("a:tr"))


def _numeric_values(values: Any, where: str) -> list[float | None]:
    if not isinstance(values, list) or not values:
        raise PptxError(f"{where}: needs `values`, a list of numbers.")
    out: list[float | None] = []
    for value in values:
        if value in (None, ""):
            out.append(None)
            continue
        try:
            out.append(float(str(value).replace(",", "")) if isinstance(value, str) else float(value))
        except (TypeError, ValueError):
            raise PptxError(f"{where}: {value!r} isn't a number.") from None
    return out


def _chart_data(op: dict[str, Any], kind: str) -> Any:
    series = op.get("series")
    if not isinstance(series, list) or not series:
        raise PptxError("needs `series`: a list of {name, values}.")
    number_format = str(op.get("number_format") or "General")
    if kind == "scatter":
        xy = XyChartData(number_format=number_format)
        for index, item in enumerate(series, 1):
            xs = _numeric_values(item.get("x"), f"series {index} x")
            ys = _numeric_values(item.get("y", item.get("values")), f"series {index} y")
            if len(xs) != len(ys):
                raise PptxError(f"Series {index} has {len(xs)} x values and {len(ys)} y values.")
            target = xy.add_series(str(item.get("name") or f"Series {index}"))
            for x, y in zip(xs, ys):
                target.add_data_point(x, y)
        return xy
    categories = op.get("categories")
    if not isinstance(categories, list) or not categories:
        raise PptxError("needs `categories`: the labels along the axis.")
    data = CategoryChartData(number_format=number_format)
    data.categories = [str(c) for c in categories]
    for index, item in enumerate(series, 1):
        if not isinstance(item, dict):
            raise PptxError(f"Series {index} isn't an object with name and values.")
        values = _numeric_values(item.get("values"), f"series {index}")
        if len(values) != len(categories):
            raise PptxError(f"Series {index} has {len(values)} values for {len(categories)} categories.")
        data.add_series(str(item.get("name") or f"Series {index}"), values)
    return data


# ── Look ─────────────────────────────────────────────────────────────


def apply_look(prs: DeckClass, *, font: Any = None, heading_font: Any = None, heading_color: Any = None,
               accent_color: Any = None) -> None:
    """Change the theme, so every layout and every slide that follows it changes too."""
    if heading_color:
        _color(heading_color)
    if accent_color:
        _color(accent_color)
    for master in prs.slide_masters:
        part = master.part.part_related_by(RT.THEME)
        root = etree.fromstring(part.blob)
        for tag, value in (("a:majorFont", heading_font), ("a:minorFont", font)):
            if value:
                latin = root.find(f".//{qn(tag)}/{qn('a:latin')}")
                if latin is not None:
                    latin.set("typeface", str(value))
        if accent_color:
            accent = root.find(f".//{qn('a:accent1')}")
            if accent is not None:
                for child in list(accent):
                    accent.remove(child)
                etree.SubElement(accent, qn("a:srgbClr"), val=str(accent_color).strip().lstrip("#").upper())
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        _THEME_FONTS.pop(id(master.part), None)
        if heading_color:
            style = master._element.find(qn("p:txStyles")).find(qn("p:titleStyle"))
            for defaults in style.iter(qn("a:defRPr")):
                for tag in ("a:solidFill", "a:schemeClr"):
                    for old in defaults.findall(qn(tag)):
                        defaults.remove(old)
                fill = etree.Element(qn("a:solidFill"))
                etree.SubElement(fill, qn("a:srgbClr"), val=str(heading_color).strip().lstrip("#").upper())
                defaults.insert(0, fill)


def _widen_blank_deck(prs: DeckClass) -> None:
    """python-pptx's blank deck is 4:3 with oversized type; a new deck gets 16:9,
    with the layouts spread to match and sizes a page of content can use."""
    old_width = prs.slide_width
    new_width = Inches(13.333)
    ratio = new_width / old_width
    prs.slide_width = new_width
    prs.slide_height = Inches(7.5)
    height_ratio = prs.slide_height / 6858000
    for master in prs.slide_masters:
        for owner in [master, *master.slide_layouts]:
            for xfrm in owner._element.iter(qn("a:xfrm")):
                off, ext = xfrm.find(qn("a:off")), xfrm.find(qn("a:ext"))
                if off is not None:
                    off.set("x", str(int(int(off.get("x")) * ratio)))
                    off.set("y", str(int(int(off.get("y")) * height_ratio)))
                if ext is not None:
                    ext.set("cx", str(int(int(ext.get("cx")) * ratio)))
                    ext.set("cy", str(int(int(ext.get("cy")) * height_ratio)))
        for layout in master.slide_layouts:
            for defaults in layout._element.iter(qn("a:defRPr")):
                if defaults.get("sz"):
                    defaults.set("sz", str(round(int(defaults.get("sz")) * 0.75 / 100) * 100))
                defaults.attrib.pop("cap", None)
            for props in layout._element.iter(*(qn(f"a:lvl{n}pPr") for n in range(1, 10))):
                if props.get("algn") == "ctr":
                    props.set("algn", "l")
        styles = master._element.find(qn("p:txStyles"))
        title = styles.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
        title.set("algn", "l")
        title.find(qn("a:defRPr")).set("sz", "3400")
        for level, size in (("a:lvl1pPr", "2400"), ("a:lvl2pPr", "2000"), ("a:lvl3pPr", "1800"),
                            ("a:lvl4pPr", "1600"), ("a:lvl5pPr", "1600")):
            props = styles.find(qn("p:bodyStyle")).find(qn(level))
            if props is not None and props.find(qn("a:defRPr")) is not None:
                props.find(qn("a:defRPr")).set("sz", size)


def build_deck(slides: list[dict[str, Any]] | None, *, edits: list[dict[str, Any]] | None = None,
               template: bytes | None = None, keep_slides: bool = False, title: str = "", author: str = "Captain",
               page_numbers: bool = False, look: dict[str, Any] | None = None,
               image_root: Callable[[str], Path | None] | None = None) -> tuple[DeckClass, Report]:
    """A new deck: the template's layouts (or a plain 16:9 set) filled from ``slides``."""
    if template is not None:
        prs = open_deck(template)
        if slides and not keep_slides:
            for slide in list(prs.slides):
                _delete_slide(prs, slide)
    else:
        prs = Presentation()
        _widen_blank_deck(prs)
    if look:
        apply_look(prs, **look)
    specs = list(slides or [])
    if template is None and not specs:
        raise PptxError("create needs `slides`: a list of {layout, title, body, …}.")
    editor = Editor(prs, image_root=image_root)
    editor.page_numbers = page_numbers
    for index, spec in enumerate(specs, 1):
        if not isinstance(spec, dict):
            raise PptxError(f"Slide {index} is not an object with `title`, `body` and so on.")
        try:
            editor.add_slide({k: v for k, v in spec.items() if k != "at"})
        except PptxError as exc:
            raise PptxError(f"Slide {index}: {exc}") from None
    core = prs.core_properties
    core.author = author
    if title:
        core.title = title
    report = editor.apply(list(edits or []))
    _drop_thumbnail(prs)
    return prs, report


# ── Checking a deck ──────────────────────────────────────────────────


def _box_of(shape: Any) -> tuple[int, int, int, int] | None:
    try:
        left, top, width, height = shape.left, shape.top, shape.width, shape.height
    except Exception:
        return None
    if None in (left, top, width, height):
        return None
    return int(left), int(top), int(width), int(height)


def _visible_box(shape: Any, slide: Any, prs: DeckClass) -> tuple[int, int, int, int] | None:
    """Where a shape's content really is: a short text in a tall box is only as tall as it is."""
    box = _box_of(shape)
    if box is None or not shape.has_text_frame:
        return box
    ctx = TextContext(shape, slide, prs)
    kind, scale, reduction = ctx.autofit()
    used = int(measure_text(ctx, box[2], scale=scale, reduction=reduction).height * EMU_PER_PT)
    if used >= box[3]:
        return box
    anchor = ctx.body_attr("anchor", "t")
    top = box[1]
    if anchor == "ctr":
        top += (box[3] - used) // 2
    elif anchor == "b":
        top += box[3] - used
    return box[0], top, box[2], used


def _visible_content(shape: Any) -> str | None:
    """"text", "table", "chart" or "picture"; None for what's only decoration."""
    if getattr(shape, "has_table", False) and shape.has_table:
        return "table"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "chart"
    if shape._element.tag == qn("p:pic"):
        return "picture"
    if shape.has_text_frame and _shape_text(shape):
        return "text"
    return None


def _table_height(shape: Any, slide: Any) -> int:
    """A row grows to hold its text, whatever the file says its height is."""
    table = shape.table
    font = _theme_fonts(slide.slide_layout.slide_master)[1]
    total = 0
    for r, row in enumerate(table.rows):
        tallest = int(row.height)
        for c, column in enumerate(table.columns):
            cell = table.cell(r, c)
            if cell.is_spanned:
                continue
            room = (int(column.width) - _CELL_MARGIN_X) / EMU_PER_PT
            lines = 0
            size = 18.0
            for para in cell._tc.txBody.findall(A_P):
                run = para.find(A_R)
                r_pr = run.find(A_RPR) if run is not None else para.find(A_END)
                if r_pr is not None and r_pr.get("sz"):
                    size = int(r_pr.get("sz")) / 100
                lines += _wrap_count([(paragraph_text(para), size, font, False)], room, room)
            tallest = max(tallest, int(lines * size * LINE_HEIGHT * EMU_PER_PT + _CELL_MARGIN_Y))
        total += tallest
    return total


def _table_width(shape: Any) -> int:
    """What the columns add up to: the frame's own width isn't what's drawn."""
    return sum(int(column.width) for column in shape.table.columns)


def fit_issues(prs: DeckClass, slide: Any, number: int) -> list[Issue]:
    """What doesn't fit, or looks unfinished, on a slide."""
    issues: list[Issue] = []
    slide_w, slide_h = prs.slide_width, prs.slide_height
    where = f"Slide {number}"
    contents: list[tuple[Any, str, tuple[int, int, int, int]]] = []
    for shape in slide.shapes:
        box = _box_of(shape)
        ref = f"{where} [{shape.shape_id}]"
        kind = placeholder_kind(shape)
        if kind in ("date", "footer", "slide number"):
            continue
        if shape.is_placeholder and shape.has_text_frame and not _shape_text(shape):
            issues.append(Issue("warn", "empty_placeholder", ref,
                                f'an empty {kind} placeholder: fill it or delete it, or people opening the file see '
                                '"Click to add".'))
            continue
        if box is not None:
            left, top, width, height = box
            if left + width <= 0 or top + height <= 0 or left >= slide_w or top >= slide_h:
                issues.append(Issue("warn", "off_slide", ref, f"{shape.name!r} is entirely off the slide."))
                continue
            margin = 18288  # a fiftieth of an inch of tolerance
            if left < -margin or top < -margin or left + width > slide_w + margin or top + height > slide_h + margin:
                if _visible_content(shape):
                    issues.append(Issue("warn", "off_slide", ref,
                                        f"{shape.name!r} runs off the slide ({inches(left)}, {inches(top)} size "
                                        f"{inches(width)} × {inches(height)} in on a {inches(slide_w)} × "
                                        f"{inches(slide_h)} in slide)."))
        if shape.has_text_frame and _shape_text(shape) and box is not None:
            issues += _text_issues(shape, slide, prs, ref, box)
        if getattr(shape, "has_table", False) and shape.has_table and box is not None:
            real = _table_height(shape, slide)
            if box[1] + real > slide_h + 18288:
                issues.append(Issue("warn", "overflow", ref,
                                    f"the table comes to about {inches(real)} in tall and runs off the bottom of the "
                                    "slide. Cut rows, split it, or use a smaller font_size."))
            wide = _table_width(shape)
            if box[0] + wide > slide_w + 18288:
                issues.append(Issue("warn", "overflow", ref,
                                    f"the table's columns come to {inches(wide)} in and run off the right of the "
                                    "slide. Narrow them or drop a column."))
        content = _visible_content(shape)
        if content is not None:
            shown = _visible_box(shape, slide, prs)
            if getattr(shape, "has_table", False) and shape.has_table and shown is not None:
                shown = (shown[0], shown[1], _table_width(shape), _table_height(shape, slide))
            if shown is not None:
                contents.append((shape, content, shown))
    issues += _overlaps(contents, where, slide_w, slide_h)
    if shape_has_no_title(slide):
        issues.append(Issue("warn", "no_title", where,
                            "no title: outline view, screen readers and navigation show it as a blank."))
    return issues


def shape_has_no_title(slide: Any) -> bool:
    if slide.slide_layout.name.casefold() == "blank":
        return False
    title = slide.shapes.title
    return title is None or not _shape_text(title)


def _text_issues(shape: Any, slide: Any, prs: DeckClass, ref: str, box: tuple[int, int, int, int]) -> list[Issue]:
    issues: list[Issue] = []
    ctx = TextContext(shape, slide, prs)
    kind, scale, reduction = ctx.autofit()
    fit = measure_text(ctx, box[2], scale=scale, reduction=reduction)
    room = box[3] / EMU_PER_PT
    if kind == "grow":
        bottom = box[1] + int(fit.height * EMU_PER_PT)
        if bottom > prs.slide_height + 18288:
            issues.append(Issue("warn", "overflow", ref,
                                f"the text grows to about {fit.height / 72:.1f} in and runs off the bottom of the slide."))
    elif fit.height > room * 1.03:
        if kind == "shrink" and scale >= 1.0:
            hint = "PowerPoint shrinks it only when someone edits the text"
        else:
            hint = "shorten it, split the slide, or set a smaller font_size"
        issues.append(Issue("warn", "overflow", ref,
                            f"about {fit.lines} lines at {fit.largest:g}pt need {fit.height / 72:.1f} in and the box is "
                            f"{room / 72:.1f} in: {hint}."))
    if fit.smallest and fit.smallest < 10:
        issues.append(Issue("warn", "small_text", ref, f"text at {fit.smallest:g}pt is hard to read on a slide."))
    return issues


def _overlaps(contents: list[tuple[Any, str, tuple[int, int, int, int]]], where: str,
              slide_w: int, slide_h: int) -> list[Issue]:
    issues: list[Issue] = []
    slide_area = slide_w * slide_h
    for i, (a, kind_a, box_a) in enumerate(contents):
        for b, kind_b, box_b in contents[i + 1:]:
            x = min(box_a[0] + box_a[2], box_b[0] + box_b[2]) - max(box_a[0], box_b[0])
            y = min(box_a[1] + box_a[3], box_b[1] + box_b[3]) - max(box_a[1], box_b[1])
            if x <= 45720 or y <= 45720:
                continue
            area_a, area_b = box_a[2] * box_a[3], box_b[2] * box_b[3]
            smaller = min(area_a, area_b)
            if not smaller or x * y < smaller * 0.03:
                continue
            # A picture filling most of the slide is a background.
            if (kind_a == "picture" and area_a > slide_area * 0.7) or (kind_b == "picture" and area_b > slide_area * 0.7):
                continue
            level = "info" if "picture" in (kind_a, kind_b) else "warn"
            issues.append(Issue(level, "overlap", f"{where} [{a.shape_id}] and [{b.shape_id}]",
                                f"{kind_a} and {kind_b} overlap by about {inches(x)} × {inches(y)} in."))
    return issues


_LEFTOVER = re.compile(r"lorem ipsum|click to add|\bxxx+\b|\btbd\b|\btodo\b|\[insert\b|<insert\b|\bplaceholder\b", re.I)


def check_deck(prs: DeckClass) -> tuple[list[Issue], dict[str, int]]:
    """What a careful reviewer looks for first. Deterministic: no model involved."""
    issues: list[Issue] = []
    titles: dict[str, list[int]] = {}
    fonts: set[str] = set()
    missing_alt: list[str] = []
    for number, slide in enumerate(prs.slides, 1):
        where = f"Slide {number}"
        issues += fit_issues(prs, slide, number)
        title = slide.shapes.title
        text = _shape_text(title).casefold() if title is not None else ""
        if text:
            titles.setdefault(text, []).append(number)
        words = 0
        for shape in _walk(slide.shapes):
            ref = f"{where} [{shape.shape_id}]"
            content = _visible_content(shape)
            if shape.has_text_frame:
                body = _shape_text(shape)
                if _LEFTOVER.search(body):
                    issues.append(Issue("warn", "leftover_text", ref, f"looks like leftover text: {body[:80]!r}."))
                if placeholder_kind(shape) not in ("date", "footer", "slide number"):
                    words += len(body.split())
                ctx = TextContext(shape, slide, prs)
                for para in ctx.tx_body.findall(A_P) if ctx.tx_body is not None else []:
                    p_pr = para.find(A_PPR)
                    level = int(p_pr.get("lvl", "0")) if p_pr is not None else 0
                    for run in para.findall(A_R):
                        fonts.add(ctx.font(level, run.find(A_RPR)))
            if content in ("picture", "chart", "table") and not _alt_text(shape):
                missing_alt.append(ref)
            if content == "picture":
                issues += _picture_issues(shape, ref)
        if words > 120:
            issues.append(Issue("info", "wordy", where, f"{words} words: a lot to read on one slide."))
        if slide._element.get("show") == "0":
            issues.append(Issue("info", "hidden", where, "the slide is hidden, but it's still in the file."))
        notes = notes_text(slide)
        if notes and _LEFTOVER.search(notes):
            issues.append(Issue("info", "leftover_text", f"{where} notes", f"looks like leftover text: {notes[:80]!r}."))
    for text, numbers in titles.items():
        if len(numbers) > 1:
            issues.append(Issue("info", "same_title", f"Slides {', '.join(map(str, numbers))}",
                                f"all titled {text[:60]!r}: hard to tell apart when navigating."))
    if len(fonts) > 3:
        issues.append(Issue("info", "many_fonts", "deck", f"{len(fonts)} fonts in use: {', '.join(sorted(fonts)[:8])}."))
    if missing_alt:
        issues.append(Issue("info", "no_alt_text", ", ".join(missing_alt[:4]) + (" …" if len(missing_alt) > 4 else ""),
                            f"{len(missing_alt)} picture(s), chart(s) or table(s) have no alt text for screen readers."))
    package = prs.part.package
    parts = [str(p.partname) for p in package.iter_parts()]
    if any(name.startswith(("/ppt/comments/", "/ppt/modernComments/")) for name in parts):
        issues.append(Issue("info", "comments", "deck", "it holds reviewers' comments: anyone you send it to can read them."))
    linked = [rel for slide in prs.slides for rel in slide.part.rels.values()
              if rel.is_external and rel.reltype != RT.HYPERLINK]
    if linked:
        issues.append(Issue("info", "links", "deck", f"{len(linked)} linked file(s) or picture(s) outside the deck "
                                                     "won't travel with it."))
    order = {"error": 0, "warn": 1, "info": 2}
    issues.sort(key=lambda i: order.get(i.level, 3))
    return issues, {"slides": len(prs.slides)}


def _picture_issues(shape: Any, ref: str) -> list[Issue]:
    try:
        px_w, px_h = shape.image.size
    except Exception:
        return []
    if not px_w or not px_h or not shape.width or not shape.height:
        return []
    shown_w = px_w * (1 - shape.crop_left - shape.crop_right)
    shown_h = px_h * (1 - shape.crop_top - shape.crop_bottom)
    if shown_w <= 0 or shown_h <= 0:
        return []
    drawn = shape.width / shape.height
    actual = shown_w / shown_h
    issues: list[Issue] = []
    if abs(drawn / actual - 1) > 0.03:
        issues.append(Issue("warn", "stretched_picture", ref,
                            f"the picture is stretched by {abs(drawn / actual - 1) * 100:.0f}%: its shape doesn't match "
                            "the image."))
    dpi = shown_w / (shape.width / EMU_PER_INCH)
    if dpi < 72:
        issues.append(Issue("info", "low_resolution", ref,
                            f"the picture has only about {dpi:.0f} pixels per inch at this size and will look soft."))
    return issues


# ── Rendering ────────────────────────────────────────────────────────


def find_renderer() -> str | None:
    """LibreOffice, where it's installed. Everything else here works without it."""
    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


def without_outside_links(data: bytes) -> bytes:
    """A copy with no linked pictures or objects: opening a stranger's deck on a
    server must not make the server fetch whatever it names."""
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for info in source.infolist():
            raw = source.read(info.filename)
            if info.filename.endswith(".rels"):
                root = etree.fromstring(raw)
                for rel in root:
                    if rel.get("TargetMode") == "External" and not rel.get("Type", "").endswith("/hyperlink"):
                        rel.set("Target", "about:blank")
                raw = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
            target.writestr(info, raw)
    return out.getvalue()


def render_pages(data: bytes, renderer: str, *, dpi: int = 110, timeout: float = 120.0) -> list[bytes]:
    """Each slide as a PNG, through LibreOffice and a PDF."""
    with tempfile.TemporaryDirectory(prefix="pptx-render-") as folder:
        root = Path(folder)
        deck = root / "deck.pptx"
        deck.write_bytes(without_outside_links(data))
        try:
            subprocess.run(
                [renderer, f"-env:UserInstallation=file://{root}/profile", "--headless", "--norestore",
                 "--convert-to", "pdf", "--outdir", str(root), str(deck)],
                check=True, capture_output=True, timeout=timeout)
        except subprocess.TimeoutExpired:
            raise PptxError(f"Rendering took over {timeout:g} seconds and was stopped.") from None
        except (subprocess.CalledProcessError, OSError) as exc:
            raise PptxError(f"LibreOffice couldn't render the deck ({type(exc).__name__}).") from None
        pdf = root / "deck.pdf"
        if not pdf.is_file():
            raise PptxError("LibreOffice produced no PDF for this deck.")
        return _pdf_to_pngs(pdf, dpi)


def _pdf_to_pngs(pdf: Path, dpi: int) -> list[bytes]:
    try:
        import pypdfium2

        document = pypdfium2.PdfDocument(str(pdf))
        images = []
        for index in range(len(document)):
            buffer = io.BytesIO()
            document[index].render(scale=dpi / 72).to_pil().save(buffer, format="PNG")
            images.append(buffer.getvalue())
        document.close()
        return images
    except ImportError:
        pass
    converter = shutil.which("pdftoppm")
    if converter is None:
        raise PptxError("Turning the PDF into pictures needs pypdfium2 or pdftoppm; neither is here.")
    subprocess.run([converter, "-png", "-r", str(dpi), str(pdf), str(pdf.with_suffix(""))],
                   check=True, capture_output=True, timeout=120)
    pages = sorted(pdf.parent.glob("deck-*.png"), key=lambda p: int(re.search(r"(\d+)\.png$", p.name).group(1)))
    return [page.read_bytes() for page in pages]

