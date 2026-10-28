"""Data library — every datastore table the caller's agents keep, profiled,
summarised and browsable.

The kiosk's "Data": the File library's sibling for the agents' datastores.
Each process agent keeps its tables in one SQLite file
(``<fd-data>/<slug>/data/home-config-parent/.captain-claw/datastore.db``:
``_ds_tables`` / ``_ds_columns`` describe the user tables, stored as
``ds_<name>``). Flight Deck opens it read-only from disk — stopped agents
count — and for every new or changed table (row count, last change, highest
row id or columns moved) works out a profile of each column (how full it is,
distinct values, ranges, the most common values, whether it holds links,
emails or dates) and asks the owner's fast tier for a title, a summary, a
category, tags, a few key facts and a line on each column. Results live in
``data_library.db``; the rows themselves are always read live, so the table
view is never stale.

Scans run like the File library's (one per user, started when the dialog
opens or by Rescan), on the same agents (the caller's own, minus the workers
Flight Deck runs itself), and spend from the same daily token allowance.
"""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import re
import sqlite3
import threading
import time
import urllib.parse
from contextlib import closing
from pathlib import Path
from typing import Any

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response

from captain_claw.flight_deck.auth import get_current_user
from captain_claw.flight_deck.file_library import (
    ScanJob,
    _ERROR_RETRY_SECONDS,
    _MAX_FAILURE_STREAK,
    _data_dir,
    _quota_spent,
    _resolve_tier,
    call_model,
    caller_agents,
    clean_tags,
    reply_json,
)
from captain_claw.logging import get_logger

log = get_logger(__name__)

router = APIRouter(prefix="/fd/data-library", tags=["data-library"])

# What a table is about: (id, label). The model picks exactly one.
CATEGORIES: list[tuple[str, str]] = [
    ("people", "People & contacts"),
    ("companies", "Companies & organisations"),
    ("leads", "Leads & deals"),
    ("finance", "Finance"),
    ("events", "Events & schedules"),
    ("products", "Products & inventory"),
    ("tasks", "Tasks & projects"),
    ("research", "Research data"),
    ("content", "Content & media"),
    ("metrics", "Metrics & logs"),
    ("other", "Other"),
]
_CATEGORY_IDS = {c for c, _ in CATEGORIES}

TABLE_PREFIX = "ds_"
# The datastore's hidden creator columns (shared agents) — never shown.
_HIDDEN_COLUMNS = {"_id", "_created_by", "_created_by_name"}
_MAX_TABLES_PER_AGENT = 200
_MAX_SUMMARIES_PER_SCAN = 150
# A profile reads at most this many rows of a table.
_PROFILE_ROWS = 50_000
_SAMPLE_HEAD, _SAMPLE_TAIL = 12, 3
_CELL_CHARS = 120
# What search keeps of a table's own text.
_SEARCH_ROWS = 400
_SEARCH_CHARS = 40_000
_PROMPT_CHARS = 9000
_MAX_TOKENS = 1200
_MAX_PAGE = 500
_MAX_EXPORT_ROWS = 200_000

# Values that only stand in for a missing one: they don't count as filled.
_PLACEHOLDERS = ("unknown", "tbd", "unknown/tbd", "n/a", "na", "none", "null", "-", "—", "?", "not found")
_URL_RE = re.compile(r"^https?://", re.I)
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}")


# ── Reading an agent's datastore ──────────────────────────────────────


def datastore_path(slug: str) -> Path | None:
    """The agent's datastore file, when it has one."""
    path = datastore_file(_data_dir() / slug)
    return path if path.is_file() else None


def datastore_file(agent_dir: Path) -> Path:
    """Where the agent in *agent_dir* keeps its datastore: the default under
    its home, or where its own config moves it — but only ever inside the
    agent's folder."""
    import yaml
    agent_dir = agent_dir.resolve()
    home = agent_dir / "data" / "home-config-parent"
    path = home / ".captain-claw" / "datastore.db"
    for cfg_file in (home / ".captain-claw" / "config.yaml", agent_dir / "config.yaml"):
        try:
            cfg = yaml.safe_load(cfg_file.read_text()) or {}
        except Exception:
            continue
        section = cfg.get("datastore") if isinstance(cfg, dict) else None
        raw = str(section.get("path") or "").strip() if isinstance(section, dict) else ""
        if raw:
            p = Path(str(home) + raw[1:]) if raw.startswith("~") else Path(raw)
            p = (p if p.is_absolute() else home / p).resolve()
            if p.is_relative_to(agent_dir):
                path = p
            break
    return path


def _casefold(v: Any) -> str | None:
    return None if v is None else str(v).casefold()


def _connect(path: Path) -> sqlite3.Connection:
    conn = sqlite3.connect(f"file:{urllib.parse.quote(str(path))}?mode=ro", uri=True,
                           timeout=5, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.create_function("cf", 1, _casefold, deterministic=True)
    return conn


def _q(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def _columns(c: sqlite3.Connection, table: str) -> list[dict]:
    """A table's visible columns in the datastore's order, with their types."""
    internal = TABLE_PREFIX + table
    physical = [r[1] for r in c.execute(f"PRAGMA table_info({_q(internal)})")]
    declared = {r[0]: (r[1], r[2]) for r in c.execute(
        "SELECT col_name, col_type, position FROM _ds_columns WHERE table_name = ?", (table,))}
    cols = [{"name": n, "type": declared.get(n, ("text", 0))[0] or "text"}
            for n in physical if n not in _HIDDEN_COLUMNS]
    cols.sort(key=lambda col: declared.get(col["name"], ("", 10_000))[1])
    return cols


def _cell(v: Any, limit: int = _CELL_CHARS) -> Any:
    if isinstance(v, bytes):
        return f"<binary, {len(v)} bytes>"
    if isinstance(v, str) and len(v) > limit:
        return v[:limit] + "…"
    return v


def _profile_column(c: sqlite3.Connection, internal: str, col: dict) -> dict:
    q = _q(col["name"])
    src = f"(SELECT {q} AS v FROM {_q(internal)} LIMIT {_PROFILE_ROWS})"
    placeholders = ", ".join("'" + x.replace("'", "''") + "'" for x in _PLACEHOLDERS)
    filled_sql = (f"v IS NOT NULL AND TRIM(CAST(v AS TEXT)) <> '' "
                  f"AND LOWER(TRIM(CAST(v AS TEXT))) NOT IN ({placeholders})")
    filled, distinct, numeric = c.execute(
        f"SELECT COUNT(CASE WHEN {filled_sql} THEN 1 END), COUNT(DISTINCT v), "
        f"COUNT(CASE WHEN typeof(v) IN ('integer', 'real') THEN 1 END) FROM {src}").fetchone()
    blank = c.execute(
        f"SELECT COUNT(*) FROM {src} WHERE v IS NOT NULL AND TRIM(CAST(v AS TEXT)) <> '' AND NOT ({filled_sql})"
    ).fetchone()[0]
    p: dict[str, Any] = {"name": col["name"], "type": col["type"], "filled": filled,
                         "distinct": distinct, "kind": "text"}
    if blank:
        p["placeholders"] = blank
    if col["type"] == "boolean":
        yes = c.execute(f"SELECT COUNT(*) FROM {src} WHERE v IN (1, '1', 'true', 'True', 'yes')").fetchone()[0]
        p.update(kind="boolean", true=yes, false=max(0, filled - yes))
        return p
    if filled and (col["type"] in ("integer", "real") or numeric >= 0.9 * filled):
        mn, mx, avg, total = c.execute(
            f"SELECT MIN(v), MAX(v), AVG(v), SUM(v) FROM {src} WHERE typeof(v) IN ('integer', 'real')").fetchone()
        p.update(kind="number", min=mn, max=mx, avg=round(avg, 4) if avg is not None else None, sum=total)
    else:
        sample = [str(r[0]) for r in c.execute(
            f"SELECT v FROM {src} WHERE {filled_sql} LIMIT 60")]
        if sample:
            share = lambda rx: sum(1 for s in sample if rx.match(s.strip())) / len(sample)  # noqa: E731
            if share(_URL_RE) >= 0.8:
                p["kind"] = "url"
            elif share(_EMAIL_RE) >= 0.8:
                p["kind"] = "email"
            elif col["type"] in ("date", "datetime") or share(_DATE_RE) >= 0.8:
                p["kind"] = "date"
                mn, mx = c.execute(f"SELECT MIN(v), MAX(v) FROM {src} WHERE {filled_sql}").fetchone()
                p.update(min=mn, max=mx)
            elif col["type"] == "json":
                p["kind"] = "json"
        avg_len = c.execute(f"SELECT AVG(LENGTH(CAST(v AS TEXT))) FROM {src} WHERE {filled_sql}").fetchone()[0]
        p["avg_length"] = round(avg_len, 1) if avg_len is not None else None
    # The most common values — only when something repeats, and for numbers
    # and dates only when they take a handful of values (not a continuum).
    top = [(_cell(r[0], 80), r[1]) for r in c.execute(
        f"SELECT CAST(v AS TEXT), COUNT(*) AS n FROM {src} WHERE v IS NOT NULL AND TRIM(CAST(v AS TEXT)) <> '' "
        "GROUP BY CAST(v AS TEXT) ORDER BY n DESC, 1 LIMIT 6")]
    few = distinct <= max(10, filled // 4)
    if (top and top[0][1] > 1 and p["kind"] not in ("url", "email")
            and (few or p["kind"] not in ("number", "date"))):
        p["top"] = [{"value": v, "count": n} for v, n in top if n > 1]
    return p


def scan_agent(slug: str, known: dict[str, str]) -> list[dict] | None:
    """The agent's tables now. A table whose fingerprint differs from
    *known* (table → fingerprint) also gets its profile, sample rows and
    search text. None when the agent has no datastore."""
    path = datastore_path(slug)
    if path is None:
        return None
    out: list[dict] = []
    with closing(_connect(path)) as c:
        try:
            meta = c.execute("SELECT name, updated_at FROM _ds_tables ORDER BY name").fetchall()
        except sqlite3.DatabaseError:
            return None
        physical = {r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type = 'table'")}
        for m in meta[:_MAX_TABLES_PER_AGENT]:
            name, internal = m["name"], TABLE_PREFIX + m["name"]
            if internal not in physical:
                continue
            try:
                cols = _columns(c, name)
                count, max_rowid = c.execute(f"SELECT COUNT(*), MAX(rowid) FROM {_q(internal)}").fetchone()
            except sqlite3.DatabaseError:
                continue
            sig = ",".join(f"{col['name']}:{col['type']}" for col in cols)
            fp = f"{count}:{m['updated_at']}:{max_rowid}:{hashlib.sha1(sig.encode()).hexdigest()[:8]}"
            t: dict[str, Any] = {"table": name, "columns": cols, "row_count": int(count or 0),
                                 "updated_at": m["updated_at"] or "", "fingerprint": fp}
            if known.get(name) != fp:
                try:
                    t.update(_details(c, internal, cols, int(count or 0)))
                except sqlite3.DatabaseError as e:
                    log.debug("data library: profile failed", table=name, error=str(e))
            out.append(t)
    return out


def _details(c: sqlite3.Connection, internal: str, cols: list[dict], count: int) -> dict:
    names = ", ".join(_q(col["name"]) for col in cols) or "rowid"
    head = [dict(r) for r in c.execute(f"SELECT {names} FROM {_q(internal)} ORDER BY rowid LIMIT {_SAMPLE_HEAD}")]
    tail: list[dict] = []
    if count > _SAMPLE_HEAD + _SAMPLE_TAIL:
        tail = [dict(r) for r in c.execute(
            f"SELECT {names} FROM {_q(internal)} ORDER BY rowid DESC LIMIT {_SAMPLE_TAIL}")][::-1]
    sample = [{k: _cell(v) for k, v in r.items()} for r in head + tail]
    parts: list[str] = []
    size = 0
    for r in c.execute(f"SELECT {names} FROM {_q(internal)} ORDER BY rowid LIMIT {_SEARCH_ROWS}"):
        line = " · ".join(str(v) for v in r if v is not None and not isinstance(v, bytes))
        parts.append(line)
        size += len(line)
        if size >= _SEARCH_CHARS:
            break
    return {
        "profile": [_profile_column(c, internal, col) for col in cols],
        "sample": sample,
        "search_text": "\n".join(parts)[:_SEARCH_CHARS],
    }


# ── Store ─────────────────────────────────────────────────────────────


class DataStore:
    """``data_library.db``: one row per agent table."""

    _JSON = ("columns", "profile", "sample", "tags", "insights", "column_notes")

    def __init__(self, db_path: Path) -> None:
        db_path.parent.mkdir(parents=True, exist_ok=True)
        self._lock = threading.RLock()
        self._conn = sqlite3.connect(str(db_path), check_same_thread=False)
        self._conn.row_factory = sqlite3.Row
        with self._lock:
            self._conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS library_tables (
                    id            TEXT PRIMARY KEY,
                    agent_slug    TEXT NOT NULL,
                    table_name    TEXT NOT NULL,
                    columns       TEXT NOT NULL DEFAULT '[]',
                    row_count     INTEGER NOT NULL DEFAULT 0,
                    updated_at    TEXT NOT NULL DEFAULT '',
                    fingerprint   TEXT NOT NULL DEFAULT '',
                    profile       TEXT NOT NULL DEFAULT '[]',
                    sample        TEXT NOT NULL DEFAULT '[]',
                    search_text   TEXT NOT NULL DEFAULT '',
                    status        TEXT NOT NULL DEFAULT 'pending',
                    title         TEXT NOT NULL DEFAULT '',
                    summary       TEXT NOT NULL DEFAULT '',
                    category      TEXT NOT NULL DEFAULT '',
                    tags          TEXT NOT NULL DEFAULT '[]',
                    insights      TEXT NOT NULL DEFAULT '[]',
                    column_notes  TEXT NOT NULL DEFAULT '{}',
                    error         TEXT NOT NULL DEFAULT '',
                    model         TEXT NOT NULL DEFAULT '',
                    tokens        INTEGER NOT NULL DEFAULT 0,
                    summarized_at REAL,
                    added_at      REAL NOT NULL DEFAULT 0
                );
                CREATE INDEX IF NOT EXISTS idx_data_agent ON library_tables(agent_slug, status);
                """
            )
            self._conn.commit()

    @staticmethod
    def table_id(slug: str, table: str) -> str:
        return hashlib.sha1(f"{slug}\0{table}".encode()).hexdigest()[:20]

    def _row(self, r: sqlite3.Row) -> dict:
        d = dict(r)
        for k in self._JSON:
            try:
                d[k] = json.loads(d.get(k) or ("{}" if k == "column_notes" else "[]"))
            except (TypeError, ValueError):
                d[k] = {} if k == "column_notes" else []
        return d

    def fingerprints(self, slug: str) -> dict[str, str]:
        with self._lock:
            return {r[0]: r[1] for r in self._conn.execute(
                "SELECT table_name, fingerprint FROM library_tables WHERE agent_slug = ?", (slug,))}

    def sync_agent(self, slug: str, tables: list[dict]) -> None:
        """New and changed tables become pending (with their fresh profile);
        dropped ones go."""
        now = time.time()
        with self._lock:
            have = self.fingerprints(slug)
            seen: set[str] = set()
            for t in tables:
                seen.add(t["table"])
                changed = have.get(t["table"]) != t["fingerprint"]
                if not changed:
                    continue
                details = (json.dumps(t.get("profile", []), ensure_ascii=False, default=str),
                           json.dumps(t.get("sample", []), ensure_ascii=False, default=str),
                           t.get("search_text", ""))
                if t["table"] not in have:
                    self._conn.execute(
                        "INSERT OR REPLACE INTO library_tables (id, agent_slug, table_name, columns, row_count, "
                        "updated_at, fingerprint, profile, sample, search_text, status, added_at) "
                        "VALUES (?,?,?,?,?,?,?,?,?,?, 'pending', ?)",
                        (self.table_id(slug, t["table"]), slug, t["table"], json.dumps(t["columns"]),
                         t["row_count"], t["updated_at"], t["fingerprint"], *details, now))
                else:
                    # Changed: the old summary shows until the new one is in.
                    self._conn.execute(
                        "UPDATE library_tables SET columns = ?, row_count = ?, updated_at = ?, fingerprint = ?, "
                        "profile = ?, sample = ?, search_text = ?, status = 'pending', error = '' "
                        "WHERE agent_slug = ? AND table_name = ?",
                        (json.dumps(t["columns"]), t["row_count"], t["updated_at"], t["fingerprint"],
                         *details, slug, t["table"]))
            for name in [n for n in have if n not in seen]:
                self._conn.execute(
                    "DELETE FROM library_tables WHERE agent_slug = ? AND table_name = ?", (slug, name))
            self._conn.commit()

    def drop_agents_except(self, slugs: set[str]) -> None:
        with self._lock:
            for (slug,) in self._conn.execute("SELECT DISTINCT agent_slug FROM library_tables").fetchall():
                if slug not in slugs:
                    self._conn.execute("DELETE FROM library_tables WHERE agent_slug = ?", (slug,))
            self._conn.commit()

    def drop_agent(self, slug: str) -> None:
        """Forget one agent's tables: it was removed for good."""
        with self._lock:
            self._conn.execute("DELETE FROM library_tables WHERE agent_slug = ?", (slug,))
            self._conn.commit()

    def pending(self, slugs: list[str], limit: int, retry_errors_before: float) -> list[dict]:
        if not slugs:
            return []
        with self._lock:
            rows = self._conn.execute(
                f"SELECT * FROM library_tables WHERE agent_slug IN ({','.join('?' * len(slugs))}) "
                "AND (status = 'pending' OR (status = 'error' AND COALESCE(summarized_at, 0) < ?)) "
                "ORDER BY updated_at DESC LIMIT ?", (*slugs, retry_errors_before, limit)).fetchall()
        return [self._row(r) for r in rows]

    def get(self, table_id: str) -> dict | None:
        with self._lock:
            r = self._conn.execute("SELECT * FROM library_tables WHERE id = ?", (table_id,)).fetchone()
        return self._row(r) if r else None

    def save_result(self, table_id: str, fingerprint: str, **fields: Any) -> None:
        for k in self._JSON:
            if k in fields and not isinstance(fields[k], str):
                fields[k] = json.dumps(fields[k], ensure_ascii=False)
        cols = ", ".join(f"{k} = ?" for k in fields)
        with self._lock:
            self._conn.execute(
                f"UPDATE library_tables SET {cols} WHERE id = ? AND fingerprint = ?",
                (*fields.values(), table_id, fingerprint))
            self._conn.commit()

    def counts(self, slugs: list[str]) -> dict[str, int]:
        out = {"total": 0, "done": 0, "pending": 0, "error": 0, "skipped": 0}
        if not slugs:
            return out
        with self._lock:
            for status, n in self._conn.execute(
                    f"SELECT status, COUNT(*) FROM library_tables WHERE agent_slug IN ({','.join('?' * len(slugs))}) "
                    "GROUP BY status", slugs):
                out[status] = out.get(status, 0) + int(n)
                out["total"] += int(n)
        return out

    def all_for(self, slugs: list[str]) -> list[dict]:
        if not slugs:
            return []
        with self._lock:
            rows = self._conn.execute(
                f"SELECT * FROM library_tables WHERE agent_slug IN ({','.join('?' * len(slugs))})", slugs).fetchall()
        return [self._row(r) for r in rows]


_STORE: DataStore | None = None
_STORE_LOCK = threading.Lock()


def get_store() -> DataStore:
    global _STORE
    with _STORE_LOCK:
        if _STORE is None:
            _STORE = DataStore(_data_dir() / "data_library.db")
        return _STORE


# ── The model ─────────────────────────────────────────────────────────


_SYSTEM = """You describe data tables that AI agents keep for the person they work for.
You get a table's name, its columns with statistics, and sample rows.
Reply with ONLY a JSON object, no prose, no code fence:
{"title": "...", "summary": "...", "category": "...", "tags": ["..."], "insights": ["..."], "columns": {"column_name": "..."}}

- title: a clear human title for what the table holds, at most 70 characters.
- summary: 2-3 sentences, at most 420 characters: what one row is, what the
  table covers, and what it is useful for. No "This table..." filler.
- category: exactly one of: %s
- tags: 2-5 short lowercase tags (topics, companies, places).
- insights: 2-4 short facts drawn from the statistics and rows (totals,
  ranges, the most common values, notable gaps such as mostly empty columns),
  each at most 140 characters. Use only numbers you were given.
- columns: for every column, at most 60 characters on what it holds.
Write in the language the data is written in."""


def _fmt(v: Any) -> str:
    if isinstance(v, float):
        return f"{v:,.2f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        return f"{v:,}"
    return str(v)


def _profile_line(p: dict, rows: int) -> str:
    bits = [f"{p['filled']}/{rows} filled", f"{p['distinct']} distinct"]
    if p.get("placeholders"):
        bits.append(f"{p['placeholders']} placeholder values like unknown/TBD")
    if p["kind"] == "number":
        bits.append(f"min {_fmt(p.get('min'))}, max {_fmt(p.get('max'))}, avg {_fmt(p.get('avg'))}, sum {_fmt(p.get('sum'))}")
    elif p["kind"] == "date":
        bits.append(f"from {p.get('min')} to {p.get('max')}")
    elif p["kind"] == "boolean":
        bits.append(f"{p.get('true', 0)} true, {p.get('false', 0)} false")
    elif p["kind"] in ("url", "email"):
        bits.append(f"{p['kind']}s")
    if p.get("top"):
        bits.append("most common: " + ", ".join(f"{t['value']} ({t['count']})" for t in p["top"][:5]))
    return f"- {p['name']} ({p['type']}): " + "; ".join(bits)


def _prompt(t: dict, agent_name: str) -> str:
    cols = [c["name"] for c in t["columns"]]
    lines = [f"Table: {t['table_name']}", f"Agent: {agent_name}",
             f"Rows: {t['row_count']} · Last changed: {(t['updated_at'] or '')[:10]}", "", "Columns:"]
    lines += [_profile_line(p, t["row_count"]) for p in t["profile"]]
    lines += ["", f"Sample rows (first {_SAMPLE_HEAD}" + (f", last {_SAMPLE_TAIL}" if t["row_count"] > _SAMPLE_HEAD + _SAMPLE_TAIL else "") + "), tab-separated:",
              "\t".join(cols)]
    for r in t["sample"]:
        lines.append("\t".join("" if r.get(c) is None else str(r.get(c)).replace("\t", " ").replace("\n", " ") for c in cols))
    text = "\n".join(lines)
    return text if len(text) <= _PROMPT_CHARS else text[:_PROMPT_CHARS] + "\n[… truncated]"


def _humanize(name: str) -> str:
    s = re.sub(r"[_-]+", " ", name)
    s = re.sub(r"\s*\b\d{8}\b", "", s).strip()
    return (s[:1].upper() + s[1:]) if s else name


async def summarize_table(t: dict, tier: dict, owner_id: str, agent_name: str) -> dict:
    """The columns to store for one table (never raises for an empty one)."""
    if t["row_count"] == 0:
        cols = ", ".join(c["name"] for c in t["columns"][:12])
        return {"status": "skipped", "title": _humanize(t["table_name"]),
                "summary": f"Empty table — no rows yet. Columns: {cols}." if cols else "Empty table.",
                "category": "other", "tags": [], "insights": [], "column_notes": {},
                "error": "", "model": "", "tokens": 0, "summarized_at": time.time()}
    content, tokens = await call_model(
        tier, _SYSTEM % ", ".join(c for c, _ in CATEGORIES), _prompt(t, agent_name), owner_id, _MAX_TOKENS)
    data = reply_json(content)
    category = str(data.get("category") or "").strip().lower()
    names = {c["name"] for c in t["columns"]}
    notes_raw = data.get("columns") if isinstance(data.get("columns"), dict) else {}
    insights_raw = data.get("insights") if isinstance(data.get("insights"), list) else []
    return {
        "status": "done",
        "title": " ".join(str(data.get("title") or "").split())[:120] or _humanize(t["table_name"]),
        "summary": " ".join(str(data.get("summary") or "").split())[:600],
        "category": category if category in _CATEGORY_IDS else "other",
        "tags": clean_tags(data.get("tags")),
        "insights": [" ".join(str(i).split())[:200] for i in insights_raw if str(i).strip()][:5],
        "column_notes": {k: " ".join(str(v).split())[:100] for k, v in notes_raw.items() if k in names},
        "error": "",
        "model": str(tier.get("model") or ""),
        "tokens": tokens,
        "summarized_at": time.time(),
    }


# ── Scans ─────────────────────────────────────────────────────────────


_JOBS: dict[str, ScanJob] = {}


def _job_for(uid: str) -> ScanJob:
    return _JOBS.setdefault(uid, ScanJob())


async def run_scan(user: dict, job: ScanJob, retry_errors: bool = False) -> None:
    uid = str(user.get("id") or "")
    store = get_store()
    try:
        agents = caller_agents(user)
        from captain_claw.flight_deck.server import _load_process_registry_strict
        try:
            store.drop_agents_except(set(_load_process_registry_strict().keys()))
        except (OSError, ValueError):
            pass
        for slug in agents:
            try:
                tables = await asyncio.to_thread(scan_agent, slug, store.fingerprints(slug))
            except sqlite3.Error as e:
                # Locked or unreadable right now: its rows stay as they were.
                log.debug("data library: datastore unreadable", agent=slug, error=str(e))
                continue
            store.sync_agent(slug, tables or [])
        todo = store.pending(list(agents), _MAX_SUMMARIES_PER_SCAN,
                             time.time() if retry_errors else time.time() - _ERROR_RETRY_SECONDS)
        if not todo:
            return
        tier = await _resolve_tier(uid)
        job.phase, job.total = "summarizing", len(todo)
        if tier is None:
            job.message = "No model is set up to write summaries — add one in your model settings."
            return
        if await _quota_spent(uid):
            job.message = "Today's token allowance is used up — summaries continue tomorrow."
            return
        sem = asyncio.Semaphore(3)
        streak = 0
        stop = asyncio.Event()

        async def one(t: dict) -> None:
            nonlocal streak
            async with sem:
                if stop.is_set():
                    return
                if await _quota_spent(uid):
                    job.message = "Today's token allowance ran out — the rest continue tomorrow."
                    stop.set()
                    return
                job.current = t["table_name"]
                try:
                    result = await summarize_table(t, tier, uid, agents.get(t["agent_slug"], {}).get("name", t["agent_slug"]))
                    streak = 0
                except Exception as e:  # noqa: BLE001 — recorded on the row
                    streak += 1
                    job.failed += 1
                    store.save_result(t["id"], t["fingerprint"], status="error",
                                      error=str(e)[:300], summarized_at=time.time())
                    if streak >= _MAX_FAILURE_STREAK:
                        job.message = f"Stopped: the model keeps failing ({str(e)[:160]})."
                        stop.set()
                    return
                store.save_result(t["id"], t["fingerprint"], **result)
                job.done += 1

        await asyncio.gather(*(one(t) for t in todo))
    except Exception as e:  # noqa: BLE001 — a scan never takes FD down
        log.warning("data library scan failed", user=uid, error=str(e))
        job.message = f"The scan failed: {str(e)[:200]}"
    finally:
        job.running, job.phase, job.current = False, "idle", ""
        job.finished_at = time.time()
        job.task = None


def start_scan(user: dict, retry_errors: bool = False) -> ScanJob:
    job = _job_for(str(user.get("id") or ""))
    if job.running:
        return job
    job.running, job.phase = True, "scanning"
    job.total = job.done = job.failed = 0
    job.current = job.message = ""
    job.started_at, job.finished_at = time.time(), 0.0
    job.task = asyncio.create_task(run_scan(user, job, retry_errors))
    return job


# ── Rows → what the UI shows ──────────────────────────────────────────


def _item(t: dict, agents: dict[str, dict], terms: list[str] | None = None, full: bool = False) -> dict:
    a = agents.get(t["agent_slug"], {})
    item = {
        "id": t["id"],
        "agent_slug": t["agent_slug"],
        "agent_name": a.get("name", t["agent_slug"]),
        "agent_running": bool(a.get("running")),
        "table_name": t["table_name"],
        "columns": t["columns"],
        "row_count": t["row_count"],
        "updated_at": t["updated_at"],
        "status": t["status"],
        "title": t["title"],
        "summary": t["summary"],
        "category": t["category"],
        "tags": t["tags"],
        "insights": t["insights"],
        "column_notes": t["column_notes"],
        "error": t["error"],
        "model": t["model"],
        "summarized_at": t["summarized_at"],
        "snippet": _snippet(t, terms) if terms else "",
    }
    if full:
        item["profile"] = t["profile"]
    return item


def _haystack(t: dict) -> str:
    cols = " ".join(c["name"] for c in t["columns"])
    return " ".join([t["table_name"], t["title"], t["summary"], " ".join(t["tags"]), t["category"],
                     cols, " ".join(t["insights"]), " ".join(t["column_notes"].values()),
                     t["search_text"]]).casefold()


def _snippet(t: dict, terms: list[str]) -> str:
    """The row text around the first term the card doesn't already show."""
    visible = " ".join([t["table_name"], t["title"], t["summary"], " ".join(t["tags"]),
                        " ".join(c["name"] for c in t["columns"])]).casefold()
    text = t["search_text"]
    low = text.casefold()
    for term in terms:
        if term in visible:
            continue
        i = low.find(term)
        if i < 0:
            continue
        a = text.rfind("\n", 0, i) + 1
        b = text.find("\n", i)
        b = len(text) if b < 0 else b
        a, b = max(a, i - 90), min(b, i + len(term) + 130)
        return ("…" if a else "") + " ".join(text[a:b].split()) + ("…" if b < len(text) else "")
    return ""


# sort → (key, newest/largest first)
_SORTS = {
    "recent": (lambda t: t["updated_at"] or "", True),
    "name": (lambda t: (t["title"] or t["table_name"]).lower(), False),
    "rows": (lambda t: int(t["row_count"] or 0), True),
    "agent": (lambda t: (t["agent_slug"], (t["title"] or t["table_name"]).lower()), False),
}


# ── Routes ────────────────────────────────────────────────────────────


def _status_payload(user: dict) -> dict:
    agents = caller_agents(user)
    return {"job": _job_for(str(user.get("id") or "")).view(),
            "counts": get_store().counts(list(agents)), "agents": len(agents)}


@router.get("/status")
async def data_status(user: dict = Depends(get_current_user)) -> dict:
    return _status_payload(user)


@router.post("/scan")
async def data_scan(retry_errors: bool = False, user: dict = Depends(get_current_user)) -> dict:
    """Look for new and changed tables and summarise them, in the background."""
    start_scan(user, retry_errors)
    return _status_payload(user)


@router.get("/tables")
async def data_tables(
    q: str = "", agent: str = "", category: str = "", sort: str = "recent",
    limit: int = 60, offset: int = 0, user: dict = Depends(get_current_user),
) -> dict:
    """The caller's tables, searched (names, summaries, columns and the data
    itself) and filtered, with each filter's counts under the others."""
    agents = caller_agents(user)
    terms = [t for t in q.casefold().split() if t][:8]
    rows = [t for t in get_store().all_for(list(agents))
            if not terms or all(term in _haystack(t) for term in terms)]

    def keep(t: dict, skip: str = "") -> bool:
        return ((skip == "agent" or not agent or t["agent_slug"] == agent)
                and (skip == "category" or not category or (t["category"] or "pending") == category))

    agent_counts: dict[str, int] = {}
    category_counts: dict[str, int] = {}
    for t in rows:
        if keep(t, "agent"):
            agent_counts[t["agent_slug"]] = agent_counts.get(t["agent_slug"], 0) + 1
        if keep(t, "category"):
            k = t["category"] or "pending"
            category_counts[k] = category_counts.get(k, 0) + 1
    key, newest_first = _SORTS.get(sort, _SORTS["recent"])
    shown = sorted((t for t in rows if keep(t)), key=key, reverse=newest_first)
    limit = max(1, min(limit, 1000))
    page = shown[max(0, offset):max(0, offset) + limit]
    return {
        "items": [_item(t, agents, terms) for t in page],
        "total": len(shown),
        "rows_total": sum(int(t["row_count"] or 0) for t in shown),
        "facets": {
            "agents": sorted(
                ({"id": s, "label": agents.get(s, {}).get("name", s), "count": n} for s, n in agent_counts.items()),
                key=lambda x: (-x["count"], x["label"].lower())),
            "categories": [{"id": c, "label": label, "count": category_counts[c]}
                           for c, label in [*CATEGORIES, ("pending", "Not summarised yet")]
                           if category_counts.get(c)],
        },
    }


def _owned(table_id: str, user: dict) -> tuple[dict, dict[str, dict]]:
    t = get_store().get(table_id)
    agents = caller_agents(user)
    if t is None or t["agent_slug"] not in agents:
        raise HTTPException(404, "table not found")
    return t, agents


@router.get("/tables/{table_id}")
async def data_table(table_id: str, user: dict = Depends(get_current_user)) -> dict:
    """One table with its column profile."""
    t, agents = _owned(table_id, user)
    return _item(t, agents, full=True)


def _open_table(t: dict) -> tuple[sqlite3.Connection, str, list[dict]]:
    path = datastore_path(t["agent_slug"])
    if path is None:
        raise HTTPException(404, "the agent's datastore is gone")
    c = _connect(path)
    internal = TABLE_PREFIX + t["table_name"]
    if not c.execute("SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?", (internal,)).fetchone():
        c.close()
        raise HTTPException(404, "the table is gone")
    return c, internal, _columns(c, t["table_name"])


def _read_rows(t: dict, limit: int, offset: int, q: str, sort: str, desc: bool) -> dict:
    c, internal, cols = _open_table(t)
    with closing(c):
        names = [col["name"] for col in cols]
        where: list[str] = []
        args: list[Any] = []
        for term in [x for x in q.casefold().split() if x][:6]:
            pat = "%" + term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
            where.append("(" + " OR ".join(f"cf({_q(n)}) LIKE ? ESCAPE '\\'" for n in names) + ")")
            args += [pat] * len(names)
        w = f"WHERE {' AND '.join(where)}" if where and names else ""
        if sort in names:
            col = next(x for x in cols if x["name"] == sort)
            coll = " COLLATE NOCASE" if col["type"] in ("text", "date", "datetime", "json") else ""
            order = f"({_q(sort)} IS NULL), {_q(sort)}{coll} {'DESC' if desc else 'ASC'}, rowid"
        else:
            order = "rowid DESC" if desc else "rowid"
        total = c.execute(f"SELECT COUNT(*) FROM {_q(internal)} {w}", args).fetchone()[0]
        select = ", ".join(["rowid AS __rowid"] + [_q(n) for n in names])
        rows = c.execute(f"SELECT {select} FROM {_q(internal)} {w} ORDER BY {order} LIMIT ? OFFSET ?",
                         (*args, limit, offset)).fetchall()
        return {
            "columns": cols,
            "rows": [{k: _plain(v) for k, v in dict(r).items()} for r in rows],
            "total": int(total),
        }


def _plain(v: Any) -> Any:
    """A cell as the table view and the exports show it: a blob by its size."""
    return f"<binary, {len(v)} bytes>" if isinstance(v, bytes) else v


@router.get("/tables/{table_id}/rows")
async def data_rows(
    table_id: str, limit: int = 50, offset: int = 0, q: str = "", sort: str = "", dir: str = "asc",
    user: dict = Depends(get_current_user),
) -> dict:
    """The table's rows, live from the agent's datastore: searched across
    every column, sorted by one, a page at a time."""
    t, _agents = _owned(table_id, user)
    try:
        return await asyncio.to_thread(
            _read_rows, t, max(1, min(limit, _MAX_PAGE)), max(0, offset), q, sort, dir == "desc")
    except sqlite3.Error as e:
        raise HTTPException(503, f"The datastore can't be read right now: {e}")


def _export_rows(c: sqlite3.Connection, internal: str, names: list[str], neutralize: bool = False):
    """What every export writes: the visible columns of each row, in the
    table's own order, up to the export cap.

    *neutralize* (the CSV export) does what the agent's own export does
    (J13): in a row a member of a shared agent created, a text cell that
    starts like a formula gets a leading ``'``, so a spreadsheet never runs
    it. The owner's rows, and a table with no creator column, stay as they
    are. The workbook needs none (every text cell is written as text) and
    nothing runs JSON."""
    if not names:
        return []
    stamped = neutralize and "_created_by" in {
        r[1] for r in c.execute(f"PRAGMA table_info({_q(internal)})")}
    select = [_q(n) for n in names] + (['"_created_by"'] if stamped else [])
    rows = c.execute(f"SELECT {', '.join(select)} FROM {_q(internal)} "
                     f"ORDER BY rowid LIMIT {_MAX_EXPORT_ROWS}")
    if not stamped:
        return rows
    from captain_claw.datastore import creator_dict, neutralize_rows

    return (neutralize_rows([list(r)[:-1]], [creator_dict(r[-1], "")], "members")[0] for r in rows)


def _csv(t: dict) -> bytes:
    c, internal, cols = _open_table(t)
    with closing(c):
        buf = io.StringIO()
        w = csv.writer(buf)
        names = [col["name"] for col in cols]
        w.writerow(names)
        for r in _export_rows(c, internal, names, neutralize=True):
            w.writerow(["" if v is None else _plain(v) for v in r])
        return ("﻿" + buf.getvalue()).encode("utf-8")


def _json(t: dict) -> bytes:
    c, internal, cols = _open_table(t)
    with closing(c):
        names = [col["name"] for col in cols]
        data = [{n: _plain(v) for n, v in zip(names, r)} for r in _export_rows(c, internal, names)]
    return json.dumps(data, ensure_ascii=False, indent=2).encode("utf-8")


# The most characters an Excel cell holds; a longer one makes Excel "repair"
# the whole workbook.
_XLSX_CELL_CHARS = 32_767


def _xlsx(t: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Font

    c, internal, cols = _open_table(t)
    with closing(c):
        names = [col["name"] for col in cols]
        wb = Workbook(write_only=True)
        ws = wb.create_sheet(t["table_name"][:31])  # Excel's cap on a sheet name
        ws.freeze_panes = "A2"
        bold = Font(bold=True)

        def cell(v: Any, font: Font | None = None) -> Any:
            v = _plain(v)
            if not isinstance(v, str):
                return v
            out = WriteOnlyCell(ws, value=ILLEGAL_CHARACTERS_RE.sub("", v)[:_XLSX_CELL_CHARS])
            # Text stays text: openpyxl would store "=…" as a formula Excel runs.
            out.data_type = "s"
            if font is not None:
                out.font = font
            return out

        ws.append([cell(n, bold) for n in names])
        for r in _export_rows(c, internal, names):
            ws.append([cell(v) for v in r])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# format → (writer, media type)
_EXPORTS = {
    "csv": (_csv, "text/csv; charset=utf-8"),
    "json": (_json, "application/json; charset=utf-8"),
    "xlsx": (_xlsx, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
}


@router.get("/tables/{table_id}/export")
async def data_export(table_id: str, format: str = "csv", user: dict = Depends(get_current_user)):
    """The whole table as CSV (opens straight in Excel), JSON (one object
    per row) or an Excel workbook."""
    fmt = format.lower()
    if fmt not in _EXPORTS:
        raise HTTPException(400, f"Unsupported format: {format}. Use csv, json or xlsx.")
    t, _agents = _owned(table_id, user)
    write, media_type = _EXPORTS[fmt]
    try:
        body = await asyncio.to_thread(write, t)
    except sqlite3.Error as e:
        raise HTTPException(503, f"The datastore can't be read right now: {e}")
    return Response(body, media_type=media_type, headers={
        "Content-Disposition": f'attachment; filename="{t["table_name"]}.{fmt}"',
        "X-Content-Type-Options": "nosniff",
    })


@router.post("/tables/{table_id}/summarize")
async def data_resummarize(table_id: str, user: dict = Depends(get_current_user)) -> dict:
    """Profile and summarise one table again, now."""
    t, agents = _owned(table_id, user)
    uid = str(user.get("id") or "")
    if await _quota_spent(uid):
        raise HTTPException(429, "Today's token allowance is used up.")
    tier = await _resolve_tier(uid)
    if tier is None:
        raise HTTPException(409, "No model is set up to write summaries.")
    store = get_store()
    # Fresh numbers first: the summary describes the table as it is now
    # (only this table is profiled again).
    known = store.fingerprints(t["agent_slug"])
    known.pop(t["table_name"], None)
    try:
        tables = await asyncio.to_thread(scan_agent, t["agent_slug"], known)
    except sqlite3.Error as e:
        raise HTTPException(503, f"The datastore can't be read right now: {e}")
    for x in tables or []:
        if x["table"] == t["table_name"] and "profile" in x:
            # Same fingerprint: sync leaves the row alone, so store the fresh profile here.
            store.save_result(t["id"], t["fingerprint"], profile=x["profile"], sample=x["sample"],
                              search_text=x["search_text"])
    store.sync_agent(t["agent_slug"], tables or [])
    fresh = store.get(table_id)
    if fresh is None:
        raise HTTPException(404, "the table is gone")
    try:
        result = await summarize_table(fresh, tier, uid, agents[fresh["agent_slug"]]["name"])
    except Exception as e:  # noqa: BLE001
        store.save_result(fresh["id"], fresh["fingerprint"], status="error",
                          error=str(e)[:300], summarized_at=time.time())
        raise HTTPException(502, f"The model failed: {str(e)[:200]}")
    store.save_result(fresh["id"], fresh["fingerprint"], **result)
    return _item(store.get(table_id) or fresh, agents, full=True)
