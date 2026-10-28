"""The spreadsheets of an agent's "way of working" setup, as datastore tables (see agent_preset).

A spreadsheet document (.xlsx, .csv) reaches the agent twice: its text as
Markdown in ``knowledge/``, like any document, and each sheet as a table in
the agent's own datastore — no model reads it — so prices, plans and codes
are exact lookups (the datastore tool's query and sql) instead of text to
search. A table is named from the document and its sheet (``price_list_plans``;
a CSV, or a workbook with one sheet, from the document alone) and never
after a table the agent already keeps for itself.

Values keep their kind: a number is a number, a date an ISO date, and text
stays text — a code like ``00123`` keeps its zeros. The header is the first
row about as wide as the sheet, so a title above it is left out.

A running agent gets its tables through its datastore REST
(``/api/datastore/tables…``); one that is off finds them in
``~/.captain-claw/pending-tables.json`` under its HOME and makes them when it
starts (``datastore.import_pending_tables``). Each table is made afresh —
dropped first if it is there — so a table given twice is the same table.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import json
import math
import os
import re
import sqlite3
import unicodedata
import urllib.parse
from collections.abc import Callable, Iterable, Sequence
from contextlib import closing
from pathlib import Path
from typing import Any

from captain_claw.logging import get_logger

log = get_logger(__name__)

TABLE_EXTS = (".xlsx", ".csv")
MAX_ROWS = 100_000          # the datastore's own max_rows_per_table
MAX_SHEETS = 20             # tables from one workbook
MAX_COLUMNS = 200
MAX_TABLES = 50             # the datastore's own max_tables, unless the agent's config says
_HEADER_ROWS = 20           # the header is looked for in a sheet's first rows
_NAME_MAX = 60
_BATCH = 2000               # rows in one REST call
_FOLD = str.maketrans({"đ": "dj", "Đ": "Dj", "ß": "ss", "æ": "ae", "Æ": "Ae", "ø": "o", "Ø": "O",
                       "ł": "l", "Ł": "L", "œ": "oe", "Œ": "Oe"})
_INT = re.compile(r"-?(?:0|[1-9]\d{0,17})")
_REAL = re.compile(r"-?(?:0|[1-9]\d*)\.\d+")
_DATE = re.compile(r"\d{4}-\d{2}-\d{2}")
_DATETIME = re.compile(r"\d{4}-\d{2}-\d{2}[ T]\d{2}:\d{2}(?::\d{2}(?:\.\d+)?)?")


def is_spreadsheet(name: str) -> bool:
    return Path(name or "").suffix.lower() in TABLE_EXTS


# ── Names ─────────────────────────────────────────────────────────────


def _safe(text: Any) -> str:
    """*text* as a datastore name — ``[a-z0-9_]``, as the datastore makes it
    (so it keeps what it is given) — or "" when nothing of it is left."""
    from captain_claw.datastore import DatastoreManager
    plain = unicodedata.normalize("NFKD", str(text or "").translate(_FOLD)).encode("ascii", "ignore").decode()
    if not re.search(r"[A-Za-z0-9]", plain):
        return ""
    return DatastoreManager._safe_name(plain)


def unique(name: str, taken: set[str]) -> str:
    out, n = name, 2
    while out in taken:
        out = f"{name}_{n}"
        n += 1
    taken.add(out)
    return out


def table_name(doc: str, sheet: str | None) -> str:
    """A table's name from its document and sheet (``price_list_2026_plans``)."""
    from captain_claw.datastore import DatastoreManager
    name = _safe(f"{Path(doc).stem} {sheet or ''}") or "spreadsheet"
    # Clipped, it is still a name the datastore keeps as it is.
    return DatastoreManager._safe_name(name[:_NAME_MAX].strip("_"))


# ── Reading a spreadsheet ─────────────────────────────────────────────


def _xlsx_value(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, bool):
        return int(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time() else v.isoformat(sep=" ")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, dt.timedelta):
        return str(v)
    if isinstance(v, int):
        return v if abs(v) < 2 ** 63 else str(v)   # SQLite's INTEGER
    if isinstance(v, float):
        return v if math.isfinite(v) else None
    v = str(v).strip()
    return v or None


def _csv_value(v: str) -> Any:
    v = (v or "").strip()
    if not v:
        return None
    if _INT.fullmatch(v):
        return int(v)
    if _REAL.fullmatch(v):
        return float(v)
    return v


def _filled(row: Sequence) -> int:
    return sum(v is not None for v in row)


def _width(row: Sequence) -> int:
    return max((i + 1 for i, v in enumerate(row) if v is not None), default=0)


def _type(values: list) -> str:
    seen = [v for v in values if v is not None]
    if not seen:
        return "text"
    if all(isinstance(v, int) for v in seen):
        return "integer"
    if all(isinstance(v, (int, float)) for v in seen):
        return "real"
    if all(isinstance(v, str) and _DATE.fullmatch(v) for v in seen):
        return "date"
    if all(isinstance(v, str) and (_DATE.fullmatch(v) or _DATETIME.fullmatch(v)) for v in seen):
        return "datetime"
    return "text"


def _table(rows: Iterable[Sequence], sheet: str | None, value: Callable[[Any], Any]) -> dict | None:
    """One sheet as a table: its header (the first row about as wide as the
    widest of the first rows, so a title above it is left out), the rows
    under it that have anything in them, and a type per column. None when
    it has no header with data under it."""
    it = iter(rows)
    head: list[list] = []
    for raw in it:
        vals = [value(v) for v in raw]
        if _filled(vals):
            head.append(vals)
            if len(head) >= _HEADER_ROWS:
                break
    if not head:
        return None
    widest = max(_filled(r) for r in head)
    at = next(i for i, r in enumerate(head) if _filled(r) >= max(1, math.ceil(widest * 0.6)))
    header, data = head[at], head[at + 1:]
    cut = False
    for raw in it:
        vals = [value(v) for v in raw]
        if not _filled(vals):
            continue
        if len(data) >= MAX_ROWS:
            cut = True
            break
        data.append(vals)
    if not data:
        return None
    width = min(MAX_COLUMNS, max(_width(r) for r in [header, *data]))
    data = [[r[i] if i < len(r) else None for i in range(width)] for r in data]
    labels = [header[i] if i < len(header) else None for i in range(width)]
    # A column with neither a heading nor a value is padding.
    keep = [i for i in range(width) if labels[i] is not None or any(r[i] is not None for r in data)]
    taken: set[str] = set()
    names = [unique(_safe(labels[i]) or f"column_{i + 1}", taken) for i in keep]
    rows_out = [[r[i] for i in keep] for r in data]
    columns = [{"name": n, "type": _type([r[j] for r in rows_out])} for j, n in enumerate(names)]
    return {"sheet": sheet, "columns": columns, "rows": rows_out, "cut": cut}


def _decode(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _csv_rows(path: Path) -> Iterable[list[str]]:
    text = _decode(path.read_bytes())
    try:
        dialect: Any = csv.Sniffer().sniff(text[:65536], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return csv.reader(io.StringIO(text), dialect)


def read_tables(path: Path) -> tuple[list[dict], list[str]]:
    """Each sheet of the spreadsheet at *path* that has a header with data
    under it (``sheet`` — None for a CSV — ``columns``, ``rows`` in the
    columns' order, ``cut``), and what was left out, said for the person."""
    ext = path.suffix.lower()
    notes: list[str] = []
    try:
        if ext == ".csv":
            got = _table(_csv_rows(path), None, _csv_value)
            return ([got] if got else []), notes
        if ext != ".xlsx":
            return [], notes
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            out: list[dict] = []
            for ws in wb.worksheets:
                if len(out) >= MAX_SHEETS:
                    notes.append(f"only its first {MAX_SHEETS} sheets with data became tables")
                    break
                got = _table(ws.iter_rows(values_only=True), ws.title, _xlsx_value)
                if got:
                    out.append(got)
            return out, notes
        finally:
            wb.close()
    except Exception as exc:
        log.warning("agent preset: a spreadsheet couldn't be read as tables", file=str(path), error=str(exc))
        return [], [f"it couldn't be read as a table ({exc})"]


def count_tables(path: Path) -> int:
    """How many tables the spreadsheet at *path* gives, roughly — its sheets
    with at least two rows (a header and data) — without reading it all."""
    def enough(rows: Iterable[Sequence]) -> bool:
        n = 0
        for raw in rows:
            if any(v not in (None, "") and str(v).strip() for v in raw):
                n += 1
                if n >= 2:
                    return True
        return False

    try:
        if path.suffix.lower() == ".csv":
            return int(enough(_csv_rows(path)))
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            return min(MAX_SHEETS, sum(enough(ws.iter_rows(values_only=True)) for ws in wb.worksheets))
        finally:
            wb.close()
    except Exception:
        return 0


# ── The agent's datastore ─────────────────────────────────────────────


def existing_tables(agent_dir: Path) -> set[str]:
    """The tables the agent's datastore has (read-only from disk: it may be
    off); none when it has no datastore yet."""
    from captain_claw.flight_deck.data_library import datastore_file
    path = datastore_file(agent_dir)
    if not path.is_file():
        return set()
    try:
        with closing(sqlite3.connect(f"file:{urllib.parse.quote(str(path))}?mode=ro", uri=True, timeout=5)) as c:
            return {str(r[0]) for r in c.execute("SELECT name FROM _ds_tables")}
    except sqlite3.Error:
        return set()


def max_tables(agent_dir: Path) -> int:
    """How many tables the agent's datastore takes: its config's
    ``datastore.max_tables``, else the datastore's own default."""
    import yaml
    cfg_file = agent_dir / "data" / "home-config-parent" / ".captain-claw" / "config.yaml"
    try:
        section = (yaml.safe_load(cfg_file.read_text(encoding="utf-8")) or {}).get("datastore") or {}
        return int(section.get("max_tables") or MAX_TABLES)
    except Exception:
        return MAX_TABLES


async def post_tables(port: int, token: str, drop: list[str], tables: list[dict]) -> tuple[set[str], set[str]]:
    """Through the agent's datastore REST: drop *drop*, then make each of
    *tables* afresh. The names dropped (one it doesn't have counts) and the
    tables made."""
    import httpx
    params = {"token": token} if token else {}
    base = f"http://127.0.0.1:{port}/api/datastore/tables"
    dropped: set[str] = set()
    made: set[str] = set()
    async with httpx.AsyncClient(timeout=httpx.Timeout(120.0, connect=5.0), trust_env=False) as client:
        async def gone(name: str) -> bool:
            resp = await client.delete(f"{base}/{name}", params=params)
            return resp.status_code < 300 or (resp.status_code == 400 and "not found" in resp.text.lower())

        for name in drop:
            try:
                if await gone(name):
                    dropped.add(name)
                else:
                    log.warning("agent preset: a table wasn't dropped", port=port, table=name)
            except Exception as exc:
                log.warning("agent preset: dropping a table failed", port=port, table=name, error=str(exc))
        for t in tables:
            name = t["name"]
            try:
                if not await gone(name):
                    log.warning("agent preset: a table in the way wasn't dropped", port=port, table=name)
                    continue
                resp = await client.post(base, params=params, json={"name": name, "columns": t["columns"]})
                if resp.status_code >= 300:
                    log.warning("agent preset: a table wasn't made", port=port, table=name, error=resp.text[:300])
                    continue
                cols = [c["name"] for c in t["columns"]]
                for i in range(0, len(t["rows"]), _BATCH):
                    resp = await client.post(f"{base}/{name}/rows", params=params, json={
                        "rows": [dict(zip(cols, row)) for row in t["rows"][i:i + _BATCH]]})
                    if resp.status_code >= 300:
                        raise RuntimeError(resp.text[:300])
                made.add(name)
            except Exception as exc:
                log.warning("agent preset: filling a table failed", port=port, table=name, error=str(exc))
    return dropped, made


# ── Tables waiting for an agent that is off ───────────────────────────


def pending_path(agent_dir: Path) -> Path:
    """Where tables wait for an agent that is off: ``~/.captain-claw`` under
    its HOME, which it reads when it starts (``datastore.import_pending_tables``)."""
    return agent_dir / "data" / "home-config-parent" / ".captain-claw" / "pending-tables.json"


def _read_pending(path: Path) -> dict:
    try:
        got = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}
    except (OSError, ValueError):
        got = {}
    got = got if isinstance(got, dict) else {}
    return {"drop": [n for n in got.get("drop") or [] if isinstance(n, str)],
            "tables": [t for t in got.get("tables") or [] if isinstance(t, dict) and t.get("name")]}


def _write_pending(path: Path, waiting: dict) -> None:
    if not waiting["drop"] and not waiting["tables"]:
        path.unlink(missing_ok=True)
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f"{path.name}.tmp")
    tmp.write_text(json.dumps(waiting, ensure_ascii=False, default=str), encoding="utf-8")
    os.replace(tmp, path)   # the agent never reads half a file


def leave_tables(agent_dir: Path, drop: list[str], tables: list[dict]) -> None:
    """Leave tables for the agent to make itself when it next starts — with
    what already waits: a table dropped now no longer waits, one given again
    waits as it is now."""
    if not drop and not tables:
        return
    path = pending_path(agent_dir)
    waiting = _read_pending(path)
    new = {t["name"] for t in tables}
    waiting["tables"] = [t for t in waiting["tables"] if t["name"] not in set(drop) | new] + [
        {"name": t["name"], "doc": t.get("doc") or "", "sheet": t.get("sheet"),
         "columns": t["columns"], "rows": t["rows"]} for t in tables]
    waiting["drop"] = list(dict.fromkeys([*waiting["drop"], *drop]))
    _write_pending(path, waiting)


def forget_pending(agent_dir: Path, names: set[str]) -> None:
    """Take *names* out of what waits for the agent's next start — the
    running agent has just been given them, so a stale drop or table there
    would undo that."""
    path = pending_path(agent_dir)
    if not names or not path.is_file():
        return
    waiting = _read_pending(path)
    waiting["drop"] = [n for n in waiting["drop"] if n not in names]
    waiting["tables"] = [t for t in waiting["tables"] if t["name"] not in names]
    _write_pending(path, waiting)
